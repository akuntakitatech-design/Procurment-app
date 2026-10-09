"""Penulis file export server-side Pusat Laporan: Excel (openpyxl write-only/streaming) & PDF (reportlab).

Keduanya menerima hasil `report_center.run()` yang SAMA dengan respon JSON UI (kolom, baris, total, filter) sehingga
isi UI = Excel = PDF. Kolom harga sudah dibuang sebelum sampai di sini bila pengguna tidak berizin.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NUM_FMT = {"qty": "#,##0.####", "money": "#,##0.00", "int": "#,##0"}
HEAD_FILL = PatternFill("solid", fgColor="E7ECF2")
THIN = Side(style="thin", color="B8C2CC")


def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def meta_lines(res):
    """Baris kop yang identik untuk Excel & PDF."""
    m = res["meta"]
    lines = [m["company"], m["title"], f"Kelompok: {m['group_title']}"]
    if m.get("date_basis"):
        lines.append(f"Dasar tanggal: {m['date_basis']}")
    lines.append("Filter: " + ("; ".join(f"{f['label']}: {f['value']}" for f in res["filters_applied"]) or "Semua data"))
    lines.append(f"Jumlah baris: {res['total_rows']:,}".replace(",", "."))
    lines.append(f"Dicetak oleh: {m['generated_by']} - {m['generated_at']} WIB")
    if not res["price_visible"]:
        lines.append("Catatan: kolom harga/nilai disembunyikan sesuai hak akses Anda.")
    return lines


def to_xlsx(res) -> BytesIO:
    cols = res["columns"]
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(res["meta"]["sheet"][:31] or "Laporan")
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(8, int(c["width"]))
    head = meta_lines(res)
    for i, line in enumerate(head):
        cell = WriteOnlyCell(ws, value=line)
        cell.font = Font(bold=i < 2, size=13 if i == 1 else 10)
        ws.append([cell])
    ws.append([])
    hdr_row = len(head) + 2
    hdr = []
    for c in cols:
        cell = WriteOnlyCell(ws, value=c["label"])
        cell.font, cell.fill = Font(bold=True), HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        hdr.append(cell)
    ws.append(hdr)
    ws.freeze_panes = f"A{hdr_row + 1}"
    fmts = [NUM_FMT.get(c["type"]) for c in cols]
    for r in res["rows"]:
        out = []
        for c, fmt in zip(cols, fmts):
            v = r.get(c["key"])
            if fmt:
                cell = WriteOnlyCell(ws, value=_num(v))
                cell.number_format = fmt
            else:
                cell = WriteOnlyCell(ws, value="" if v is None else str(v))
            out.append(cell)
        ws.append(out)
    last = hdr_row + len(res["rows"])
    if any(c["total"] for c in cols):
        tot = []
        for i, (c, fmt) in enumerate(zip(cols, fmts)):
            v = res["totals"].get(c["key"]) if c["total"] else ("TOTAL" if i == 0 else None)
            cell = WriteOnlyCell(ws, value=v)
            cell.font = Font(bold=True)
            if fmt and c["total"]:
                cell.number_format = fmt
            tot.append(cell)
        ws.append(tot)
    ws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(max(1, len(cols)))}{max(last, hdr_row)}"
    ps = wb.create_sheet("Parameter")
    ps.append(["Parameter", "Nilai"])
    for f in res["filters_applied"]:
        ps.append([f["label"], f["value"]])
    for k, lb in (("report_key", "Kode laporan"), ("generated_at", "Dibuat (WIB)"), ("generated_by", "Dibuat oleh")):
        ps.append([lb, res["meta"][k]])
    ps.append(["Jumlah baris", res["total_rows"]])
    ps.append(["Kolom harga", "Ditampilkan" if res["price_visible"] else "Disembunyikan (hak akses)"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _fmt(v, typ):
    if v is None or v == "":
        return ""
    if typ in NUM_FMT and isinstance(v, (int, float)):
        s = f"{v:,.2f}" if typ == "money" else (f"{v:,.0f}" if typ == "int" else f"{v:,.4f}".rstrip("0").rstrip("."))
        return s.replace(",", "#").replace(".", ",").replace("#", ".")
    return str(v)


def to_pdf(res) -> BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    cols = res["columns"]
    page = landscape(A3 if len(cols) > 14 else A4)   # laporan lebar -> A3 agar kolom tetap terbaca
    margin = 10 * mm
    avail = page[0] - 2 * margin
    wsum = sum(max(6, c["width"]) for c in cols) or 1
    widths = [avail * max(6, c["width"]) / wsum for c in cols]
    fs = 6.5 if len(cols) > 18 else 7.5
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=fs, leading=fs + 1.6)
    cell_r = ParagraphStyle("cellr", parent=cell, alignment=2)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    meta = meta_lines(res)

    def esc(s):
        return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    data = [[Paragraph(esc(c["label"]), head) for c in cols]]
    for r in res["rows"]:
        data.append([Paragraph(esc(_fmt(r.get(c["key"]), c["type"])), cell_r if c["type"] in NUM_FMT else cell)
                     for c in cols])
    has_total = any(c["total"] for c in cols)
    if has_total:
        data.append([Paragraph(esc(_fmt(res["totals"].get(c["key"]), c["type"]) if c["total"] else ("TOTAL" if i == 0 else "")),
                               ParagraphStyle("t", parent=cell_r if c["type"] in NUM_FMT else cell, fontName="Helvetica-Bold"))
                     for i, c in enumerate(cols)])
    tbl = LongTable(data, colWidths=widths, repeatRows=1)
    style = [("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B8C2CC")),
             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7ECF2")),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]
    if has_total:
        style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F3F5F8")))
    tbl.setStyle(TableStyle(style))

    class NumberedCanvas(rl_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._pages = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            n = len(self._pages)
            for st in self._pages:
                self.__dict__.update(st)
                self.setFont("Helvetica", 7)
                self.drawString(margin, 6 * mm, f"{res['meta']['title']} - {res['meta']['generated_at']} WIB")
                self.drawRightString(page[0] - margin, 6 * mm, f"Hal {self._pageNumber}/{n}")
                super().showPage()
            super().save()

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=margin, rightMargin=margin, topMargin=margin,
                            bottomMargin=12 * mm, title=res["meta"]["title"], author=res["meta"]["company"],
                            subject=f"rows={res['total_rows']}")
    story = [Paragraph(esc(meta[0]), ParagraphStyle("c", fontName="Helvetica-Bold", fontSize=10)),
             Paragraph(esc(meta[1]), ParagraphStyle("t1", fontName="Helvetica-Bold", fontSize=13, leading=16))]
    story += [Paragraph(esc(x), ParagraphStyle("m", fontName="Helvetica", fontSize=8, leading=10)) for x in meta[2:]]
    story += [Spacer(1, 4 * mm), tbl]
    doc.build(story, canvasmaker=NumberedCanvas)
    buf.seek(0)
    return buf
