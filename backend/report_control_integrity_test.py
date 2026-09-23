"""Disposable integration test for Report & Control division/warehouse scope."""
import os
import sys
import time
import uuid
from datetime import datetime, timezone
from io import BytesIO

import requests
from openpyxl import load_workbook

API = os.environ.get("TEST_API_URL", "http://report-control-api:8000/api").rstrip("/")
ADMIN_EMAIL = os.environ["ADMIN_EMAIL"]
ADMIN_PASSWORD = os.environ["ADMIN_PASSWORD"]


def check(ok, msg):
    if not ok:
        raise AssertionError(msg)
    print(f"PASS: {msg}")


def wait_api():
    base = API[:-4] if API.endswith("/api") else API
    for _ in range(90):
        try:
            if requests.get(f"{base}/docs", timeout=2).status_code == 200:
                return
        except requests.RequestException:
            pass
        time.sleep(1)
    raise RuntimeError("Backend test tidak siap")


def login(email, password):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=15)
    check(r.status_code == 200, f"Login {email} berhasil")
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {r.json()['token']}"})
    return s


def master(s, name, payload):
    r = s.post(f"{API}/master/{name}", json=payload, timeout=15)
    check(r.status_code == 200, f"Master {name} berhasil dibuat")
    return r.json()


def create_adjustment(s, div, wh, item, qty):
    r = s.post(f"{API}/adjustments", json={
        "date": datetime.now(timezone.utc).date().isoformat(),
        "warehouse_id": wh["id"], "division_id": div["id"],
        "adj_type": "Saldo Test", "reason": "REPORT CONTROL TEST",
        "lines": [{"item_id": item["id"], "adjustment": qty, "reason": "Saldo awal test"}],
    }, timeout=15)
    check(r.status_code == 200, f"Stok test {item['code']} berhasil dibuat")
    return r.json()


def create_mro(s, div, wh, item, note):
    r = s.post(f"{API}/mro", json={
        "date": datetime.now(timezone.utc).date().isoformat(),
        "need_date": datetime.now(timezone.utc).date().isoformat(),
        "division_id": div["id"], "requester": "Report Tester", "department": "Workshop",
        "default_warehouse_id": wh["id"], "submitted": True, "notes": note,
        "lines": [{"item_id": item["id"], "qty": 5, "unit": "pcs", "warehouse_id": wh["id"]}],
    }, timeout=15)
    check(r.status_code == 200, f"MRO {note} berhasil dibuat")
    return r.json()


def create_mi(s, div, wh, item, unit, note):
    r = s.post(f"{API}/mi", json={
        "date": datetime.now(timezone.utc).date().isoformat(),
        "division_id": div["id"], "default_warehouse_id": wh["id"],
        "receiver": "Report Tester", "department": "Workshop", "source_type": "Direct", "notes": note,
        "lines": [{"item_id": item["id"], "qty": 2, "unit": "pcs", "warehouse_id": wh["id"], "unit_id": unit["id"]}],
    }, timeout=15)
    check(r.status_code == 200, f"MI {note} berhasil dibuat")
    return r.json()


def main():
    wait_api()
    admin = login(ADMIN_EMAIL, ADMIN_PASSWORD)
    run = uuid.uuid4().hex[:8]
    common = f"RPT{run}"

    div_a = master(admin, "divisions", {"code": f"RA{run}", "name": f"Report A {run}"})
    div_b = master(admin, "divisions", {"code": f"RB{run}", "name": f"Report B {run}"})
    wh_a = master(admin, "warehouses", {"code": f"RWA{run}", "name": f"Gudang Report A {run}", "division_id": div_a["id"]})
    wh_b = master(admin, "warehouses", {"code": f"RWB{run}", "name": f"Gudang Report B {run}", "division_id": div_b["id"]})
    item_a = master(admin, "items", {"code": f"RIA{run}", "name": f"{common} Item A", "unit": "pcs", "division_id": div_a["id"]})
    item_b = master(admin, "items", {"code": f"RIB{run}", "name": f"{common} Item B", "unit": "pcs", "division_id": div_b["id"]})
    unit_a = master(admin, "units", {"code": f"RUA{run}", "name": f"{common} Unit A", "plate_no": f"A-{run}", "division_id": div_a["id"]})
    unit_b = master(admin, "units", {"code": f"RUB{run}", "name": f"{common} Unit B", "plate_no": f"B-{run}", "division_id": div_b["id"]})

    create_adjustment(admin, div_a, wh_a, item_a, 20)
    create_adjustment(admin, div_b, wh_b, item_b, 20)
    mro_a = create_mro(admin, div_a, wh_a, item_a, f"{common}-MRO-A")
    mro_b = create_mro(admin, div_b, wh_b, item_b, f"{common}-MRO-B")
    mi_a = create_mi(admin, div_a, wh_a, item_a, unit_a, f"{common}-MI-A")
    mi_b = create_mi(admin, div_b, wh_b, item_b, unit_b, f"{common}-MI-B")

    # Sanity: global admin sees both divisions before limited-user assertions.
    admin_trace = admin.get(f"{API}/reports/mro-traceability", timeout=20)
    check(admin_trace.status_code == 200, "Admin dapat membaca report traceability")
    admin_ids = {x.get("mro_id") for x in admin_trace.json()}
    check(mro_a["id"] in admin_ids and mro_b["id"] in admin_ids, "Data report dua divisi tersedia untuk admin")

    limited_email = f"report-limited-{run}@example.test"
    limited_password = f"Report-{run}-Pass!"
    user = admin.post(f"{API}/users", json={
        "email": limited_email, "password": limited_password, "name": "Report Limited",
        "role": "warehouse", "scope": "limited", "divisions": [div_a["id"]], "warehouses": [wh_a["id"]],
        "permissions": ["view", "export"], "is_active": True,
    }, timeout=15)
    check(user.status_code == 200, "User report terbatas berhasil dibuat")
    limited = login(limited_email, limited_password)

    dash = limited.get(f"{API}/dashboard", timeout=20)
    check(dash.status_code == 200, "Dashboard user terbatas dapat dibaca")
    dd = dash.json()
    check(dd.get("mro_open") == 1, "Dashboard hanya menghitung MRO divisi sendiri")
    check(dd.get("stock", {}).get("total") == 1, "Dashboard stok hanya menghitung gudang/divisi sendiri")
    check(dd.get("counts", {}).get("items") == 1, "Dashboard master item hanya menghitung divisi sendiri")

    premium = limited.get(f"{API}/dashboard-premium", timeout=20)
    check(premium.status_code == 200, "Dashboard premium user terbatas dapat dibaca")
    pd = premium.json()
    check(pd.get("financial_visible") is False, "Nilai pembelian tersembunyi tanpa izin harga")
    check(all(x.get("entity_id") != mro_b["id"] for x in (pd.get("recent") or [])), "Aktivitas premium tidak bocor dari divisi lain")

    pos = limited.get(f"{API}/inventory/position", timeout=15)
    check(pos.status_code == 200, "Inventory position dapat dibaca")
    pos_rows = pos.json()
    check(any(x.get("item_id") == item_a["id"] for x in pos_rows), "Inventory position memuat barang divisi sendiri")
    check(all(x.get("item_id") != item_b["id"] for x in pos_rows), "Inventory position menyembunyikan barang divisi lain")

    ledger = limited.get(f"{API}/inventory/ledger", timeout=15)
    check(ledger.status_code == 200, "Stock ledger dapat dibaca")
    ledger_rows = ledger.json()
    check(any(x.get("item_id") == item_a["id"] for x in ledger_rows), "Stock ledger memuat barang divisi sendiri")
    check(all(x.get("item_id") != item_b["id"] for x in ledger_rows), "Stock ledger menyembunyikan barang divisi lain")

    trace_a = limited.get(f"{API}/traceability/{mro_a['id']}", timeout=15)
    check(trace_a.status_code == 200, "Traceability MRO divisi sendiri dapat dibuka")
    trace_b = limited.get(f"{API}/traceability/{mro_b['id']}", timeout=15)
    check(trace_b.status_code == 403, "Traceability MRO divisi lain ditolak")

    trace = limited.get(f"{API}/reports/mro-traceability", timeout=20)
    check(trace.status_code == 200, "Report MRO traceability user terbatas dapat dibaca")
    trace_rows = trace.json()
    trace_ids = {x.get("mro_id") for x in trace_rows}
    check(mro_a["id"] in trace_ids and mro_b["id"] not in trace_ids, "Report traceability mengikuti scope divisi")
    own_row = next(x for x in trace_rows if x.get("mro_id") == mro_a["id"])
    check(own_row.get("po_gross") is None and own_row.get("po_total") is None, "Nilai finansial report disamarkan tanpa izin harga")

    export = limited.get(f"{API}/reports/mro-traceability/export.xlsx", timeout=30)
    check(export.status_code == 200 and export.content[:2] == b"PK", "Export Excel traceability berhasil")
    wb = load_workbook(BytesIO(export.content), read_only=True, data_only=True)
    ws = wb.active
    cell_text = "\n".join(str(c.value or "") for row in ws.iter_rows() for c in row)
    check(str(mro_a.get("no")) in cell_text and str(mro_b.get("no")) not in cell_text, "Export Excel tidak membocorkan divisi lain")

    lead = limited.get(f"{API}/reports/lead-time", timeout=20)
    check(lead.status_code == 200, "Lead time dapat dibaca")
    lead_nos = {x.get("mro_no") for x in lead.json()}
    check(mro_a.get("no") in lead_nos and mro_b.get("no") not in lead_nos, "Lead time mengikuti scope divisi")

    usage = limited.get(f"{API}/reports/unit-usage", timeout=20)
    check(usage.status_code == 200, "Pemakaian unit dapat dibaca")
    usage_rows = usage.json()
    check(any(x.get("unit_id") == unit_a["id"] for x in usage_rows), "Pemakaian unit divisi sendiri muncul")
    check(all(x.get("unit_id") != unit_b["id"] for x in usage_rows), "Pemakaian unit divisi lain tersembunyi")

    search = limited.get(f"{API}/search", params={"q": common}, timeout=20)
    check(search.status_code == 200, "Global search dapat digunakan")
    search_rows = search.json()
    blocked_ids = {mro_b["id"], item_b["id"], unit_b["id"], mi_b["id"]}
    check(not any(x.get("id") in blocked_ids for x in search_rows), "Global search tidak membocorkan divisi lain")
    check(any(x.get("id") in {mro_a["id"], item_a["id"], unit_a["id"], mi_a["id"]} for x in search_rows), "Global search tetap menemukan data divisi sendiri")

    audit = limited.get(f"{API}/audit", timeout=20)
    check(audit.status_code == 200, "Audit log dapat dibaca")
    audit_rows = audit.json()
    check(any(x.get("entity_id") == mro_a["id"] for x in audit_rows), "Audit log divisi sendiri terlihat")
    check(all(x.get("entity_id") != mro_b["id"] for x in audit_rows), "Audit log divisi lain tersembunyi")

    admin_notif = admin.get(f"{API}/notifications", timeout=15)
    check(admin_notif.status_code == 200, "Admin dapat membaca notifikasi test")
    notif_b = next((x for x in admin_notif.json() if str(x.get("division")) == str(div_b["id"]) and x.get("read") is False), None)
    check(notif_b is not None, "Notifikasi divisi B tersedia untuk pengujian")

    own_notif = limited.get(f"{API}/notifications", timeout=15)
    check(own_notif.status_code == 200, "Notifikasi user terbatas dapat dibaca")
    check(all(x.get("division") in (None, div_a["id"]) for x in own_notif.json()), "Daftar notifikasi mengikuti scope divisi")

    foreign_read = limited.post(f"{API}/notifications/{notif_b['id']}/read", timeout=15)
    check(foreign_read.status_code == 403, "User tidak dapat menandai notifikasi divisi lain")
    read_all = limited.post(f"{API}/notifications/read-all", timeout=15)
    check(read_all.status_code == 200, "Read-all notifikasi divisi sendiri berhasil")
    admin_notif_after = admin.get(f"{API}/notifications", timeout=15).json()
    b_after = next(x for x in admin_notif_after if x.get("id") == notif_b["id"])
    check(b_after.get("read") is False, "Read-all tidak mengubah notifikasi divisi lain")

    print("\nRESULT: PASS - Report & Control konsisten, finansial terlindungi, dan tidak bocor lintas divisi.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print(f"\nRESULT: FAIL - {type(exc).__name__}: {exc}", file=sys.stderr)
        raise
