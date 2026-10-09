"""Penyesuaian Stok multi-gudang per item — create/validasi/atomik/edit/hapus/legacy/akses/permission/draft lampiran/
Import Excel. Tenant QA sementara (paket business) — tidak menyentuh tenant lain / tidak reset DB.

Gudang A, B, D = divisi Asset; Gudang C = divisi Operasional; Gudang X = nonaktif.
"""
import datetime as _dt
import io
import json
import os
import sys
import uuid
from pathlib import Path

from openpyxl import Workbook, load_workbook

import receipt_control_test as T
import transfer_multi_warehouse_test as TM
from vendor_invoice_test import PW, U

call, check, API = T.call, T.check, T.API
u = uuid.uuid4().hex[:6]
m, stock, opening, db_conn = TM.m, TM.stock, TM.opening, TM.db_conn
_ENV = T.test_env()
STORAGE_LOCAL = (_ENV.get("STORAGE_DRIVER") or "local").lower() == "local"
LOCAL_PATH = Path(_ENV.get("STORAGE_LOCAL_PATH") or "/app/data/uploads")
TODAY = _dt.date.today().isoformat()


def sql(q, args=()):
    with db_conn() as cx, cx.cursor() as cur:
        cur.execute(q, args)
        return cur.fetchall()


def ledger(doc_id):
    _, led = call("GET", "inventory/ledger?limit=3000")
    return sorted((x.get("item_name"), x.get("warehouse_name"), float(x.get("qty_in") or 0), float(x.get("qty_out") or 0))
                  for x in led if x.get("doc_id") == doc_id and not x.get("is_reversal") and not x.get("reversed"))


def raw_ledger(doc_id):
    return [json.loads(d) for (d,) in sql("SELECT doc FROM stock_ledger WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.doc_id'))=%s ORDER BY id", (doc_id,))]


def vrows(doc_no=None):
    _, vl = call("GET", "reports/valuation-ledger")
    rows = vl.get("rows", []) if isinstance(vl, dict) else vl
    return [x for x in rows if doc_no is None or x.get("doc_no") == doc_no]


def wh_value():
    out = {}
    for x in vrows():
        k = f"{x.get('item_name')}@{x.get('warehouse_name')}"
        out[k] = round(out.get(k, 0.0) + float(x.get("value_in") or 0) - float(x.get("value_out") or 0), 4)
    return {k: v for k, v in out.items() if abs(v) > 1e-9}


def att_rows(entity_id):
    return [tuple(r) for r in sql("SELECT JSON_UNQUOTE(JSON_EXTRACT(doc,'$.entity')), JSON_UNQUOTE(JSON_EXTRACT(doc,'$.storage_path')) FROM attachments "
                                  "WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.entity_id'))=%s ORDER BY id", (entity_id,))]


def obj_exists(path):
    return (LOCAL_PATH / path).exists() if STORAGE_LOCAL and path else None


def n_ledger():
    return sql("SELECT COUNT(*) FROM stock_ledger")[0][0]


def upload(sess, draft_id, name, entity="adjustment_draft"):
    return sess.post(f"{API}/attachments", files={"file": (name, io.BytesIO(b"%PDF-1.4 " + name.encode()), "application/pdf")},
                     data={"entity": entity, "entity_id": draft_id, "category": "Lampiran Penyesuaian"})


def xlsx(rows, cols):
    wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return b


def main():
    TM.register(); other_hdr = dict(T.S.headers)
    T.S.headers.pop("Authorization", None); TM.register()
    asset = m("divisions", {"code": f"AS{u}", "name": "Asset"})
    ops = m("divisions", {"code": f"OP{u}", "name": "Operasional"})
    cat = m("item_categories", {"code": f"KC{u}", "name": "Umum"})
    uom = m("uoms", {"code": f"PC{u}", "name": "Pcs"})
    W = {k: m("warehouses", {"code": f"W{k}{u}", "name": f"Gudang {k}", "is_active": True,
                             "division_id": (ops if k == "C" else asset)["id"]}) for k in "ABCDX"}
    call("PUT", f"master/warehouses/{W['X']['id']}", {**W["X"], "is_active": False})
    P = {k: m("projects", {"code": f"P{k}{u}", "name": f"Project P{k}"}) for k in "123"}
    UN = {k: m("units", {"code": f"U{k}{u}", "name": f"Unit U{k}", "is_active": True}) for k in "12"}
    ref = {"category_id": cat["id"], "division_id": asset["id"], "base_uom_id": uom["id"], "unit": "PCS", "is_active": True}
    I = {n: m("items", {"code": f"I{n}{u}", "name": f"Item {n}", **ref}) for n in "ABC"}
    # Saldo awal via payload LAMA (gudang header saja) -> kompatibel ke belakang
    opening(W["A"]["id"], asset["id"], I["A"]["id"], 20, 1000)
    opening(W["B"]["id"], asset["id"], I["A"]["id"], 10, 1200)
    opening(W["D"]["id"], asset["id"], I["C"]["id"], 8, 3000)
    S = lambda n, w: stock(I[n]["id"], W[w]["id"])  # noqa: E731
    KEYS = ("A@A", "A@B", "B@C", "B@A", "C@D", "C@B")
    snap = lambda: {k: S(*k.split("@")) for k in KEYS}  # noqa: E731
    dlt = lambda a, b: {k: round(b[k] - a[k], 6) for k in a if abs(b[k] - a[k]) > 1e-9}  # noqa: E731
    head = {"date": TODAY, "division_id": asset["id"], "warehouse_id": W["A"]["id"], "project_id": P["1"]["id"],
            "adj_type": "Koreksi", "reason": "UAT multi gudang", "notes": "ADJ MULTI"}

    def L(item, qty, wh=None, p="__", un=None, cost=None, reason="QA"):
        ln = {"item_id": I[item]["id"], "adjustment": qty, "reason": reason}
        if wh: ln["warehouse_id"] = W[wh]["id"]
        if p != "__": ln["project_id"] = P[p]["id"] if p else None
        if un: ln["unit_id"] = UN[un]["id"]
        if cost is not None: ln["approved_unit_cost"] = cost
        return ln

    def n_docs():
        return len(call("GET", "adjustments")[1])

    # ================= A. Create multi gudang =================
    s0 = snap()
    sc, a1 = call("POST", "adjustments", {**head, "lines": [L("A", 5, "A", "1", "1"), L("A", -3, "B", "2"), L("B", 4, "C", "3", "2", cost=500),
                                                             L("C", -2, "D", None)]})
    check("A1. Create multi gudang (+A, -B, +C, -D) -> 200", sc == 200 and len(a1.get("lines", [])) == 4, a1)
    s1 = snap()
    check("A2. Stok per gudang LINE: A@A +5, A@B -3, B@C +4, C@D -2", dlt(s0, s1) == {"A@A": 5, "A@B": -3, "B@C": 4, "C@D": -2}, dlt(s0, s1))
    check("A3. Ledger per LINE (bukan gudang header)", ledger(a1["id"]) == sorted([("Item A", "Gudang A", 5, 0), ("Item A", "Gudang B", 0, 3),
                                                                            ("Item B", "Gudang C", 4, 0), ("Item C", "Gudang D", 0, 2)]), ledger(a1["id"]))
    ln = {(x["item_name"], x["warehouse_name"]): x for x in a1["lines"]}
    check("A4. Detail per baris: gudang/project/unit sesuai line", (ln[("Item A", "Gudang A")]["project_name"], ln[("Item A", "Gudang A")]["unit_name"]) == ("Project P1", "Unit U1")
          and ln[("Item A", "Gudang B")]["project_name"] == "Project P2" and ln[("Item B", "Gudang C")]["unit_name"] == "Unit U2" and ln[("Item C", "Gudang D")]["project_name"] is None, a1["lines"])
    check("A5. Stok Sebelum/Sesudah per gudang LINE", (ln[("Item A", "Gudang A")]["before"], ln[("Item A", "Gudang A")]["after"]) == (20, 25)
          and (ln[("Item A", "Gudang B")]["before"], ln[("Item A", "Gudang B")]["after"]) == (10, 7) and (ln[("Item B", "Gudang C")]["before"], ln[("Item B", "Gudang C")]["after"]) == (0, 4))
    hd = sql("SELECT JSON_UNQUOTE(JSON_EXTRACT(doc,'$.warehouse_source')), JSON_EXTRACT(doc,'$.line_warehouse_ids') FROM adjustments WHERE id=%s", (a1["id"],))[0]
    check("A6. Header: warehouse_source=line + line_warehouse_ids dihitung server", hd[0] == "line" and sorted(json.loads(hd[1])) == sorted(W[k]["id"] for k in "ABCD"), hd)
    rl = raw_ledger(a1["id"])
    pj = {(r["item_id"], r["warehouse_id"]): (r.get("project_id"), r.get("unit_id")) for r in rl}
    check("A7. Ledger membawa project/unit per baris", pj[(I["A"]["id"], W["A"]["id"])] == (P["1"]["id"], UN["1"]["id"]) and pj[(I["B"]["id"], W["C"]["id"])] == (P["3"]["id"], UN["2"]["id"]), pj)
    vv = {(x["item_name"], x["warehouse_name"]): (float(x.get("value_in") or 0), float(x.get("value_out") or 0)) for x in vrows(a1["no"])}
    check("A8. Valuasi per gudang LINE: + pakai MWA gudang (A@A 5x1000), - pakai MWA gudang (A@B 3x1200), + tanpa avg pakai biaya disetujui (4x500)",
          vv == {("Item A", "Gudang A"): (5000, 0), ("Item A", "Gudang B"): (0, 3600), ("Item B", "Gudang C"): (2000, 0), ("Item C", "Gudang D"): (0, 6000)}, vv)
    sc, lst = call("GET", "adjustments")
    row = next(x for x in lst if x["id"] == a1["id"])
    check("A9. List: ringkasan gudang dari LINE ('Gudang A +3') + multi_warehouse", row["warehouse_name"] == "Gudang A +3" and row["multi_warehouse"] is True and row["line_count"] == 4, row)

    # B. Payload lama: baris tanpa gudang -> Gudang Default; project key absen -> Project Default
    sc, b1 = call("POST", "adjustments", {**head, "lines": [L("A", 1)]})
    check("B1. Payload lama (gudang header saja) -> baris memakai Gudang Default + Project Default", sc == 200 and b1["lines"][0]["warehouse_name"] == "Gudang A"
          and b1["lines"][0]["project_name"] == "Project P1" and b1["lines"][0]["legacy_header_warehouse"] is False, b1)
    sc, b2 = call("POST", "adjustments", {**head, "lines": [L("A", 1, "A", None)]})
    check("B2. Override project kosong di baris dihormati (tidak ditimpa default)", sc == 200 and b2["lines"][0]["project_id"] is None, b2["lines"])

    # ================= C. Validasi sebelum penulisan =================
    def blocked(name, body, code=400, needle=None):
        nd, nl, sn = n_docs(), n_ledger(), snap()
        sc_, r_ = call("POST", "adjustments", body)
        ok = sc_ == code and (needle is None or needle.lower() in str(r_.get("detail")).lower())
        check(name, ok and n_docs() == nd and n_ledger() == nl and snap() == sn, (sc_, r_))

    blocked("C1. Qty 0 -> 400 tanpa penulisan", {**head, "lines": [L("A", 0, "A")]}, 400, "tidak boleh 0")
    blocked("C2. Minus melebihi stok gudang LINE (A@B 7) -> 400 tanpa penulisan", {**head, "lines": [L("A", 1, "A"), L("A", -8, "B")]}, 400, "melebihi stok gudang b")
    blocked("C3. Barang+gudang sama diakumulasi (-4 & -4, stok 7) -> 400", {**head, "lines": [L("A", -4, "B"), L("A", -4, "B")]}, 400, "baris 2")
    sc, c4 = call("POST", "adjustments", {**head, "lines": [L("C", 3, "B", cost=2500), L("C", -3, "B")]})
    check("C4. Urutan berurutan +3 lalu -3 di gudang tanpa stok -> 200 (simulasi per urutan posting)", sc == 200, c4)
    blocked("C5. Urutan -3 lalu +3 (stok 0) -> 400 (tidak lolos karena net)", {**head, "lines": [L("C", -3, "B"), L("C", 3, "B", cost=2500)]}, 400, "baris 1")
    blocked("C6. Qty + tanpa rata-rata & tanpa biaya (baris 2) -> 400 atomik (baris 1 tidak terposting)",
            {**head, "lines": [L("A", 2, "A"), L("B", 2, "D")]}, 400, "approved unit cost wajib")
    blocked("C7. Biaya beda rata-rata tanpa alasan -> 400", {**head, "lines": [L("A", 2, "A", cost=999, reason="")]}, 400, "alasan wajib")
    blocked("C8. Gudang baris nonaktif -> 400", {**head, "lines": [L("A", 1, "X")]}, 400, "tidak ditemukan atau tidak aktif")
    blocked("C9. Gudang baris id tenant lain/tidak ada -> 400", {**head, "lines": [{**L("A", 1), "warehouse_id": str(uuid.uuid4())}]}, 400, "gudang")
    blocked("C10. Project baris tidak ada -> 400", {**head, "lines": [{**L("A", 1, "A"), "project_id": str(uuid.uuid4())}]}, 400, "project")
    blocked("C11. Header tanpa Gudang Default -> 400", {**head, "warehouse_id": None, "lines": [L("A", 1, "A")]}, 400, "gudang adjustment wajib")

    # ================= E. Edit =================
    sc, e1 = call("POST", "adjustments", {**head, "lines": [L("A", -3, "B", "1"), L("A", 2, "A", "1")]})
    se, we = snap(), wh_value()
    nl = n_ledger(); lines_before = sql("SELECT doc FROM adjustment_lines WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.adjustment_id'))=%s", (e1["id"],))
    sc, r = call("PUT", f"transactions/adjustment/{e1['id']}", {**head, "lines": [L("A", -50, "B")]})
    check("E1. Edit tidak valid (-50 > stok+kredit reversal) -> 400 SEBELUM reversal: stok/ledger/baris tidak berubah",
          sc == 400 and snap() == se and n_ledger() == nl and sql("SELECT doc FROM adjustment_lines WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.adjustment_id'))=%s", (e1["id"],)) == lines_before, (sc, r))
    sc, r = call("PUT", f"transactions/adjustment/{e1['id']}", {**head, "lines": [L("B", 3, "D")]})
    check("E2. Edit dengan qty + tanpa rata-rata & tanpa biaya -> 400 sebelum reversal", sc == 400 and snap() == se and n_ledger() == nl, (sc, r))
    avail_b = S("A", "B")
    sc, r = call("PUT", f"transactions/adjustment/{e1['id']}", {**head, "lines": [L("A", -(avail_b + 3), "B", "2"), L("A", 4, "D", "3", "1", cost=1100)]})
    se2 = snap()
    check("E3. Edit memakai kredit reversal (minus = stok + 3 lama) -> 200", sc == 200, (sc, r))
    check("E3. Reversal baris ASLI (A@B +3, A@A -2) lalu posting baris BARU (A@B -(stok+3), A@D +4)",
          dlt(se, se2) == {"A@A": -2, "A@B": -avail_b} and S("A", "D") == 4, (dlt(se, se2), S("A", "D")))
    check("E4. Ledger aktif hanya baris baru", ledger(e1["id"]) == sorted([("Item A", "Gudang B", 0, avail_b + 3), ("Item A", "Gudang D", 4, 0)]), ledger(e1["id"]))
    sc, d = call("GET", f"adjustments/{e1['id']}")
    dl = {x["warehouse_name"]: x for x in d["lines"]}
    check("E5. Detail setelah edit: gudang/project/unit/before/after baris baru", dl["Gudang D"]["project_name"] == "Project P3" and dl["Gudang D"]["unit_name"] == "Unit U1"
          and (dl["Gudang D"]["before"], dl["Gudang D"]["after"]) == (0, 4) and dl["Gudang B"]["after"] == 0, d["lines"])
    hd = sql("SELECT JSON_EXTRACT(doc,'$.line_warehouse_ids') FROM adjustments WHERE id=%s", (e1["id"],))[0][0]
    check("E6. line_warehouse_ids dihitung ulang server-side saat edit", sorted(json.loads(hd)) == sorted([W["B"]["id"], W["D"]["id"]]), hd)
    wv = wh_value()
    check("E7. Valuasi edit: A@D masuk 4x1100, tanpa nilai ganda di A@A", wv.get("Item A@Gudang D") == 4400 and wv.get("Item A@Gudang A") == round(we.get("Item A@Gudang A", 0) - 2000, 4), (we, wv))

    # ================= F. Hapus =================
    sc, f1 = call("POST", "adjustments", {**head, "lines": [L("A", -1, "A"), L("C", -1, "D")]})
    sf = snap()
    sc, r = call("DELETE", f"transactions/adjustment/{f1['id']}")
    check("F1. Hapus -> reversal ke gudang LINE asli (A@A +1, C@D +1)", sc == 200 and dlt(sf, snap()) == {"A@A": 1, "C@D": 1}, (sc, r, dlt(sf, snap())))
    sc, f2 = call("POST", "adjustments", {**head, "lines": [L("B", 2, "A", cost=700)]})
    call("POST", "adjustments", {**head, "lines": [L("B", -2, "A")]})  # stok hasil f2 terpakai habis
    sc, cap = call("GET", f"transactions/adjustment/{f2['id']}/capability")
    sc2, r = call("DELETE", f"transactions/adjustment/{f2['id']}")
    check("F2. Guard existing: stok hasil penyesuaian + sudah terpakai -> edit/hapus dikunci (409)", cap.get("can_delete") is False and sc2 == 409, (cap, sc2))

    # ================= G. Legacy =================
    sc, lg = call("POST", "adjustments", {**head, "warehouse_id": W["B"]["id"], "project_id": P["2"]["id"], "lines": [L("A", 2, "B", "2", cost=1200), L("A", -1, "B", "2")]})
    for x in lg["lines"]:
        sql("UPDATE adjustment_lines SET doc = JSON_REMOVE(doc, '$.warehouse_id', '$.project_id', '$.unit_id') WHERE id=%s", (x["id"],))
    sql("UPDATE adjustments SET doc = JSON_REMOVE(doc, '$.warehouse_source', '$.line_warehouse_ids') WHERE id=%s", (lg["id"],))
    led_before = raw_ledger(lg["id"])
    sc, d = call("GET", f"adjustments/{lg['id']}")
    check("G1. Legacy: detail fallback header per baris (Gudang B, Project P2) + flag legacy", sc == 200 and all(x["warehouse_name"] == "Gudang B" and x["project_name"] == "Project P2"
          and x["legacy_header_warehouse"] for x in d["lines"]) and d["legacy_header_warehouse"] is True, d["lines"])
    raw = sql("SELECT JSON_EXTRACT(doc,'$.warehouse_id') FROM adjustment_lines WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.adjustment_id'))=%s", (lg["id"],))
    check("G2. Tanpa backfill: line legacy tetap tanpa gudang setelah dibaca; ledger historis identik", all(r_[0] is None for r_ in raw) and raw_ledger(lg["id"]) == led_before, raw)
    sc, lst = call("GET", "adjustments")
    check("G3. List legacy: gudang fallback header", next(x for x in lst if x["id"] == lg["id"])["warehouse_name"] == "Gudang B")
    # Mixed legacy: satu baris punya gudang line, satu tidak
    sc, mx = call("POST", "adjustments", {**head, "lines": [L("A", 1, "D"), L("A", 1, "A")]})
    sql("UPDATE adjustment_lines SET doc = JSON_REMOVE(doc, '$.warehouse_id') WHERE id=%s", (mx["lines"][1]["id"],))
    sql("UPDATE adjustments SET doc = JSON_REMOVE(doc, '$.warehouse_source', '$.line_warehouse_ids') WHERE id=%s", (mx["id"],))
    sc, d = call("GET", f"adjustments/{mx['id']}")
    check("G4. Mixed legacy diputuskan per baris: line1 Gudang D (line), line2 Gudang A (header)",
          [(x["warehouse_name"], x["legacy_header_warehouse"]) for x in d["lines"]] == [("Gudang D", False), ("Gudang A", True)], d["lines"])
    sl = snap()
    sc, r = call("DELETE", f"transactions/adjustment/{lg['id']}")
    check("G5. Hapus legacy: reversal dari ledger asli (Gudang B): A@B -1", sc == 200 and dlt(sl, snap()) == {"A@B": -1}, (sc, r, dlt(sl, snap())))
    sl = snap(); ad0 = S("A", "D")
    sc, r = call("PUT", f"transactions/adjustment/{mx['id']}", {**head, "lines": [L("A", 1, "B", cost=1200), L("A", 1, "A")]})
    check("G6. Edit legacy campuran: reversal ledger asli (D, A) lalu posting baris baru (B, A)", sc == 200 and dlt(sl, snap()) == {"A@B": 1} and S("A", "D") == ad0 - 1, (sc, r, dlt(sl, snap())))

    # ================= H. Akses / tenant =================
    email = f"adj_{u}@example.com"
    _, usr = call("POST", "users", {"email": email, "password": PW, "name": "QA Adj Asset", "role": "warehouse"}, 200)
    call("PUT", f"access/users/{usr['id']}", {"overrides": {"adjustment.create": "allow", "adjustment.post": "allow"}, "division_override": {"mode": "selected", "divisions": [asset["id"]]}}, 200)
    C = U(email)
    sc, lst = C("GET", "adjustments")
    ids = {x["id"] for x in lst} if isinstance(lst, list) else set()
    check("H1. User divisi Asset: dokumen berisi gudang divisi lain (C) TIDAK tampil di list (meski header Gudang A)", a1["id"] not in ids and b1["id"] in ids, (sc, len(ids)))
    sc, r = C("GET", f"adjustments/{a1['id']}")
    check("H2. User divisi Asset: detail dokumen lintas divisi -> 403", sc == 403, sc)
    check("H3. Multi gudang dalam satu divisi (B & D, Asset) tetap terlihat", e1["id"] in ids)
    sn = snap()
    sc, r = C("POST", "adjustments", {**head, "lines": [L("A", 1, "A"), L("B", 1, "C", cost=500)]})
    check("H4. User Asset: create dengan baris gudang divisi lain (C) -> 403, tanpa penulisan", sc == 403 and snap() == sn, (sc, r))
    sc, own = C("POST", "adjustments", {**head, "lines": [L("A", 1, "A", cost=1000)]})
    check("H5. User Asset dengan izin adjustment: create satu gudang dalam scope -> 200", sc == 200, (sc, own))
    sc, d = C("GET", f"adjustments/{own['id']}")
    check("H6. User tanpa izin Lihat Harga: API detail tanpa approved_unit_cost/cost_overridden", sc == 200 and all("approved_unit_cost" not in x and "cost_overridden" not in x for x in d["lines"]), d.get("lines"))
    sc, d = call("GET", f"adjustments/{own['id']}")
    check("H7. Admin (berizin) tetap melihat biaya", d["lines"][0].get("approved_unit_cost") == 1000, d["lines"][0])
    sc, r = C("PUT", f"transactions/adjustment/{own['id']}", {**head, "lines": [L("B", 1, "C", cost=500)]})
    check("H8. User Asset: edit memindah baris ke gudang divisi lain -> 403", sc == 403, (sc, r))
    mine = dict(T.S.headers); T.S.headers.clear(); T.S.headers.update(other_hdr)
    xs = [call("GET", f"adjustments/{a1['id']}")[0], call("PUT", f"transactions/adjustment/{a1['id']}", {**head, "lines": [L("A", 1, "A")]})[0],
          call("DELETE", f"transactions/adjustment/{a1['id']}")[0]]
    sc_x, r_x = call("POST", "adjustments", {**head, "lines": [L("A", 1, "A")]})
    T.S.headers.clear(); T.S.headers.update(mine)
    check("H9. Tenant lain: detail/edit/hapus -> 404; create pakai gudang tenant lain -> 400", all(c in (403, 404) for c in xs) and sc_x in (400, 403, 404), (xs, sc_x))
    E2 = f"adjnp_{u}@example.com"
    call("POST", "users", {"email": E2, "password": PW, "name": "QA No Adj", "role": "warehouse"}, 200)
    NP = U(E2)
    sc, _ = NP("POST", "attachment-drafts", {"module": "adjustment"})
    sc2, _ = NP("POST", "adjustments", {**head, "lines": [L("A", 1, "A")]})
    check("H10. User tanpa izin Penyesuaian Stok: draft lampiran & create -> 403 (adjustment.create = stock_adjustment)", sc == 403 and sc2 == 403, (sc, sc2))

    # ================= D. Draft multi-lampiran =================
    sc, dr = call("POST", "attachment-drafts", {"module": "adjustment"})
    check("D1. Draft Adjustment dibuat sebelum dokumen ada", sc == 200 and dr.get("module") == "adjustment", dr)
    ups = [upload(T.S, dr["id"], f"bukti-{i}.pdf") for i in (1, 2)]
    check("D1. Multi-lampiran (2 file) diunggah sebagai draft", all(x.status_code == 200 for x in ups) and len(att_rows(dr["id"])) == 2)
    paths = [p for _, p in att_rows(dr["id"])]
    nd = n_docs()
    sc, r = call("POST", "adjustments", {**head, "lines": [L("A", -999, "A")], "attachment_draft_id": dr["id"]})
    st = sql("SELECT JSON_EXTRACT(doc,'$.bound_to') FROM attachment_drafts WHERE id=%s", (dr["id"],))[0][0]
    check("D2. Posting gagal -> dokumen tidak terbentuk, draft tidak terikat, 2 file tetap draft", sc == 400 and n_docs() == nd and st in (None, "null") and len(att_rows(dr["id"])) == 2, (sc, st))
    sc, ok1 = call("POST", "adjustments", {**head, "lines": [L("A", -1, "A")], "attachment_draft_id": dr["id"]})
    fin = att_rows(ok1["id"])
    check("D3. Retry sukses -> tepat 2 lampiran final (path sama, tanpa duplikat)", sc == 200 and sorted(p for _, p in fin) == sorted(paths) and all(e == "adjustment" for e, _ in fin)
          and len(att_rows(dr["id"])) == 0 and n_docs() == nd + 1, fin)
    sc, att = call("GET", f"attachments?entity=adjustment&entity_id={ok1['id']}")
    check("D3. Lampiran muncul di detail Adjustment", sc == 200 and len(att) == 2, att)
    sc, r = call("POST", "adjustments", {**head, "lines": [L("A", -1, "A")], "attachment_draft_id": dr["id"]})
    check("D4. Replay draft ter-bind -> 409 (tanpa dokumen baru)", sc == 409 and n_docs() == nd + 1, (sc, r))
    sc, r = call("POST", "loans", {"date": TODAY, "division_id": asset["id"], "from_warehouse_id": W["A"]["id"], "to_warehouse_id": W["B"]["id"],
                                   "lines": [{"item_id": I["A"]["id"], "qty": 1, "uom_id": uom["id"]}], "attachment_draft_id": dr["id"]})
    check("D5. Isolasi modul: draft Adjustment tidak bisa dipakai Pinjam Barang (404)", sc == 404, (sc, r))
    sc, dl_ = call("POST", "attachment-drafts", {"module": "loan"})
    up_l = upload(T.S, dl_["id"], "loan.pdf", entity="loan_draft")
    up_x = upload(T.S, dl_["id"], "x.pdf")  # entity adjustment_draft dengan id draft loan
    check("D5. Draft Loan tetap berfungsi & tidak bisa diisi sebagai adjustment_draft", up_l.status_code == 200 and up_x.status_code == 404, (up_l.status_code, up_x.status_code))
    call("DELETE", f"attachment-drafts/{dl_['id']}")
    # Cross-user & cross-tenant
    sc, dA = call("POST", "attachment-drafts", {"module": "adjustment"})
    upA = upload(T.S, dA["id"], "rahasia-adj.pdf").json(); pathA = att_rows(dA["id"])[0][1]
    cu = [C("GET", f"attachments?entity=adjustment_draft&entity_id={dA['id']}")[0], upload(C.S, dA["id"], "b.pdf").status_code,
          C("DELETE", f"attachments/{upA['id']}")[0], C.S.get(f"{API}/attachments/{upA['id']}/download").status_code,
          C("DELETE", f"attachment-drafts/{dA['id']}")[0], C("POST", "adjustments", {**head, "lines": [L("A", 1, "A", cost=1000)], "attachment_draft_id": dA["id"]})[0]]
    check("D6. User lain (tenant sama): lihat/upload/hapus/download/batal/bind draft -> 404", all(c == 404 for c in cu) and len(att_rows(dA["id"])) == 1, cu)
    T.S.headers.clear(); T.S.headers.update(other_hdr)
    ct = [call("GET", f"attachments?entity=adjustment_draft&entity_id={dA['id']}"), call("DELETE", f"attachment-drafts/{dA['id']}")]
    upX = upload(T.S, dA["id"], "x.pdf"); dlX = T.S.get(f"{API}/attachments/{upA['id']}/download")
    bx = call("POST", "adjustments", {**head, "lines": [L("A", 1)], "attachment_draft_id": dA["id"]})
    T.S.headers.clear(); T.S.headers.update(mine)
    codes = [ct[0][0], ct[1][0], upX.status_code, dlX.status_code]
    blob = str(ct) + upX.text + dlX.text + str(bx)
    check("D7. Tenant lain: lihat/batal/upload/download draft -> 404; bind ditolak", all(c == 404 for c in codes) and bx[0] in (400, 403, 404), (codes, bx[0]))
    check("D7. Tidak bocor filename/storage path/tenant/user", "rahasia-adj" not in blob and pathA not in blob and usr["id"] not in blob and "tenant_id" not in blob)
    existed = obj_exists(pathA)
    sc, r = call("DELETE", f"attachment-drafts/{dA['id']}")
    check("D8. Batal: row draft + metadata lampiran + object storage terhapus", sc == 200 and len(att_rows(dA["id"])) == 0
          and sql("SELECT COUNT(*) FROM attachment_drafts WHERE id=%s", (dA["id"],))[0][0] == 0 and ((existed is True and obj_exists(pathA) is False) if STORAGE_LOCAL else True), (existed, obj_exists(pathA)))
    # Cleanup > 24 jam
    old = (_dt.datetime.now(_dt.timezone.utc) - _dt.timedelta(hours=25)).isoformat()
    sc, dO = call("POST", "attachment-drafts", {"module": "adjustment"}); upload(T.S, dO["id"], "lama.pdf"); pO = att_rows(dO["id"])[0][1]
    sc, dP = call("POST", "attachment-drafts", {"module": "adjustment"}); upload(T.S, dP["id"], "proses.pdf")
    fake = str(uuid.uuid4())
    sql("UPDATE attachment_drafts SET doc = JSON_SET(doc, '$.created_at', %s) WHERE id IN (%s, %s, %s)", (old, dO["id"], dP["id"], dr["id"]))
    sql("UPDATE attachment_drafts SET doc = JSON_SET(doc, '$.bound_to', %s, '$.binding', true, '$.claimed_at', %s) WHERE id=%s",
        (fake, _dt.datetime.now(_dt.timezone.utc).isoformat(), dP["id"]))
    call("POST", "attachment-drafts", {"module": "adjustment"})  # memicu cleanup tenant aktif
    check("D9. Cleanup > 24 jam: draft orphan + lampiran + object terhapus", sql("SELECT COUNT(*) FROM attachment_drafts WHERE id=%s", (dO["id"],))[0][0] == 0
          and len(att_rows(dO["id"])) == 0 and (obj_exists(pO) is False if STORAGE_LOCAL else True))
    check("D9. Draft yang SEDANG diproses (klaim baru) dilindungi dari cleanup", sql("SELECT COUNT(*) FROM attachment_drafts WHERE id=%s", (dP["id"],))[0][0] == 1 and len(att_rows(dP["id"])) == 1)
    check("D9. Draft final + 2 lampiran final aman dari cleanup", sql("SELECT COUNT(*) FROM attachment_drafts WHERE id=%s", (dr["id"],))[0][0] == 1
          and len(att_rows(ok1["id"])) == 2 and all(obj_exists(p) in (True, None) for p in paths))

    # ================= K. Import Excel =================
    _, tpl = None, T.S.get(f"{API}/excel/template/adjustment")
    wb = load_workbook(io.BytesIO(tpl.content)); hdrs = [c.value for c in wb.active[1]]
    guide = " ".join(str(c.value) for row in wb["Petunjuk"].iter_rows() for c in row if c.value)
    check("K1. Template baru memuat kolom opsional line_warehouse_code + petunjuk", "line_warehouse_code" in hdrs and "line_warehouse_code" in guide and "line_warehouse_code" not in
          [x for x in (call("GET", "excel/datasets")[1]["transactions"]) if x["key"] == "adjustment"][0]["required"], hdrs)
    OLD = ["batch_ref", "date", "warehouse_code", "division_code", "project_code", "adj_type", "reason", "notes", "item_code", "adjustment", "line_reason"]
    NEW = OLD + ["line_warehouse_code"]
    base = {"date": TODAY, "warehouse_code": W["A"]["code"], "division_code": asset["code"], "adj_type": "Koreksi", "reason": "Import"}

    def imp(rows, cols, sess=None):
        r_ = (sess or T.S).post(f"{API}/excel/import/adjustment", files={"file": ("adj.xlsx", xlsx(rows, cols), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        return r_.status_code, r_.json()

    sn = snap()
    sc, r = imp([{**base, "batch_ref": "OLD1", "item_code": I["A"]["code"], "adjustment": 1}], OLD)
    check("K2. Template LAMA (tanpa kolom baru) tetap bisa diimpor -> Gudang Default", sc == 200 and r.get("ok") and dlt(sn, snap()) == {"A@A": 1}
          and r["created_docs"][0]["lines"][0]["warehouse"] == "Gudang A", (sc, r))
    sn = snap()
    sc, r = imp([{**base, "batch_ref": "MW1", "item_code": I["A"]["code"], "adjustment": 2, "line_warehouse_code": W["B"]["code"]},
                 {**base, "batch_ref": "MW1", "item_code": I["C"]["code"], "adjustment": -1, "line_warehouse_code": W["D"]["code"]},
                 {**base, "batch_ref": "MW1", "item_code": I["A"]["code"], "adjustment": 1, "line_warehouse_code": ""}], NEW)
    check("K3. Multi gudang + kolom kosong: B +2, D -1, kosong -> Gudang Default A +1 (1 dokumen)", sc == 200 and r.get("ok") and dlt(sn, snap()) == {"A@B": 2, "C@D": -1, "A@A": 1}
          and len(r["created_docs"]) == 1, (sc, r, dlt(sn, snap())))
    check("K3. Preview hasil import menampilkan gudang masing-masing barang", [x["warehouse"] for x in r["created_docs"][0]["lines"]] == ["Gudang B", "Gudang D", "Gudang A"], r["created_docs"])
    nd, sn = n_docs(), snap()
    sc, r = imp([{**base, "batch_ref": "BAD", "item_code": I["A"]["code"], "adjustment": 1, "line_warehouse_code": "TIDAK-ADA"},
                 {**base, "batch_ref": "BAD", "item_code": I["A"]["code"], "adjustment": 1, "line_warehouse_code": W["X"]["code"]}], NEW)
    errs = " | ".join(r.get("errors") or [])
    check("K4. Kode gudang baris tidak ada / nonaktif -> error per nomor baris Excel, tanpa penulisan & tanpa fallback default",
          r.get("ok") is False and "Baris 2: line_warehouse_code 'TIDAK-ADA' tidak ditemukan" in errs and f"Baris 3: line_warehouse_code '{W['X']['code']}' tidak aktif" in errs
          and n_docs() == nd and snap() == sn, errs)
    sc, r = imp([{**base, "batch_ref": "SC1", "item_code": I["A"]["code"], "adjustment": 1, "line_warehouse_code": W["C"]["code"]}], NEW, sess=C.S)
    errs = " | ".join(r.get("errors") or [])
    check("K5. User Asset: kode gudang baris divisi lain -> error baris Excel (di luar cakupan)", r.get("ok") is False and "Baris 2" in errs and "cakupan" in errs and snap() == sn, (sc, errs))

    # ================= M. Kegagalan engine DI TENGAH posting (setelah validasi lolos) =================
    # Baris 1 (item baru K1 +1 di A) terposting; baris 2 (K2 -2 di B, backdate sebelum saldo ada) ditolak engine
    # -> seluruh dokumen dibatalkan (reversal baris 1), draft lampiran tetap draft, retry sukses tanpa duplikasi.
    today = _dt.date.today()
    K1 = m("items", {"code": f"K1{u}", "name": "Item K1", **ref}); K2 = m("items", {"code": f"K2{u}", "name": "Item K2", **ref})
    call("POST", "adjustments", {**head, "date": (today - _dt.timedelta(days=3)).isoformat(), "adj_type": "Opening",
                                 "lines": [{"item_id": K2["id"], "adjustment": 4, "warehouse_id": W["B"]["id"], "approved_unit_cost": 300, "reason": "QA"}]}, 200)
    sc, dM = call("POST", "attachment-drafts", {"module": "adjustment"})
    upload(T.S, dM["id"], "m-1.pdf"); upload(T.S, dM["id"], "m-2.pdf"); pathsM = sorted(p for _, p in att_rows(dM["id"]))
    nd, sn, nl0 = n_docs(), snap(), n_ledger()
    lines_m = [{"item_id": K1["id"], "adjustment": 1, "warehouse_id": W["A"]["id"], "approved_unit_cost": 100, "reason": "QA"},
               {"item_id": K2["id"], "adjustment": -2, "warehouse_id": W["B"]["id"], "reason": "QA"}]
    sc, r = call("POST", "adjustments", {**head, "date": (today - _dt.timedelta(days=5)).isoformat(), "lines": lines_m, "attachment_draft_id": dM["id"]})
    k1_rows = sql("SELECT JSON_UNQUOTE(JSON_EXTRACT(doc,'$.doc_type')) FROM stock_ledger WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.item_id'))=%s", (K1["id"],))
    check("M1. Gagal engine di baris 2 -> error, dokumen TIDAK terbentuk", sc >= 400 and n_docs() == nd, (sc, r))
    check("M2. Tanpa partial posting: stok K1@A = 0, K2@B = 4, gudang lain tidak berubah", stock(K1["id"], W["A"]["id"]) == 0 and stock(K2["id"], W["B"]["id"]) == 4 and snap() == sn)
    check("M3. Baris 1 sempat terposting lalu dibalik engine reversal (bukan pre-check)", len(k1_rows) >= 2 and any("Reversal" in (x[0] or "") for x in k1_rows), k1_rows)
    check("M4. Tidak ada header/line Adjustment yatim", sql("SELECT COUNT(*) FROM adjustment_lines WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.item_id'))=%s", (K1["id"],))[0][0] == 0)
    stM = sql("SELECT JSON_EXTRACT(doc,'$.bound_to') FROM attachment_drafts WHERE id=%s", (dM["id"],))[0][0]
    check("M5. Draft dilepas (bound_to null) + 2 file tetap draft", stM in (None, "null") and len(att_rows(dM["id"])) == 2, stM)
    nl1 = n_ledger()
    sc, okM = call("POST", "adjustments", {**head, "lines": lines_m, "attachment_draft_id": dM["id"]})
    finM = att_rows(okM.get("id")) if sc == 200 else []
    check("M6. Retry sukses -> 1 dokumen, 2 lampiran final (path sama), tanpa duplikat", sc == 200 and n_docs() == nd + 1 and sorted(p for _, p in finM) == pathsM, (sc, okM, finM))
    check("M7. Retry: stok tepat sekali (K1@A +1, K2@B -2) & ledger aktif 2 baris", stock(K1["id"], W["A"]["id"]) == 1 and stock(K2["id"], W["B"]["id"]) == 2
          and len(ledger(okM["id"])) == 2 and n_ledger() - nl1 == 2, (n_ledger() - nl1, ledger(okM["id"])))
    vm = {(x["item_name"], x["warehouse_name"]): (float(x.get("value_in") or 0), float(x.get("value_out") or 0)) for x in vrows(okM["no"])}
    check("M8. Valuasi retry konsisten: K1@A masuk 1x100, K2@B keluar 2x300 (MWA gudang line)", vm == {("Item K1", "Gudang A"): (100, 0), ("Item K2", "Gudang B"): (0, 600)}, vm)
    check("M9. Audit trail: create hanya untuk dokumen sukses", sql("SELECT COUNT(*) FROM audit_logs WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.entity_id'))=%s AND JSON_UNQUOTE(JSON_EXTRACT(doc,'$.action'))='create'", (okM["id"],))[0][0] == 1)
    _ = nl0

    # ================= N. Manipulasi line_warehouse_ids / warehouse_source dari klien =================
    sc, n1 = call("POST", "adjustments", {**head, "line_warehouse_ids": [W["C"]["id"]], "warehouse_source": "legacy", "lines": [L("A", 1, "A", cost=1000)]})
    hn = sql("SELECT JSON_UNQUOTE(JSON_EXTRACT(doc,'$.warehouse_source')), JSON_EXTRACT(doc,'$.line_warehouse_ids') FROM adjustments WHERE id=%s", (n1["id"],))[0]
    check("N1. Create: line_warehouse_ids/warehouse_source dari klien diabaikan (dihitung server dari LINE)", sc == 200 and hn[0] == "line" and json.loads(hn[1]) == [W["A"]["id"]], hn)
    sc, lst = C("GET", "adjustments")
    check("N2. User Asset tetap melihat dokumen (gudang LINE sebenarnya A), bukan C palsu", n1["id"] in {x["id"] for x in lst}, sc)
    sn = snap()
    sc, r = C("POST", "adjustments", {**head, "line_warehouse_ids": [W["A"]["id"]], "lines": [L("B", 1, "C", cost=500)]})
    check("N3. User Asset: baris gudang C + line_warehouse_ids palsu [A] -> 403 tanpa penulisan", sc == 403 and snap() == sn, (sc, r))
    sc, r = call("PUT", f"transactions/adjustment/{n1['id']}", {**head, "line_warehouse_ids": [W["A"]["id"]], "lines": [L("A", -1, "B")]})
    hn = sql("SELECT JSON_EXTRACT(doc,'$.line_warehouse_ids') FROM adjustments WHERE id=%s", (n1["id"],))[0][0]
    check("N4. Edit: line_warehouse_ids dihitung ulang server-side (B), bukan nilai klien (A)", sc == 200 and json.loads(hn) == [W["B"]["id"]], (sc, r, hn))
    sc, r = C("PUT", f"transactions/adjustment/{n1['id']}", {**head, "line_warehouse_ids": [W["A"]["id"]], "lines": [L("B", 1, "C", cost=500)]})
    check("N5. User Asset: edit ke gudang C + ids palsu -> 403", sc == 403, (sc, r))

    failed = [n for n, ok in T.RESULTS if not ok]
    print(f"\n{len(T.RESULTS) - len(failed)}/{len(T.RESULTS)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    os.environ.setdefault("PYTHONUNBUFFERED", "1")
    main()
