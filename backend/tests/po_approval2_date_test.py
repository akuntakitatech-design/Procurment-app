"""Tanggal Approval 2 pada Daftar PO — E2E (tenant uji, alur approval existing) + unit (fake DB).

`approval2_at` (WIB 'YYYY-MM-DD HH:MM:SS') hanya untuk Status Dokumen = Approved (Disetujui), bersumber dari task Approval 2
final (approval_tasks seq=2 Approved, acted_at); bukan tanggal PO / updated_at; tanpa riwayat -> None (tanpa perkiraan).
Sort daftar `sort=approval2_at&dir=asc|desc` (tanpa tanggal selalu di akhir); export txn-po mengikuti urutan & memuat kolomnya.
Workflow approval, status & aturan transaksi tidak berubah.
"""
import asyncio
import io
import sys
import time
from datetime import datetime, timedelta, timezone

import po_approval2_batch_test as A
import po_document_status_test as S
import receipt_control_test as T

call, check = T.call, T.check
API = T.API
WIB = timezone(timedelta(hours=7))


def wib(iso):
    return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(WIB).strftime("%Y-%m-%d %H:%M:%S")


def paged(qs=""):
    sc, d = call("GET", f"po?page=1&page_size=100{qs}")
    return sc, (d.get("items", []) if isinstance(d, dict) else [])


def a2_task_acted(po_id):
    """acted_at task Approval 2 (approval_tasks seq=2) dari riwayat detail PO (po_approvals mirror existing)."""
    d = S.detail({"id": po_id})
    l2 = [a for a in d.get("approvals") or [] if int(a.get("seq") or 0) == 2]
    return l2[0].get("acted_at") if l2 else None, d


def main():
    M = T.setup()
    a1, a2 = A.mk_user("l1", [M["div"]["id"]]), A.mk_user("l2", [M["div"]["id"]])
    A.set_ov(a2, {"upload_attachment": "allow"})
    call("PUT", "settings/approval_modules", {"modules": {
        "mro": {"enabled": False, "levels": []}, "ro": {"enabled": False, "levels": []},
        "po": {"enabled": True, "levels": [{"level": 1, "email": a1.email}, {"level": 2, "email": a2.email}]}}}, 200)

    pA, pB, pC, pW, pR, pDr = (A.mk_po(M, [A.line(M, "item", 10, 1000 + i)]) for i in range(6))
    for p in (pA, pB, pC, pW, pR):
        call("POST", f"po/{p['id']}/submit", {}, 200)
    for p in (pA, pB, pC, pW):
        a1("POST", f"approvals/{A.task(a1, p['id'], 1)['id']}/approve", {"note": "L1 OK"})
    a1("POST", f"approvals/{A.task(a1, pR['id'], 1)['id']}/reject", {"reason": "revisi"})

    sc, el = a2("GET", "approval2/po/eligible")
    tid = {r["po_id"]: r["approval_task_id"] for r in el or []}
    sc, B = a2("POST", "approval2/po/batches", {"title": "Uji Tgl A2", "submission_date": datetime.now(WIB).date().isoformat(),
                                                "approval_task_ids": [tid[p["id"]] for p in (pA, pB, pC, pW)]})
    check("Batch Approval 2 dibuat", sc == 200 and B.get("id"), (sc, B))
    a2("POST", f"approval2/po/batches/{B['id']}/submit", {})
    sc, _ = a2.upload("po_approval2_batch", B["id"], "bukti.png", A.PNG, "image/png")
    it = {i["po_id"]: i["id"] for i in B["items"]}
    # Urutan approve sengaja C -> A -> B (berbeda dari urutan dibuat/nomor PO) dengan jeda > 1 detik
    for p in (pC, pA, pB):
        sc, d = a2("POST", f"approval2/po/batches/{B['id']}/approve", {"item_ids": [it[p["id"]]]})
        check(f"Approve Terpilih Approval 2 {p.get('no')}", sc == 200, (sc, d))
        time.sleep(1.2)

    rows = {r["id"]: r for r in paged()[1]}
    arr = {r["id"]: r for r in call("GET", "po")[1] or []}
    # ---------------------------------------------------------------- tanggal & sumber
    for p in (pA, pB, pC):
        acted, d = a2_task_acted(p["id"])
        r = rows.get(p["id"]) or {}
        check(f"{d.get('no')}: Disetujui + approval2_at = acted_at Approval 2 (WIB) — list paged & array sama",
              r.get("document_status") == "Approved" and acted and r.get("approval2_at") == wib(acted) == arr[p["id"]].get("approval2_at"),
              (r.get("document_status"), r.get("approval2_at"), acted))
    check("Konversi WIB: 10 karakter pertama = tanggal WIB acted_at",
          all(rows[p["id"]]["approval2_at"][:10] == wib(a2_task_acted(p["id"])[0])[:10] for p in (pA, pB, pC)))
    # ---------------------------------------------------------------- status lain tetap, tanpa tanggal
    want = {pW["id"]: "Waiting Approval 2", pR["id"]: "Rejected", pDr["id"]: "Draft"}
    check("Status selain Disetujui tetap (Menunggu Approval 2 / Ditolak / Draft) & tanpa approval2_at",
          all(rows[k].get("document_status") == v and rows[k].get("approval2_at") is None for k, v in want.items()),
          {rows[k].get("no"): (rows[k].get("document_status"), rows[k].get("approval2_at")) for k in want})

    # ---------------------------------------------------------------- sorting
    approved_order = sorted((pA, pB, pC), key=lambda p: rows[p["id"]]["approval2_at"])
    exp_asc = [p["id"] for p in approved_order]
    check("Urutan tanggal = urutan approve sebenarnya C -> A -> B (bukan urutan dibuat/nomor A -> B -> C), stempel naik tegas",
          exp_asc == [pC["id"], pA["id"], pB["id"]] and rows[pC["id"]]["approval2_at"] < rows[pA["id"]]["approval2_at"] < rows[pB["id"]]["approval2_at"],
          [(rows[i]["no"], rows[i]["approval2_at"]) for i in exp_asc])
    no_of = {k: v["no"] for k, v in rows.items()}
    sc, asc = paged("&sort=approval2_at&dir=asc")
    sc2, desc = paged("&sort=approval2_at&dir=desc")
    ia, idd = [r["id"] for r in asc], [r["id"] for r in desc]
    check("sort=approval2_at asc: terlama -> terbaru, PO tanpa tanggal di akhir", sc == 200 and ia[:3] == exp_asc
          and all(r.get("approval2_at") is None for r in asc[3:]), [(r["no"], r.get("approval2_at")) for r in asc])
    check("sort=approval2_at desc: terbaru -> terlama, PO tanpa tanggal tetap di akhir", sc2 == 200 and idd[:3] == exp_asc[::-1]
          and all(r.get("approval2_at") is None for r in desc[3:]), [(r["no"], r.get("approval2_at")) for r in desc])
    sc, f = paged("&sort=approval2_at&dir=asc&q=" + pA["no"])
    check("sort + pencarian tetap berfungsi", sc == 200 and [r["id"] for r in f] == [pA["id"]], [r["no"] for r in f])

    # ---------------------------------------------------------------- export P6 mengikuti urutan & memuat kolom
    sc, rj = call("GET", "report-center/txn-po?sort=approval2_at&dir=asc")
    check("Export JSON txn-po sort=approval2_at: 200, urutan = daftar, kolom 'Tgl Approval 2 (WIB)'",
          sc == 200 and [r["no"] for r in rj["rows"]] == [no_of[i] for i in ia]
          and [r["approval2_at"] for r in rj["rows"]][:3] == [rows[i]["approval2_at"] for i in exp_asc]
          and any(c["key"] == "approval2_at" and c["label"] == "Tgl Approval 2 (WIB)" for c in rj["columns"]), (sc, rj.get("detail")))
    rx = T.S.get(f"{API}/report-center/txn-po/export.xlsx?sort=approval2_at&dir=desc")
    from openpyxl import load_workbook
    vals = [list(v) for v in load_workbook(io.BytesIO(rx.content), read_only=True).worksheets[0].iter_rows(values_only=True)] if rx.status_code == 200 else []
    hi = next((i for i, v in enumerate(vals) if v and str(v[0]).startswith("No.")), None)
    ci = vals[hi].index("Tgl Approval 2 (WIB)") if hi is not None and "Tgl Approval 2 (WIB)" in vals[hi] else None
    body = [v for v in vals[hi + 1:] if v and v[0] not in (None, "TOTAL")] if hi is not None else []
    check("Excel txn-po sort desc: kolom Tgl Approval 2 berisi stempel WIB, urutan terbaru -> terlama",
          ci is not None and [b[ci] for b in body[:3]] == [rows[i]["approval2_at"] for i in exp_asc[::-1]] and all(not b[ci] for b in body[3:]),
          (rx.status_code, ci))
    sp = T.S.get(f"{API}/report-center/txn-po/export.pdf?sort=approval2_at&dir=asc")
    check("PDF txn-po sort=approval2_at -> 200", sp.status_code == 200 and sp.content[:4] == b"%PDF", sp.status_code)

    # ---------------------------------------------------------------- Closed (diterima penuh) -> badge lain, tanpa tanggal
    xb = S.detail(pB)
    call("POST", "do", T.do_body(M, xb, 10, "supX"), 200)
    r = {x["id"]: x for x in paged()[1]}[pB["id"]]
    check("Diterima penuh -> Status Dokumen Closed (tetap seperti sekarang), approval2_at tidak ditampilkan",
          r.get("document_status") == "Closed" and r.get("approval2_at") is None, (r.get("document_status"), r.get("approval2_at")))

    # ---------------------------------------------------------------- 1 tahap & tanpa approval (legacy tanpa riwayat)
    call("PUT", "settings/approval_modules", {"modules": {
        "mro": {"enabled": False, "levels": []}, "ro": {"enabled": False, "levels": []},
        "po": {"enabled": True, "levels": [{"level": 1, "email": a1.email}]}}}, 200)
    pG = A.mk_po(M, [A.line(M, "item", 2, 5000)])
    call("POST", f"po/{pG['id']}/submit", {}, 200)
    a1("POST", f"approvals/{A.task(a1, pG['id'], 1)['id']}/approve", {"note": "ok"})
    call("PUT", "settings/approval_modules", {"modules": {
        "mro": {"enabled": False, "levels": []}, "ro": {"enabled": False, "levels": []},
        "po": {"enabled": False, "levels": []}}}, 200)
    pH = A.mk_po(M, [A.line(M, "item", 2, 5000)])
    call("POST", f"po/{pH['id']}/submit", {}, 200)
    rows = {r["id"]: r for r in paged()[1]}
    check("Approval 1 tahap: Disetujui tanpa tanggal Approval 2", rows[pG["id"]].get("document_status") == "Approved"
          and rows[pG["id"]].get("approval2_at") is None, rows[pG["id"]].get("approval2_at"))
    check("PO tanpa riwayat approval (approval nonaktif / legacy): Disetujui tanpa tanggal, tanpa perkiraan",
          rows[pH["id"]].get("document_status") == "Approved" and rows[pH["id"]].get("approval2_at") is None, rows[pH["id"]].get("approval2_at"))
    check("Workflow approval tidak berubah: PO.status internal Approved untuk PO Disetujui",
          all(S.detail(p)["status"] == "Approved" for p in (pA, pC, pG, pH)))

    unit_tests()
    ok = sum(1 for _, c in T.RESULTS if c)
    print(f"\n{ok}/{len(T.RESULTS)} passed")
    return 0 if ok == len(T.RESULTS) else 1


def unit_tests():
    sys.path.insert(0, __file__.rsplit("/tests/", 1)[0])
    import receipt_control_layer as R
    check("Unit wib_stamp: UTC -> WIB (lintas tanggal)", R.wib_stamp("2026-10-10T20:30:05.123+00:00") == "2026-10-11 03:30:05")
    check("Unit wib_stamp: naive dianggap UTC (now_iso)", R.wib_stamp("2026-10-10T01:00:00") == "2026-10-10 08:00:00")
    check("Unit wib_stamp: kosong/tidak valid -> None", [R.wib_stamp(v) for v in (None, "", "bukan")] == [None, None, None])
    AP = {"document_status": "Approved"}
    pos = [{"id": i, **AP} for i in ("eng", "mirror", "none", "pend", "dup", "nolvl2", "noacted")] + \
          [{"id": "wa2", "document_status": "Waiting Approval 2"}, {"id": "closed", "document_status": "Closed"}]
    ok_l1 = {"seq": 1, "status": "Approved", "acted_at": "2026-10-09T03:00:00+00:00"}
    db = S._Db(
        approval_tasks=[
            {"module": "po", "document_id": "eng", **ok_l1}, {"module": "po", "document_id": "eng", "seq": 2, "status": "Approved",
                                                             "acted_at": "2026-10-10T18:00:00+00:00"},
            {"module": "po", "document_id": "pend", **ok_l1}, {"module": "po", "document_id": "pend", "seq": 2, "status": "Pending"},
            {"module": "po", "document_id": "dup", "seq": 2, "status": "Approved", "acted_at": "2026-10-10T01:00:00+00:00"},
            {"module": "po", "document_id": "dup", "seq": 2, "status": "Approved", "acted_at": "2026-10-11T01:00:00+00:00"},
            {"module": "po", "document_id": "nolvl2", **ok_l1},
            {"module": "po", "document_id": "noacted", "seq": 2, "status": "Approved", "acted_at": None},
            {"module": "po", "document_id": "wa2", **ok_l1}, {"module": "po", "document_id": "wa2", "seq": 2, "status": "Approved",
                                                             "acted_at": "2026-10-10T18:00:00+00:00"},
            {"module": "po", "document_id": "closed", "seq": 2, "status": "Approved", "acted_at": "2026-10-10T18:00:00+00:00"},
        ],
        po_approvals=[{"po_id": "mirror", **ok_l1}, {"po_id": "mirror", "seq": 2, "status": "Approved", "acted_at": "2026-10-08T23:30:00+00:00"},
                      {"po_id": "eng", "seq": 2, "status": "Approved", "acted_at": "2000-01-01T00:00:00+00:00"}])
    m = asyncio.run(R.po_approval2_at_map(S._Srv(db), pos))
    check("Unit: engine approval_tasks seq=2 final -> WIB", m.get("eng") == "2026-10-11 01:00:00", m.get("eng"))
    check("Unit: mirror po_approvals hanya dipakai bila PO tanpa task (legacy engine lama)", m.get("mirror") == "2026-10-09 06:30:00", m.get("mirror"))
    check("Unit: PO legacy tanpa riwayat -> tidak ada tanggal", "none" not in m)
    check("Unit: siklus belum final (task Pending) -> tidak ada tanggal", "pend" not in m)
    check("Unit: riwayat ambigu (2 task seq=2) -> tidak ada tanggal (tanpa menebak)", "dup" not in m)
    check("Unit: 1 tahap (tanpa seq=2) / acted_at kosong -> tidak ada tanggal", "nolvl2" not in m and "noacted" not in m)
    check("Unit: status selain Approved (Menunggu Approval 2 / Closed) -> tidak ada tanggal", "wa2" not in m and "closed" not in m)


if __name__ == "__main__":
    sys.exit(main())
