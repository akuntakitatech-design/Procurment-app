"""Stock Opname hardening — workflow, Freeze/Live, posting atomik + retry, valuasi, redaksi harga, import, koreksi.

Jalankan (backend aktif di :8001, DB MariaDB TERISOLASI dari backend/.env):
  /root/.venv/bin/python backend/tests/stock_opname_workflow_test.py
Kegagalan baris tengah disimulasikan dengan TRIGGER MariaDB pada tabel item_warehouse (tidak ada hook di kode aplikasi).
"""
import concurrent.futures as cf
import datetime as _dt
import io
import json
import sys
import uuid

import requests
from openpyxl import load_workbook

import receipt_control_test as T
import transfer_multi_warehouse_test as TM
from vendor_invoice_test import PW

call, check, API = T.call, T.check, T.API
m, stock, opening, db_conn = TM.m, TM.stock, TM.opening, TM.db_conn
u = uuid.uuid4().hex[:6]
TODAY = _dt.date.today().isoformat()
TOMORROW = (_dt.date.today() + _dt.timedelta(days=1)).isoformat()


def sql(q, args=()):
    with db_conn() as cx, cx.cursor() as cur:
        cur.execute(q, args)
        return cur.fetchall()


def vsum(doc_id, item_id):
    rows = sql("SELECT doc FROM valuation_ledger WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.doc_id'))=%s", (doc_id,))
    tot = 0.0
    for (d,) in rows:
        x = json.loads(d)
        if x.get("item_id") == item_id:
            tot += float(x.get("value_in") or 0) - float(x.get("value_out") or 0)
    return round(tot, 4)


def live_entries(doc_id):
    return [json.loads(d) for (d,) in sql("SELECT doc FROM stock_ledger WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.doc_id'))=%s", (doc_id,))
            if not json.loads(d).get("is_reversal") and not json.loads(d).get("reversed")]


def pool(item_id, wh):
    r = sql("SELECT doc FROM item_warehouse WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.item_id'))=%s AND JSON_UNQUOTE(JSON_EXTRACT(doc,'$.warehouse_id'))=%s", (item_id, wh))
    d = json.loads(r[0][0]) if r else {}
    return round(float(d.get("current_stock") or 0), 6), round(float(d.get("total_value") or 0), 4)


def OL():
    _, r = call("GET", "opname?page_size=500")
    return r.get("rows", r.get("items", [])) if isinstance(r, dict) else r


def main():
    TM.register(); other_hdr = dict(T.S.headers)
    T.S.headers.pop("Authorization", None); TM.register()
    div = m("divisions", {"code": f"AS{u}", "name": "Asset"})
    cat = m("item_categories", {"code": f"KC{u}", "name": "Umum"})
    pcs = m("uoms", {"code": f"PC{u}", "name": "Pcs"})
    box = m("uoms", {"code": f"BX{u}", "name": "Box", "symbol": f"BX{u}"})
    W = {k: m("warehouses", {"code": f"W{k}{u}", "name": f"Gudang {k}", "is_active": True, "division_id": div["id"]}) for k in "ABC"}
    ref = {"category_id": cat["id"], "division_id": div["id"], "base_uom_id": pcs["id"], "unit": "PCS", "is_active": True}
    I = {n: m("items", {"code": f"I{n}{u}", "name": f"Item {n}", **ref, **({"uoms": [{"uom_id": box["id"], "factor": 6}]} if n == "A" else {})})
         for n in "ABCD"}
    for n, c in (("A", 1000), ("B", 2000), ("C", 3000)):
        opening(W["A"]["id"], div["id"], I[n]["id"], 10, c)
    opening(W["B"]["id"], div["id"], I["A"]["id"], 10, 1000)
    opening(W["C"]["id"], div["id"], I["A"]["id"], 5, 1000)
    S = lambda n, w: stock(I[n]["id"], W[w]["id"])  # noqa: E731
    adj = lambda w, n, q, d=TODAY, cost=None: call("POST", "adjustments", {  # noqa: E731
        "date": d, "warehouse_id": W[w]["id"], "division_id": div["id"], "adj_type": "Koreksi", "reason": "QA",
        "lines": [{"item_id": I[n]["id"], "adjustment": q, "reason": "QA", "warehouse_id": W[w]["id"], **({"approved_unit_cost": cost} if cost else {})}]})

    # ===== A. Create, lock 1 gudang 1 opname aktif, konkurensi
    sc, o = call("POST", "opname", {"date": TODAY, "warehouse_id": W["A"]["id"], "division_id": div["id"], "mode": "freeze", "scope": "all", "counter_name": "Budi"})
    check("A1. Create Opname freeze -> 200, 3 baris snapshot, Counting, v2", sc == 200 and o.get("summary", {}).get("total") == 3 and o.get("status") == "Counting" and o.get("workflow_version") == 2, (sc, o))
    did, no = o["id"], o["no"]
    sc, _ = call("POST", "opname", {"date": TODAY, "warehouse_id": W["A"]["id"], "division_id": div["id"], "mode": "live"})
    check("A2. Opname kedua di gudang yang sama (mode lain) -> 409", sc == 409)
    hdr = dict(T.S.headers); cookies = T.S.cookies.get_dict()

    def mk():
        s = requests.Session(); s.headers.update(hdr); s.cookies.update(cookies)
        rr = s.post(f"{API}/opname", json={"date": TODAY, "warehouse_id": W["B"]["id"], "division_id": div["id"], "mode": "live"})
        return rr.status_code, (rr.json() if rr.status_code == 200 else None)
    with cf.ThreadPoolExecutor(6) as ex:
        outs = list(ex.map(lambda _: mk(), range(6)))
    codes = [c for c, _ in outs]
    check("A3. 6 create bersamaan di Gudang B -> tepat 1 sukses, sisanya 409", codes.count(200) == 1 and codes.count(409) == 5, codes)
    live = next(dd for c, dd in outs if c == 200)
    check("A4. Daftar Opname memuat dokumen baru", any(x.get("id") == live["id"] for x in OL()), len(OL()))

    # ===== B. Freeze memblokir mutasi gudang tsb (DO/MI/Transfer/Adjustment/item-warehouse), gudang lain tetap jalan
    s0 = (S("A", "A"), S("B", "A"), S("C", "A"))
    sc, r = adj("A", "C", 1)
    check("B1. Adjustment di gudang Freeze -> 409 + pesan nomor Opname", sc == 409 and no in str(r) and "Mutasi stok tidak dapat dilakukan" in str(r), (sc, r))
    sc, r = call("POST", "transfers", {"date": TODAY, "division_id": div["id"], "from_warehouse_id": W["A"]["id"], "to_warehouse_id": W["C"]["id"],
                                       "lines": [{"item_id": I["A"]["id"], "qty": 1, "from_warehouse_id": W["A"]["id"], "to_warehouse_id": W["C"]["id"]}]})
    check("B2. Transfer keluar dari gudang Freeze -> ditolak, tanpa penulisan", sc in (400, 409) and S("A", "C") == 5, (sc, r))
    sc, r = call("POST", "transfers", {"date": TODAY, "division_id": div["id"], "from_warehouse_id": W["C"]["id"], "to_warehouse_id": W["A"]["id"],
                                       "lines": [{"item_id": I["A"]["id"], "qty": 1, "from_warehouse_id": W["C"]["id"], "to_warehouse_id": W["A"]["id"]}]})
    check("B3. Transfer MASUK ke gudang Freeze -> ditolak; gudang asal tidak berubah (atomik)", sc in (400, 409) and S("A", "C") == 5, (sc, r))
    sc, r = call("POST", "item-warehouse", {"item_id": I["A"]["id"], "warehouse_id": W["A"]["id"], "current_stock": 99})
    # current_stock klien selalu dibuang guard existing; ubah min/max bukan mutasi stok -> saldo wajib tetap.
    check("B4. Set saldo langsung (item-warehouse) di gudang Freeze -> saldo tidak berubah", sc in (200, 409) and S("A", "A") == 10, (sc, r))
    check("B5. Stok gudang Freeze tidak berubah", (S("A", "A"), S("B", "A"), S("C", "A")) == s0)
    sc, _ = adj("C", "A", 1)
    check("B6. Gudang lain (tanpa Freeze) tetap bisa mutasi", sc == 200 and S("A", "C") == 6)

    # ===== C. Penghitungan: kosong vs 0, UOM, alasan, status ignored, redaksi harga
    lines = {r["item_code"]: r for r in call("GET", f"opname/{did}")[1]["lines"]}
    LA, LB, LC = (lines[f"I{n}{u}"]["id"] for n in "ABC")
    check("C1. Harga Moving Average tampil di semua baris (admin)", all(lines[f"I{n}{u}"].get("avg_cost") == c for n, c in (("A", 1000), ("B", 2000), ("C", 3000))), lines)
    sc, r = call("PUT", f"opname/{did}/count", {"status": "Posted", "lines": [{"line_id": LA, "qty": 2, "uom_id": box["id"]}, {"line_id": LC, "qty": 0}]})
    _, d = call("GET", f"opname/{did}")
    la = next(x for x in d["lines"] if x["id"] == LA); lc = next(x for x in d["lines"] if x["id"] == LC)
    check("C2. Status dari payload diabaikan (tetap Counting)", d["status"] == "Counting")
    check("C3. 2 Box -> 12 pcs dasar, satuan & faktor tersimpan", la["counted"] == 12 and la["counted_qty"] == 2 and la["conversion_factor"] == 6, la)
    check("C4. 0 = sudah dihitung (selisih -10), kosong = Belum Dihitung", lc["status"] == "minus" and d["summary"]["uncounted"] == 1 and d["summary"]["counted"] == 2, d["summary"])
    sc, _ = call("PUT", f"opname/{did}/count", {"lines": [{"line_id": LC, "clear": True}]})
    check("C5. Kosongkan hitung -> kembali Belum Dihitung", call("GET", f"opname/{did}")[1]["summary"]["uncounted"] == 2)
    sc, _ = call("PUT", f"opname/{did}/count", {"lines": [{"line_id": LA, "qty": -1}]})
    check("C6. Qty negatif ditolak", sc == 400)
    sc, _ = call("PUT", f"opname/{did}/count", {"lines": [{"line_id": LA, "qty": 1, "uom_id": W["A"]["id"]}]})
    check("C7. Satuan bukan milik barang ditolak", sc == 400)
    call("PUT", f"opname/{did}/count", {"lines": [{"line_id": LB, "qty": 7}, {"line_id": LC, "qty": 10}]})
    sc, r = call("POST", f"opname/{did}/workflow", {"action": "review"})
    check("C8. Review ditolak bila selisih belum beralasan", sc == 400 and "alasan" in str(r), r)
    # user tanpa izin harga
    sc, su = call("POST", "users", {"email": f"staff{u}@qa.test", "password": PW, "name": "Staff QA", "role": "warehouse"})
    call("PUT", f"access/users/{su['id']}", {"overrides": {"opname.view": "allow", "opname.edit": "allow", "view_purchase_price": "deny"}})
    ss = requests.Session(); ss.post(f"{API}/auth/login", json={"email": f"staff{u}@qa.test", "password": PW})
    sd = ss.get(f"{API}/opname/{did}").json()
    flat = json.dumps(sd)
    check("C9. API redaksi harga untuk user tanpa view_purchase_price (baris + ringkasan)",
          all(k not in flat for k in ('"avg_cost"', '"unit_price"', '"value"', '"net_value"', '"surplus_value"')), [k for k in ("avg_cost", "unit_price", "net_value") if k in flat])
    pr = ss.get(f"{API}/opname/{did}/print-data").json()
    check("C10. Data cetak juga diredaksi", "net_value" not in json.dumps(pr) and "avg_cost" not in json.dumps(pr))
    sc = ss.put(f"{API}/opname/{did}/count", json={"lines": [{"line_id": LB, "approved_unit_cost": 1}]}).status_code
    check("C11. User tanpa izin harga mengisi Harga Satuan -> 403", sc == 403)
    sc = ss.post(f"{API}/opname/{did}/post").status_code
    check("C12. User tanpa izin post -> 403", sc == 403)

    # ===== D. Tambah barang belum tercatat + harga khusus (tanpa average)
    sc, _ = call("POST", f"opname/{did}/add-item", {"item_id": I["D"]["id"]})
    check("D1. Tambah barang tanpa alasan -> 400", sc == 400)
    sc, r = call("POST", f"opname/{did}/add-item", {"item_id": I["D"]["id"], "reason": "Ditemukan di rak 3"})
    LD = r.get("line_id")
    check("D2. Tambah barang dengan alasan -> 200", sc == 200 and LD)
    sc, _ = call("POST", f"opname/{did}/add-item", {"item_id": I["D"]["id"], "reason": "dup"})
    check("D3. Barang duplikat -> 409", sc == 409)
    call("PUT", f"opname/{did}/count", {"lines": [{"line_id": LA, "reason": "Salah catat"}, {"line_id": LB, "reason": "Rusak"}, {"line_id": LD, "qty": 1, "reason": "Temuan"}]})
    sc, r = call("POST", f"opname/{did}/workflow", {"action": "review"})
    check("D4. Surplus tanpa average wajib Harga Satuan -> review ditolak", sc == 400 and "Harga Satuan" in str(r), r)
    sc, _ = call("PUT", f"opname/{did}/count", {"lines": [{"line_id": LD, "approved_unit_cost": 500}]})
    check("D5. Harga khusus tanpa alasan -> 400", sc == 400)
    call("PUT", f"opname/{did}/count", {"lines": [{"line_id": LD, "approved_unit_cost": 500, "cost_reason": "Harga beli terakhir"}]})
    sm = call("GET", f"opname/{did}")[1]["summary"]
    check("D6. Ringkasan server: surplus 2.500, shortage -6.000, bersih -3.500",
          (sm["plus"], sm["minus"], sm["match"], sm["surplus_value"], sm["shortage_value"], sm["net_value"]) == (2, 1, 1, 2500, -6000, -3500), sm)

    # ===== E. Workflow + approval
    check("E1. Review -> 200", call("POST", f"opname/{did}/workflow", {"action": "review"})[0] == 200)
    check("E2. Submit -> Waiting Approval", call("POST", f"opname/{did}/submit")[1].get("status") == "Waiting Approval")
    sc, _ = call("PUT", f"opname/{did}/count", {"lines": [{"line_id": LB, "qty": 8}]})
    check("E3. Hasil hitung terkunci saat Waiting Approval (409)", sc == 409)
    check("E4. Reject tanpa alasan -> 400", call("POST", f"opname/{did}/workflow", {"action": "reject"})[0] == 400)
    sc, r = call("POST", f"opname/{did}/workflow", {"action": "reject", "reason": "Cek ulang Item B"})
    check("E5. Reject beralasan -> kembali Counting + history", sc == 200 and r["status"] == "Counting" and r["history"][-1]["action"] == "reject")
    call("POST", f"opname/{did}/workflow", {"action": "review"}); call("POST", f"opname/{did}/submit")
    sc, _ = adj("A", "C", 1)
    check("E6. Freeze tetap aktif setelah reject/return", sc == 409)

    # ===== F. Posting gagal di baris tengah -> rollback total; retry aman; tanpa double posting
    before = {n: pool(I[n]["id"], W["A"]["id"]) for n in "ABCD"}
    sql("DROP TRIGGER IF EXISTS qa_opname_fail")
    sql("CREATE TRIGGER qa_opname_fail BEFORE UPDATE ON item_warehouse FOR EACH ROW BEGIN "
        "IF JSON_UNQUOTE(JSON_EXTRACT(NEW.doc,'$.item_id')) = %s THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='QA fault item tengah'; END IF; END", (I["B"]["id"],))
    try:
        sc, r = call("POST", f"opname/{did}/post")
    finally:
        sql("DROP TRIGGER IF EXISTS qa_opname_fail")
    _, d = call("GET", f"opname/{did}")
    after = {n: pool(I[n]["id"], W["A"]["id"]) for n in "ABCD"}
    check("F1. Posting gagal -> error + status kembali Waiting Approval", sc >= 400 and d["status"] == "Waiting Approval" and d.get("last_post_error"), (sc, r, d.get("status")))
    check("F2. Stok + nilai persediaan semua barang tidak berubah (tanpa partial)", after == before, (before, after))
    check("F3. Tidak ada movement aktif tersisa dari dokumen", live_entries(did) == [], live_entries(did))
    check("F4. Freeze tetap aktif setelah posting gagal", adj("A", "C", 1)[0] == 409)
    sc, r = call("POST", f"opname/{did}/post")
    check("F5. Retry posting -> Posted", sc == 200 and r.get("status") == "Posted", (sc, r))
    check("F6. Stok final A12 B7 C10 D1", (S("A", "A"), S("B", "A"), S("C", "A"), S("D", "A")) == (12, 7, 10, 1))
    ents = live_entries(did)
    check("F7. Tepat 3 movement aktif (A,B,D), tanpa duplikasi", sorted(e["item_id"] for e in ents) == sorted([I["A"]["id"], I["B"]["id"], I["D"]["id"]]), len(ents))
    check("F8. Valuasi: A +2.000 @avg, B -6.000 @avg, D +500 @harga disetujui",
          (vsum(did, I["A"]["id"]), vsum(did, I["B"]["id"]), vsum(did, I["D"]["id"])) == (2000, -6000, 500))
    p = call("GET", f"opname/{did}")[1]
    check("F9. Nilai posted di dokumen = valuation ledger (bersih -3.500)", p["summary"]["net_value"] == -3500, p["summary"])
    check("F10. Posting ulang -> 409 (tidak double)", call("POST", f"opname/{did}/post")[0] == 409 and len(live_entries(did)) == 3)
    sc, _ = adj("A", "C", 1)
    check("F11. Freeze dilepas setelah Posted", sc == 200)

    # ===== G. Koreksi Posted via reversal + validasi
    all_lines = call("GET", f"opname/{did}")[1]["lines"]
    body = {"lines": [{"id": x["id"], "counted": (8 if x["id"] == LB else x["counted"]), "reason": x.get("reason")} for x in all_lines]}
    check("G1. Koreksi dengan baris tidak lengkap -> 400", call("PUT", f"transactions/opname/{did}", {"lines": body["lines"][:2]})[0] == 400)
    check("G2. Koreksi ganti gudang -> 400", call("PUT", f"transactions/opname/{did}", {**body, "warehouse_id": W["B"]["id"]})[0] == 400)
    sc, r = call("PUT", f"transactions/opname/{did}", body)
    check("G3. Koreksi B 7->8 via reversal -> stok B 8, ledger aktif tetap 3", sc == 200 and S("B", "A") == 8 and len(live_entries(did)) == 3, (sc, r))
    check("G4. Valuasi B setelah koreksi -4.000", vsum(did, I["B"]["id"]) == -4000, vsum(did, I["B"]["id"]))

    # ===== H. Live: rekonsiliasi kronologis + tanda hitung ulang
    lid_live = live["id"]
    call("PUT", f"transactions/opname/{lid_live}", {"date": TOMORROW})
    LL = call("GET", f"opname/{lid_live}")[1]["lines"][0]["id"]
    call("PUT", f"opname/{lid_live}/count", {"lines": [{"line_id": LL, "qty": 9, "reason": "Hilang"}]})
    adj("B", "A", 5)  # tanggal hari ini (<= waktu hitung), diinput SETELAH hitung
    ln = call("GET", f"opname/{lid_live}")[1]["lines"][0]
    check("H1. Mutasi backdate/same-day setelah hitung -> baris ditandai perlu hitung ulang", ln.get("needs_recount") is True, ln)
    sc, r = call("POST", f"opname/{lid_live}/workflow", {"action": "review"})
    check("H2. Review diblokir selama ada baris perlu hitung ulang", sc == 400 and "dihitung ulang" in str(r))
    call("PUT", f"opname/{lid_live}/count", {"lines": [{"line_id": LL, "qty": 14}]})
    ln = call("GET", f"opname/{lid_live}")[1]["lines"][0]
    check("H3. Hitung ulang: stok sistem saat hitung 15, selisih -1", ln["system_qty"] == 15 and ln["variance"] == -1 and not ln["needs_recount"], ln)
    adj("B", "A", 3, d=TOMORROW)  # mutasi SETELAH waktu hitung (kronologis)
    ln = call("GET", f"opname/{lid_live}")[1]["lines"][0]
    check("H4. Mutasi bertanggal setelah hitung tidak menandai ulang; selisih tetap -1", not ln["needs_recount"] and ln["variance"] == -1, ln)
    call("POST", f"opname/{lid_live}/workflow", {"action": "review"}); call("POST", f"opname/{lid_live}/submit")
    setst = "UPDATE opname SET doc=JSON_SET(doc,'$.status','Posting','$.posting_started_at',%s) WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.id'))=%s"
    sql(setst, (_dt.datetime.now(_dt.timezone.utc).isoformat(), lid_live))
    check("H4b. Posting sedang berjalan (belum basi) -> 409, tidak diproses ganda", call("POST", f"opname/{lid_live}/post")[0] == 409 and S("A", "B") == 18)
    sql(setst, ("2000-01-01T00:00:00+00:00", lid_live))  # simulasi crash di tengah posting (status Posting tertinggal)
    sc, _ = call("POST", f"opname/{lid_live}/post")
    check("H5. Pemulihan posting basi + Posting Live: stok 18 + (-1) = 17 (bukan fisik-snapshot = +4), tanpa double", sc == 200 and S("A", "B") == 17 and len(live_entries(lid_live)) == 1, (sc, S("A", "B")))

    # ===== I. Cancel melepas lock; delete dokumen aktif melepas lock
    sc, c1 = call("POST", "opname", {"date": TOMORROW, "warehouse_id": W["C"]["id"], "division_id": div["id"], "mode": "freeze"})
    check("I1. Cancel tanpa alasan -> 400", call("POST", f"opname/{c1['id']}/workflow", {"action": "cancel"})[0] == 400)
    sc, r = call("POST", f"opname/{c1['id']}/workflow", {"action": "cancel", "reason": "Salah gudang"})
    check("I2. Cancel -> Cancelled, freeze dilepas", sc == 200 and r["status"] == "Cancelled" and adj("C", "A", 1, d=TOMORROW)[0] == 200)
    sc, c2 = call("POST", "opname", {"date": TOMORROW, "warehouse_id": W["C"]["id"], "division_id": div["id"], "mode": "live"})
    check("I3. Setelah Cancelled, Opname baru di gudang sama boleh", sc == 200)
    sc, _ = call("PUT", f"transactions/opname/{c2['id']}", {"warehouse_id": W["A"]["id"]})
    check("I4. Gudang dikunci setelah snapshot (edit gudang -> 400)", sc == 400)

    # ===== J. Template + Import ke dokumen yang sama (preview -> commit), tidak posting
    r = T.S.get(f"{API}/opname/{c2['id']}/template.xlsx")
    wb = load_workbook(io.BytesIO(r.content)); ws = wb.active
    check("J1. Template dari snapshot: No. Opname, gudang, header kolom", ws["B1"].value == c2["no"] and ws.cell(row=6, column=6).value == "Stok Fisik" and ws.max_row >= 7)
    ws.cell(row=7, column=6, value=4); ws.cell(row=7, column=7, value="Rusak 2")
    ws.append([99, f"NOPE{u}", "x", "PCS", 0, 1, "", ""]); ws.append([100, ws.cell(row=7, column=2).value, "", "PCS", 0, 2, "", ""])
    b = io.BytesIO(); wb.save(b)
    pv = T.S.post(f"{API}/opname/{c2['id']}/import?mode=preview", files={"file": ("h.xlsx", b.getvalue())}).json()
    check("J2. Preview: error kode tidak dikenal + duplikat dengan nomor baris Excel", any("Baris 8" in e and "tidak ditemukan" in e for e in pv["errors"]) and any("Baris 9" in e and "duplikat" in e for e in pv["errors"]), pv)
    rc = T.S.post(f"{API}/opname/{c2['id']}/import?mode=commit", files={"file": ("h.xlsx", b.getvalue())})
    check("J3. Commit dengan error -> 400, tidak ada perubahan", rc.status_code == 400 and call("GET", f"opname/{c2['id']}")[1]["summary"]["counted"] == 0)
    ws.delete_rows(8, 2); b = io.BytesIO(); wb.save(b)
    rc = T.S.post(f"{API}/opname/{c2['id']}/import?mode=commit", files={"file": ("h.xlsx", b.getvalue())}).json()
    _, d2 = call("GET", f"opname/{c2['id']}")
    check("J4. Commit valid -> dokumen SAMA ter-update, status tetap Counting", rc.get("applied") == 1 and d2["summary"]["counted"] == 1 and d2["status"] == "Counting" and len([x for x in OL() if x.get("warehouse_id") == W["C"]["id"]]) == 2, rc)
    ws["B1"] = "OPN/LAIN"; b = io.BytesIO(); wb.save(b)
    pv = T.S.post(f"{API}/opname/{c2['id']}/import?mode=preview", files={"file": ("h.xlsx", b.getvalue())}).json()
    check("J5. Nomor Opname berbeda -> error", any("Nomor Opname" in e for e in pv["errors"]))

    # ===== K. Delete Posted -> reversal; legacy; tenant isolation
    sc, r = call("DELETE", f"transactions/opname/{did}")
    check("K1. Hapus Posted: reversal penuh atau dikunci bila stok sudah dipakai",
          (sc == 200 and (S("A", "A"), S("B", "A"), S("D", "A")) == (10, 10, 0)) or (sc == 409 and S("B", "A") == 8), (sc, r))
    lg = str(uuid.uuid4())
    call("POST", "opname", {"date": TODAY, "warehouse_id": W["C"]["id"], "division_id": div["id"]})  # 409 (aktif) - abaikan
    sql("INSERT INTO opname (pk, doc, created_at, updated_at) SELECT %s, JSON_REMOVE(JSON_SET(doc,'$.id',%s,'$.no','OPN-LEGACY','$.status','Posted'),'$.workflow_version','$.history'), NOW(), NOW() FROM opname WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.id'))=%s", (lg[:36], lg, c2["id"]))
    sc, r = call("GET", f"opname/{lg}")
    check("K2. Dokumen legacy (tanpa workflow_version) tetap terbaca", sc == 200 and r.get("legacy") is True, (sc, str(r)[:200]))
    sql("DELETE FROM opname WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.id'))=%s", (lg,))
    mine = dict(T.S.headers); T.S.headers.clear(); T.S.headers.update(other_hdr); jar = T.S.cookies.copy(); T.S.cookies.clear()
    sc, _ = call("GET", f"opname/{did}")
    T.S.headers.clear(); T.S.headers.update(mine); T.S.cookies.update(jar)
    check("K3. Tenant lain tidak dapat membaca Opname (404)", sc in (403, 404), sc)

    # ===== L. Lampiran draft Opname: multi-file -> terikat ke dokumen final; isolasi tenant; draft sekali pakai
    def upload(draft_id, name):
        return T.S.post(f"{API}/attachments", files={"file": (name, io.BytesIO(b"%PDF-1.4 " + name.encode()), "application/pdf")},
                        data={"entity": "opname_draft", "entity_id": draft_id, "category": "Lampiran Stock Opname"})
    sc, dr = call("POST", "attachment-drafts", {"module": "opname"})
    check("L1. Draft lampiran modul opname -> 200", sc == 200 and dr.get("id"), (sc, dr))
    ups = [upload(dr["id"], n) for n in ("foto-rak.pdf", "ba-hitung.pdf")]
    check("L2. Multi-file upload ke draft opname -> 2x 200", all(x.status_code == 200 for x in ups), [x.status_code for x in ups])
    T.S.headers.clear(); T.S.headers.update(other_hdr); jar = T.S.cookies.copy(); T.S.cookies.clear()
    o_list = call("GET", f"attachments?entity=opname_draft&entity_id={dr['id']}")
    o_up = upload(dr["id"], "asing.pdf").status_code
    o_del = call("DELETE", f"attachments/{ups[0].json()['id']}")[0]
    o_dl = T.S.get(f"{API}/attachments/{ups[0].json()['id']}/download").status_code
    T.S.headers.clear(); T.S.headers.update(mine); T.S.cookies.update(jar)
    # DELETE lampiran lintas tenant: endpoint lampiran existing (generik, semua modul) menjawab 200 tanpa efek bila id tidak
    # ada di tenant pemanggil -> yang diuji adalah efeknya: lampiran pemilik tetap utuh & bisa diunduh pemilik.
    still = call("GET", f"attachments?entity=opname_draft&entity_id={dr['id']}")[1]
    own_dl = T.S.get(f"{API}/attachments/{ups[0].json()['id']}/download").status_code
    check("L3. Tenant lain tidak dapat membaca/unggah/hapus/unduh lampiran draft opname",
          (o_list[0] in (403, 404) or (o_list[0] == 200 and o_list[1] == [])) and o_up in (400, 403, 404) and o_dl in (403, 404)
          and o_del in (200, 403, 404) and len(still) == 2 and own_dl == 200,
          (o_list, o_up, o_del, o_dl, len(still), own_dl))
    sc, od = call("POST", "opname", {"date": TODAY, "warehouse_id": W["A"]["id"], "division_id": div["id"], "mode": "live", "attachment_draft_id": dr["id"]})
    check("L4. Create Opname dengan draft lampiran -> 200", sc == 200, (sc, od))
    sc, att = call("GET", f"attachments?entity=opname&entity_id={od.get('id')}")
    check("L5. 2 lampiran draft terikat ke dokumen Opname final", sc == 200 and len(att) == 2, (sc, att))
    sc, r = call("POST", "opname", {"date": TODAY, "warehouse_id": W["B"]["id"], "division_id": div["id"], "mode": "live", "attachment_draft_id": dr["id"]})
    check("L6. Draft yang sudah terpakai tidak dapat diklaim ulang (ditolak)", sc in (400, 409), (sc, r))
    sc, ob = call("POST", "opname", {"date": TODAY, "warehouse_id": W["B"]["id"], "division_id": div["id"], "mode": "live"})
    check("L7. Klaim ulang gagal tidak meninggalkan lock gudang (create normal -> 200)", sc == 200, (sc, ob))
    for x in (od, ob):
        call("POST", f"opname/{x.get('id')}/workflow", {"action": "cancel", "reason": "QA lampiran"})

    # ===== M. Approve & Post bersamaan: klaim CAS dapat memicu deadlock InnoDB (1213) pada transaksi kalah;
    #          wajib dijawab 409 (bukan 500), tepat 1 posting, tanpa movement ganda.
    sc, pm = call("POST", "opname", {"date": TOMORROW, "warehouse_id": W["B"]["id"], "division_id": div["id"], "mode": "freeze", "scope": "all"})
    pl = call("GET", f"opname/{pm['id']}")[1]["lines"]
    tgt = next(x for x in pl if float(x["system_qty"]) >= 1)
    call("PUT", f"opname/{pm['id']}/count", {"lines": [{"line_id": x["id"], "qty": float(x["system_qty"]) - (1 if x["id"] == tgt["id"] else 0),
                                                        **({"reason": "QA paralel"} if x["id"] == tgt["id"] else {})} for x in pl]})
    call("POST", f"opname/{pm['id']}/workflow", {"action": "review"})
    check("M1. Dokumen siap posting (Waiting Approval)", call("POST", f"opname/{pm['id']}/submit")[1].get("status") == "Waiting Approval")
    pq0, pv0 = pool(tgt["item_id"], W["B"]["id"])
    ph, pc = dict(T.S.headers), T.S.cookies.get_dict()
    bar = __import__("threading").Barrier(6)

    def par_post(_):
        s = requests.Session(); s.headers.update(ph); s.cookies.update(pc)
        bar.wait()
        rr = s.post(f"{API}/opname/{pm['id']}/post", timeout=120)
        return rr.status_code, (str(rr.json().get("detail"))[:160] if rr.status_code != 200 else "Posted")
    with cf.ThreadPoolExecutor(6) as ex:
        pouts = list(ex.map(par_post, range(6)))
    pcodes = [c for c, _ in pouts]
    check("M2. 6 Approve & Post bersamaan -> tepat 1x200, 5x409, tanpa 500", pcodes.count(200) == 1 and pcodes.count(409) == 5,
          [(c, d) for c, d in pouts if c != 409] or pcodes)
    check("M3. Tepat 1 movement aktif, status Posted (tanpa double posting)",
          len(live_entries(pm["id"])) == 1 and call("GET", f"opname/{pm['id']}")[1].get("status") == "Posted", len(live_entries(pm["id"])))
    pq1, pv1 = pool(tgt["item_id"], W["B"]["id"])
    vrows = [json.loads(d) for (d,) in sql("SELECT doc FROM valuation_ledger WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.doc_id'))=%s", (pm["id"],))]
    check("M4. Stok pool turun TEPAT 1 & nilai persediaan = nilai ledger valuasi tunggal (tanpa ledger ganda)",
          abs((pq0 - pq1) - 1) < 1e-6 and len(vrows) == 1 and abs(round(pv1 - pv0, 4) - vsum(pm["id"], tgt["item_id"])) < 0.01,
          (pq0, pq1, pv0, pv1, len(vrows)))

    failed = [n for n, ok in T.RESULTS if not ok]
    print(f"\n{len(T.RESULTS) - len(failed)}/{len(T.RESULTS)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
