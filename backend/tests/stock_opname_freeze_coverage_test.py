"""Stock Opname Freeze — cakupan SEMUA jalur mutasi (DO, MI, Loan, Return, Transfer, Adjustment, edit/delete+reversal,
set saldo langsung, import Excel saldo/limit, hapus master) + gudang lain tetap jalan + posting Opname sah.

Jalankan (backend aktif di :8001, DB MariaDB TERISOLASI dari backend/.env):
  /root/.venv/bin/python backend/tests/stock_opname_freeze_coverage_test.py
"""
import datetime as _dt
import io
import json
import sys
import uuid

import receipt_control_test as rc
import transfer_multi_warehouse_test as TM
from openpyxl import load_workbook

call, check, API, S = rc.call, rc.check, rc.API, rc.S
TODAY = _dt.datetime.now(_dt.timezone.utc).date().isoformat()  # konvensi tanggal = default server now_iso() (UTC)
u = uuid.uuid4().hex[:6]


def sql(q, args=()):
    with TM.db_conn() as cx, cx.cursor() as cur:
        cur.execute(q, args)
        return cur.fetchall()


def iw(item_id, wh):
    r = sql("SELECT doc FROM item_warehouse WHERE JSON_UNQUOTE(JSON_EXTRACT(doc,'$.item_id'))=%s AND JSON_UNQUOTE(JSON_EXTRACT(doc,'$.warehouse_id'))=%s", (item_id, wh))
    return json.loads(r[0][0]) if r else None


def snap(item_id, wh):
    d = iw(item_id, wh) or {}
    return (round(float(d.get("current_stock") or 0), 6), round(float(d.get("total_value") or 0), 4), d.get("min_stock"))


def ledger_count():
    return sql("SELECT COUNT(*) FROM stock_ledger")[0][0], sql("SELECT COUNT(*) FROM valuation_ledger")[0][0]


def xlsx_import(dataset, rows):
    r = S.get(f"{API}/excel/template/{dataset}")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Data"]
    headers = [c.value for c in ws[1]]
    ws.delete_rows(2, ws.max_row)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    b = io.BytesIO()
    wb.save(b)
    rr = S.post(f"{API}/excel/import/{dataset}", files={"file": ("x.xlsx", b.getvalue())})
    return rr.status_code, (rr.json() if rr.headers.get("content-type", "").startswith("application/json") else rr.text)


def main():
    M = rc.setup()
    wh, wh2, item, item2, div = M["wh"]["id"], M["wh2"]["id"], M["item"]["id"], M["item2"]["id"], M["div"]["id"]
    sc, it3 = call("POST", "master/items", {"code": f"IZ{u}", "name": "Barang Nol", "unit": "PCS", "is_active": True, **rc.item_ref(M)})
    sc, it4 = call("POST", "master/items", {"code": f"IY{u}", "name": "Barang Saldo Awal", "unit": "PCS", "is_active": True, **rc.item_ref(M)})
    for w, i, q in ((wh, item, 50), (wh2, item, 20), (wh, item2, 30)):
        call("POST", "adjustments", {"date": TODAY, "warehouse_id": w, "division_id": div, "reason": "seed",
                                     "lines": [{"item_id": i, "adjustment": q, "approved_unit_cost": 1000}]}, 200)
    call("POST", "item-warehouse", {"item_id": it3["id"], "warehouse_id": wh, "min_stock": 1, "max_stock": 5}, 200)
    po, _, _ = rc.make_po(M, "supX", 10)
    sc, mI = call("POST", "mro", {"no": f"MRO-F-{u}", "division_id": div, "requester": "Budi", "submitted": True,
                                  "lines": [{"item_id": item, "qty": 3, "warehouse_id": wh, "project_id": M["pa"]["id"]}]}, 200)
    call("POST", f"mro/{mI['id']}/submit", {})
    ml = mI["lines"][0]
    mi_line = {"item_id": item, "qty": 1, "warehouse_id": wh, "mro_id": mI["id"], "mro_line_id": ml["id"],
               "sources": [{"mro_id": mI["id"], "line_id": ml["id"], "qty": 1}]}
    # dokumen Posted sebelum Freeze (untuk uji edit/delete/return saat Freeze)
    sc, adj_old = call("POST", "adjustments", {"date": TODAY, "warehouse_id": wh, "division_id": div, "reason": "lama",
                                               "lines": [{"item_id": item2, "adjustment": 2, "approved_unit_cost": 1000}]}, 200)
    sc, trf_old = call("POST", "transfers", {"date": TODAY, "division_id": div, "from_warehouse_id": wh2, "to_warehouse_id": wh,
                                             "lines": [{"item_id": item, "qty": 1}]}, 200)
    sc, loan = call("POST", "loans", {"date": TODAY, "division_id": div, "project_id": M["pa"]["id"], "from_warehouse_id": wh,
                                      "to_warehouse_id": wh2, "due_date": TODAY, "requester": "Budi",
                                      "lines": [{"item_id": item, "qty": 2, "uom_id": M["uom"]["id"], "from_warehouse_id": wh,
                                                 "to_warehouse_id": wh2, "project_id": M["pa"]["id"]}]}, 200)
    loan = call("GET", f"loans/{loan['id']}")[1]
    # dokumen Posted tambahan (sebelum Freeze) untuk jalur edit/delete DO, MI, Loan, Loan Return
    sc, do_old = call("POST", "do", rc.do_body(M, po, 2, "supX"), 200)
    sc, mi_old = call("POST", "mi", {"source_type": "MRO", "lines": [mi_line]}, 200)
    call("POST", f"loans/{loan['id']}/return", {"date": TODAY, "lines": [{"loan_line_id": loan["lines"][0]["id"], "qty": 1}]}, 200)
    ret_old = call("GET", f"loans/{loan['id']}/returns")[1][0]
    sc, loan2 = call("POST", "loans", {"date": TODAY, "division_id": div, "project_id": M["pa"]["id"], "from_warehouse_id": wh,
                                       "to_warehouse_id": wh2, "due_date": TODAY, "requester": "Budi",
                                       "lines": [{"item_id": item, "qty": 1, "uom_id": M["uom"]["id"], "from_warehouse_id": wh,
                                                  "to_warehouse_id": wh2, "project_id": M["pa"]["id"]}]}, 200)

    # ===== Freeze gudang A
    sc, op = call("POST", "opname", {"date": TODAY, "warehouse_id": wh, "division_id": div, "mode": "freeze", "scope": "all"})
    check("Z0. Opname Freeze Gudang A dibuat", sc == 200 and op.get("status") == "Counting", (sc, op))
    no = op["no"]
    s_a = {i: snap(i, wh) for i in (item, item2, it3["id"])}
    s_b = snap(item, wh2)
    led0 = ledger_count()

    def frozen(sc, r):
        return sc == 409 and no in json.dumps(r) and "Mutasi stok tidak dapat dilakukan" in json.dumps(r)

    sc, r = call("POST", "do", rc.do_body(M, po, 3, "supX"))
    check("Z1. DO (penerimaan) ke gudang Freeze -> 409 + nomor Opname", frozen(sc, r), (sc, r))
    sc, r = call("POST", "mi", {"source_type": "MRO", "lines": [mi_line]})
    check("Z2. MI (pengeluaran) dari gudang Freeze -> 409", frozen(sc, r), (sc, r))
    sc, r = call("POST", "loans", {"date": TODAY, "division_id": div, "project_id": M["pa"]["id"], "from_warehouse_id": wh,
                                   "to_warehouse_id": wh2, "due_date": TODAY, "requester": "Budi",
                                   "lines": [{"item_id": item, "qty": 1, "uom_id": M["uom"]["id"], "from_warehouse_id": wh,
                                              "to_warehouse_id": wh2, "project_id": M["pa"]["id"]}]})
    check("Z3. Loan dari gudang Freeze -> 409", frozen(sc, r), (sc, r))
    sc, r = call("POST", f"loans/{loan['id']}/return", {"date": TODAY, "lines": [{"loan_line_id": loan["lines"][0]["id"], "qty": 1}]})
    check("Z4. Loan Return ke gudang Freeze -> 409", frozen(sc, r), (sc, r))
    sc, r = call("POST", "adjustments", {"date": TODAY, "warehouse_id": wh2, "division_id": div, "reason": "multi",
                                         "lines": [{"item_id": item, "adjustment": 1, "approved_unit_cost": 1000, "warehouse_id": wh2},
                                                   {"item_id": item, "adjustment": 1, "approved_unit_cost": 1000, "warehouse_id": wh}]})
    check("Z5. Adjustment multi-gudang (1 baris gudang Freeze) -> 409, gudang lain tidak berubah (atomik)",
          frozen(sc, r) and snap(item, wh2) == s_b, (sc, r, snap(item, wh2), s_b))
    sc, r = call("DELETE", f"transactions/adjustment/{adj_old['id']}")
    check("Z6. Hapus Adjustment Posted (reversal) di gudang Freeze -> 409", frozen(sc, r), (sc, r))
    # Pengecualian posting Opname TIDAK dapat dipalsukan lewat payload/jalur modul 'opname'
    sc, r = call("POST", "adjustments", {"date": TODAY, "warehouse_id": wh, "division_id": div, "reason": "forge",
                                         "opname_id": op["id"], "doc_id": op["id"], "source_type": "opname", "module": "opname",
                                         "lines": [{"item_id": item, "adjustment": 1, "approved_unit_cost": 1000, "warehouse_id": wh,
                                                    "opname_id": op["id"]}]})
    check("Z5b. Payload memalsukan opname_id/doc_id/module=opname -> tetap 409", frozen(sc, r) and snap(item, wh)[:2] == s_a[item][:2], (sc, r))
    sc, r = call("DELETE", f"transactions/opname/{adj_old['id']}")
    check("Z5c. Reversal dokumen lain lewat path modul 'opname' -> ditolak, saldo tetap",
          sc != 200 and snap(item2, wh)[:2] == s_a[item2][:2], (sc, r))
    # Transfer BARU: gudang Freeze sebagai ASAL maupun TUJUAN -> 409 sebelum mutasi pertama; gudang lawan tidak berubah
    led_t = ledger_count()
    sc, r = call("POST", "transfers", {"date": TODAY, "division_id": div, "from_warehouse_id": wh, "to_warehouse_id": wh2,
                                       "lines": [{"item_id": item, "qty": 1}]})
    check("Z7a. Transfer baru DARI gudang Freeze -> 409, gudang tujuan & ledger tidak berubah",
          frozen(sc, r) and snap(item, wh2) == s_b and snap(item, wh)[:2] == s_a[item][:2] and ledger_count() == led_t, (sc, r))
    sc, r = call("POST", "transfers", {"date": TODAY, "division_id": div, "from_warehouse_id": wh2, "to_warehouse_id": wh,
                                       "lines": [{"item_id": item, "qty": 1}]})
    check("Z7b. Transfer baru KE gudang Freeze -> 409, gudang asal & ledger tidak berubah (atomik)",
          frozen(sc, r) and snap(item, wh2) == s_b and snap(item, wh)[:2] == s_a[item][:2] and ledger_count() == led_t, (sc, r))
    sc, r = call("PUT", f"transactions/transfer/{trf_old['id']}", {"date": TODAY, "division_id": div, "from_warehouse_id": wh2,
                                                                     "to_warehouse_id": wh, "lines": [{"item_id": item, "qty": 2}]})
    check("Z7. Edit Transfer Posted yang menyentuh gudang Freeze -> 409", frozen(sc, r), (sc, r))
    sc, r = call("DELETE", f"transactions/transfer/{trf_old['id']}")
    check("Z8. Hapus Transfer Posted (gudang tujuan Freeze) -> 409, gudang asal tidak berubah",
          frozen(sc, r) and snap(item, wh2) == s_b, (sc, r))
    # jalur edit/delete + reversal dokumen Posted lain yang menyentuh gudang Freeze (pre-check sebelum mutasi pertama)
    sc, r = call("PUT", f"transactions/adjustment/{adj_old['id']}", {"date": TODAY, "warehouse_id": wh, "division_id": div, "reason": "edit",
                                                                       "lines": [{"item_id": item2, "adjustment": 3, "approved_unit_cost": 1000}]})
    check("Z6b. Edit Adjustment Posted di gudang Freeze -> 409", frozen(sc, r), (sc, r))
    sc, r = call("DELETE", f"transactions/do/{do_old['id']}")
    check("Z1b. Hapus DO Posted (reversal penerimaan) gudang Freeze -> 409", frozen(sc, r), (sc, r))
    sc, r = call("PUT", f"transactions/do/{do_old['id']}", rc.do_body(M, po, 1, "supX"))
    check("Z1c. Edit DO Posted gudang Freeze -> 409", frozen(sc, r), (sc, r))
    sc, r = call("DELETE", f"transactions/mi/{mi_old['id']}")
    check("Z2b. Hapus MI Posted (reversal pengeluaran) gudang Freeze -> 409", frozen(sc, r), (sc, r))
    sc, r = call("DELETE", f"transactions/loan/{loan2['id']}")
    check("Z3b. Hapus Loan Posted (gudang asal Freeze) -> 409, gudang tujuan tidak berubah", frozen(sc, r) and snap(item, wh2) == s_b, (sc, r))
    sc, r = call("PUT", f"loan-returns/{ret_old['id']}", {"date": TODAY, "notes": "edit saat freeze",
                                                          "lines": [{"loan_line_id": loan["lines"][0]["id"], "qty": 2}]})
    check("Z4b. Edit Loan Return (masuk ke gudang Freeze) -> 409, gudang peminjam tidak berubah", frozen(sc, r) and snap(item, wh2) == s_b, (sc, r))
    sc, r = call("DELETE", f"loan-returns/{ret_old['id']}")
    check("Z4c. Hapus Loan Return (reversal gudang Freeze) -> 409", frozen(sc, r) and snap(item, wh2) == s_b, (sc, r))
    sc, r = call("POST", "item-warehouse", {"item_id": item, "warehouse_id": wh, "current_stock": 999, "min_stock": 1, "max_stock": 9})
    # current_stock dari klien selalu dibuang guard existing (item_warehouse_guard_layer) -> saldo tidak pernah berubah.
    check("Z9. Set saldo langsung (item-warehouse) gudang Freeze -> saldo tidak berubah", snap(item, wh)[:2] == s_a[item][:2], (sc, r, snap(item, wh)))
    sc, r = call("POST", "item-warehouse", {"item_id": item, "warehouse_id": wh, "min_stock": 3, "max_stock": 90})
    check("Z10. Ubah min/max saja (stok tidak berubah) saat Freeze -> diizinkan", sc == 200 and snap(item, wh)[2] == 3, (sc, r))
    sc, r = xlsx_import("stock_limits", [{"item_code": M["item"]["code"], "warehouse_code": M["wh2"]["code"], "min_stock": 7, "max_stock": 70},
                                         {"item_code": M["item"]["code"], "warehouse_code": M["wh"]["code"], "min_stock": 1, "max_stock": 10, "current_stock": 500}])
    check("Z11. Import Excel limit stok berisi saldo gudang Freeze -> 409, baris lain TIDAK ikut terimport",
          frozen(sc, r) and snap(item, wh2)[2] != 7, (sc, r, snap(item, wh2)))
    sc, r = xlsx_import("stock_limits", [{"item_code": M["item"]["code"], "warehouse_code": M["wh"]["code"], "min_stock": 4, "max_stock": 40}])
    check("Z12. Import Excel limit tanpa saldo (min/max) gudang Freeze -> diizinkan", sc == 200 and snap(item, wh)[2] == 4, (sc, r))
    sc, r = xlsx_import("opening_inventory", [{"item_code": it4["code"], "item_name": it4["name"], "uom_code": M["uom"]["code"],
                                              "warehouse_code": M["wh"]["code"], "qty": 5, "purchase_price": 1000}])
    check("Z13. Import Saldo Awal ke gudang Freeze -> 409, tanpa pool/ledger baru", frozen(sc, r) and iw(it4["id"], wh) is None, (sc, r))
    sc, r = call("POST", "valuation/opening", {"item_id": item2, "warehouse_id": wh, "opening_avg_cost": 1234, "cutoff_date": TODAY})
    check("Z14. Opening valuation pada pool gudang Freeze -> ditolak, nilai pool tidak berubah",
          sc in (400, 409) and snap(item2, wh) == s_a[item2], (sc, r))
    sc, r = call("DELETE", f"master/items/{it3['id']}")
    check("Z15. Hapus master barang yang memiliki baris gudang Freeze -> 409", sc == 409 and iw(it3["id"], wh) is not None, (sc, r))
    sc, r = call("POST", "master-bulk-delete/items", {"ids": [it3["id"]]})
    check("Z15b. Hapus massal master barang (baris gudang Freeze) -> ditolak, baris gudang tetap ada",
          sc in (400, 409) and iw(it3["id"], wh) is not None, (sc, r))
    sc, r = call("POST", "item-warehouse/clear-minmax", {"item_id": it3["id"], "warehouse_id": wh})
    check("Z15c. Kosongkan min/max (bukan mutasi stok) saat Freeze -> diizinkan, saldo/nilai tetap",
          sc == 200 and snap(it3["id"], wh)[:2] == s_a[it3["id"]][:2], (sc, r))
    sc, dr = call("POST", "valuation/replay/dry-run", {"warehouse_id": wh})
    check("Z16. Valuation replay DRY-RUN (read-only) tetap diizinkan saat Freeze", sc == 200, (sc, str(dr)[:200]))
    check("Z17. Saldo + nilai gudang Freeze tidak berubah (kecuali min/max)",
          all(snap(i, wh)[:2] == s_a[i][:2] for i in s_a), ({i: snap(i, wh) for i in s_a}, s_a))
    sc, r = call("POST", "adjustments", {"date": TODAY, "warehouse_id": wh2, "division_id": div, "reason": "lain",
                                         "lines": [{"item_id": item, "adjustment": 1, "approved_unit_cost": 1000}]})
    check("Z18. Gudang lain (tanpa Freeze) tetap bisa mutasi", sc == 200 and snap(item, wh2)[0] == s_b[0] + 1, (sc, r))
    led1 = ledger_count()
    check("Z19. Ledger hanya bertambah 1 movement (mutasi gudang lain) — tidak ada ledger yatim",
          (led1[0] - led0[0], led1[1] - led0[1]) == (1, 1), (led0, led1))

    # ===== Posting Opname sah tetap jalan, Freeze dilepas
    d = call("GET", f"opname/{op['id']}")[1]
    call("PUT", f"opname/{op['id']}/count", {"lines": [{"line_id": l["id"], "qty": l["system_qty"]} for l in d["lines"]]})
    call("POST", f"opname/{op['id']}/workflow", {"action": "review"})
    call("POST", f"opname/{op['id']}/submit")
    sc, r = call("POST", f"opname/{op['id']}/post")
    check("Z20. Posting Opname pemilik Freeze -> Posted", sc == 200 and r.get("status") == "Posted", (sc, r))
    sc, r = call("POST", "do", rc.do_body(M, po, 3, "supX"))
    check("Z21. Setelah Posted, DO ke gudang A kembali normal", sc == 200, (sc, r))
    sc, r = call("POST", f"loans/{loan['id']}/return", {"date": TODAY, "lines": [{"loan_line_id": loan["lines"][0]["id"], "qty": 1}]})
    check("Z22. Setelah Posted, Loan Return kembali normal", sc == 200, (sc, r))
    sc, r = call("DELETE", f"transactions/adjustment/{adj_old['id']}")
    check("Z23. Setelah Posted, hapus Adjustment (reversal) kembali normal", sc == 200, (sc, r))
    sc, r = call("DELETE", f"transactions/loan/{loan2['id']}")
    check("Z24. Setelah Posted, hapus Loan kembali normal", sc == 200, (sc, r))
    sc, r = call("DELETE", f"transactions/mi/{mi_old['id']}")
    check("Z25. Setelah Posted, hapus MI kembali normal", sc == 200, (sc, r))

    failed = [n for n, ok in rc.RESULTS if not ok]
    print(f"\n{len(rc.RESULTS) - len(failed)}/{len(rc.RESULTS)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
