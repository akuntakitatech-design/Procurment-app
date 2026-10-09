"""P2a Reporting Operasional Procurement — Register MRO/RO/PO/DO/MI, Lead Time, Pemakaian per Unit/Proyek, MRO
Traceability (filter barang/kategori), dan konsolidasi 5 card Pusat Laporan.

Hanya pada backend + DB test terisolasi (DB *itest*; tenant QA via /saas/register) — BUKAN preview/sandbox/production.
Membuktikan: parsial & multi-referensi (1 PO -> 2 DO, 1 MRO -> 2 RO) tanpa nilai berganda, pembatalan MI (reversal) tidak
dihitung, Nilai PO / Nilai DO / HPP MI terpisah dan konsisten dengan MRO Traceability & Rekap HPP MI, filter periode WIB/
supplier/proyek/status/pencarian, scope divisi, redaksi harga server-side (JSON/Excel/PDF/total), export = seluruh hasil
filter, data besar (pagination + export tanpa pemotongan), endpoint lama tetap, read-only.
"""
import io
import os
import sys
import uuid
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

import receipt_control_test as T
from master_item_stock_test import mk_user

call, check, API = T.call, T.check, T.API
TODAY = datetime.now(ZoneInfo("Asia/Jakarta")).date()
EPS = 0.01


def near(a, b, eps=EPS):
    return abs(float(a or 0) - float(b or 0)) <= eps


def rc(caller, key, qs=""):
    return caller("GET", f"report-center/{key}?page_size=500&{qs}")


def get_bin(u, path):
    s = u.S if u is not None else T.S
    r = s.get(f"{API}/{path}")
    return r.status_code, r.content, r.headers


def xlsx_rows(content, first_header):
    vals = [list(r) for r in load_workbook(io.BytesIO(content), read_only=True).worksheets[0].iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(vals) if r and r[0] == first_header), None)
    return (vals[hi], [r for r in vals[hi + 1:] if r and r[0] and r[0] != "TOTAL"]) if hi is not None else ([], [])


def main():
    dbn = (os.environ.get("TEST_DATABASE_URL") or "").rsplit("/", 1)[-1].split("?")[0]
    if not ("itest" in dbn or dbn.endswith("_test")):
        print("Menolak berjalan: hanya untuk backend + DB test terisolasi (scripts/run_regression_itest.sh, DB *itest*).")
        sys.exit(2)
    M = T.setup()
    u = uuid.uuid4().hex[:5]
    unit = call("POST", "master/units", {"code": f"UN{u}", "name": "Dump Truck P2", "plate_no": f"B {u} QA", "is_active": True}, 200)[1]
    dB = call("POST", "master/divisions", {"code": f"DB{u}", "name": f"Divisi B {u}"}, 200)[1]

    # ---- Alur 1: MRO1 10 -> RO1 -> PO1 (Supplier X, Project A) -> DO1 4 + DO2 6 -> MI1 3, MI2 2 (MI2 dibatalkan)
    po1, mro1_no, ro1_no = T.make_po(M, "supX", 10, project="pa")
    sc, do1 = call("POST", "do", T.do_body(M, po1, 4, "supX"))
    sc2, do2 = call("POST", "do", T.do_body(M, po1, 6, "supX"))
    check("setup: 1 PO -> 2 DO parsial (4 + 6)", sc == 200 and sc2 == 200, (sc, do1, sc2, do2))
    sc, mros = call("GET", "mro")
    mros = mros if isinstance(mros, list) else (mros or {}).get("rows") or (mros or {}).get("items") or []
    mro1 = next(m for m in mros if m.get("no") == mro1_no)
    sc, mro1d = call("GET", f"mro/{mro1['id']}")
    ml = mro1d["lines"][0]

    def mi(q):
        return call("POST", "mi", {"division_id": M["div"]["id"], "source_type": "MRO", "default_warehouse_id": M["wh"]["id"],
                                   "receiver": "Operator", "lines": [{
                                       "item_id": M["item"]["id"], "qty": q, "uom_id": M["uom"]["id"], "conversion_factor": 1,
                                       "warehouse_id": M["wh"]["id"], "project_id": M["pa"]["id"], "unit_id": unit["id"],
                                       "mro_id": mro1["id"], "mro_line_id": ml["id"],
                                       "sources": [{"mro_id": mro1["id"], "line_id": ml["id"], "qty": q, "base_qty": q}]}]})
    sc, mi1 = mi(3)
    sc2, mi2 = mi(2)
    check("setup: MI1 3 & MI2 2 dari MRO1 (Unit, Project A)", sc == 200 and sc2 == 200, (sc, mi1, sc2, mi2))
    sc, dl = call("DELETE", f"transactions/mi/{mi2.get('id')}")
    check("setup: MI2 dibatalkan (hapus -> reversal ledger)", sc == 200, (sc, dl))

    # ---- Alur 2: MRO2 10 -> RO2a 6 + RO2b 4 (multi-referensi) -> PO2 dari RO2a (Supplier Y, Project B)
    sc, mro2 = call("POST", "mro", {"no": f"MRO-B{u}", "division_id": M["div"]["id"], "requester": "Sari", "submitted": True,
                                    "lines": [{"item_id": M["item"]["id"], "qty": 10, "warehouse_id": M["wh"]["id"], "project_id": M["pb"]["id"]}]}, 200)
    call("POST", f"mro/{mro2['id']}/submit", {})
    m2l = mro2["lines"][0]
    ros = []
    for q in (6, 4):
        sc, ro = call("POST", "ro", {"division_id": M["div"]["id"], "submitted": True, "lines": [{
            "item_id": M["item"]["id"], "qty": q, "warehouse_id": M["wh"]["id"], "project_id": M["pb"]["id"],
            "sources": [{"mro_id": mro2["id"], "line_id": m2l["id"], "qty": q}]}]}, 200)
        call("POST", f"ro/{ro['id']}/submit", {})
        ros.append(ro)
    sc, po2 = call("POST", "po", {"supplier_id": M["supY"]["id"], "division_id": M["div"]["id"], "lines": [{
        "item_id": M["item"]["id"], "qty": 6, "price": 2000, "warehouse_id": M["wh"]["id"], "project_id": M["pb"]["id"],
        "sources": [{"ro_id": ros[0]["id"], "line_id": ros[0]["lines"][0]["id"], "qty": 6}]}]}, 200)
    check("setup: MRO2 -> 2 RO (6 + 4) -> PO2 (6, Supplier Y)", bool(po2.get("id")), po2)
    sc, po1d = call("GET", f"po/{po1['id']}")
    sc, po2d = call("GET", f"po/{po2['id']}")

    NP = mk_user("warehouse", divs=[M["div"]["id"]], overrides={"export": "allow"})
    UB = mk_user("manager", divs=[dB["id"]], overrides={"export": "allow"})
    sc, iw0 = call("GET", "item-warehouse")

    # ------------------------------------------------------------------ C. katalog 5 card
    sc, cat = call("GET", "report-center/catalog")
    g = {x["key"]: x for x in cat.get("groups", [])}
    check("C1 5 card kategori berurutan + judul", [x["key"] for x in cat["groups"]] == ["persediaan", "procurement", "warehouse", "spk", "hutang"]
          and [x["title"] for x in cat["groups"]] == ["Persediaan & Nilai Persediaan", "Procurement", "Warehouse", "SPK & Kontrak Vendor", "Invoice & Hutang"],
          [x["title"] for x in cat.get("groups", [])])
    check("C2 card Procurement: MRO Traceability + 5 register + Lead Time + Pemakaian + P2b outstanding/rekap (tanpa duplikat)",
          [r["key"] for r in g["procurement"]["reports"]] == ["mro-traceability", "register-mro", "register-ro", "register-po", "register-do",
                                                            "register-mi", "lead-time", "pemakaian-barang", "outstanding-mro",
                                                            "outstanding-ro", "outstanding-po", "outstanding-do", "rekap-pembelian",
                                                            "rekap-nilai-do"], [r["key"] for r in g["procurement"]["reports"]])
    allkeys = [r["key"] for x in cat["groups"] for r in x["reports"]]
    check("C3 tidak ada laporan ganda di seluruh card", len(allkeys) == len(set(allkeys)), allkeys)
    check("C4 belum tersedia ditandai fase (bukan laporan aktif); halaman modul terkait ber-path",
          all(p.get("phase") for x in cat["groups"] for p in x["planned"]) and g["warehouse"]["reports"] == []
          and any(lk["to"] == "/traceability" for lk in g["procurement"]["links"]), g["warehouse"])
    sc, cn = NP("GET", "report-center/catalog")
    hl = {lk["to"] for x in cn.get("groups", []) for lk in x["links"]}
    check("C5 link modul mengikuti permission (staff gudang tanpa invoice.view: tanpa link Invoice)",
          not any(t.startswith("/invoice") for t in hl), sorted(hl))

    # ------------------------------------------------------------------ R. register
    sc, rm = rc(call, "register-mro")
    by = {r["no"]: r for r in rm.get("rows", [])}
    check("R1 register MRO: MRO1 Partial (MI net 3 dari 10; MI2 batal), MRO2 Open", by.get(mro1_no, {}).get("status") == "Partial"
          and by.get(mro2["no"], {}).get("status") == "Open", {k: v.get("status") for k, v in by.items()})
    sc, rr = rc(call, "register-ro")
    byr = {r["no"]: r for r in rr.get("rows", [])}
    check("R2 register RO: RO2a Fully Ordered, RO2b Open, RO1 Fully Ordered; Ref. MRO terisi",
          byr.get(ros[0]["no"], {}).get("status") == "Fully Ordered" and byr.get(ros[1]["no"], {}).get("status") == "Open"
          and byr.get(ro1_no, {}).get("status") == "Fully Ordered" and mro2["no"] in byr.get(ros[1]["no"], {}).get("mro_refs", ""),
          {k: (v.get("status"), v.get("mro_refs")) for k, v in byr.items()})
    sc, rp = rc(call, "register-po")
    byp = {r["no"]: r for r in rp.get("rows", [])}
    check("R3 register PO: 1 baris per PO (2 DO tidak menggandakan), Total = detail PO, Closed / Diterima Penuh (label existing)",
          len([r for r in rp.get("rows", []) if r["no"] == po1d["no"]]) == 1 and near(byp[po1d["no"]]["grand_total"], po1d["grand_total"])
          and near(byp[po1d["no"]]["gross"], 10000) and byp[po1d["no"]]["receipt_status"] == "Diterima Penuh" and byp[po1d["no"]]["status"] == "Closed",
          byp.get(po1d["no"]))
    check("R4 total register PO = Σ grand_total detail PO", near(rp["totals"]["grand_total"], float(po1d["grand_total"]) + float(po2d["grand_total"])),
          (rp.get("totals"), po1d.get("grand_total"), po2d.get("grand_total")))
    sc, rd = rc(call, "register-do")
    byd = {r["no"]: r for r in rd.get("rows", [])}
    sc, tr = rc(call, "mro-traceability", f"q={mro1_no}")
    check("R5 register DO: DO1 4.000 + DO2 6.000 = Nilai DO MRO Traceability (DPP × qty diterima); Ref. PO",
          near(byd[do1["no"]]["do_value"], 4000) and near(byd[do2["no"]]["do_value"], 6000)
          and near(rd["totals"]["do_value"], tr["totals"]["do_total"]) and po1d["no"] in byd[do1["no"]]["po_refs"],
          (byd.get(do1["no"]), rd.get("totals"), tr.get("totals")))
    sc, rmi = rc(call, "register-mi")
    sc2, hpp = rc(call, "hpp-mi")
    check("R6 register MI: MI1 tampil (HPP 3 × 1.000), MI2 batal tidak tampil; Σ HPP = Rekap HPP MI (P1)",
          mi1["no"] in {r["no"] for r in rmi["rows"]} and mi2["no"] not in {r["no"] for r in rmi["rows"]}
          and near(rmi["totals"]["hpp"], 3000) and near(rmi["totals"]["hpp"], hpp["totals"]["hpp"]), (rmi.get("totals"), hpp.get("totals")))
    check("R7 Nilai PO, Nilai DO, HPP MI terpisah (kolom berbeda, tidak dijumlahkan)", "do_value" not in rp["totals"] and "hpp" not in rd["totals"]
          and "grand_total" not in rmi["totals"], "")
    sc, f1 = rc(call, "register-po", f"supplier_id={M['supY']['id']}")
    sc2, f2 = rc(call, "register-po", f"q={po1d['no']}")
    sc3, f3 = rc(call, "register-ro", "status=Open")
    sc4, f4 = rc(call, "register-mro", f"project_id={M['pb']['id']}")
    check("R8 filter supplier / pencarian / status / proyek", [r["no"] for r in f1["rows"]] == [po2d["no"]] and [r["no"] for r in f2["rows"]] == [po1d["no"]]
          and ros[1]["no"] in {r["no"] for r in f3["rows"]} and ro1_no not in {r["no"] for r in f3["rows"]}
          and [r["no"] for r in f4["rows"]] == [mro2["no"]], ([r["no"] for r in f1.get("rows", [])], [r["no"] for r in f4.get("rows", [])]))
    tmr = (TODAY + timedelta(days=1)).isoformat()
    sc, fs = rc(call, "register-po", "status=Closed")
    sc2, fr = rc(call, "register-po", "receipt_status=Belum Diterima")
    check("R8b filter status PO (kelompok dasar) & status penerimaan (label existing)", [r["no"] for r in fs["rows"]] == [po1d["no"]]
          and po2d["no"] in {r["no"] for r in fr["rows"]} and po1d["no"] not in {r["no"] for r in fr["rows"]},
          ([r["no"] for r in fs.get("rows", [])], [r["no"] for r in fr.get("rows", [])]))
    sc, hu = rc(call, "hpp-mi", "group_by=unit")
    check("R8c Rekap HPP MI per Unit (P1): reversal MI2 mengurangi unit asal (tanpa baris '(Tanpa Unit)' negatif)",
          [(r["group_name"], r["hpp"]) for r in hu["rows"]] == [("Dump Truck P2", 3000.0)], hu.get("rows"))
    sc, fp = rc(call, "register-po", f"date_from={tmr}")
    sc2, fq = rc(call, "register-po", f"date_to={TODAY.isoformat()}")
    check("R9 periode tanggal dokumen WIB: mulai besok -> 0; s/d hari ini -> semua PO", sc == 200 and fp["total_rows"] == 0 and fq["total_rows"] == rp["total_rows"], (fp.get("total_rows"), fq.get("total_rows")))
    check("R10 drill-down ke dokumen sumber", (rp["rows"][0].get("_drill") or {}).get("to", "").startswith("/po/")
          and (rmi["rows"][0].get("_drill") or {}).get("to", "").startswith("/mi/"), rp["rows"][0].get("_drill"))

    # ------------------------------------------------------------------ L. lead time & P. pemakaian
    sc, lt = rc(call, "lead-time")
    l1 = next((r for r in lt.get("rows", []) if r["mro_no"] == mro1_no), {})
    l2 = next((r for r in lt.get("rows", []) if r["mro_no"] == mro2["no"]), {})
    check("L1 lead time MRO1: tanggal RO/PO/DO/MI pertama terisi, hari ≥ 0, supplier X, proyek A",
          all(l1.get(k) for k in ("ro_date", "po_date", "do_date", "mi_date")) and all((l1.get(k) or 0) >= 0 for k in ("mro_to_ro", "ro_to_po", "po_to_do", "do_to_mi", "mro_to_mi"))
          and l1.get("supplier") == "Supplier X" and l1.get("project") == "Project A" and l1.get("status") == "Partial", l1)
    check("L2 lead time MRO2: RO & PO ada, DO/MI belum (hari kosong), satu baris per MRO × barang walau 2 RO",
          l2.get("ro_date") and l2.get("po_date") and l2.get("do_date") is None and l2.get("po_to_do") is None
          and len([r for r in lt["rows"] if r["mro_no"] == mro2["no"]]) == 1, l2)
    sc, ls = rc(call, "lead-time", f"supplier_id={M['supY']['id']}")
    sc2, lp = rc(call, "lead-time", f"project_id={M['pa']['id']}&status=Partial")
    check("L3 lead time filter supplier & proyek+status", [r["mro_no"] for r in ls["rows"]] == [mro2["no"]] and [r["mro_no"] for r in lp["rows"]] == [mro1_no],
          ([r["mro_no"] for r in ls.get("rows", [])], [r["mro_no"] for r in lp.get("rows", [])]))
    sc, pu = rc(call, "pemakaian-barang")
    sc2, pp = rc(call, "pemakaian-barang", "group_by=project")
    sc3, mu = rc(call, "mutasi-persediaan")
    check("P1 pemakaian per unit: Dump Truck 3 (MI2 batal net 0), HPP 3.000 = Rekap HPP = Mutasi Pemakaian MI",
          [(r["group_name"], r["qty"]) for r in pu["rows"]] == [("Dump Truck P2", 3.0)] and near(pu["totals"]["hpp"], hpp["totals"]["hpp"])
          and near(pu["totals"]["qty"], mu["totals"]["q_mi"]) and pu["rows"][0]["base_unit"], pu.get("rows"))
    check("P2 pemakaian per proyek: Project A 3", [(r["group_name"], r["qty"]) for r in pp["rows"]] == [("Project A", 3.0)], pp.get("rows"))
    sc, pf = rc(call, "pemakaian-barang", f"unit_id={unit['id']}&date_from={tmr}")
    check("P3 filter periode mulai besok -> 0", sc == 200 and pf["total_rows"] == 0, pf.get("total_rows"))
    sc, ti = rc(call, "mro-traceability", f"item_id={M['item']['id']}")
    sc2, tc = rc(call, "mro-traceability", f"category_id={M['cat']['id']}")
    check("T1 MRO Traceability: filter barang & kategori baru (baris = barang tsb)", sc == 200 and ti["total_rows"] >= 2
          and {r["item_code"] for r in ti["rows"]} == {M["item"]["code"]} and tc["total_rows"] >= ti["total_rows"], (ti.get("total_rows"), tc.get("total_rows")))

    # ------------------------------------------------------------------ S. scope & redaksi
    sc, nb = rc(UB, "register-mro")
    sc2, nb2 = rc(UB, "lead-time")
    check("S1 user Divisi B: MRO/lead time Divisi A tidak tampil", sc == 200 and mro1_no not in {r["no"] for r in nb["rows"]}
          and mro1_no not in {r["mro_no"] for r in nb2["rows"]}, (nb.get("total_rows"), nb2.get("total_rows")))
    sc, _ = rc(UB, "register-po", f"division_id={M['div']['id']}")
    check("S2 user Divisi B filter Divisi A -> 403", sc == 403, sc)
    sc, np_ = rc(NP, "register-po")
    ck = {c["key"] for c in np_.get("columns", [])}
    check("S3 tanpa izin harga: register PO tanpa Bruto/Diskon/DPP/Pajak/Total (kolom, baris, total)",
          sc == 200 and not ({"gross", "discount", "dpp", "tax", "grand_total"} & ck) and all(not ({"gross", "grand_total"} & set(r)) for r in np_["rows"])
          and not ({"gross", "grand_total"} & set(np_["totals"])) and np_["total_rows"] == rp["total_rows"], ck)
    sc, nd = rc(NP, "register-do")
    sc2, nm = rc(NP, "register-mi")
    sc3, npu = rc(NP, "pemakaian-barang")
    check("S4 tanpa izin harga: Nilai DO, HPP MI, HPP pemakaian disembunyikan; qty tetap", "do_value" not in {c["key"] for c in nd["columns"]}
          and "hpp" not in {c["key"] for c in nm["columns"]} and "hpp" not in npu["totals"] and near(npu["totals"]["qty"], 3), npu.get("totals"))
    sc, content, _ = get_bin(NP, "report-center/register-po/export.xlsx")
    hdr, body = xlsx_rows(content, "No. Dokumen") if sc == 200 else ([], [])
    check("S5 Excel register PO staff: tanpa kolom nilai", sc == 200 and "Total PO" not in hdr and "DPP" not in hdr and len(body) == np_["total_rows"], hdr)
    sc, pdf, _ = get_bin(NP, "report-center/register-do/export.pdf")
    txt = "".join(" ".join(pg.extract_text() or "" for pg in PdfReader(io.BytesIO(pdf)).pages).split()) if sc == 200 else ""
    check("S6 PDF register DO staff: tanpa 'Nilai DO', dengan catatan redaksi", sc == 200 and "NilaiDO" not in txt and do1["no"].replace(" ", "") in txt, sc)

    # ------------------------------------------------------------------ E. export & data besar
    sc, content, hdrs = get_bin(None, "report-center/register-do/export.xlsx")
    hdr, body = xlsx_rows(content, "No. Dokumen") if sc == 200 else ([], [])
    check("E1 Excel register DO = JSON (baris & X-Report-Rows), kolom Nilai DO ada", sc == 200 and len(body) == rd["total_rows"]
          and hdrs.get("X-Report-Rows") == str(rd["total_rows"]) and "Nilai DO (DPP × Qty Diterima)" in hdr, (len(body), rd.get("total_rows")))
    for i in range(120):
        call("POST", "mro", {"no": f"MRO-BULK{u}-{i:03d}", "division_id": M["div"]["id"], "requester": "Bulk", "submitted": True,
                             "lines": [{"item_id": M["item"]["id"], "qty": 1, "warehouse_id": M["wh"]["id"]}]})
    sc, big = call("GET", "report-center/register-mro?page_size=50&page=3&q=BULK")
    check("E2 data besar: 120 MRO -> total 120, halaman 3 = 20 baris, TOTAL atas seluruh baris", sc == 200 and big["total_rows"] == 120
          and len(big["rows"]) == 20 and big["pages"] == 3 and near(big["totals"]["line_count"], 120), (big.get("total_rows"), len(big.get("rows", []))))
    sc, content, hdrs = get_bin(None, "report-center/register-mro/export.xlsx?q=BULK")
    hdr, body = xlsx_rows(content, "No. Dokumen") if sc == 200 else ([], [])
    sc2, pdf, _ = get_bin(None, "report-center/register-mro/export.pdf?q=BULK")
    rd_ = PdfReader(io.BytesIO(pdf)) if sc2 == 200 else None
    t2 = "".join(" ".join(pg.extract_text() or "" for pg in rd_.pages).split()) if rd_ else ""
    check("E3 export data besar: Excel 120 baris (bukan halaman aktif), PDF multi-halaman memuat baris pertama & terakhir",
          sc == 200 and len(body) == 120 and rd_ and len(rd_.pages) >= 2 and f"MRO-BULK{u}-000" in t2 and f"MRO-BULK{u}-119" in t2,
          (len(body), len(rd_.pages) if rd_ else None))

    # ------------------------------------------------------------------ K. kompatibilitas & read-only
    codes = [call("GET", p)[0] for p in ("reports/lead-time", "reports/unit-usage", "reports/mro-traceability", "mro", "po", "do", "mi")]
    check("K1 endpoint lama tetap 200 (lead-time, unit-usage, mro-traceability, daftar dokumen)", all(c == 200 for c in codes), codes)
    sc, iw1 = call("GET", "item-warehouse")
    snap = lambda iw: sorted((x.get("item_id"), x.get("warehouse_id"), x.get("current_stock"), x.get("total_value")) for x in iw or [])  # noqa: E731
    check("K2 read-only: stok & nilai pool tidak berubah oleh laporan/export", snap(iw0) == snap(iw1), "")

    failed = [nm_ for nm_, ok in T.RESULTS if not ok]
    print(f"\n{len(T.RESULTS) - len(failed)}/{len(T.RESULTS)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
