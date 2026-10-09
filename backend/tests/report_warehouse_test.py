"""P3 Reporting Warehouse (Transfer, Pinjam & Pengembalian, Penyesuaian, Stock Opname) pada Pusat Laporan.

Hanya pada backend + DB test terisolasi (DB *itest*; tenant QA via /saas/register) — BUKAN preview/sandbox/production.
Membuktikan: register = agregasi detail (tanpa hitung ganda); gudang/proyek/unit per baris; transfer/pengembalian yang dihapus
(reversal) tidak dihitung; outstanding pinjaman per BARIS dengan partial return & cut-off; umur & keterlambatan; nilai
dari valuation ledger (= detail dokumen existing); opname Posted vs Counting vs Cancelled (nilai hanya Posted); scope divisi;
redaksi harga server-side (JSON/Excel/PDF); export = seluruh hasil filter & paritas JSON = Excel = PDF; read-only.
"""
import io
import os
import sys
import uuid
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

import receipt_control_test as T
import transfer_multi_warehouse_test as TM
from master_item_stock_test import mk_user

call, check, API = T.call, T.check, T.API
TODAY = datetime.now(ZoneInfo("Asia/Jakarta")).date()
D0 = TODAY - timedelta(days=30)
d = lambda n: (D0 + timedelta(days=n)).isoformat()  # noqa: E731
u = uuid.uuid4().hex[:5]


def near(a, b, eps=0.01):
    return abs(float(a or 0) - float(b or 0)) <= eps


def rc(caller, key, qs=""):
    return caller("GET", f"report-center/{key}?page_size=500&{qs}")


def rows_of(res):
    return [r for r in (res or {}).get("rows", []) if not r.get("_kind")]


def xl(sess, key, qs=""):
    r = sess.get(f"{API}/report-center/{key}/export.xlsx?{qs}")
    if not r.ok:
        return r.status_code, [], [], []
    vals = [list(v) for v in load_workbook(io.BytesIO(r.content), read_only=True).worksheets[0].iter_rows(values_only=True)]
    hi = next((i for i, v in enumerate(vals) if v and v[0] in ("No. Dokumen", "No. Pengembalian")), None)
    if hi is None:
        return r.status_code, [], [], []
    body = [v for v in vals[hi + 1:] if v and v[0] not in (None, "TOTAL")]
    tot = next((v for v in vals[hi + 1:] if v and v[0] == "TOTAL"), [])
    return r.status_code, [str(x) for x in vals[hi]], body, tot


def pdf_text(sess, key, qs=""):
    r = sess.get(f"{API}/report-center/{key}/export.pdf?{qs}")
    return r.status_code, (" ".join(" ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(r.content)).pages).split()) if r.ok else "")


def main():
    dbn = (os.environ.get("TEST_DATABASE_URL") or "").rsplit("/", 1)[-1].split("?")[0]
    if not ("itest" in dbn or dbn.endswith("_test")):
        print("Menolak berjalan: hanya untuk backend + DB test terisolasi (scripts/run_regression_itest.sh, DB *itest*).")
        sys.exit(2)
    TM.register()
    m = TM.m
    dA, dB = m("divisions", {"code": f"DA{u}", "name": f"Asset {u}"}), m("divisions", {"code": f"DB{u}", "name": f"Ops {u}"})
    cat, uom = m("item_categories", {"code": f"KC{u}", "name": "Umum"}), m("uoms", {"code": f"PC{u}", "name": "Pcs", "symbol": "pcs"})
    W = {k: m("warehouses", {"code": f"W{k}{u}", "name": f"Gudang {k} {u}", "is_active": True, "division_id": (dB if k in "CD" else dA)["id"]})
         for k in "ABCD"}
    pa, pb = m("projects", {"code": f"PA{u}", "name": f"Proyek A {u}"}), m("projects", {"code": f"PB{u}", "name": f"Proyek B {u}"})
    unit = m("units", {"code": f"UN{u}", "name": f"Excavator {u}", "is_active": True})
    ref = {"category_id": cat["id"], "division_id": dA["id"], "base_uom_id": uom["id"], "unit": "PCS", "is_active": True}
    X, Y = m("items", {"code": f"X{u}", "name": f"Bearing {u}", **ref}), m("items", {"code": f"Y{u}", "name": f"Oli {u}", **ref})

    def opening(wh, div, item, qty, cost, day):
        return call("POST", "adjustments", {"date": day, "warehouse_id": W[wh]["id"], "division_id": div["id"], "adj_type": "Opening",
                                            "reason": "Saldo awal QA", "lines": [{"item_id": item["id"], "adjustment": qty, "reason": "Saldo awal",
                                                                                  "approved_unit_cost": cost}]}, 200)[1]
    opening("A", dA, X, 50, 1000, d(0)); opening("A", dA, Y, 20, 500, d(0)); opening("C", dB, X, 10, 2000, d(0))

    def tl(item, qty, frm, to, prj=None, un=None):
        return {"item_id": item["id"], "qty": qty, "uom_id": uom["id"], "from_warehouse_id": W[frm]["id"], "to_warehouse_id": W[to]["id"],
                "project_id": prj["id"] if prj else None, "unit_id": un["id"] if un else None}
    sc1, tr1 = call("POST", "transfers", {"date": d(1), "division_id": dA["id"], "from_warehouse_id": W["A"]["id"], "to_warehouse_id": W["B"]["id"],
                                          "lines": [tl(X, 5, "A", "B", pa, unit), tl(Y, 3, "A", "B", pb)]})
    sc2, tr2 = call("POST", "transfers", {"date": d(2), "division_id": dA["id"], "lines": [tl(X, 2, "A", "B")]})
    sc3, trc = call("POST", "transfers", {"date": d(2), "division_id": dB["id"], "lines": [tl(X, 1, "C", "D")]})
    sc4, tr3 = call("POST", "transfers", {"date": d(3), "division_id": dA["id"], "lines": [tl(X, 1, "A", "B")]})
    check("setup: opening + TR1 (2 baris) + TR2 + TRC (divisi B) + TR3", all(s == 200 for s in (sc1, sc2, sc3, sc4)), (sc1, sc2, sc3, sc4, tr1))

    NP = mk_user("warehouse", divs=[dA["id"]], overrides={"export": "allow"})
    UB = mk_user("manager", divs=[dB["id"]], overrides={"export": "allow"})
    by_no = lambda res, no: [r for r in rows_of(res) if r.get("no") == no]  # noqa: E731

    # ------------------------------------------------------------------ L. Pinjam & pengembalian
    lh = {"division_id": dA["id"], "from_warehouse_id": W["A"]["id"], "to_warehouse_id": W["B"]["id"], "requester": "Budi QA"}
    sc1, L1 = call("POST", "loans", {**lh, "date": d(4), "due_date": d(10), "lines": [tl(X, 10, "A", "B", pa), tl(Y, 4, "A", "B", pa)]})
    sc2, L2 = call("POST", "loans", {**lh, "date": d(5), "due_date": (TODAY + timedelta(days=10)).isoformat(), "lines": [tl(X, 3, "A", "B")]})
    sc3, L3 = call("POST", "loans", {**lh, "date": d(5), "lines": [tl(Y, 2, "A", "B")]})
    # L4 multi-gudang per baris: X A->B, Y B->A (header A->B); pengembalian kemudian DIEDIT (reversal + posting ulang)
    sc4l, L4 = call("POST", "loans", {**lh, "date": d(5), "lines": [tl(X, 2, "A", "B"), tl(Y, 1, "B", "A")]})
    check("setup: 4 pinjaman dibuat (L4 multi-gudang per baris)", sc1 == 200 and sc2 == 200 and sc3 == 200 and sc4l == 200,
          (sc1, L1, sc2, L2, sc3, L3, sc4l, L4))
    ll = {x["item_id"]: x["id"] for x in call("GET", f"loans/{L1['id']}")[1]["lines"]}
    l3 = call("GET", f"loans/{L3['id']}")[1]["lines"][0]["id"]
    l4 = {x["item_id"]: x["id"] for x in call("GET", f"loans/{L4['id']}")[1]["lines"]}
    s1, _ = call("POST", f"loans/{L1['id']}/return", {"date": d(6), "lines": [{"loan_line_id": ll[X["id"]], "qty": 4}]})
    s2, _ = call("POST", f"loans/{L3['id']}/return", {"date": d(7), "lines": [{"loan_line_id": l3, "qty": 2}]})
    s3, _ = call("POST", f"loans/{L4['id']}/return", {"date": d(7), "lines": [{"loan_line_id": l4[X["id"]], "qty": 2},
                                                                             {"loan_line_id": l4[Y["id"]], "qty": 1}]})
    s4, _ = call("POST", f"loans/{L1['id']}/return", {"date": d(12), "lines": [{"loan_line_id": ll[X["id"]], "qty": 6}, {"loan_line_id": ll[Y["id"]], "qty": 4}]})
    rets = {r["date"][:10]: r for r in call("GET", f"loans/{L1['id']}/returns")[1]}
    check("setup: L1 (X10,Y4 jatuh tempo D+10) R1 X4 (D+6), R2 X6+Y4 (D+12); L2 X3 jatuh tempo +10 hari; L3 Y2 tanpa jatuh tempo, return dihapus",
          all(s == 200 for s in (sc1, sc2, sc3, s1, s2, s3, s4)) and len(rets) == 2, (sc1, sc2, sc3, s1, s2, s3, s4, list(rets)))
    R1, R2 = rets.get(d(6), {}).get("no"), rets.get(d(12), {}).get("no")
    # ------------------------------------------------------------------ A. Penyesuaian
    sx, sy = TM.stock(X["id"], W["A"]["id"]), TM.stock(Y["id"], W["A"]["id"])
    sa1, A1 = call("POST", "adjustments", {"date": d(13), "division_id": dA["id"], "warehouse_id": W["A"]["id"], "project_id": pa["id"],
                                           "adj_type": "Koreksi", "reason": "Koreksi header", "lines": [
                                               {"item_id": X["id"], "adjustment": 5, "reason": "Temuan", "unit_id": unit["id"]},
                                               {"item_id": Y["id"], "adjustment": -1, "reason": "Rusak"}]})
    sa2, A2 = call("POST", "adjustments", {"date": d(13), "division_id": dA["id"], "warehouse_id": W["B"]["id"], "adj_type": "Koreksi",
                                           "reason": "Hilang", "lines": [{"item_id": X["id"], "adjustment": -2}]})
    check("setup: ADJ1 (A: X +5, Y −1) + ADJ2 (B: X −2)", sa1 == 200 and sa2 == 200, (A1, A2))
    sc, ad = rc(call, "adjustment-detail")
    a1 = {r["item_code"]: r for r in by_no(ad, A1["no"])}
    check("A1 detail: stok sebelum/sesudah = snapshot baris, +/−, alasan baris, proyek & unit, nilai = MWA (5×1.000, −1×500)",
          sc == 200 and near(a1[X["code"]]["before"], sx) and near(a1[X["code"]]["after"], sx + 5) and near(a1[X["code"]]["plus"], 5)
          and near(a1[Y["code"]]["minus"], 1) and near(a1[Y["code"]]["after"], sy - 1) and a1[X["code"]]["reason"] == "Temuan"
          and a1[Y["code"]]["reason"] == "Rusak" and a1[X["code"]]["project"] == pa["name"] and a1[X["code"]]["unit"] == unit["name"]
          and near(a1[X["code"]]["value"], 5000) and near(a1[Y["code"]]["value"], -500) and a1[X["code"]]["status"] == "Posted", a1)
    sc, ag = rc(call, "adjustment-register", f"date_from={d(13)}")
    g = (by_no(ag, A1["no"]) or [{}])[0]
    ad13 = rc(call, "adjustment-detail", f"date_from={d(13)}")[1]
    check("A2 register ADJ1: +5, −1, net 4, nilai 4.500; total register = total detail",
          near(g.get("plus"), 5) and near(g.get("minus"), 1) and near(g.get("adjustment"), 4) and near(g.get("value"), 4500)
          and near(ag["totals"]["value"], ad13["totals"]["value"]) and near(ag["totals"]["plus"], ad13["totals"]["plus"]), (g, ag.get("totals")))
    fm = sorted((r["no"], r["item_code"]) for r in rows_of(rc(call, "adjustment-detail", f"date_from={d(13)}&direction=minus")[1]))
    fb = [r["no"] for r in rows_of(rc(call, "adjustment-register", f"warehouse_id={W['B']['id']}")[1])]
    check("A3 filter arah (−) & gudang", fm == sorted([(A1["no"], Y["code"]), (A2["no"], X["code"])]) and fb == [A2["no"]], (fm, fb))

    # ------------------------------------------------------------------ O. Stock opname
    sc, o1 = call("POST", "opname", {"date": d(14), "warehouse_id": W["A"]["id"], "division_id": dA["id"], "mode": "live", "scope": "all"})
    ol = {r["item_code"]: r for r in call("GET", f"opname/{o1['id']}")[1]["lines"]}
    sysx, sysy = TM.stock(X["id"], W["A"]["id"]), TM.stock(Y["id"], W["A"]["id"])
    call("PUT", f"opname/{o1['id']}/count", {"lines": [{"line_id": ol[X["code"]]["id"], "qty": sysx + 2, "reason": "Salah catat"},
                                                       {"line_id": ol[Y["code"]]["id"], "qty": sysy - 1, "reason": "Rusak"}]})
    call("POST", f"opname/{o1['id']}/workflow", {"action": "review"}); call("POST", f"opname/{o1['id']}/submit")
    sp, po = call("POST", f"opname/{o1['id']}/post")
    sc2, o2 = call("POST", "opname", {"date": d(14), "warehouse_id": W["B"]["id"], "division_id": dA["id"], "mode": "live", "scope": "all"})
    sc3, o3 = call("POST", "opname", {"date": d(14), "warehouse_id": W["C"]["id"], "division_id": dB["id"], "mode": "live", "scope": "all"})
    sc4, _ = call("POST", f"opname/{o3['id']}/workflow", {"action": "cancel", "reason": "Salah gudang"})
    check("setup: OP1 gudang A Posted (X +2, Y −1); OP2 gudang B Counting; OP3 gudang C Cancelled",
          sc == 200 and sp == 200 and (po or {}).get("status") == "Posted" and sc2 == 200 and sc3 == 200 and sc4 == 200, (sp, po, sc2, sc3, sc4))
    # pembatalan (reversal, tanggal hari ini) setelah seluruh posting bertanggal mundur
    sc5, _ = call("DELETE", f"transactions/transfer/{tr3.get('id')}")
    r3 = call("GET", f"loans/{L3['id']}/returns")[1][0]
    sc6, _ = call("DELETE", f"loan-returns/{r3['id']}")
    r4 = call("GET", f"loans/{L4['id']}/returns")[1][0]
    sc7, _ = call("PUT", f"loan-returns/{r4['id']}", {"lines": [{"loan_line_id": l4[X["id"]], "qty": 1}]})
    check("setup: TR3 dihapus, pengembalian L3 dihapus, pengembalian L4 diedit X2+Y1 -> X1 (reversal engine existing)",
          sc5 == 200 and sc6 == 200 and sc7 == 200, (sc5, sc6, sc7))

    # ------------------------------------------------------------------ T. Transfer
    sc, td = rc(call, "transfer-detail")
    t1 = {r["item_code"]: r for r in by_no(td, tr1["no"])}
    check("T1 detail TR1: gudang asal/tujuan, proyek & unit per baris, qty, satuan dasar, status, nilai = MWA (5×1.000, 3×500)",
          sc == 200 and len(t1) == 2 and t1[X["code"]]["from_wh"] == W["A"]["name"] and t1[X["code"]]["to_wh"] == W["B"]["name"]
          and t1[X["code"]]["project"] == pa["name"] and t1[X["code"]]["unit"] == unit["name"] and t1[Y["code"]]["project"] == pb["name"]
          and near(t1[X["code"]]["qty"], 5) and t1[X["code"]]["base_unit"] == "pcs" and t1[X["code"]]["status"] == "Posted"
          and near(t1[X["code"]]["value"], 5000) and near(t1[Y["code"]]["value"], 1500), t1)
    check("T2 transfer yang dihapus (reversal) tidak muncul", not by_no(td, tr3.get("no")), [r["no"] for r in rows_of(td)])
    sc, tg = rc(call, "transfer-register")
    g1 = (by_no(tg, tr1["no"]) or [{}])[0]
    check("T3 register = agregasi detail: TR1 2 baris, qty 8, nilai 6.500; total register = total detail (tanpa hitung ganda out/in)",
          sc == 200 and g1.get("line_count") == 2 and near(g1.get("qty"), 8) and near(g1.get("value"), 6500)
          and near(tg["totals"]["qty"], td["totals"]["qty"]) and near(tg["totals"]["qty"], 5 + 3 + 2 + 1)
          and near(tg["totals"]["value"], td["totals"]["value"]), (g1, tg.get("totals"), td.get("totals")))
    fw = rows_of(rc(call, "transfer-detail", f"warehouse_id={W['D']['id']}")[1])
    fi = rows_of(rc(call, "transfer-detail", f"item_id={Y['id']}")[1])
    fp = rows_of(rc(call, "transfer-detail", f"project_id={pb['id']}")[1])
    fd = rows_of(rc(call, "transfer-register", f"division_id={dB['id']}")[1])
    fpd = rows_of(rc(call, "transfer-register", f"date_from={d(2)}&date_to={d(2)}")[1])
    check("T4 filter gudang (tujuan D) / barang / proyek / divisi / periode",
          [r["no"] for r in fw] == [trc["no"]] and [r["item_code"] for r in fi] == [Y["code"]] and [r["item_code"] for r in fp] == [Y["code"]]
          and [r["no"] for r in fd] == [trc["no"]] and sorted(r["no"] for r in fpd) == sorted([tr2["no"], trc["no"]]), (fw, fi, fp, fd, fpd))
    ub = rows_of(UB("GET", "report-center/transfer-register?page_size=500")[1])
    check("T5 scope divisi: user Divisi B hanya melihat transfer divisi B", [r["no"] for r in ub] == [trc["no"]], [r["no"] for r in ub])
    check("T6 drill-down ke halaman Transfer (?open=id)", (g1.get("_drill") or {}).get("to") == f"/transfer?open={tr1['id']}", g1.get("_drill"))


    def ld(qs=""):
        res = rc(call, "loan-detail", qs)[1]
        return {(r["no"], r["item_code"]): r for r in rows_of(res)}, res
    a, res_now = ld()
    x1, y1, x2, y3 = a.get((L1["no"], X["code"]), {}), a.get((L1["no"], Y["code"]), {}), a.get((L2["no"], X["code"]), {}), a.get((L3["no"], Y["code"]), {})
    check("L1 baris X L1 (hari ini): 10 dipinjam, kembali 10 (4 + 6, parsial 2 dokumen), outstanding 0, Completed, kembali penuh D+12, "
          "umur 8 hari, Kembali Terlambat 2 hari, ref R1 & R2",
          near(x1.get("qty"), 10) and near(x1.get("returned"), 10) and near(x1.get("outstanding"), 0) and x1.get("status") == "Completed"
          and x1.get("full_return_date") == d(12) and x1.get("age") == 8 and x1.get("late") == "Kembali Terlambat" and x1.get("late_days") == 2
          and R1 in (x1.get("return_refs") or "") and R2 in (x1.get("return_refs") or ""), x1)
    check("L2 baris Y L1 Completed (R2), L2 Open Belum Jatuh Tempo umur 25 hari, L3 pengembalian dihapus -> Open outstanding 2 Tanpa Jatuh Tempo",
          y1.get("status") == "Completed" and near(y1.get("returned"), 4) and x2.get("status") == "Open" and x2.get("late") == "Belum Jatuh Tempo"
          and x2.get("age") == 25 and y3.get("status") == "Open" and near(y3.get("outstanding"), 2) and y3.get("late") == "Tanpa Jatuh Tempo"
          and y3.get("return_refs") == "-", (y1, x2, y3))
    x4, y4 = a.get((L4["no"], X["code"]), {}), a.get((L4["no"], Y["code"]), {})
    check("L2b L4 multi-gudang per baris & pengembalian diedit: X (A->B) kembali 1 sisa 1 Partial Returned umur 25; "
          "Y (B->A) kembali 0 (net reversal) sisa 1 Open; nilai sisa = harga pokok saat pinjam",
          x4.get("from_wh") == W["A"]["name"] and x4.get("to_wh") == W["B"]["name"] and y4.get("from_wh") == W["B"]["name"]
          and y4.get("to_wh") == W["A"]["name"] and near(x4.get("returned"), 1) and near(x4.get("outstanding"), 1)
          and x4.get("status") == "Partial Returned" and x4.get("age") == 25 and near(y4.get("returned"), 0) and near(y4.get("outstanding"), 1)
          and y4.get("status") == "Open" and y4.get("return_refs") == "-" and near(x4.get("outstanding_value"), 1000)
          and near(y4.get("outstanding_value"), 500), (x4, y4))
    b6, _ = ld(f"date_to={d(6)}")
    check("L2c cut-off D+6 (sebelum pengembalian L4): X L4 kembali 0 Open, umur 1",
          near(b6.get((L4["no"], X["code"]), {}).get("returned"), 0) and b6.get((L4["no"], X["code"]), {}).get("status") == "Open"
          and b6.get((L4["no"], X["code"]), {}).get("age") == 1, b6.get((L4["no"], X["code"])))
    b, _ = ld(f"date_to={d(8)}")
    bx = b.get((L1["no"], X["code"]), {})
    check("L3 cut-off D+8: X L1 kembali 4 (R2 belum), outstanding 6, Partial Returned 40%, Belum Jatuh Tempo, umur 4",
          near(bx.get("returned"), 4) and near(bx.get("outstanding"), 6) and bx.get("status") == "Partial Returned" and near(bx.get("pct"), 40)
          and bx.get("late") == "Belum Jatuh Tempo" and bx.get("age") == 4 and bx.get("return_refs") == R1, bx)
    c, cres = ld(f"date_to={d(11)}&late=Terlambat")
    check("L4 cut-off D+11 filter Terlambat: hanya 2 baris L1 (X sisa 6, Y sisa 4), 1 hari; nilai sisa = sisa × harga pokok saat pinjam",
          sorted(c) == sorted([(L1["no"], X["code"]), (L1["no"], Y["code"])]) and all(r["late_days"] == 1 for r in c.values())
          and near(c[(L1["no"], X["code"])]["outstanding_value"], 6000) and near(c[(L1["no"], Y["code"])]["outstanding_value"], 2000)
          and near(cres["totals"]["outstanding"], 10) and near(cres["totals"]["outstanding_value"], 8000), (c, cres.get("totals")))
    sc, lr = rc(call, "loan-register", f"date_to={d(11)}")
    g = (by_no(lr, L1["no"]) or [{}])[0]
    check("L5 register L1 @D+11: qty 14, kembali 4, outstanding 10, Partial Returned, Terlambat 1 hari, nilai pinjam 12.000",
          sc == 200 and near(g.get("qty"), 14) and near(g.get("returned"), 4) and near(g.get("outstanding"), 10) and g.get("status") == "Partial Returned"
          and g.get("late") == "Terlambat" and g.get("late_days") == 1 and near(g.get("loan_value"), 12000), g)
    st = {r["no"]: r["status"] for r in rows_of(rc(call, "loan-register")[1])}
    check("L6 register hari ini: L1 Completed, L2 Open, L3 Open, L4 Partial Returned; filter status Completed -> hanya L1",
          st.get(L1["no"]) == "Completed" and st.get(L2["no"]) == "Open" and st.get(L3["no"]) == "Open" and st.get(L4["no"]) == "Partial Returned"
          and [r["no"] for r in rows_of(rc(call, "loan-register", "status=Completed")[1])] == [L1["no"]], st)
    sc, rr = rc(call, "loan-return-register")
    rk = sorted((r["return_no"], r["item_code"], r["qty"], r["value"]) for r in rows_of(rr))
    check("L7 register pengembalian: R1 X4 (4.000), R2 X6 (6.000) + Y4 (2.000), R4 hasil edit X1 (1.000; Y net 0 tidak muncul); "
          "return dihapus tidak muncul; arah Peminjam -> Pemberi",
          sc == 200 and rk == sorted([(R1, X["code"], 4.0, 4000.0), (R2, X["code"], 6.0, 6000.0), (R2, Y["code"], 4.0, 2000.0),
                                      (r4["no"], X["code"], 1.0, 1000.0)])
          and all(r["from_wh"] == W["B"]["name"] and r["to_wh"] == W["A"]["name"] for r in rows_of(rr)), rk)
    check("L8 periode register pengembalian = tanggal pengembalian (D+12 -> hanya R2)",
          {r["return_no"] for r in rows_of(rc(call, "loan-return-register", f"date_from={d(12)}")[1])} == {R2})

    sc, od = rc(call, "opname-detail")
    o1r = {r["item_code"]: r for r in by_no(od, o1["no"])}
    det = {r["item_code"]: r for r in call("GET", f"opname/{o1['id']}")[1]["lines"]}
    check("O1 detail OP1 Posted: stok sistem/fisik/selisih/hasil/alasan = detail opname existing; nilai selisih = valuation ledger (+2.000, −500)",
          sc == 200 and near(o1r[X["code"]]["variance"], 2) and near(o1r[Y["code"]]["variance"], -1) and o1r[X["code"]]["line_status"] == "Lebih"
          and o1r[Y["code"]]["line_status"] == "Kurang" and o1r[X["code"]]["reason"] == "Salah catat"
          and near(o1r[X["code"]]["system_qty"], det[X["code"]].get("system_qty")) and near(o1r[X["code"]]["counted"], det[X["code"]].get("counted"))
          and near(o1r[X["code"]]["value"], 2000) and near(o1r[Y["code"]]["value"], -500) and near(o1r[X["code"]]["value"], det[X["code"]].get("value"))
          and o1r[X["code"]]["approved_by"] not in (None, "-"), (o1r, det))
    o2r = by_no(od, o2["no"])
    check("O2 OP2 Counting: belum dihitung, tanpa nilai selisih (hanya Posted bernilai)",
          o2r and all(r["status"] == "Counting" and r["value"] is None and r["line_status"] == "Belum Dihitung" for r in o2r), o2r)
    sc, og = rc(call, "opname-register")
    g = (by_no(og, o1["no"]) or [{}])[0]
    check("O3 register OP1: 2 barang, selisih +2/−1, net 1, nilai lebih 2.000, kurang 500, net 1.500, disetujui; total nilai = Posted saja",
          g.get("item_count") == 2 and near(g.get("plus_qty"), 2) and near(g.get("minus_qty"), 1) and near(g.get("variance"), 1)
          and near(g.get("plus_value"), 2000) and near(g.get("minus_value"), 500) and near(g.get("value"), 1500) and g.get("approved_by") not in (None, "-")
          and near(og["totals"]["value"], 1500) and (by_no(og, o2["no"]) or [{}])[0].get("value") is None, (g, og.get("totals")))
    check("O4 filter status: Cancelled -> OP3, Counting -> OP2, Posted -> OP1",
          [r["no"] for r in rows_of(rc(call, "opname-register", "status=Cancelled")[1])] == [o3["no"]]
          and [r["no"] for r in rows_of(rc(call, "opname-register", "status=Counting")[1])] == [o2["no"]]
          and [r["no"] for r in rows_of(rc(call, "opname-register", "status=Posted")[1])] == [o1["no"]])
    check("O5 scope divisi: user Divisi B hanya melihat OP3 (gudang C)",
          [r["no"] for r in rows_of(UB("GET", "report-center/opname-register?page_size=500")[1])] == [o3["no"]])

    # ------------------------------------------------------------------ S. harga, export, paritas
    PRICE = {"transfer-detail": "value", "loan-detail": "outstanding_value", "loan-register": "outstanding_value",
             "adjustment-detail": "value", "opname-register": "value", "loan-return-register": "value"}
    for key, col in PRICE.items():
        sc, j = rc(NP, key)
        hdr = xl(NP.S, key)[1]
        _, pt = pdf_text(NP.S, key)
        _, pa_ = pdf_text(T.S, key)
        lbl = {"value": ("Nilai",), "outstanding_value": ("Nilai Sisa",)}[col]
        check(f"S1 {key}: tanpa izin harga -> kolom nilai dibuang di JSON/Excel/PDF/total (kontrol: admin memuat)",
              sc == 200 and col not in [c["key"] for c in j["columns"]] and col not in (j.get("totals") or {})
              and not any(h.startswith(lbl) for h in hdr) and not any(x in pt for x in ("Nilai Sisa", "Nilai Transfer", "Nilai Penyesuaian",
                                                                                         "Nilai Selisih", "Nilai Pengembalian", "Harga Satuan"))
              and any(x in pa_ for x in ("Nilai Sisa", "Nilai Transfer", "Nilai Penyesuaian", "Nilai Selisih", "Nilai Pengembalian")), (key, hdr))
    for key, qs in (("transfer-detail", ""), ("loan-detail", f"date_to={d(11)}"), ("adjustment-register", ""), ("opname-detail", "")):
        _, j = rc(call, key, qs)
        sc, hdr, body, tot = xl(T.S, key, qs)
        _, pt = pdf_text(T.S, key, qs)
        cols = [c for c in j["columns"]]
        tcol = next(c for c in cols if c["total"] and c["type"] == "qty")
        ix = hdr.index(tcol["label"])
        nos = {r["no"] for r in rows_of(j)}
        check(f"S2 paritas {key}: Excel = seluruh {len(rows_of(j))} baris JSON, TOTAL {tcol['label']} sama, PDF memuat nomor dokumen & TOTAL",
              sc == 200 and len(body) == len(rows_of(j)) == j["total_rows"] and [str(v[0]) for v in body] == [str(r["no"]) for r in rows_of(j)]
              and near(tot[ix], j["totals"][tcol["key"]]) and all(n in pt for n in nos) and "TOTAL" in pt, (key, len(body), j.get("total_rows")))
    sc, p1 = rc(call, "transfer-detail", "page_size=2&page=2")
    check("S3 pagination: halaman 2 ukuran 2, total_rows & totals = seluruh hasil filter", p1.get("page") == 2 and len(p1.get("rows", [])) <= 2
          and p1.get("total_rows") == td["total_rows"] and near(p1["totals"]["qty"], td["totals"]["qty"]), (p1.get("page"), p1.get("total_rows")))
    sc, sq = rc(call, "transfer-detail", f"q={Y['code']}")
    check("S4 pencarian kode barang", sc == 200 and {r["item_code"] for r in rows_of(sq)} == {Y["code"]}, [r["item_code"] for r in rows_of(sq)])
    sc, cat_ = call("GET", "report-center/catalog")
    wg = next(g for g in cat_["groups"] if g["key"] == "warehouse")
    check("S5 card Warehouse: 9 laporan aktif, tanpa placeholder 'Segera tersedia'",
          len(wg["reports"]) == 9 and wg["planned"] == [], [r["key"] for r in wg["reports"]])
    failed = [nm_ for nm_, ok in T.RESULTS if not ok]
    print(f"\n{len(T.RESULTS) - len(failed)}/{len(T.RESULTS)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
