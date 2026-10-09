"""P0 Reporting Foundation — keamanan valuation-summary/ledger + Pusat Laporan (registry, builder bersama, export).

Tenant QA sementara (/saas/register) pada backend test terisolasi — BUKAN production. Membuktikan:
  V. /reports/valuation-summary & /reports/valuation-ledger: cakupan divisi + penugasan gudang server-side, warehouse_id di
     luar cakupan -> 403, tanpa izin harga -> 403, bentuk respon kompatibel, filter tanggal efektif WIB, read-only.
  R. /report-center: katalog 5 kelompok; pagination tidak memengaruhi total; paritas JSON (semua halaman) = Excel = PDF;
     paritas dengan endpoint lama MRO Traceability; scope divisi; redaksi kolom harga di JSON/Excel/PDF; izin export;
     validasi filter; pencarian; batas export tegas (Excel 100.000 / PDF 5.000) tanpa pemotongan; audit export; cut-off WIB.
"""
import io
import os
import sys
import uuid
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

import receipt_control_test as T
from master_item_stock_test import mk_user

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from fastapi import HTTPException  # noqa: E402

import reporting.report_center as RC  # noqa: E402
from reporting.exporters import to_pdf, to_xlsx  # noqa: E402
from report_trace_detail import _join_refs as RTDJ  # noqa: E402
from reporting.scope import local_day  # noqa: E402
from valuation_report_scope_layer import effective_day  # noqa: E402

call, check, API = T.call, T.check, T.API
TODAY = datetime.now(ZoneInfo("Asia/Jakarta")).date()
PRICE_KEYS = {"po_gross", "po_discount", "po_dpp", "po_tax", "po_total", "do_total"}
PRICE_LABELS = {"Nilai PO", "Diskon PO", "DPP PO", "Pajak PO", "Total PO", "Total DO (DPP x Qty Diterima)"}
SUMMARY_KEYS = {"item_id", "item_name", "item_code", "warehouse_id", "warehouse_name", "qty_on_hand", "base_uom", "avg_cost",
                "inventory_value"}


def adj(wh, div, item, qty, cost, date):
    line = {"item_id": item, "adjustment": qty, "reason": "QA report center"}
    if qty > 0:
        line["approved_unit_cost"] = cost
    return call("POST", "adjustments", {"date": date, "warehouse_id": wh, "division_id": div,
                                        "adj_type": "Opening" if qty > 0 else "Koreksi", "reason": "QA report center",
                                        "notes": "REPORT CENTER TEST", "lines": [line]})


def get_bin(u, path):
    s = u.S if u is not None else T.S
    r = s.get(f"{API}/{path}")
    return r.status_code, r.content, r.headers


def xlsx_table(content):
    ws = load_workbook(io.BytesIO(content), read_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Kode Barang")
    header = [h for h in rows[hi] if h is not None]
    body = [r[:len(header)] for r in rows[hi + 1:] if r and any(v is not None for v in r)]
    total = body[-1] if body and body[-1][0] == "TOTAL" else None
    return header, (body[:-1] if total else body), total, rows[:hi]


def all_pages(caller, key, qs=""):
    rows, page, first = [], 1, None
    while True:
        sc, r = caller("GET", f"report-center/{key}?page_size=1&page={page}{qs}")
        if sc != 200:
            return sc, r, rows
        first = first or r
        rows += r["rows"]
        if page >= r["pages"]:
            return sc, first, rows
        page += 1


def main():
    M = T.setup()
    u = uuid.uuid4().hex[:5]
    dA = M["div"]
    dB = call("POST", "master/divisions", {"code": f"RB{u}", "name": f"Divisi RB {u}"}, 200)[1]
    wA1 = call("POST", "master/warehouses", {"code": f"R1{u}", "name": f"Gudang RA1 {u}", "division_id": dA["id"], "is_active": True}, 200)[1]
    wA2 = call("POST", "master/warehouses", {"code": f"R2{u}", "name": f"Gudang RA2 {u}", "division_id": dA["id"], "is_active": True}, 200)[1]
    wB = call("POST", "master/warehouses", {"code": f"RB{u}", "name": f"Gudang RB {u}", "division_id": dB["id"], "is_active": True}, 200)[1]
    ref = {"unit": "PCS", "base_uom_id": M["uom"]["id"], "category_id": M["cat"]["id"], "is_active": True}
    X = call("POST", "master/items", {"code": f"RX{u}", "name": "Barang RX", "division_id": dA["id"], **ref}, 200)[1]
    Z = call("POST", "master/items", {"code": f"RZ{u}", "name": "Barang RZ", "division_id": dB["id"], **ref}, 200)[1]
    for wh, dv, it, q, c, dt in ((wA1, dA, X, 10, 1000, "2026-08-10"), (wB, dB, Z, 3, 2000, "2026-09-01"),
                                 (wA2, dA, X, 4, 500, "2026-09-15"), (wA1, dA, X, -5, None, "2026-09-28")):
        sc, r = adj(wh["id"], dv["id"], it["id"], q, c, dt)
        check(f"setup: penyesuaian {it['name']} {q:+} @ {dt}", sc == 200, (sc, r))
    T.make_po(M, "supX", 5)
    T.make_po(M, "supY", 2, item="item2")
    sc, mroB = call("POST", "mro", {"no": f"MRO-RB{u}", "division_id": dB["id"], "requester": "Budi", "submitted": True,
                                    "lines": [{"item_id": Z["id"], "qty": 7, "warehouse_id": wB["id"], "project_id": M["pa"]["id"]}]}, 200)
    check("setup: MRO Divisi B", sc == 200, (sc, mroB))

    A = mk_user("manager", divs=[dA["id"]], overrides={"export": "allow"})
    B = mk_user("manager", divs=[dB["id"]], overrides={"export": "deny"})   # juga dipakai untuk uji tanpa izin export
    AW = mk_user("warehouse", divs=[dA["id"]], warehouses=[wA2["id"]], overrides={"view_purchase_price": "allow"})
    NP = mk_user("warehouse", divs=[dA["id"]], overrides={"export": "allow"})
    NX = B   # batas paket starter: 5 pengguna per tenant

    sc, iw0 = call("GET", "item-warehouse")
    snap0 = sorted((x.get("item_id"), x.get("warehouse_id"), x.get("current_stock"), x.get("avg_cost"), x.get("total_value")) for x in iw0 or [])

    # ------------------------------------------------------------------ V. valuation-summary / ledger
    sc, s0 = call("GET", "reports/valuation-summary")
    whs0 = {r["warehouse_id"] for r in s0.get("rows", [])}
    check("V1 admin (global): summary memuat gudang Divisi A & B", sc == 200 and {wA1["id"], wA2["id"], wB["id"]} <= whs0, whs0)
    check("V2 bentuk respon kompatibel (method/rows/total_value; kunci baris sama)", set(s0) == {"method", "rows", "total_value"}
          and all(set(r) == SUMMARY_KEYS for r in s0["rows"]), list(s0))
    sc, sa = A("GET", "reports/valuation-summary")
    wa = {r["warehouse_id"] for r in sa.get("rows", [])}
    check("V3 user Divisi A: hanya gudang Divisi A (tanpa gudang B)", sc == 200 and wa == {wA1["id"], wA2["id"]}, wa)
    check("V4 total_value user A = Σ baris A = 7.000", abs(sa.get("total_value", 0) - sum(r["inventory_value"] for r in sa["rows"])) < 1e-6
          and abs(sa.get("total_value", 0) - 7000) < 0.01, sa.get("total_value"))
    sc, _ = A("GET", f"reports/valuation-summary?warehouse_id={wB['id']}")
    check("V5 user A minta gudang Divisi B -> 403", sc == 403, sc)
    sc, _ = A("GET", f"reports/valuation-ledger?warehouse_id={wB['id']}")
    check("V6 user A ledger gudang Divisi B -> 403", sc == 403, sc)
    sc, sb = B("GET", "reports/valuation-summary")
    check("V7 user Divisi B: hanya gudang B (nilai 6.000)", sc == 200 and {r["warehouse_id"] for r in sb["rows"]} == {wB["id"]}
          and abs(sb["total_value"] - 6000) < 0.01, sb)
    sc, sw = AW("GET", "reports/valuation-summary")
    check("V8 user ditugaskan gudang RA2 saja: hanya RA2", sc == 200 and {r["warehouse_id"] for r in sw.get("rows", [])} == {wA2["id"]}, (sc, sw))
    sc, _ = AW("GET", f"reports/valuation-summary?warehouse_id={wA1['id']}")
    check("V9 user gudang RA2 minta RA1 (divisi sama, tidak ditugaskan) -> 403", sc == 403, sc)
    sc, la = A("GET", "reports/valuation-ledger")
    names = {r.get("warehouse_name") for r in la.get("rows", [])}
    check("V10 ledger user A: tanpa entri gudang Divisi B", sc == 200 and wB["name"] not in names and wA1["name"] in names, names)
    sc, lb = B("GET", "reports/valuation-ledger")
    check("V11 ledger user B: hanya gudang B", sc == 200 and {r.get("warehouse_name") for r in lb["rows"]} == {wB["name"]}, lb)
    sc, lp = A("GET", "reports/valuation-ledger?date_from=2026-09-01&date_to=2026-09-30")
    days = sorted(r.get("txn_date") for r in lp.get("rows", []))
    check("V12 filter tanggal = tanggal efektif transaksi (09-15, 09-28; tanpa 08-10)", sc == 200 and days == ["2026-09-15", "2026-09-28"], days)
    for nm, path in (("summary", "reports/valuation-summary"), ("ledger", "reports/valuation-ledger")):
        sc, _ = NP("GET", path)
        check(f"V13 tanpa izin harga: valuation-{nm} -> 403", sc == 403, sc)
    check("V14 tanggal efektif WIB: 2026-09-30T18:30Z -> 2026-10-01; tanggal murni tetap",
          effective_day({"txn_at": "2026-09-30T18:30:00+00:00"}) == "2026-10-01" and effective_day({"txn_at": "2026-09-30"}) == "2026-09-30"
          and effective_day({"at": "2026-09-30T16:59:59+00:00"}) == "2026-09-30", "")

    # ------------------------------------------------------------------ R. Pusat Laporan
    sc, cat = call("GET", "report-center/catalog")
    gk = [g["key"] for g in cat.get("groups", [])]
    check("R1 katalog: 5 kelompok berurutan + MRO Traceability di Procurement",
          sc == 200 and gk == ["persediaan", "procurement", "warehouse", "spk", "hutang"]
          and any(r["key"] == "mro-traceability" for r in cat["groups"][1]["reports"]), gk)
    sc, first, rows = all_pages(call, "mro-traceability")
    n = first.get("total_rows", -1) if isinstance(first, dict) else -1
    check("R2 pagination: total_rows = jumlah baris semua halaman (page_size=1) & >= 3", sc == 200 and n == len(rows) and n >= 3 and first["pages"] == n, (n, len(rows)))
    tot = first.get("totals", {})
    check("R3 total = Σ seluruh baris (bukan halaman aktif)", abs(tot.get("request", -1) - sum(r["request"] or 0 for r in rows)) < 1e-9
          and abs(tot.get("po_total", -1) - round(sum(r["po_total"] or 0 for r in rows), 2)) < 0.01, tot)
    sc, legacy = call("GET", "reports/mro-traceability")
    check("R4 kompatibel: endpoint lama MRO Traceability tetap 200 & jumlah baris = Pusat Laporan", sc == 200 and isinstance(legacy, list) and len(legacy) == n, (sc, len(legacy or [])))
    sc, content, hdr = get_bin(None, "report-center/mro-traceability/export.xlsx")
    header, body, trow, head = xlsx_table(content)
    cols = first["columns"]
    check("R5 Excel: header = kolom JSON, jumlah baris = total_rows, X-Report-Rows benar",
          sc == 200 and header == [c["label"] for c in cols] and len(body) == n and hdr.get("x-report-rows") == str(n), (sc, len(body)))
    order = {r["mro_no"] + r["item_code"]: r for r in rows}
    ok = True
    for b in body:
        rec = dict(zip([c["key"] for c in cols], b))
        j = order.get((rec["mro_no"] or "") + (rec["item_code"] or ""))
        for c in cols:
            jv, xv = (j or {}).get(c["key"]), rec[c["key"]]
            if c["type"] in ("qty", "money", "int"):
                ok &= j is not None and abs(float(jv or 0) - float(xv or 0)) < 1e-6
            else:
                ok &= j is not None and (jv or "") == (xv or "")
    check("R6 paritas Excel = JSON per sel (semua baris & kolom)", ok and len(order) == n, "")
    check("R7 Excel: baris TOTAL = totals JSON", trow is not None and all(
        abs(float(trow[i] or 0) - tot[c["key"]]) < 1e-6 for i, c in enumerate(cols) if c["total"]), trow)
    check("R8 Excel: kop memuat judul, filter, pencetak", any("MRO Traceability" in str(h[0]) for h in head if h)
          and any(str(h[0]).startswith("Dicetak oleh") for h in head if h), head[:3])
    sc, pdf, hdr = get_bin(None, "report-center/mro-traceability/export.pdf")
    rd = PdfReader(io.BytesIO(pdf))
    text = " ".join(p.extract_text() or "" for p in rd.pages)
    flat = "".join(text.split())   # sel sempit dapat memenggal teks ke baris berikut
    check("R9 PDF: valid, metadata rows = total_rows, setiap No. MRO tercetak, nomor halaman",
          sc == 200 and pdf[:4] == b"%PDF" and (rd.metadata or {}).get("/Subject") == f"rows={n}"
          and all(r["mro_no"] in flat for r in rows) and f"Hal1/{len(rd.pages)}" in flat, (sc, len(pdf)))
    sc, ra = A("GET", "report-center/mro-traceability?page_size=500")
    sc2, rb = B("GET", "report-center/mro-traceability?page_size=500")
    check("R10 scope divisi: user A tanpa MRO Divisi B; user B hanya MRO Divisi B",
          sc == 200 and sc2 == 200 and mroB["no"] not in {r["mro_no"] for r in ra["rows"]}
          and {r["mro_no"] for r in rb["rows"]} == {mroB["no"]}, ({r["mro_no"] for r in rb.get("rows", [])},))
    def rowset(rs):
        return sorted(((r.get("mro_no") or ""), (r.get("item_code") or "")) for r in rs)
    for nm, usr, rc in (("A", A, ra), ("B", B, rb)):
        sc, lg = usr("GET", "reports/mro-traceability")
        lg_rows = [{"mro_no": RTDJ(x.get("mro_refs"), "no"), "item_code": x.get("item_code")} for x in (lg or [])]
        check(f"R10b tanpa perluasan akses: user {nm} endpoint lama = Pusat Laporan (baris identik)",
              sc == 200 and rowset(lg_rows) == rowset(rc["rows"]), (len(lg or []), rc.get("total_rows")))
    sc, xa, _ = get_bin(A, "report-center/mro-traceability/export.xlsx")
    hxa, bxa, _, _ = xlsx_table(xa)
    ino = hxa.index("No. MRO")
    sc2, pa, _ = get_bin(A, "report-center/mro-traceability/export.pdf")
    fa = "".join(" ".join(pg.extract_text() or "" for pg in PdfReader(io.BytesIO(pa)).pages).split())
    check("R10c scope divisi di Excel & PDF user A: tanpa MRO Divisi B, baris = JSON user A",
          sc == 200 and sc2 == 200 and len(bxa) == ra["total_rows"] and mroB["no"] not in {b[ino] for b in bxa}
          and mroB["no"] not in fa and all(r["mro_no"] in fa for r in ra["rows"]), (len(bxa), ra.get("total_rows")))
    sc, _ = A("GET", f"report-center/mro-traceability?division_id={dB['id']}")
    check("R11 user A filter Divisi B -> 403", sc == 403, sc)
    sc, rf = call("GET", f"report-center/mro-traceability?division_id={dB['id']}&page_size=500")
    check("R12 filter divisi (admin) = hanya MRO Divisi B + label filter", sc == 200 and {r["mro_no"] for r in rf["rows"]} == {mroB["no"]}
          and any(f["key"] == "division_id" and f["value"] == dB["name"] for f in rf["filters_applied"]), rf.get("filters_applied"))
    sc, rn = NP("GET", "report-center/mro-traceability?page_size=500")
    keys = {c["key"] for c in rn.get("columns", [])}
    check("R13 tanpa izin harga: kolom & nilai harga tidak dikirim (JSON)", sc == 200 and not (keys & PRICE_KEYS)
          and all(not (set(r) & PRICE_KEYS) for r in rn["rows"]) and rn["price_visible"] is False and not (set(rn["totals"]) & PRICE_KEYS), keys & PRICE_KEYS)
    sc, content, _ = get_bin(NP, "report-center/mro-traceability/export.xlsx")
    h2, b2, _, head2 = xlsx_table(content)
    check("R14 tanpa izin harga: Excel tanpa kolom harga + catatan redaksi", sc == 200 and not (set(h2) & PRICE_LABELS)
          and len(b2) == rn["total_rows"] and any("disembunyikan" in str(h[0]) for h in head2 if h), set(h2) & PRICE_LABELS)
    sc, pdf2, _ = get_bin(NP, "report-center/mro-traceability/export.pdf")
    t2 = "".join(" ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf2)).pages).split())
    check("R15 tanpa izin harga: PDF tanpa label harga", sc == 200 and not any(lb in t2 for lb in ("NilaiPO", "DPPPO", "TotalPO", "PajakPO")), "")
    for f in ("xlsx", "pdf"):
        sc, _, _ = get_bin(NX, f"report-center/mro-traceability/export.{f}")
        check(f"R16 tanpa izin export: export.{f} -> 403", sc == 403, sc)
    sc, _, _ = get_bin(NX, "reports/mro-traceability/export.xlsx")
    check("R16b paritas izin export: endpoint lama export.xlsx juga 403 untuk user tanpa izin export", sc == 403, sc)
    sc, _, _ = get_bin(A, "reports/mro-traceability/export.xlsx")
    check("R16c paritas izin export: endpoint lama export.xlsx 200 untuk user dengan izin export", sc == 200, sc)
    for qs, exp, nm in (("date_from=2026-13-01", 400, "tanggal invalid"), ("date_from=2026-10-02&date_to=2026-10-01", 400, "awal > akhir"),
                        ("status=Bogus", 400, "status tidak dikenal"), ("page=999&page_size=1", 200, "halaman > max dijepit")):
        sc, r = call("GET", f"report-center/mro-traceability?{qs}")
        check(f"R17 validasi filter: {nm} -> {exp}", sc == exp, (sc, r))
    sc, r = call("GET", "report-center/tidak-ada")
    check("R18 laporan tidak dikenal -> 404", sc == 404, sc)
    target = rows[0]["mro_no"]
    sc, rq = call("GET", f"report-center/mro-traceability?q={target}&page_size=500")
    check("R19 pencarian No. MRO -> hanya baris MRO tsb", sc == 200 and rq["total_rows"] >= 1 and all(r["mro_no"] == target for r in rq["rows"]), rq.get("total_rows"))
    yday = (TODAY - timedelta(days=1)).isoformat()
    sc, rt = call("GET", f"report-center/mro-traceability?date_from={TODAY.isoformat()}&date_to={TODAY.isoformat()}")
    sc2, ry = call("GET", f"report-center/mro-traceability?date_to={yday}")
    check("R20 cut-off WIB: MRO hari ini masuk periode hari ini, tidak masuk s/d kemarin", sc == 200 and rt["total_rows"] == n and ry["total_rows"] == 0,
          (rt.get("total_rows"), ry.get("total_rows")))
    sc, st = call("GET", "report-center/mro-traceability?status=Open&page_size=500")
    check("R21 filter status Open", sc == 200 and st["total_rows"] >= 1 and all(r["status"] == "Open" for r in st["rows"]), st.get("total_rows"))
    sc, au = call("GET", "audit?entity=report&limit=50")
    items = au.get("items", au) if isinstance(au, dict) else au
    ex = [a for a in items or [] if a.get("action") == "export" and a.get("entity_id") == "mro-traceability"]
    check("R22 audit export tercatat (format & jumlah baris)", len(ex) >= 2 and any((a.get("after") or {}).get("rows") == n for a in ex), len(ex))

    # ------------------------------------------------------------------ L. batas export & exporter (in-process)
    def fake(rows_n):
        return {"total_rows": rows_n}
    for fmt, lim in (("xlsx", 100_000), ("pdf", 5_000)):
        try:
            RC.check_limit(fake(lim), fmt); ok_at = True
        except HTTPException:
            ok_at = False
        try:
            RC.check_limit(fake(lim + 1), fmt); over = None
        except HTTPException as e:
            over = e
        check(f"L1 batas {fmt}: {lim:,} baris diizinkan, {lim + 1:,} -> 422 dengan pesan 'Tidak ada data yang dipotong'",
              ok_at and over is not None and over.status_code == 422 and "dipotong" in over.detail, getattr(over, "detail", None))
    empty = {"meta": {"report_key": "x", "title": "Uji", "group_title": "G", "date_basis": "", "company": "PT", "sheet": "Uji",
                      "generated_by": "QA", "generated_at": "01-01-2026 00:00"},
             "columns": [{"key": "a", "label": "A", "type": "qty", "price": False, "total": True, "width": 10}],
             "rows": [], "totals": {"a": 0}, "total_rows": 0, "price_visible": True, "filters_applied": []}
    check("L2 exporter tahan data kosong (xlsx & pdf)", to_xlsx(empty).getvalue()[:2] == b"PK" and to_pdf(empty).getvalue()[:4] == b"%PDF", "")
    big_cols = [{"key": "item_code", "label": "Kode Barang", "type": "text", "price": False, "total": False, "width": 14},
                {"key": "qty", "label": "Qty", "type": "qty", "price": False, "total": True, "width": 10},
                {"key": "val", "label": "Nilai PO", "type": "money", "price": True, "total": True, "width": 14}]
    big_rows = [{"item_code": f"BRG-{i:04d}", "qty": 1.5, "val": 1234.5} for i in range(400)]
    big = {**empty, "columns": big_cols, "rows": big_rows, "totals": {"qty": 600.0, "val": 493800.0}, "total_rows": 400}
    rd = PdfReader(io.BytesIO(to_pdf(big).getvalue()), strict=True)
    pages = ["".join((pg.extract_text() or "").split()) for pg in rd.pages]
    np_ = len(pages)
    check("L4 PDF multi-halaman valid (strict), header tabel berulang, nomor halaman Hal x/N, semua baris",
          np_ >= 2 and all(f"Hal{i + 1}/{np_}" in t for i, t in enumerate(pages)) and all("KodeBarang" in t for t in pages)
          and all(f"BRG-{i:04d}" in "".join(pages) for i in range(400)), np_)
    check("L5 PDF total di halaman terakhir dengan format id-ID", "TOTAL" in pages[-1] and "493.800,00" in pages[-1] and "600" in pages[-1], pages[-1][-200:])
    red = {**big, "columns": big_cols[:2], "rows": [{k: r[k] for k in ("item_code", "qty")} for r in big_rows],
           "totals": {"qty": 600.0}, "price_visible": False}
    rt = "".join(" ".join(pg.extract_text() or "" for pg in PdfReader(io.BytesIO(to_pdf(red).getvalue()), strict=True).pages).split())
    check("L6 PDF tanpa izin harga: tanpa kolom/nilai harga + catatan redaksi", "NilaiPO" not in rt and "493.800" not in rt
          and "disembunyikan" in rt, "")
    check("L3 local_day WIB (UTC 17:00 = besok WIB)", local_day("2026-09-30T17:00:00+00:00") == "2026-10-01"
          and local_day("2026-09-30T16:59:59Z") == "2026-09-30", "")

    sc, iw1 = call("GET", "item-warehouse")
    snap1 = sorted((x.get("item_id"), x.get("warehouse_id"), x.get("current_stock"), x.get("avg_cost"), x.get("total_value")) for x in iw1 or [])
    check("Z1 read-only: stok/avg/nilai pool tidak berubah setelah laporan & export dibaca", snap0 == snap1, "")

    failed = [nm for nm, ok in T.RESULTS if not ok]
    print(f"\n{len(T.RESULTS) - len(failed)}/{len(T.RESULTS)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
