"""P2b Outstanding & Nilai Procurement + konsolidasi navigasi Laporan (Riwayat Pergerakan Stok).

Hanya pada backend + DB test terisolasi (DB *itest*; tenant QA via /saas/register) — BUKAN preview/sandbox/production.
Membuktikan: outstanding per ALOKASI baris (parsial, multi-referensi 1 MRO -> 2 RO dan 1 PO -> 2 DO, pembatalan MI/DO
mengembalikan outstanding) tanpa double counting; paritas dengan monitor MRO existing; komitmen PO vs realisasi DO (DPP basis,
PPN & total terpisah, 1 baris PO dihitung sekali walau multi-DO); pengelompokan Supplier/Barang/Proyek/Periode; scope divisi;
redaksi harga server-side (JSON/Excel/PDF); export = seluruh hasil filter; read-only.
"""
import io
import os
import sys
import uuid
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

import receipt_control_test as T
from master_item_stock_test import mk_user

call, check, API = T.call, T.check, T.API
TODAY = datetime.now(ZoneInfo("Asia/Jakarta")).date()


def near(a, b, eps=0.01):
    return abs(float(a or 0) - float(b or 0)) <= eps


def rc(caller, key, qs=""):
    return caller("GET", f"report-center/{key}?page_size=500&{qs}")


def rows_of(res):
    return [r for r in (res or {}).get("rows", []) if not r.get("_kind")]


def main():
    dbn = (os.environ.get("TEST_DATABASE_URL") or "").rsplit("/", 1)[-1].split("?")[0]
    if not ("itest" in dbn or dbn.endswith("_test")):
        print("Menolak berjalan: hanya untuk backend + DB test terisolasi (scripts/run_regression_itest.sh, DB *itest*).")
        sys.exit(2)
    M = T.setup()
    u = uuid.uuid4().hex[:5]
    dB = call("POST", "master/divisions", {"code": f"DB{u}", "name": f"Divisi B {u}"}, 200)[1]

    # Alur A: MRO-A 10 -> RO 10 -> PO-A 10 (Supplier X, Project A, @1000) -> DO1 4 + DO2 3 -> MI 2 + MI 1 (MI 1 dibatalkan)
    poA, mroA_no, roA_no = T.make_po(M, "supX", 10, project="pa", mro_no=f"MRO-A{u}")
    sc1, do1 = call("POST", "do", T.do_body(M, poA, 4, "supX"))
    sc2, do2 = call("POST", "do", T.do_body(M, poA, 3, "supX"))
    check("setup: PO-A -> DO1 4 + DO2 3", sc1 == 200 and sc2 == 200, (do1, do2))
    sc, mros = call("GET", "mro")
    mros = mros if isinstance(mros, list) else (mros or {}).get("rows") or (mros or {}).get("items") or []
    mroA = next(m for m in mros if m.get("no") == mroA_no)
    ml = call("GET", f"mro/{mroA['id']}")[1]["lines"][0]

    def mi(q):
        return call("POST", "mi", {"division_id": M["div"]["id"], "source_type": "MRO", "default_warehouse_id": M["wh"]["id"],
                                   "receiver": "Operator", "lines": [{
                                       "item_id": M["item"]["id"], "qty": q, "uom_id": M["uom"]["id"], "conversion_factor": 1,
                                       "warehouse_id": M["wh"]["id"], "project_id": M["pa"]["id"], "mro_id": mroA["id"],
                                       "mro_line_id": ml["id"], "sources": [{"mro_id": mroA["id"], "line_id": ml["id"], "qty": q, "base_qty": q}]}]})
    sc1, mi1 = mi(2)
    sc2, mi2 = mi(1)
    sc3, _ = call("DELETE", f"transactions/mi/{mi2.get('id')}")
    check("setup: MI 2 + MI 1, MI kedua dibatalkan (reversal)", sc1 == 200 and sc2 == 200 and sc3 == 200)

    # Alur B: MRO-B 10 -> RO-B1 6 + RO-B2 4 (multi-referensi) -> PO-B 6 dari RO-B1 (Supplier Y, Project B, @2000, approved)
    sc, mroB = call("POST", "mro", {"no": f"MRO-B{u}", "division_id": M["div"]["id"], "requester": "Sari", "submitted": True,
                                    "lines": [{"item_id": M["item"]["id"], "qty": 10, "warehouse_id": M["wh"]["id"], "project_id": M["pb"]["id"]}]}, 200)
    call("POST", f"mro/{mroB['id']}/submit", {})
    ros = []
    for q in (6, 4):
        sc, ro = call("POST", "ro", {"division_id": M["div"]["id"], "submitted": True, "lines": [{
            "item_id": M["item"]["id"], "qty": q, "warehouse_id": M["wh"]["id"], "project_id": M["pb"]["id"],
            "sources": [{"mro_id": mroB["id"], "line_id": mroB["lines"][0]["id"], "qty": q}]}]}, 200)
        call("POST", f"ro/{ro['id']}/submit", {})
        ros.append(ro)
    sc, poB = call("POST", "po", {"supplier_id": M["supY"]["id"], "division_id": M["div"]["id"], "lines": [{
        "item_id": M["item"]["id"], "qty": 6, "price": 2000, "warehouse_id": M["wh"]["id"], "project_id": M["pb"]["id"],
        "sources": [{"ro_id": ros[0]["id"], "line_id": ros[0]["lines"][0]["id"], "qty": 6}]}]}, 200)
    call("POST", f"po/{poB['id']}/submit", {})
    if call("GET", f"po/{poB['id']}")[1].get("status") != "Approved":
        call("POST", f"po/{poB['id']}/approve", {})
    # Alur C: PO-C draft (belum disubmit) -> bukan komitmen
    sc, poC = call("POST", "po", {"supplier_id": M["supY"]["id"], "division_id": M["div"]["id"], "lines": [{
        "item_id": M["item"]["id"], "qty": 5, "price": 9999, "warehouse_id": M["wh"]["id"], "project_id": M["pb"]["id"],
        "sources": [{"ro_id": ros[1]["id"], "line_id": ros[1]["lines"][0]["id"], "qty": 4}]}]})
    check("setup: MRO-B -> 2 RO (6 + 4) -> PO-B 6 approved; PO-C draft", bool(poB.get("id")), (poB, sc, poC))
    NP = mk_user("warehouse", divs=[M["div"]["id"]], overrides={"export": "allow"})
    UB = mk_user("manager", divs=[dB["id"]], overrides={"export": "allow"})
    by_no = lambda res, no: [r for r in rows_of(res) if r.get("no") == no]  # noqa: E731

    # ------------------------------------------------------------------ O. Outstanding
    sc, om = rc(call, "outstanding-mro")
    a, b = (by_no(om, mroA_no) or [{}])[0], (by_no(om, mroB["no"]) or [{}])[0]
    check("O1 Outstanding MRO-A = 10 − MI 2 (MI batal tidak dihitung) = 8, Sebagian 20%, ref MI",
          sc == 200 and a.get("qty") == 10 and near(a.get("done"), 2) and near(a.get("outstanding"), 8) and a.get("status") == "Sebagian"
          and near(a.get("pct"), 20) and (mi1.get("no") or "?") in (a.get("refs") or "") and (mi2.get("no") or "?") not in (a.get("refs") or ""), a)
    check("O2 Outstanding MRO-B (2 RO, belum ada MI) = 10, Belum Diproses, umur 0 hari", near(b.get("outstanding"), 10)
          and b.get("status") == "Belum Diproses" and b.get("age") == 0, b)
    mon = call("GET", f"mro/{mroA['id']}")[1]["lines"][0]["monitor"]
    check("O3 paritas dengan monitor MRO existing (qty MI, outstanding)", near(mon["qty_mi"], a.get("done")) and near(mon["outstanding"], a.get("outstanding")), (mon, a))
    sc, orr = rc(call, "outstanding-ro")
    rb = [r for r in rows_of(orr) if r.get("no") in (ros[0]["no"], ros[1]["no"])]
    rod = call("GET", f"ro/{ros[1]['id']}")[1]
    st2 = (rod.get("lines") or [{}])[0]
    b2 = next((r for r in rb if r["no"] == ros[1]["no"]), {})
    check("O4 Outstanding RO: RO-B1 (6 dipesan penuh) tidak tampil; RO-B2 qty 4, outstanding = qty − alokasi PO (= status RO existing)",
          all(r["no"] != ros[0]["no"] for r in rb) and b2.get("qty") == 4 and near(b2.get("outstanding") + b2.get("done"), 4)
          and near(st2.get("outstanding"), b2.get("outstanding")) and near(st2.get("ordered"), b2.get("done")), (b2, st2.get("outstanding"), st2.get("ordered")))
    sc, ora = rc(call, "outstanding-ro", "show=all")
    r1 = [r for r in rows_of(ora) if r.get("no") == ros[0]["no"]]
    check("O5 show=all menampilkan RO-B1 Selesai 100%", r1 and r1[0]["status"] == "Selesai" and near(r1[0]["pct"], 100), r1)
    sc, op = rc(call, "outstanding-po")
    pa_, pb_ = (by_no(op, poA["no"]) or [{}])[0], (by_no(op, poB["no"]) or [{}])[0]
    check("O6 Outstanding PO-A = 10 − (DO1 4 + DO2 3) = 3, 70%, ref 2 DO, Supplier X; PO-B 6 belum diterima",
          near(pa_.get("outstanding"), 3) and near(pa_.get("pct"), 70) and do1.get("no", "?") in pa_.get("refs", "") and do2.get("no", "?") in pa_.get("refs", "")
          and pa_.get("supplier") == "Supplier X" and near(pb_.get("outstanding"), 6), (pa_, pb_))
    check("O7 PO draft bukan komitmen -> tidak muncul di Outstanding PO", not by_no(op, poC.get("no")), poC.get("no"))
    sc, od = rc(call, "outstanding-do")
    d_ = (by_no(od, mroA_no) or [{}])[0]
    check("O8 Outstanding DO→MI MRO-A: diterima 7 − diserahkan 2 = 5 (paritas monitor qty_received)",
          near(d_.get("qty"), 7) and near(d_.get("done"), 2) and near(d_.get("outstanding"), 5) and near(mon["qty_received"], 7), (d_, mon))
    sc, ost = rc(call, "outstanding-po", f"supplier_id={M['supX']['id']}")
    check("O9 filter supplier", all(r.get("supplier") == "Supplier X" for r in rows_of(ost)) and by_no(ost, poA["no"]), rows_of(ost)[:2])
    sc, oy = rc(call, "outstanding-mro", f"date_to={(TODAY - timedelta(days=1)).isoformat()}")
    check("O10 cut-off/periode: Tanggal Akhir kemarin -> dokumen hari ini tidak dihitung", not by_no(oy, mroA_no), len(rows_of(oy)))

    # ------------------------------------------------------------------ R. Rekap Pembelian & Rekap Nilai DO
    sc, rp = rc(call, "rekap-pembelian", "group_by=supplier")
    g = {r["group"]: r for r in rows_of(rp)}
    x, y = g.get("Supplier X", {}), g.get("Supplier Y", {})
    check("R1 Rekap Pembelian per supplier: X qty 10 DPP 10.000 diterima 7.000 sisa 3.000; Y DPP 12.000 (PO draft tidak dihitung)",
          x.get("po_count") == 1 and near(x.get("qty"), 10) and near(x.get("dpp"), 10000) and near(x.get("dpp_received"), 7000)
          and near(x.get("dpp_open"), 3000) and near(y.get("dpp"), 12000) and near(y.get("qty"), 6), (x, y))
    check("R2 kolom bruto/diskon/PPN/total = snapshot PO existing (tanpa pajak: total = DPP, diskon 0)",
          near(x.get("gross"), 10000) and near(x.get("discount"), 0) and near(x.get("tax"), 0) and near(x.get("total"), 10000), x)
    tot = (rp.get("totals") or {})
    check("R3 TOTAL = seluruh hasil filter", near(tot.get("dpp"), sum(r["dpp"] for r in rows_of(rp))), tot)
    for gb, expect in (("item", None), ("project", {"Project A": 10000, "Project B": 12000}), ("month", {TODAY.isoformat()[:7]: 22000}),
                       ("year", {TODAY.isoformat()[:4]: 22000})):
        sc, rg = rc(call, "rekap-pembelian", f"group_by={gb}")
        gg = {r["group"]: r["dpp"] for r in rows_of(rg)}
        ok = sc == 200 and (near(sum(gg.values()), 22000) if expect is None else all(near(gg.get(k), v) for k, v in expect.items()))
        check(f"R4 pengelompokan {gb} konsisten (total DPP 22.000)", ok, gg)
    sc, rpf = rc(call, "rekap-pembelian", f"group_by=supplier&project_id={M['pa']['id']}")
    check("R5 filter kombinasi (proyek A) hanya Supplier X", [r["group"] for r in rows_of(rpf)] == ["Supplier X"], rows_of(rpf))
    sc, rd = rc(call, "rekap-nilai-do", "group_by=supplier")
    dx = {r["group"]: r for r in rows_of(rd)}.get("Supplier X", {})
    check("R6 Rekap Nilai DO X: 2 DO, 1 PO, qty 7, DPP 7.000, PPN 0, total 7.000, komitmen PO 10.000 dihitung sekali, 70%",
          dx.get("do_count") == 2 and dx.get("po_count") == 1 and near(dx.get("qty"), 7) and near(dx.get("dpp"), 7000)
          and near(dx.get("tax"), 0) and near(dx.get("total"), 7000) and near(dx.get("po_dpp"), 10000) and near(dx.get("pct"), 70), dx)

    # ------------------------------------------------------------------ T. PPN & diskon (DPP basis; PPN & total kolom terpisah)
    supZ = call("POST", "master/suppliers", {"code": f"SZ{u}", "name": "Supplier Z", "supplier_category_id": M["scat"]["id"]}, 200)[1]
    M["supZ"] = supZ
    sc, mroD = call("POST", "mro", {"no": f"MRO-D{u}", "division_id": M["div"]["id"], "requester": "Dewi", "submitted": True,
                                    "lines": [{"item_id": M["item"]["id"], "qty": 5, "warehouse_id": M["wh"]["id"], "project_id": M["pa"]["id"]}]}, 200)
    call("POST", f"mro/{mroD['id']}/submit", {})
    sc, roD = call("POST", "ro", {"division_id": M["div"]["id"], "submitted": True, "lines": [{
        "item_id": M["item"]["id"], "qty": 5, "warehouse_id": M["wh"]["id"], "project_id": M["pa"]["id"],
        "sources": [{"mro_id": mroD["id"], "line_id": mroD["lines"][0]["id"], "qty": 5}]}]}, 200)
    call("POST", f"ro/{roD['id']}/submit", {})
    sc, poD = call("POST", "po", {"supplier_id": supZ["id"], "division_id": M["div"]["id"], "lines": [{
        "item_id": M["item"]["id"], "qty": 5, "price": 1000, "discount": 500, "tax": 11, "warehouse_id": M["wh"]["id"],
        "project_id": M["pa"]["id"], "sources": [{"ro_id": roD["id"], "line_id": roD["lines"][0]["id"], "qty": 5}]}]}, 200)
    call("POST", f"po/{poD['id']}/submit", {})
    if call("GET", f"po/{poD['id']}")[1].get("status") != "Approved":
        call("POST", f"po/{poD['id']}/approve", {})
    poD = call("GET", f"po/{poD['id']}")[1]
    pl = poD["lines"][0]
    sc, doD = call("POST", "do", T.do_body(M, poD, 2, "supZ"))
    check("setup: PO-D Supplier Z 5 @1.000 diskon 500 PPN 11% (approved) -> DO 2", sc == 200 and poD.get("status") == "Approved"
          and near(pl.get("dpp"), 4500) and near(pl.get("tax_amount"), 495), (sc, poD.get("status"), pl.get("dpp"), pl.get("tax_amount")))
    sc, rz = rc(call, "rekap-pembelian", f"group_by=supplier&supplier_id={supZ['id']}")
    z = (rows_of(rz) or [{}])[0]
    check("T1 Rekap Pembelian = snapshot baris PO existing: bruto 5.000, diskon 500, DPP 4.500, PPN 495, total 4.995, DPP diterima 1.800",
          sc == 200 and len(rows_of(rz)) == 1 and near(z.get("gross"), 5000) and near(z.get("discount"), 500) and near(z.get("dpp"), pl.get("dpp"))
          and near(z.get("tax"), pl.get("tax_amount")) and near(z.get("total"), pl.get("total")) and near(z.get("total"), 4995)
          and near(z.get("dpp_received"), 1800) and near(z.get("dpp_open"), 2700), (z, pl.get("dpp"), pl.get("tax_amount"), pl.get("total")))
    sc, dz = rc(call, "rekap-nilai-do", f"group_by=supplier&supplier_id={supZ['id']}")
    dzz = (rows_of(dz) or [{}])[0]
    check("T2 Rekap Nilai DO basis DPP: 2/5 × DPP 4.500 = 1.800; PPN 198 & total 1.998 terpisah; komitmen PO 4.500; 40%",
          sc == 200 and near(dzz.get("dpp"), 1800) and near(dzz.get("tax"), 198) and near(dzz.get("total"), 1998)
          and near(dzz.get("po_dpp"), 4500) and near(dzz.get("pct"), 40), dzz)
    sc, oz = rc(call, "outstanding-po", f"supplier_id={supZ['id']}")
    check("T3 Outstanding PO-D = 5 − 2 = 3 (40%)", [round(r["outstanding"], 6) for r in rows_of(oz)] == [3] and near(rows_of(oz)[0]["pct"], 40), rows_of(oz))
    tj = (rc(call, "rekap-nilai-do", "group_by=supplier")[1] or {}).get("totals") or {}
    xr = T.S.get(f"{API}/report-center/rekap-nilai-do/export.xlsx?group_by=supplier")
    xv = [list(v) for v in load_workbook(io.BytesIO(xr.content), read_only=True).worksheets[0].iter_rows(values_only=True)] if xr.ok else []
    hrow = next((v for v in xv if v and v[0] == "Kelompok"), [])
    trow = next((v for v in xv if v and v[0] == "TOTAL"), [])
    ix = {str(h): i for i, h in enumerate(hrow) if h}
    pdf = T.S.get(f"{API}/report-center/rekap-nilai-do/export.pdf?group_by=supplier")
    ptxt = " ".join(" ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf.content)).pages).split()) if pdf.ok else ""
    check("T4 konsistensi JSON = Excel (TOTAL DPP/PPN/Total) = PDF (baris Supplier Z & TOTAL)", xr.ok and pdf.ok and trow
          and near(trow[ix["DPP Diterima (Sebelum PPN)"]], tj.get("dpp")) and near(trow[ix["PPN"]], tj.get("tax"))
          and near(trow[ix["Total Termasuk PPN"]], tj.get("total")) and "Supplier Z" in ptxt and "TOTAL" in ptxt, (trow, tj))

    # ------------------------------------------------------------------ S. scope, harga, export
    sc, ub = rc(UB, "outstanding-mro")
    check("S1 scope divisi: user Divisi B tidak melihat MRO divisi A", sc == 200 and not by_no(ub, mroA_no), (sc, len(rows_of(ub))))
    scs = [rc(UB, k, q) for k, q in (("rekap-pembelian", "group_by=supplier"), ("rekap-nilai-do", "group_by=supplier"), ("outstanding-po", ""))]
    check("S1b scope divisi rekap/outstanding: user Divisi B tidak melihat PO/DO Supplier X/Y/Z divisi A",
          all(s == 200 for s, _ in scs) and not any(r.get("group") in ("Supplier X", "Supplier Y", "Supplier Z") or r.get("supplier") in ("Supplier X", "Supplier Y", "Supplier Z")
                                                     for _, j in scs for r in rows_of(j)), [len(rows_of(j)) for _, j in scs])
    sc, npg = rc(NP, "rekap-pembelian", "group_by=supplier")
    gx = {r["group"]: r for r in rows_of(npg)}
    check("S1c user divisi A (non-global) melihat rekap Supplier X qty 10 & Supplier Z qty 5 (visibilitas existing)",
          sc == 200 and near(gx.get("Supplier X", {}).get("qty"), 10) and near(gx.get("Supplier Z", {}).get("qty"), 5), list(gx))
    sc, npj = rc(NP, "rekap-pembelian", "group_by=supplier")
    cols = [c["key"] for c in (npj or {}).get("columns", [])]
    check("S2 tanpa view_purchase_price: kolom nilai dibuang server-side (qty tetap)", sc == 200 and "dpp" not in cols and "qty" in cols
          and all("dpp" not in r and "total" not in r for r in rows_of(npj)), (sc, cols))
    for key in ("outstanding-po", "rekap-pembelian", "rekap-nilai-do", "riwayat-pergerakan-stok"):
        r = T.S.get(f"{API}/report-center/{key}/export.xlsx")
        vals = [list(v) for v in load_workbook(io.BytesIO(r.content), read_only=True).worksheets[0].iter_rows(values_only=True)] if r.ok else []
        sc, js = rc(call, key)
        n = js.get("total_rows") if isinstance(js, dict) else None
        body = [v for v in vals if v and v[0] not in (None, "TOTAL")]
        rp_ = T.S.get(f"{API}/report-center/{key}/export.pdf")
        txt = " ".join(" ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(rp_.content)).pages).split()) if rp_.ok else ""
        check(f"S3 export {key}: Excel & PDF 200, Excel memuat seluruh {n} baris hasil filter", r.ok and rp_.ok and n is not None
              and len(body) >= n and len(txt) > 0, (r.status_code, rp_.status_code, n, len(body)))
    rn = NP.S.get(f"{API}/report-center/rekap-nilai-do/export.xlsx")
    vals = [list(v) for v in load_workbook(io.BytesIO(rn.content), read_only=True).worksheets[0].iter_rows(values_only=True)] if rn.ok else []
    hi = next((i for i, v in enumerate(vals) if v and v[0] == "Kelompok"), None)
    hdr = [str(c) for c in (vals[hi] if hi is not None else []) if c]
    check("S4 Excel tanpa izin harga: header & data tanpa kolom DPP/PPN/Total (qty tetap)", rn.ok and hdr and "Qty Diterima" in hdr
          and not any(k in h for h in hdr for k in ("DPP", "PPN", "Total", "Realisasi")) and all(len(v) <= len(hdr) + 1 for v in vals[hi + 1:] if v),
          (rn.status_code, hdr))
    pn = NP.S.get(f"{API}/report-center/rekap-pembelian/export.pdf?group_by=supplier")
    pt = " ".join(" ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pn.content)).pages).split()) if pn.ok else ""
    pa = T.S.get(f"{API}/report-center/rekap-pembelian/export.pdf?group_by=supplier")
    pat = " ".join(" ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pa.content)).pages).split()) if pa.ok else ""
    check("S5a kontrol: PDF admin memuat kolom nilai (DPP/Bruto) & angka 4.995", "Nilai Bruto" in pat and "4.995" in pat, pat[:300])
    check("S5 PDF tanpa izin harga: tanpa kolom nilai (DPP/PPN/Diskon/Bruto/Total) & tanpa angka nilai; qty tetap", pn.ok and "Qty Dipesan" in pt
          and not any(k in pt for k in ("DPP (Sebelum PPN)", "Nilai Bruto", "Diskon", "Total Termasuk PPN", "DPP Diterima", "DPP Belum Diterima",
                                        "4.995", "4.500", "12.000")), (pn.status_code, pt[:200]))

    # ------------------------------------------------------------------ P. Riwayat Pergerakan Stok (pengganti tab ledger Inventory)
    sc, rw = rc(call, "riwayat-pergerakan-stok", f"item_id={M['item']['id']}")
    ins = sum(r["qty_in"] for r in rows_of(rw)); outs = sum(r["qty_out"] for r in rows_of(rw))
    sc2, iw = call("GET", "item-warehouse")
    cur = sum(float(x.get("current_stock") or 0) for x in (iw or []) if x.get("item_id") == M["item"]["id"])
    check("P1 Riwayat Pergerakan Stok: masuk − keluar = stok Inventory saat ini (sumber stock_ledger sama)", sc == 200 and near(ins - outs, cur), (ins, outs, cur))

    # ------------------------------------------------------------------ X. pembatalan DO mengembalikan outstanding
    sc, _ = call("DELETE", f"transactions/do/{do2.get('id')}")
    sc_, op2 = rc(call, "outstanding-po")
    p2 = (by_no(op2, poA["no"]) or [{}])[0]
    sc_, rd2 = rc(call, "rekap-nilai-do", "group_by=supplier")
    dx2 = {r["group"]: r for r in rows_of(rd2)}.get("Supplier X", {})
    check("X1 DO2 dibatalkan -> Outstanding PO-A kembali 6; Rekap DO X = 1 DO, DPP 4.000",
          sc == 200 and near(p2.get("outstanding"), 6) and dx2.get("do_count") == 1 and near(dx2.get("dpp"), 4000), (sc, p2, dx2))
    sc, rl = call("GET", "reports/mro-traceability")
    check("X2 endpoint lama tetap tersedia (/api/reports/mro-traceability)", sc == 200, sc)

    failed = [nm_ for nm_, ok in T.RESULTS if not ok]
    print(f"\n{len(T.RESULTS) - len(failed)}/{len(T.RESULTS)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
