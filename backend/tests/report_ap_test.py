"""P5 Reporting Invoice, Pembayaran & Hutang Vendor pada Pusat Laporan (backend + DB test terisolasi *itest*).

Fixture lewat API Invoice Vendor existing (DO -> Invoice multi-DO / penagihan parsial -> pembayaran parsial / sekaligus ->
pembatalan pembayaran) + data legacy terkontrol (alokasi DP, invoice tanpa jatuh tempo, rekening supplier). Membuktikan:
Register (multi-DO 1 baris, DPP/PPN, posisi as-of), Monitoring Pembayaran (riwayat, kumulatif, batal, tanpa rekening/metode bayar),
Outstanding (= Dashboard per cut-off, lunas/DP tidak outstanding), ringkasan per divisi/proyek (multi-divisi tanpa hitung
ganda), Aging (bucket per cut-off WIB, tanpa jatuh tempo terpisah), Rekap Supplier (Saldo Awal + Invoice − Pembayaran
Efektif − DP = Saldo Akhir = Outstanding), rincian per bulan, izin invoice.view & cakupan divisi, hub, paritas JSON = Excel =
PDF, read-only.
"""
import io
import os
import sys
import uuid

from openpyxl import load_workbook
from pypdf import PdfReader

sys.path.insert(0, "/app/backend/tests")
sys.path.insert(0, "/app/backend")
import receipt_control_test as T
import vendor_invoice_test as V
from opening_correction_test import conn, ins

call, check, API = T.call, T.check, T.API
u = uuid.uuid4().hex[:5]
TODAY = None


def near(a, b, eps=0.01):
    return abs(float(a or 0) - float(b or 0)) <= eps


def rc(fn, key, qs=""):
    return fn("GET", f"report-center/{key}?page_size=500&{qs}")


def rows(res):
    return (res or {}).get("rows", [])


def by_no(res, k="no"):
    return {r[k]: r for r in rows(res)}


def xl(sess, key, qs, first):
    r = sess.get(f"{API}/report-center/{key}/export.xlsx?{qs}")
    if not r.ok:
        return r.status_code, [], [], []
    vals = [list(v) for v in load_workbook(io.BytesIO(r.content), read_only=True).worksheets[0].iter_rows(values_only=True)]
    hi = next((i for i, v in enumerate(vals) if v and v[0] == first), None)
    if hi is None:
        return r.status_code, [], [], []
    body = [v for v in vals[hi + 1:] if v and v[0] not in (None, "TOTAL")]
    return r.status_code, [str(x) for x in vals[hi]], body, next((v for v in vals[hi + 1:] if v and v[0] == "TOTAL"), [])


def pdf(sess, key, qs):
    r = sess.get(f"{API}/report-center/{key}/export.pdf?{qs}")
    return r.status_code, (" ".join(" ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(r.content)).pages).split()) if r.ok else "")


def sql(q, args):
    cn = conn()
    try:
        cn.cursor().execute(q, args)
    finally:
        cn.close()


def main():
    if not os.environ.get("DATABASE_URL") or "itest" not in os.environ.get("DATABASE_URL", ""):
        print("Menolak berjalan: hanya untuk DB test terisolasi (*itest*).")
        sys.exit(2)
    from datetime import datetime
    from zoneinfo import ZoneInfo
    global TODAY
    TODAY = datetime.now(ZoneInfo("Asia/Jakarta")).date().isoformat()
    M = T.setup()
    T.M = M
    V.M = M
    tenant = call("GET", "auth/me")[1].get("tenant_id")
    divA = M["div"]
    sc, divB = call("POST", "master/divisions", {"code": f"DB{u}", "name": f"Divisi B {u}"}, 200)

    # ------------------------------------------------------------------ fixture (API Invoice Vendor existing)
    _, dA1 = V.mk_do("supX", 10000)
    _, dA2 = V.mk_do("supX", 5000)
    _, dA3 = V.mk_do("supX", 8000)
    _, dA4 = V.mk_do("supX", 2000)
    _, dA5 = V.mk_do("supX", 500)
    _, dA6 = V.mk_do("supX", 1000)
    _, dY1 = V.mk_do("supY", 4000)
    M["div"] = divB
    _, dB1 = V.mk_do("supX", 500)
    M["div"] = divA

    def inv(no, sup, allocs, idate, due, tax=0, **kw):
        b = V.inv_body(f"{no}-{u}", sup, allocs, due=due, **kw)
        b.update({"invoice_date": idate, "received_date": idate, "tax_amount": tax})
        sc, r = call("POST", "vendor-invoices", b)
        assert sc == 200, (no, sc, r)
        return r

    def pay(i, day, amt, ref):
        sc, r = call("POST", f"vendor-invoices/{i['id']}/payments", {"date": day, "amount": amt, "reference": ref, "notes": f"cat {ref}"})
        assert sc == 200, (sc, r)
        return r["created_payment_id"]

    iM = inv("INV-M", "supX", [(dA1, 10_000_000), (dA2, 5_000_000)], "2026-06-01", "2026-07-01", tax=1_500_000)   # multi-DO
    iP1 = inv("INV-P1", "supX", [(dA3, 5_000_000)], "2026-09-01", "2026-10-31")                                   # DO ditagih parsial
    iP2 = inv("INV-P2", "supX", [(dA3, 3_000_000)], "2026-09-20", TODAY if TODAY >= "2026-09-20" else "2026-10-10")
    iY = inv("INV-Y", "supY", [(dY1, 4_000_000)], "2026-03-01", "2026-04-01")
    iL = inv("INV-L", "supX", [(dA4, 2_000_000)], "2026-05-01", "2026-05-31")
    iB = inv("INV-B", "supX", [(dB1, 500_000), (dA5, 500_000)], "2026-09-25", "2026-11-30")                     # multi-divisi
    iN = inv("INV-N", "supX", [(dA6, 1_000_000)], "2026-08-01", "2026-08-31")
    sql("UPDATE vendor_invoices SET doc = JSON_REMOVE(doc, '$.due_date') WHERE id = %s", [iN["id"]])           # legacy: tanpa jatuh tempo
    sql("UPDATE suppliers SET doc = JSON_SET(doc, '$.banks', JSON_EXTRACT(%s, '$')) WHERE id = %s",
        ['[{"bank_name": "BCA", "account_no": "1234567", "account_name": "PT Supplier X", "is_primary": true}]', M["supX"]["id"]])
    cn = conn()
    try:  # DP Supplier yang SUDAH dialokasikan ke INV-P1 (engine existing membaca vendor_invoice_dp_allocations)
        ins(cn, "vendor_invoice_dp_allocations", {"id": f"dpa-{u}", "tenant_id": tenant, "invoice_id": iP1["id"], "po_id": f"po-dp-{u}",
                                                  "dp_payment_id": f"dpp-{u}", "amount": 1_000_000, "created_at": "2026-09-05T03:00:00+00:00"})
    finally:
        cn.close()
    p1 = pay(iM, "2026-06-15", 5_000_000, "TRF-001")
    p2 = pay(iM, "2026-07-10", 4_000_000, "TRF-002")
    p3 = pay(iM, "2026-08-05", 3_000_000, "GIRO-003")
    sc, _ = call("POST", f"vendor-invoices/{iM['id']}/payments/{p2}/cancel", {"reason": "Transfer gagal"})
    check("fixture: pembayaran ke-2 INV-M dibatalkan (engine existing)", sc == 200, sc)
    pay(iP1, "2026-09-10", 4_000_000, "TRF-P1")      # 5 jt − DP 1 jt − 4 jt = lunas
    pay(iL, "2026-05-20", 2_000_000, "TRF-L")        # lunas sekaligus
    before = sorted((x["id"], x.get("paid_total"), x.get("dp_allocated_total"), x.get("amount")) for x in call("GET", "vendor-invoices")[1])
    n = lambda k: f"{k}-{u}"  # noqa: E731  (no invoice vendor)

    # ------------------------------------------------------------------ 1. Register Invoice
    sc, r = rc(call, "ap-invoice-register", "date_from=2026-06-01&date_to=2026-06-30")
    R = by_no(r, "invoice_no")
    m = R.get(n("INV-M"), {})
    check("R1 periode Juni: hanya INV-M; multi-DO 2 DO = SATU baris (total 15 jt tidak digandakan)",
          sc == 200 and list(R) == [n("INV-M")] and m.get("do_count") == 2 and near(r["totals"].get("amount"), 15_000_000), (sc, list(R), r.get("totals")))
    check("R2 DPP 13,5 jt + PPN 1,5 jt = Total 15 jt; No. DO kedua DO", (m.get("dpp"), m.get("tax_amount"), m.get("amount")) == (13_500_000, 1_500_000, 15_000_000)
          and dA1["no"] in m.get("do_nos", "") and dA2["no"] in m.get("do_nos", ""), m)
    check("R3 posisi per 30-06 (as-of): dibayar 5 jt, sisa 10 jt, 1 pembayaran, Dibayar Sebagian, Belum Jatuh Tempo",
          (m.get("paid"), m.get("remaining"), m.get("payment_count"), m.get("payment_status"), m.get("due_state")) ==
          (5_000_000, 10_000_000, 1, "Dibayar Sebagian", "Belum Jatuh Tempo"), m)
    sc, r = rc(call, "ap-invoice-register")
    R = by_no(r, "invoice_no")
    m, p1r, lr = R.get(n("INV-M"), {}), R.get(n("INV-P1"), {}), R.get(n("INV-L"), {})
    check("R4 kini: INV-M dibayar 8 jt (pembayaran batal tidak dihitung), sisa 7 jt, 2 pembayaran aktif",
          (m.get("paid"), m.get("remaining"), m.get("payment_count")) == (8_000_000, 7_000_000, 2), m)
    check("R5 INV-P1: DP dialokasikan 1 jt kolom terpisah, dibayar 4 jt, sisa 0 -> Lunas; INV-L lunas sekaligus",
          (p1r.get("dp_allocated"), p1r.get("paid"), p1r.get("remaining"), p1r.get("settlement_status"), lr.get("payment_status")) ==
          (1_000_000, 4_000_000, 0, "Lunas", "Lunas"), (p1r, lr))
    tot_amount = sum(x["amount"] for x in rows(r))
    check("R6 total register = Σ nilai header invoice (DO parsial 2 invoice tidak digandakan)", near(r["totals"].get("amount"), tot_amount)
          and near(R.get(n("INV-P1"), {}).get("amount", 0) + R.get(n("INV-P2"), {}).get("amount", 0), 8_000_000), r.get("totals"))
    sc, r = rc(call, "ap-invoice-register", "status=outstanding")
    check("R7 filter Belum Lunas: tanpa INV-P1 & INV-L", sc == 200 and n("INV-P1") not in by_no(r, "invoice_no") and n("INV-L") not in by_no(r, "invoice_no")
          and n("INV-M") in by_no(r, "invoice_no"), sorted(by_no(r, "invoice_no")))
    check("R8 drill register -> detail invoice existing", m and (rows(rc(call, "ap-invoice-register", f"invoice_id={iM['id']}")[1]) or [{}])[0].get("_drill", {}).get("to") == f"/invoice/{iM['id']}")

    # ------------------------------------------------------------------ 2. Monitoring Pembayaran
    sc, r = rc(call, "ap-payments", f"invoice_id={iM['id']}")
    P = rows(r)
    check("P1 riwayat INV-M: 3 pembayaran berurutan ke-1..3 (tanggal), referensi & catatan",
          sc == 200 and [x["seq"] for x in P] == [1, 2, 3] and [x["reference"] for x in P] == ["TRF-001", "TRF-002", "GIRO-003"]
          and P[0]["notes"] == "cat TRF-001", P)
    check("P2 pembayaran batal: status Dibatalkan, alasan, nilai efektif 0; kumulatif aktif 5 jt -> 8 jt",
          P and P[1]["status"] == "Dibatalkan" and P[1]["cancel_reason"] == "Transfer gagal" and P[1]["amount_effective"] == 0
          and P[1]["cumulative"] is None and P[2]["cumulative"] == 8_000_000 and near(r["totals"].get("amount_effective"), 8_000_000), P)
    check("P3 rekening/metode bayar tidak dicatat per pembayaran -> tidak ditampilkan (rekening master supplier pun tidak)",
          P and "supplier_bank" not in P[0] and not any("Rekening" in c["label"] for c in r["columns"])
          and "1234567" not in str(r), [c["label"] for c in r["columns"]])
    sc, r = rc(call, "ap-payments", f"invoice_id={iM['id']}&date_to=2026-07-31")
    P = rows(r)
    check("P4 per 31-07 (lampau): pembayaran ke-2 batal SESUDAH cut-off -> masih efektif 4 jt, kumulatif 9 jt",
          [x["seq"] for x in P] == [1, 2] and P[1]["amount_effective"] == 4_000_000 and P[1]["cumulative"] == 9_000_000, P)
    sc, r = rc(call, "ap-payments", "status=cancelled")
    check("P5 filter Dibatalkan: hanya pembayaran batal", [x["reference"] for x in rows(r)] == ["TRF-002"], rows(r))
    sc, r = rc(call, "ap-payments", f"invoice_id={iP1['id']}")
    check("P6 DP Supplier bukan pembayaran (INV-P1 hanya 1 baris pembayaran 4 jt)", [x["amount"] for x in rows(r)] == [4_000_000], rows(r))
    check("P7 drill pembayaran -> detail invoice", rows(r) and rows(r)[0]["_drill"]["to"] == f"/invoice/{iP1['id']}")

    # ------------------------------------------------------------------ 3. Outstanding (+ paritas Dashboard)
    def dash(dt):
        sc, d = call("GET", f"dashboard/control-center?date_from={dt[:8]}01&date_to={dt}")
        return ((d or {}).get("finance") or {}).get("payable") or {}

    def outst(qs):
        sc, r = rc(call, "ap-outstanding", qs)
        assert sc == 200, (sc, str(r)[:300])
        return r

    o = outst("date_to=2026-07-31")
    O = by_no(o, "invoice_no")
    check("O1 per 31-07: INV-M 6 jt (bayar ke-2 masih aktif saat itu) + INV-Y 4 jt; invoice sesudah cut-off tidak tampil",
          set(O) == {n("INV-M"), n("INV-Y")} and O[n("INV-M")]["remaining"] == 6_000_000 and near(o["totals"]["remaining"], 10_000_000), {k: v["remaining"] for k, v in O.items()})
    check("O2 paritas Dashboard 31-07: Outstanding = Sisa Hutang Supplier", near(o["totals"]["remaining"], dash("2026-07-31").get("value")), dash("2026-07-31"))
    o = outst("")
    O = by_no(o, "invoice_no")
    check("O3 kini: lunas (INV-L) & lunas via DP+bayar (INV-P1) tidak outstanding; total 16 jt",
          set(O) == {n(k) for k in ("INV-M", "INV-Y", "INV-P2", "INV-B", "INV-N")} and near(o["totals"]["remaining"], 16_000_000),
          {k: v["remaining"] for k, v in O.items()})
    check("O4 paritas Dashboard kini", near(o["totals"]["remaining"], dash(TODAY).get("value")), (o["totals"], dash(TODAY)))
    check("O5 invoice tanpa jatuh tempo -> Umur 'Tanpa Jatuh Tempo'; lewat jatuh tempo dihitung", O.get(n("INV-N"), {}).get("due_state") == "Tanpa Jatuh Tempo"
          and O.get(n("INV-M"), {}).get("overdue_amount") == 7_000_000, (O.get(n("INV-N")), O.get(n("INV-M"))))
    od = outst(f"division_id={divA['id']}")
    check("O6 filter Divisi A: invoice multi-divisi ikut (salah satu DO cocok), nilai tidak dipecah", n("INV-B") in by_no(od, "invoice_no")
          and by_no(od, "invoice_no")[n("INV-B")]["remaining"] == 1_000_000 and near(od["totals"]["remaining"], 16_000_000))
    check("O6b baris invoice multi-divisi berlabel 'Lebih dari satu divisi: ...' (nama kedua divisi), satu baris",
          str(by_no(od, "invoice_no")[n("INV-B")]["division"]).startswith("Lebih dari satu divisi: ") and divB["name"] in by_no(od, "invoice_no")[n("INV-B")]["division"]
          and sum(1 for x in rows(od) if x["invoice_no"] == n("INV-B")) == 1, by_no(od, "invoice_no")[n("INV-B")]["division"])
    ob = outst(f"division_id={divB['id']}")
    check("O7 filter Divisi B: hanya INV-B", set(by_no(ob, "invoice_no")) == {n("INV-B")}, sorted(by_no(ob, "invoice_no")))
    sc, s = rc(call, "ap-outstanding-summary", "group_by=division")
    S_ = by_no(s)
    check("O8 ringkasan per divisi: multi-divisi di kelompok tersendiri, Σ kelompok = Outstanding (tanpa hitung ganda)",
          sc == 200 and S_.get("Lebih dari satu divisi", {}).get("remaining") == 1_000_000 and near(s["totals"]["remaining"], 16_000_000),
          {k: v["remaining"] for k, v in S_.items()})
    def drill_same(res):
        out = []
        for g_ in rows(res):
            to = (g_.get("_drill") or {}).get("to", "")
            sc_, dr = rc(call, to.split("/report-center/")[1].split("?")[0], to.split("?", 1)[1] if "?" in to else "")
            out.append(sc_ == 200 and len(rows(dr)) == g_["invoice_count"] and near(dr["totals"]["remaining"], g_["remaining"]))
        return out
    same = drill_same(s)
    check("O8b drill tiap kelompok divisi -> Outstanding berisi invoice kelompok yang sama persis (jumlah & nilai)",
          len(same) == 2 and all(same), same)
    sc, s = rc(call, "ap-outstanding-summary", "group_by=supplier")
    S_ = by_no(s)
    check("O9 ringkasan per supplier: X 12 jt, Y 4 jt; drill -> Outstanding difilter supplier",
          S_.get("Supplier X", {}).get("remaining") == 12_000_000 and S_.get("Supplier Y", {}).get("remaining") == 4_000_000
          and f"supplier_id={M['supX']['id']}" in (S_.get("Supplier X", {}).get("_drill") or {}).get("to", ""), {k: v["remaining"] for k, v in S_.items()})
    sc, s = rc(call, "ap-outstanding-summary", "group_by=project")
    check("O10 ringkasan per proyek: Σ = Outstanding; drill tiap kelompok = isi kelompok", sc == 200 and near(s["totals"]["remaining"], 16_000_000)
          and all(drill_same(s)), s.get("totals"))
    sc, s = rc(call, "ap-outstanding-summary", "group_by=supplier")
    check("O11 drill ringkasan supplier = isi kelompok", all(drill_same(s)) and len(rows(s)) == 2)

    # ------------------------------------------------------------------ 4. Aging
    def ag(qs):
        sc, r = rc(call, "ap-aging", qs)
        assert sc == 200, (sc, str(r)[:300])
        return r

    a = ag("")
    A = by_no(a, "invoice_no")
    exp_m = "Lewat > 90 Hari"
    check("A1 kini: INV-M > 90 hari, INV-Y > 90 hari, INV-P2 jatuh tempo hari ini = Belum Jatuh Tempo, INV-N terpisah",
          A.get(n("INV-M"), {}).get("bucket") == exp_m and A.get(n("INV-Y"), {}).get("bucket") == exp_m
          and A.get(n("INV-P2"), {}).get("bucket") == "Belum Jatuh Tempo" and A.get(n("INV-P2"), {}).get("days_overdue") == 0
          and A.get(n("INV-N"), {}).get("bucket") == "Tanpa Jatuh Tempo" and A.get(n("INV-N"), {}).get("b_nodue") == 1_000_000,
          {k: (v["bucket"], v["days_overdue"]) for k, v in A.items()})
    t = a["totals"]
    check("A2 Σ bucket = Total Outstanding = Outstanding report", near(sum(t.get(f"b_{k}", 0) for k in ("nodue", "current", "d1_30", "d31_60", "d61_90", "d90")),
                                                                     t.get("remaining")) and near(t.get("remaining"), 16_000_000), t)
    a = ag("date_to=2026-07-31")
    A = by_no(a, "invoice_no")
    check("A3 per 31-07: INV-M 30 hari -> 1–30 (6 jt), INV-Y 121 hari -> > 90",
          (A.get(n("INV-M"), {}).get("days_overdue"), A.get(n("INV-M"), {}).get("b_d1_30"), A.get(n("INV-Y"), {}).get("bucket")) == (30, 6_000_000, exp_m), A)
    a = ag("date_to=2026-08-31")
    check("A4 per 31-08: INV-M 61 hari -> 61–90 (3 jt)", by_no(a, "invoice_no").get(n("INV-M"), {}).get("b_d61_90") == 3_000_000, by_no(a, "invoice_no").get(n("INV-M")))
    a = ag("date_to=2026-04-15")
    check("A5 per 15-04: INV-Y 14 hari -> 1–30", by_no(a, "invoice_no").get(n("INV-Y"), {}).get("bucket") == "Lewat 1–30 Hari", rows(a))
    a = ag("level=supplier")
    A = {x["supplier"]: x for x in rows(a)}
    check("A6 per supplier: X 12 jt (4 invoice), Y 4 jt; Σ = 16 jt", A.get("Supplier X", {}).get("remaining") == 12_000_000
          and A.get("Supplier X", {}).get("invoice_count") == 4 and near(a["totals"]["remaining"], 16_000_000), A)
    a = ag("bucket=nodue")
    check("A7 filter kategori Tanpa Jatuh Tempo", [x["invoice_no"] for x in rows(a)] == [n("INV-N")], rows(a))

    # ------------------------------------------------------------------ 5. Rekap per Supplier
    def recap(qs, key="ap-supplier-recap"):
        sc, r = rc(call, key, qs)
        assert sc == 200, (sc, str(r)[:300])
        return r

    def ident(x):
        return near(x["opening"] + x["invoiced"] - x["paid_period"] - x["dp_period"], x["closing"])

    r = recap("date_from=2026-07-01&date_to=2026-07-31")
    X = by_no(r).get("Supplier X", {})
    check("S1 Juli: Saldo Awal 10 jt, Invoice 0, Pembayaran Efektif 4 jt, DP 0, Saldo Akhir 6 jt",
          (X.get("opening"), X.get("invoiced"), X.get("paid_period"), X.get("dp_period"), X.get("closing")) == (10_000_000, 0, 4_000_000, 0, 6_000_000), X)
    check("S2 Juli: identitas Saldo Awal + Invoice − Bayar − DP = Saldo Akhir; Σ Saldo Akhir = Outstanding 31-07",
          all(ident(x) for x in rows(r)) and near(r["totals"]["closing"], outst("date_to=2026-07-31")["totals"]["remaining"]), rows(r))
    end = "2026-10-31"
    r = recap(f"date_from=2026-09-01&date_to={end}")
    X = by_no(r).get("Supplier X", {})
    check("S3 Sep–Okt: Saldo Awal 4 jt, Invoice 9 jt, Pembayaran Efektif 0 (bayar baru 4 jt − batal bayar lama 4 jt), DP 1 jt, Saldo Akhir 12 jt",
          (X.get("opening"), X.get("invoiced"), X.get("paid_period"), X.get("dp_period"), X.get("closing"), X.get("invoice_count")) ==
          (4_000_000, 9_000_000, 0, 1_000_000, 12_000_000, 3), X)
    check("S4 Saldo Akhir = Outstanding per cut-off yang sama; DP belum dialokasikan tidak mengurangi; identitas tiap supplier",
          all(ident(x) for x in rows(r)) and near(r["totals"]["closing"], outst(f"date_to={end}")["totals"]["remaining"]), (r["totals"], rows(r)))
    check("S5 total kumulatif: Total Tagihan − Total Bayar − Total DP = Saldo Akhir",
          near(X.get("total_billed", 0) - X.get("total_paid", 0) - X.get("total_dp", 0), X.get("closing")), X)
    r = recap("date_from=2026-05-01&date_to=2026-08-31", "ap-supplier-recap-monthly")
    xs = [x for x in rows(r) if x["no"] == "Supplier X"]
    check("S6 per bulan Mei–Agu: Saldo Akhir bulan = Saldo Awal bulan berikutnya; identitas tiap bulan",
          [x["month"] for x in xs] == ["2026-05", "2026-06", "2026-07", "2026-08"] and all(ident(x) for x in xs)
          and all(near(xs[i]["closing"], xs[i + 1]["opening"]) for i in range(len(xs) - 1)), [(x["month"], x["opening"], x["closing"]) for x in xs])
    check("S7 per bulan Agu: Saldo Akhir = Outstanding 31-08", near(sum(x["closing"] for x in rows(r) if x["month"] == "2026-08"),
                                                                       outst("date_to=2026-08-31")["totals"]["remaining"]))
    check("S8 drill rekap -> Outstanding difilter supplier", f"supplier_id={M['supX']['id']}" in (xs[0].get("_drill") or {}).get("to", ""), xs[:1])

    # ------------------------------------------------------------------ 6. Izin, cakupan divisi, hub
    wh = V.mk_user("warehouse")
    mg = V.mk_user("manager", divs=[divA["id"]])
    check("I1 tanpa invoice.view -> 403 seluruh laporan P5 (JSON & export)", all(rc(wh, k)[0] == 403 for k in AP_KEYS)
          and wh.S.get(f"{API}/report-center/ap-outstanding/export.xlsx").status_code == 403
          and rc(wh, "ap-outstanding")[1].get("detail") == "Anda tidak memiliki izin untuk melihat laporan Outstanding Hutang Vendor.")
    sc, cat = wh("GET", "report-center/catalog")
    check("I2 katalog tanpa invoice.view: tidak ada laporan hutang", sc == 200 and not next(g for g in cat["groups"] if g["key"] == "hutang")["reports"])
    sc, o = rc(mg, "ap-outstanding")
    check("I3 Manajer Divisi A (Lihat): invoice multi-divisi A+B tersembunyi (visibilitas existing), total 15 jt",
          sc == 200 and n("INV-B") not in by_no(o, "invoice_no") and near(o["totals"]["remaining"], 15_000_000) and o.get("price_visible") is True, (sc, o.get("totals")))
    check("I4 Manajer Divisi A memilih Divisi B -> 403", rc(mg, "ap-outstanding", f"division_id={divB['id']}")[0] == 403)
    sc, hb = rc(mg, "ap-payments", f"invoice_id={iB['id']}")
    sc2, ha = rc(call, "ap-payments", f"invoice_id={iB['id']}")
    lab = lambda res: next((f["value"] for f in res.get("filters_applied", []) if f["key"] == "invoice_id"), None)  # noqa: E731
    check("I4b riwayat invoice di luar cakupan: 0 baris & label filter tidak membocorkan nomor invoice; admin melihat label",
          sc == 200 and rows(hb) == [] and lab(hb) == iB["id"] and sc2 == 200 and n("INV-B") in str(lab(ha)), (lab(hb), lab(ha)))
    # Kebocoran lintas divisi: INV-B (DO Divisi A + B) di luar cakupan Manajer Divisi A di SELURUH laporan, pencarian, total,
    # pagination, export, filter invoice_id & drill. Pembayaran (lalu dibatalkan) pada INV-B agar Monitoring Pembayaran ikut diuji.
    pb = pay(iB, "2026-09-26", 100_000, "TRF-B")
    sc, _ = call("POST", f"vendor-invoices/{iB['id']}/payments/{pb}/cancel", {"reason": "uji cakupan"})
    nb = n("INV-B")
    leak = {}
    for k in AP_KEYS:
        qs = "date_from=2026-01-01&date_to=" + TODAY if k.startswith("ap-supplier-recap") else ""
        sc_a, ra = rc(call, k, qs)
        sc_m, rm = rc(mg, k, qs)
        leak[k] = (sc_a, sc_m, nb in str(ra), nb in str(rm), iB["id"] in str(rm))
    check("I3b admin melihat INV-B; Manajer Divisi A: INV-B (nomor & id) tidak ada di JSON ketujuh laporan P5",
          all(v[0] == 200 and v[1] == 200 and not v[3] and not v[4] for v in leak.values())
          and all(leak[k][2] for k in ("ap-invoice-register", "ap-payments", "ap-outstanding", "ap-aging")), leak)
    srch = {k: rc(mg, k, f"q={nb}") for k in ("ap-invoice-register", "ap-payments", "ap-outstanding", "ap-aging")}
    check("I3c pencarian nomor INV-B oleh Manajer Divisi A -> 0 baris, total 0 (admin menemukan)",
          all(sc_ == 200 and rows(r_) == [] and r_["total_rows"] == 0 and not any(v for v in r_["totals"].values()) for sc_, r_ in srch.values())
          and all(len(rows(rc(call, k, f"q={nb}")[1])) >= 1 for k in srch), {k: (v[0], v[1].get("total_rows")) for k, v in srch.items()})
    tm = {k: rc(mg, k, "date_from=2026-01-01&date_to=" + TODAY if k.startswith("ap-supplier-recap") else "")[1]["totals"].get(
        "closing" if k.startswith("ap-supplier-recap") else "remaining") for k in ("ap-outstanding", "ap-outstanding-summary", "ap-aging", "ap-supplier-recap")}
    check("I3d total Manajer Divisi A tanpa INV-B (15 jt) di Outstanding, Ringkasan, Aging, Rekap", all(near(v, 15_000_000) for v in tm.values()), tm)
    sc, pg = mg("GET", "report-center/ap-invoice-register?page_size=1&page=999")
    sc2, pa = call("GET", "report-center/ap-invoice-register?page_size=1")
    check("I3e pagination Manajer: total_rows/pages hanya invoice dalam cakupan (admin −1), halaman terakhir bukan INV-B",
          sc == 200 and pg["total_rows"] == pa["total_rows"] - 1 and pg["pages"] == pg["total_rows"] and nb not in str(pg), (pg.get("total_rows"), pa.get("total_rows")))
    ex = {}
    for k in ("ap-invoice-register", "ap-payments", "ap-outstanding", "ap-aging"):
        rx = mg.S.get(f"{API}/report-center/{k}/export.xlsx")
        txt = " ".join(str(c) for ws in load_workbook(io.BytesIO(rx.content), read_only=True).worksheets for v in ws.iter_rows(values_only=True) for c in v if c) if rx.ok else ""
        ps, ptxt = pdf(mg.S, k, "")
        ex[k] = (rx.status_code, ps, nb in txt, nb in ptxt.replace(" ", ""))
    check("I3f export Excel & PDF Manajer Divisi A: 200 tanpa INV-B", all(v[0] == 200 and v[1] == 200 and not v[2] and not v[3] for v in ex.values()), ex)
    sc, hr = rc(mg, "ap-invoice-register", f"invoice_id={iB['id']}")
    rx = mg.S.get(f"{API}/report-center/ap-payments/export.xlsx?invoice_id={iB['id']}")
    xtxt = " ".join(str(c) for v in load_workbook(io.BytesIO(rx.content), read_only=True).worksheets[0].iter_rows(values_only=True) for c in v if c) if rx.ok else ""
    check("I3g akses langsung invoice_id INV-B (JSON register & export Excel pembayaran): 0 baris, label filter = id mentah (tanpa nomor)",
          sc == 200 and rows(hr) == [] and lab(hr) == iB["id"] and rx.status_code == 200 and nb not in xtxt and "TRF-B" not in xtxt, (lab(hr), rx.status_code))
    check("I3h drill detail invoice di luar cakupan -> 403 (detail & pembayaran existing)",
          mg("GET", f"vendor-invoices/{iB['id']}")[0] == 403, mg("GET", f"vendor-invoices/{iB['id']}")[0])
    nv = V.mk_user("purchasing", overrides={"invoice.view": "deny"})  # Semua Divisi + export
    deny = {k: (rc(nv, k)[0], nv.S.get(f"{API}/report-center/{k}/export.xlsx").status_code,
                nv.S.get(f"{API}/report-center/{k}/export.pdf").status_code) for k in AP_KEYS}
    check("I6 user ber-izin export & lintas divisi tetapi invoice.view ditolak -> 403 JSON, Excel, PDF seluruh laporan P5 (bukan 200 kosong)",
          all(v == (403, 403, 403) for v in deny.values()), deny)
    sc, cat = nv("GET", "report-center/catalog")
    check("I6b katalog user tanpa invoice.view: tidak ada laporan hutang & tautan modul invoice", sc == 200
          and not next(g for g in cat["groups"] if g["key"] == "hutang")["reports"], sc)
    sc, cat = call("GET", "report-center/catalog")
    g = next(x for x in cat["groups"] if x["key"] == "hutang")
    check("I5 hub Invoice & Hutang: 7 laporan P5, tautan Status Penagihan DO & DP Supplier, tanpa placeholder; 5 kategori tetap",
          [x["key"] for x in cat["groups"]] == ["persediaan", "procurement", "warehouse", "spk", "hutang"]
          and [r["key"] for r in g["reports"]] == AP_KEYS
          and [lk["to"] for lk in g["links"]] == ["/invoice?tab=do", "/dp-supplier"] and not g["planned"], g)

    # ------------------------------------------------------------------ 7. Paritas JSON = Excel = PDF
    for key, qs, first, money in (("ap-outstanding", "date_to=2026-07-31", "No. Internal", "remaining"),
                                  ("ap-invoice-register", "", "No. Internal", "amount"),
                                  ("ap-supplier-recap", f"date_from=2026-09-01&date_to={end}", "Supplier", "closing"),
                                  ("ap-payments", "", "No. Internal Invoice", "amount_effective")):
        sc, j = rc(call, key, qs)
        xs_, hdr, body, tot = xl(T.S, key, qs, first)
        labels = [c["label"] for c in j["columns"]]
        mi = [c["key"] for c in j["columns"]].index(money)
        check(f"X {key}: Excel header = kolom JSON, baris = JSON, TOTAL = JSON", xs_ == 200 and hdr == labels and len(body) == len(rows(j))
              and near(tot[mi] if tot else None, j["totals"].get(money)), (xs_, hdr[:4], len(body), len(rows(j))))
        ps, txt = pdf(T.S, key, qs)
        check(f"X {key}: PDF 200 memuat judul & nomor dokumen", ps == 200 and j["meta"]["title"].split(" — ")[0] in txt
              and all(str(r.get("no") or "")[:8] in txt for r in rows(j)[:3]), (ps, txt[:200]))

    # ------------------------------------------------------------------ 8. Read-only
    after = sorted((x["id"], x.get("paid_total"), x.get("dp_allocated_total"), x.get("amount")) for x in call("GET", "vendor-invoices")[1])
    check("RO laporan P5 read-only: invoice/pembayaran/DP tidak berubah", before == after)

    passed = sum(1 for _, ok in T.RESULTS if ok)
    print(f"\n{passed}/{len(T.RESULTS)} passed")
    return 0 if passed == len(T.RESULTS) else 1


AP_KEYS = ["ap-invoice-register", "ap-payments", "ap-outstanding", "ap-outstanding-summary", "ap-aging", "ap-supplier-recap",
           "ap-supplier-recap-monthly"]


if __name__ == "__main__":
    sys.exit(main())
