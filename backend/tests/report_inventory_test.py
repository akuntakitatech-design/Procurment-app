"""P1 Reporting Persediaan & Nilai Persediaan (Pusat Laporan) — paritas stok, nilai, dan laporan historis.

Tenant QA sementara (/saas/register) pada backend test terisolasi (DB *itest*) — BUKAN production. Membuktikan:
  P. Posisi Stok: qty = item_warehouse (hari ini) / Σ stock_ledger (historis); nilai = Valuation Summary / Dashboard /
     Inventory per tanggal cut-off (WIB) yang sama; scope divisi/gudang; redaksi harga.
  N. Ringkasan Nilai: total per Gudang/Divisi/Kategori = Dashboard; tanpa izin harga -> 403 & tersembunyi di katalog.
  K. Kartu Stok Qty/Nilai: Saldo Awal + Masuk − Keluar (+ Selisih) = Saldo Akhir; saldo akhir = posisi per tanggal;
     barang wajib; barang/gudang di luar cakupan -> 403.
  M. Mutasi: identitas per baris (qty & nilai), Transfer/Loan internal Masuk = Keluar (tidak berganda) pada cakupan
     seluruh gudang; Σ saldo awal/akhir = posisi per tanggal; Σ Pemakaian MI = Rekap HPP MI.
  H. HPP MI = Σ nilai keluar MI valuation ledger (net reversal) per Proyek/Unit/Barang; filter proyek.
  X. Min/Max & Reorder: status & saran reorder sesuai rule yang disetujui.
  E. Export Excel/PDF = JSON (seluruh baris, total), redaksi harga di file, notice barang wajib -> 400.
  Z. Read-only: pool stok/nilai tidak berubah.
"""
import io
import os
import sys
import uuid
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

import receipt_control_test as T
from master_item_stock_test import mk_user

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import reporting.reports_inventory as RI  # noqa: E402
import stock_summary as SS  # noqa: E402

call, check, API = T.call, T.check, T.API
TODAY = datetime.now(ZoneInfo("Asia/Jakarta")).date().isoformat()
EPS = 0.01


def near(a, b, eps=EPS):
    return abs(float(a or 0) - float(b or 0)) <= eps


def rc(caller, key, qs=""):
    sc, r = caller("GET", f"report-center/{key}?page_size=500&{qs}")
    return sc, r


def get_bin(u, path):
    s = u.S if u is not None else T.S
    r = s.get(f"{API}/{path}")
    return r.status_code, r.content, r.headers


def dash_value(caller, date_to, division_id=""):
    qs = f"date_from=2026-01-01&date_to={date_to}" + (f"&division_id={division_id}" if division_id else "")
    sc, d = caller("GET", f"dashboard/control-center?{qs}")
    return (d.get("inventory") or {}).get("inventory_value") if sc == 200 else None


def adj(wh, div, item, qty, cost, date):
    line = {"item_id": item, "adjustment": qty, "reason": "QA P1"}
    if qty > 0:
        line["approved_unit_cost"] = cost
    return call("POST", "adjustments", {"date": date, "warehouse_id": wh, "division_id": div,
                                        "adj_type": "Opening" if qty > 0 else "Koreksi", "reason": "QA P1",
                                        "notes": "P1 REPORT TEST", "lines": [line]})


def main():
    db = os.environ.get("TEST_DATABASE_URL") or ""
    dbn = db.rsplit("/", 1)[-1].split("?")[0]
    if not ("itest" in dbn or dbn.endswith("_test")):
        print("Menolak berjalan: test ini membuat tenant QA dan hanya boleh dijalankan pada backend + DB test terisolasi "
              "(scripts/run_regression_itest.sh, DB *itest*), bukan preview/sandbox/production.")
        sys.exit(2)
    M = T.setup()
    u = uuid.uuid4().hex[:5]
    dA = M["div"]
    dB = call("POST", "master/divisions", {"code": f"PB{u}", "name": f"Divisi PB {u}"}, 200)[1]
    cat2 = call("POST", "master/item_categories", {"code": f"K2{u}", "name": "Kategori Dua"}, 200)[1]
    unit = call("POST", "master/units", {"code": f"UN{u}", "name": "Excavator P1", "is_active": True}, 200)[1]
    wA1 = call("POST", "master/warehouses", {"code": f"P1{u}", "name": f"Gudang PA1 {u}", "division_id": dA["id"], "is_active": True}, 200)[1]
    wA2 = call("POST", "master/warehouses", {"code": f"P2{u}", "name": f"Gudang PA2 {u}", "division_id": dA["id"], "is_active": True}, 200)[1]
    wB = call("POST", "master/warehouses", {"code": f"PB{u}", "name": f"Gudang PB {u}", "division_id": dB["id"], "is_active": True}, 200)[1]
    ref = {"unit": "PCS", "base_uom_id": M["uom"]["id"], "is_active": True}
    X = call("POST", "master/items", {"code": f"PX{u}", "name": "Barang PX", "division_id": dA["id"], "category_id": M["cat"]["id"], **ref}, 200)[1]
    Y = call("POST", "master/items", {"code": f"PY{u}", "name": "Barang PY", "division_id": dA["id"], "category_id": cat2["id"], **ref}, 200)[1]
    Z = call("POST", "master/items", {"code": f"PZ{u}", "name": "Barang PZ", "division_id": dB["id"], "category_id": M["cat"]["id"], **ref}, 200)[1]
    call("PUT", "settings/mi", {"mi_mode": "mro_plus_direct"})

    for wh, dv, it, q, c, dt in ((wA1, dA, X, 10, 1000, "2026-08-01"), (wA1, dA, Y, 20, 500, "2026-08-01"),
                                 (wB, dB, Z, 5, 2000, "2026-08-01")):
        sc, r = adj(wh["id"], dv["id"], it["id"], q, c, dt)
        check(f"setup: saldo awal {it['name']} {q} @ {dt}", sc == 200, (sc, r))

    def tl(path, qty, date, frm, to):
        return call("POST", path, {"date": date, "division_id": dA["id"], "from_warehouse_id": frm["id"], "to_warehouse_id": to["id"],
                                   "notes": "P1", "lines": [{"item_id": X["id"], "qty": qty, "uom_id": M["uom"]["id"],
                                                              "from_warehouse_id": frm["id"], "to_warehouse_id": to["id"]}]})
    sc, trf = tl("transfers", 4, "2026-08-15", wA1, wA2)
    check("setup: transfer PX PA1 -> PA2 4 @ 08-15", sc == 200, (sc, trf))
    sc, loan = tl("loans", 2, "2026-08-20", wA1, wA2)
    check("setup: loan PX PA1 -> PA2 2 @ 08-20", sc == 200, (sc, loan))
    sc, ld = call("GET", f"loans/{loan.get('id')}")
    ll = (ld.get("lines") or [{}])[0]
    sc, ret = call("POST", f"loans/{loan.get('id')}/return", {"date": "2026-09-05", "lines": [{"loan_line_id": ll.get("id"), "qty": 1}]})
    check("setup: loan return PX 1 @ 09-05", sc == 200, (sc, ret))
    sc, mi = call("POST", "mi", {"date": "2026-09-10", "division_id": dA["id"], "default_warehouse_id": wA2["id"], "receiver": "QA",
                                 "department": "Workshop", "source_type": "Direct", "notes": "P1 HPP",
                                 "lines": [{"item_id": X["id"], "qty": 3, "unit": "PCS", "warehouse_id": wA2["id"],
                                            "project_id": M["pa"]["id"], "unit_id": unit["id"]}]})
    check("setup: MI Direct PX 3 @ PA2 (Project A, Unit) 09-10", sc == 200, (sc, mi))
    sc, r = adj(wA1["id"], dA["id"], X["id"], -1, None, "2026-09-12")
    check("setup: koreksi PX PA1 -1 @ 09-12", sc == 200, (sc, r))
    sc, r = call("POST", "item-warehouse", {"item_id": Y["id"], "warehouse_id": wA1["id"], "min_stock": 25, "max_stock": 40,
                                            "current_stock": 20})
    check("setup: min/max PY @ PA1 = 25/40", sc == 200, (sc, r))

    A = mk_user("manager", divs=[dA["id"]], overrides={"export": "allow"})
    NP = mk_user("warehouse", divs=[dA["id"]], overrides={"export": "allow"})
    sc, iw0 = call("GET", "item-warehouse")
    snap0 = sorted((x.get("item_id"), x.get("warehouse_id"), x.get("current_stock"), x.get("avg_cost"), x.get("total_value")) for x in iw0 or [])
    pool = {(x["item_id"], x["warehouse_id"]): x for x in iw0 or []}

    # ------------------------------------------------------------------ P. posisi stok
    sc, cat = call("GET", "report-center/catalog")
    keys = [r["key"] for r in cat["groups"][0]["reports"]] if sc == 200 else []
    check("P0 katalog Persediaan: 7 laporan P1 + Riwayat Pergerakan Stok (P2b)", keys == ["posisi-stok", "ringkasan-nilai", "kartu-stok-qty",
          "riwayat-pergerakan-stok", "kartu-stok-nilai", "mutasi-persediaan", "hpp-mi", "min-max-reorder"], keys)
    sc, vs = call("GET", "reports/valuation-summary")
    sc, p = rc(call, "posisi-stok")
    check("P1 posisi hari ini: Σ nilai = Valuation Summary = Dashboard", sc == 200 and near(p["totals"]["value"], vs["total_value"])
          and near(p["totals"]["value"], dash_value(call, TODAY)), (p.get("totals"), vs.get("total_value"), dash_value(call, TODAY)))
    by = {(r["item_code"], r["warehouse"]): r for r in p.get("rows", [])}
    # PX PA1 = 10 − 4 (transfer) − 2 (loan) + 1 (return) − 1 (koreksi) = 4; PA2 = 4 + 2 − 1 − 3 (MI) = 2
    exp_now = {(X["code"], wA1["name"]): 4, (X["code"], wA2["name"]): 2, (Y["code"], wA1["name"]): 20, (Z["code"], wB["name"]): 5}
    wn = {wA1["id"]: wA1["name"], wA2["id"]: wA2["name"], wB["id"]: wB["name"]}
    ic = {X["id"]: X["code"], Y["id"]: Y["code"], Z["id"]: Z["code"]}
    check("P2 qty per barang × gudang hari ini = item_warehouse (PX PA1 4, PA2 2; PY 20; PZ 5)",
          all(near(by.get(k, {}).get("qty"), v) for k, v in exp_now.items())
          and all(near(by.get((ic[i], wn[w]), {}).get("qty"), x.get("current_stock")) for (i, w), x in pool.items() if i in ic and w in wn),
          {k: by.get(k, {}).get("qty") for k in exp_now})
    for d, exp in (("2026-08-31", {(X["code"], wA1["name"]): 4, (X["code"], wA2["name"]): 6}),
                   ("2026-09-10", {(X["code"], wA1["name"]): 5, (X["code"], wA2["name"]): 2})):
        sc, ph = rc(call, "posisi-stok", f"as_of={d}")
        bh = {(r["item_code"], r["warehouse"]): r for r in ph.get("rows", [])}
        check(f"P3 posisi historis {d}: qty = Σ mutasi s/d cut-off & nilai = Dashboard per tanggal",
              sc == 200 and all(near(bh.get(k, {}).get("qty"), v) for k, v in exp.items())
              and near(ph["totals"]["value"], dash_value(call, d)), ({k: bh.get(k, {}).get("qty") for k in exp}, ph.get("totals"), dash_value(call, d)))
    sc, pa = rc(A, "posisi-stok")
    check("P4 user Divisi A: tanpa gudang/barang Divisi B; Σ nilai = Dashboard user A",
          sc == 200 and wB["name"] not in {r["warehouse"] for r in pa["rows"]} and near(pa["totals"]["value"], dash_value(A, TODAY)),
          (pa.get("totals"), dash_value(A, TODAY)))
    sc, _ = rc(A, "posisi-stok", f"warehouse_id={wB['id']}")
    check("P5 user A filter gudang Divisi B -> 403", sc == 403, sc)
    sc, pn = rc(NP, "posisi-stok")
    ck = {c["key"] for c in pn.get("columns", [])}
    check("P6 tanpa izin harga: posisi tampil tanpa kolom/nilai harga", sc == 200 and not ({"avg_cost", "value"} & ck)
          and all(not ({"avg_cost", "value"} & set(r)) for r in pn["rows"]) and "value" not in pn["totals"], ck)
    sc, pw = rc(call, "posisi-stok", f"warehouse_id={wA2['id']}&category_id={M['cat']['id']}")
    check("P7 filter gudang + kategori", sc == 200 and {r["warehouse"] for r in pw["rows"]} == {wA2["name"]}
          and {r["item_code"] for r in pw["rows"]} == {X["code"]}, [(r["item_code"], r["warehouse"]) for r in pw.get("rows", [])])

    # ------------------------------------------------------------------ N. ringkasan nilai
    for gb in ("warehouse", "division", "category"):
        sc, n = rc(call, "ringkasan-nilai", f"group_by={gb}")
        check(f"N1 ringkasan per {gb}: total = posisi = Dashboard; porsi 100%", sc == 200 and near(n["totals"]["value"], p["totals"]["value"])
              and near(n["totals"]["share"], 100, 0.05), n.get("totals"))
    sc, nd = rc(call, "ringkasan-nilai", "group_by=division&as_of=2026-08-31")
    check("N2 ringkasan historis 08-31 = Dashboard per tanggal", sc == 200 and near(nd["totals"]["value"], dash_value(call, "2026-08-31")), nd.get("totals"))
    sc, na = rc(A, "ringkasan-nilai", "group_by=division")
    check("N3 user A: hanya Divisi A", sc == 200 and [r["group_name"] for r in na["rows"]] == [dA["name"]], na.get("rows"))
    sc, cn = NP("GET", "report-center/catalog")
    hidden = {r["key"] for g in cn.get("groups", []) for r in g["reports"]}
    for key in ("ringkasan-nilai", "kartu-stok-nilai", "hpp-mi"):
        sc, _ = rc(NP, key, f"item_id={X['id']}")
        check(f"N4 tanpa izin harga: {key} -> 403 & tidak tampil di katalog", sc == 403 and key not in hidden, sc)
    for key in ("posisi-stok", "kartu-stok-qty", "mutasi-persediaan", "min-max-reorder"):
        check(f"N5 tanpa izin harga: {key} tetap tampil di katalog", key in hidden, sorted(hidden))

    # ------------------------------------------------------------------ K. kartu stok
    sc, kq0 = rc(call, "kartu-stok-qty")
    check("K1 kartu stok tanpa barang: 0 baris + notice wajib pilih barang", sc == 200 and kq0["total_rows"] == 0
          and "Barang" in kq0["meta"]["notice"], kq0.get("meta"))
    sc, _, _ = get_bin(None, "report-center/kartu-stok-qty/export.xlsx")
    check("K2 export kartu stok tanpa barang -> 400", sc == 400, sc)
    sc, kq = rc(call, "kartu-stok-qty", f"item_id={X['id']}&date_from=2026-08-10&date_to=2026-09-30")
    rows = kq.get("rows", [])
    op, cl, body = rows[0], rows[-1], rows[1:-1]
    check("K3 kartu qty PX: Saldo Awal 10 (sebelum 08-10)", op.get("_kind") == "opening" and near(op["balance"], 10), op)
    check("K4 Saldo Awal + Masuk − Keluar = Saldo Akhir = posisi 09-30 (6)",
          near(op["balance"] + kq["totals"]["qty_in"] - kq["totals"]["qty_out"], cl["balance"]) and near(cl["balance"], 6)
          and cl.get("_kind") == "closing" and not cl.get("note"), (op, kq.get("totals"), cl))
    check("K5 saldo berjalan kronologis (tanggal naik) & baris terakhir = saldo akhir",
          [r["txn_date"] for r in body] == sorted(r["txn_date"] for r in body) and near(body[-1]["balance"], cl["balance"]), [r["txn_date"] for r in body])
    tin = sum(r["qty_in"] for r in body if r["category"] == "Transfer/Loan Masuk")
    tout = sum(r["qty_out"] for r in body if r["category"] == "Transfer/Loan Keluar")
    check("K6 seluruh gudang: Transfer/Loan Masuk = Keluar (internal tidak berganda)", near(tin, tout) and tin > 0, (tin, tout))
    sc, kw = rc(call, "kartu-stok-qty", f"item_id={X['id']}&warehouse_id={wA2['id']}")
    check("K7 kartu qty PX gudang PA2 (semua periode): saldo akhir = item_warehouse PA2",
          sc == 200 and near(kw["rows"][-1]["balance"], pool[(X["id"], wA2["id"])]["current_stock"]), kw.get("rows", [])[-1:])
    sc, kn = rc(call, "kartu-stok-nilai", f"item_id={X['id']}&date_from=2026-08-10&date_to=2026-09-30")
    nr = kn.get("rows", [])
    t = kn.get("totals", {})
    sc2, ph = rc(call, "posisi-stok", f"as_of=2026-09-30&item_id={X['id']}")
    check("K8 kartu nilai: Saldo Awal + Nilai Masuk − Keluar + Selisih = Saldo Akhir = posisi nilai 09-30",
          sc == 200 and near(nr[0]["balance_value"] + t["value_in"] - t["value_out"] + t["value_diff"], nr[-1]["balance_value"])
          and near(nr[-1]["balance_value"], ph["totals"]["value"]) and near(nr[0]["balance_value"], 10000), (nr[:1], t, nr[-1:], ph.get("totals")))
    check("K9 kartu nilai: selisih nilai = 0 (rantai valuation ledger utuh)", near(t["value_diff"], 0), t)
    mi_rows = [r for r in nr if r["doc_type"] == "MI"]
    check("K10 kartu nilai: MI dinilai pada rata-rata MWA ledger (3 × 1.000)", len(mi_rows) == 1 and near(mi_rows[0]["value_out"], 3000), mi_rows)
    sc, _ = rc(A, "kartu-stok-qty", f"item_id={Z['id']}")
    check("K11 user A kartu barang Divisi B -> 403", sc == 403, sc)

    # ------------------------------------------------------------------ M. mutasi
    sc, mu = rc(call, "mutasi-persediaan", "date_from=2026-08-10&date_to=2026-09-30")
    tot = mu.get("totals", {})
    ok_rows = all(near(r["q_open"] + r["q_do"] - r["q_mi"] + r["q_tl_in"] - r["q_tl_out"] + r["q_adj"] + r["q_opn"] + r["q_opening"]
                       + r["q_other"] + r["q_diff"], r["q_close"])
                  and near(r["v_open"] + r["v_do"] - r["v_mi"] + r["v_tl_in"] - r["v_tl_out"] + r["v_adj"] + r["v_opn"] + r["v_opening"]
                           + r["v_other"] + r["v_diff"], r["v_close"]) for r in mu.get("rows", []))
    check("M1 mutasi: identitas per baris qty & nilai (Saldo Awal + Mutasi + Selisih = Saldo Akhir)", sc == 200 and ok_rows and mu["rows"], mu.get("rows", [])[:2])
    check("M2 seluruh gudang: Transfer/Loan Masuk = Keluar (qty & nilai) → saldo perusahaan tidak berganda",
          near(tot["q_tl_in"], tot["q_tl_out"]) and near(tot["v_tl_in"], tot["v_tl_out"]) and tot["q_tl_in"] > 0, tot)
    sc, p0 = rc(call, "posisi-stok", "as_of=2026-08-09")
    sc, p1 = rc(call, "posisi-stok", "as_of=2026-09-30")
    check("M3 Σ saldo awal/akhir mutasi = posisi per 08-09 / 09-30 (qty & nilai)",
          near(tot["v_open"], p0["totals"]["value"]) and near(tot["v_close"], p1["totals"]["value"])
          and near(tot["q_open"], p0["totals"]["qty"]) and near(tot["q_close"], p1["totals"]["qty"]), (tot, p0.get("totals"), p1.get("totals")))
    check("M4 selisih qty & nilai = 0 dan Pemakaian MI = 3 / 3.000", near(tot["q_diff"], 0) and near(tot["v_diff"], 0)
          and near(tot["q_mi"], 3) and near(tot["v_mi"], 3000), tot)
    sc, mw = rc(call, "mutasi-persediaan", f"date_from=2026-08-10&date_to=2026-09-30&warehouse_id={wA2['id']}")
    rx = next((r for r in mw.get("rows", []) if r["item_code"] == X["code"]), {})
    check("M5 cakupan 1 gudang (PA2): Transfer/Loan Masuk 6, Keluar 1 (return), MI 3, saldo akhir 2", near(rx.get("q_tl_in"), 6)
          and near(rx.get("q_tl_out"), 1) and near(rx.get("q_mi"), 3) and near(rx.get("q_close"), 2), rx)
    sc, mn = rc(NP, "mutasi-persediaan", "date_from=2026-08-10")
    check("M6 tanpa izin harga: mutasi tanpa kolom nilai", sc == 200 and not any(c["key"].startswith("v_") for c in mn["columns"]), [c["key"] for c in mn.get("columns", [])][:8])

    # ------------------------------------------------------------------ H. HPP MI
    sc, lg = call("GET", "reports/valuation-ledger?doc_type=MI&date_from=2026-08-10&date_to=2026-09-30")
    ledger_mi = sum(float(r.get("value_out") or 0) - float(r.get("value_in") or 0) for r in lg.get("rows", []))
    for gb in ("project", "unit", "item"):
        sc, h = rc(call, "hpp-mi", f"date_from=2026-08-10&date_to=2026-09-30&group_by={gb}")
        check(f"H1 HPP MI per {gb} = Σ nilai keluar MI ledger = Mutasi Pemakaian MI (3.000)", sc == 200 and near(h["totals"]["hpp"], ledger_mi)
              and near(h["totals"]["hpp"], tot["v_mi"]) and near(ledger_mi, 3000), (h.get("totals"), ledger_mi))
    sc, hp = rc(call, "hpp-mi", "group_by=project")
    check("H2 HPP per proyek: Project A = 3.000", sc == 200 and [(r["group_name"], r["hpp"]) for r in hp["rows"]] == [(M["pa"]["name"], 3000.0)], hp.get("rows"))
    sc, hf = rc(call, "hpp-mi", f"project_id={M['pb']['id']}")
    check("H3 filter Project B: tanpa HPP", sc == 200 and hf["total_rows"] == 0, hf.get("rows"))
    sc, hs = rc(call, "hpp-mi", "group_by=spk")
    check("H4 HPP per SPK: MI tanpa alokasi SPK -> (Tanpa SPK), total sama", sc == 200 and near(hs["totals"]["hpp"], 3000)
          and hs["rows"][0]["group_name"] == "(Tanpa SPK)", hs.get("rows"))

    # ------------------------------------------------------------------ X. min/max
    sc, mm = rc(call, "min-max-reorder")
    py = next((r for r in mm.get("rows", []) if r["item_code"] == Y["code"]), {})
    check("X1 PY 20 vs Min 25/Max 40: Di bawah Min, saran reorder = 40 − 20 = 20", py.get("status") == "Di bawah Min" and near(py.get("reorder"), 20), py)
    sc, mr = rc(call, "min-max-reorder", "need=reorder")
    check("X2 filter perlu reorder: hanya baris saran > 0", sc == 200 and mr["rows"] and all(r["reorder"] > 0 for r in mr["rows"]), mr.get("rows"))
    check("X3 rule reorder: Max kosong -> Min − stok; Normal -> 0; Kosong -> Max",
          RI.reorder_qty(3, 5, 0) == 2 and RI.reorder_qty(10, 5, 20) == 0 and RI.reorder_qty(0, 5, 20) == 20, "")

    # ------------------------------------------------------------------ E. export
    sc, content, hdr = get_bin(None, "report-center/mutasi-persediaan/export.xlsx?date_from=2026-08-10&date_to=2026-09-30")
    ws = load_workbook(io.BytesIO(content), read_only=True).worksheets[0] if sc == 200 else None
    vals = [list(r) for r in ws.iter_rows(values_only=True)] if ws else []
    hi = next((i for i, r in enumerate(vals) if r and r[0] == "Kode Barang"), None)
    body = [r for r in vals[hi + 1:] if r and r[0] and r[0] != "TOTAL"] if hi is not None else []
    check("E1 Excel mutasi = seluruh baris JSON + header X-Report-Rows", sc == 200 and len(body) == mu["total_rows"]
          and hdr.get("X-Report-Rows") == str(mu["total_rows"]), (sc, len(body), mu.get("total_rows")))
    sc, pdf, _ = get_bin(None, f"report-center/kartu-stok-nilai/export.pdf?item_id={X['id']}&date_from=2026-08-10&date_to=2026-09-30")
    txt = "".join(" ".join(pg.extract_text() or "" for pg in PdfReader(io.BytesIO(pdf)).pages).split()) if sc == 200 else ""
    check("E2 PDF kartu nilai valid (Saldo Awal/Akhir, TOTAL, Hal)", sc == 200 and pdf[:4] == b"%PDF" and "SaldoAwal" in txt
          and "SaldoAkhir" in txt and "TOTAL" in txt and "Hal" in txt, sc)
    sc, content, _ = get_bin(NP, "report-center/posisi-stok/export.xlsx")
    hdrs = []
    if sc == 200:
        wv = [list(r) for r in load_workbook(io.BytesIO(content), read_only=True).worksheets[0].iter_rows(values_only=True)]
        hdrs = next((r for r in wv if r and r[0] == "Kode Barang"), [])
    check("E3 Excel posisi tanpa izin harga: tanpa kolom Rata-rata Biaya/Nilai", sc == 200 and "Nilai Persediaan" not in hdrs
          and "Rata-rata Biaya" not in hdrs and "Qty" in hdrs, hdrs)
    sc, pdf2, _ = get_bin(NP, "report-center/mutasi-persediaan/export.pdf?date_from=2026-08-10")
    t2 = "".join(" ".join(pg.extract_text() or "" for pg in PdfReader(io.BytesIO(pdf2)).pages).split()) if sc == 200 else ""
    check("E4 PDF mutasi tanpa izin harga: tanpa label nilai", sc == 200 and "NilaiSaldo" not in t2 and "QtySaldoAwal" in t2, sc)

    # ------------------------------------------------------------------ W. cut-off WIB & Z. read-only
    check("W1 tanggal efektif WIB: 2026-09-30T18:30Z -> 2026-10-01 (laporan & Inventory sama)",
          RI.eff_day({"txn_at": "2026-09-30T18:30:00+00:00"}) == "2026-10-01" == SS._txn_day({"txn_at": "2026-09-30T18:30:00+00:00"})
          and SS._txn_day({"txn_at": "2026-09-30"}) == "2026-09-30", "")
    check("W2 kategori mutasi: Reversal X = kategori X", RI.cat_of("Reversal Transfer Out") == "tl_out" and RI.cat_of("Loan Return In") == "tl_in"
          and RI.cat_of("Opening Valuation") == "opening" and RI.cat_of("???") == "other", "")
    sc, iw1 = call("GET", "item-warehouse")
    snap1 = sorted((x.get("item_id"), x.get("warehouse_id"), x.get("current_stock"), x.get("avg_cost"), x.get("total_value")) for x in iw1 or [])
    check("Z1 read-only: stok/avg/nilai pool tidak berubah setelah laporan & export", snap0 == snap1, "")

    failed = [nm for nm, ok in T.RESULTS if not ok]
    print(f"\n{len(T.RESULTS) - len(failed)}/{len(T.RESULTS)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
