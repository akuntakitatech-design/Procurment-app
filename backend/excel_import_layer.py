"""Excel import/export templates for master data and operational transactions.

Design goals:
- one downloadable XLSX template per dataset;
- human-readable reference codes, never internal IDs;
- row-level validation with clear Excel row numbers;
- transactions are grouped by `batch_ref` (one header may have many item rows);
- normal transaction endpoints are reused so numbering, audit, stock and allocation rules stay authoritative.
"""
from collections import OrderedDict, defaultdict
from copy import deepcopy
from io import BytesIO
from typing import Any

from fastapi import Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


MASTER_DATASETS = {
    "items": {
        "label": "Master Barang",
        "columns": [
            "code", "name", "alias", "category_code", "division_code", "base_uom_code",
            "brand", "part_number", "spec", "is_active",
        ],
        "required": ["code", "name", "base_uom_code"],
        "example": ["BRG-00001", "RANTAI BESI 5/8 PANJANG 8 M", "Rantai 5/8", "SPAREPART", "OPS", "PCS", "", "", "", "Ya"],
    },
    "suppliers": {
        "label": "Master Supplier",
        "columns": [
            "code", "name", "supplier_category_code", "phone", "email", "legal_name", "supplier_type",
            "address", "city", "province", "country", "npwp", "nib", "pkp", "payment_term", "currency",
            "default_tax_code", "lead_time_days", "min_order", "delivery_days", "service_area",
            "contact_name", "contact_role", "contact_phone", "contact_email",
            "bank_name", "bank_account_no", "bank_account_name", "bank_branch", "bank_currency", "is_active",
        ],
        "required": ["code", "name"],
        "example": ["SUP-00001", "PT. Profita Abadi", "SPAREPART", "", "", "PT. Profita Abadi", "Lokal", "", "", "", "Indonesia", "", "", "Ya", "Net 120", "IDR", "PPN11", 7, 0, "H+7 After PO", "", "Ibu Lidya", "Marketing", "085214080112", "", "BANK MANDIRI", "108-001-023-364-2", "PT. Profita Abadi", "", "IDR", "Ya"],
    },
    "units": {
        "label": "Master Unit / Aset",
        "columns": ["code", "name", "type", "plate_no", "asset_no", "brand", "model", "year", "division_code", "is_active"],
        "required": ["code", "name"],
        "example": ["UNT-00001", "Excavator 01", "Heavy Equipment", "", "AST-001", "CAT", "320", 2024, "OPS", "Ya"],
    },
    "stock_limits": {
        "label": "Stok Minimal / Maksimal",
        "columns": ["item_code", "warehouse_code", "min_stock", "max_stock", "current_stock"],
        "required": ["item_code", "warehouse_code"],
        "example": ["BRG-00001", "GDG-00001", 5, 20, ""],
    },
}


TRANSACTION_DATASETS = {
    "mro": {
        "label": "MRO — Permintaan Material",
        "path": "/api/mro",
        "columns": ["batch_ref", "date", "need_date", "division_code", "requester", "department", "default_warehouse_code", "default_project_code", "default_unit_code", "spk", "notes", "item_code", "qty", "uom_code", "warehouse_code", "project_code", "unit_code", "line_notes"],
        "required": ["batch_ref", "date", "item_code", "qty"],
    },
    "ro": {
        "label": "RO — Permintaan Pembelian",
        "path": "/api/ro",
        "columns": ["batch_ref", "date", "division_code", "requester", "department", "default_warehouse_code", "default_project_code", "default_unit_code", "spk", "notes", "item_code", "qty", "uom_code", "warehouse_code", "project_code", "unit_code", "line_notes", "source_mro_no"],
        "required": ["batch_ref", "date", "item_code", "qty"],
    },
    "po": {
        "label": "PO — Pesanan Pembelian",
        "path": "/api/po",
        "columns": ["batch_ref", "date", "supplier_code", "buyer_contact_code", "division_code", "payment_term", "delivery_term", "currency", "eta", "default_warehouse_code", "default_project_code", "default_unit_code", "spk", "tax_inclusive", "supplier_notes", "internal_notes", "item_code", "qty", "uom_code", "warehouse_code", "project_code", "unit_code", "price", "discount", "tax_code", "line_notes", "source_ro_no"],
        "required": ["batch_ref", "date", "supplier_code", "item_code", "qty"],
    },
    "do": {
        "label": "DO — Penerimaan Barang",
        "path": "/api/do",
        "columns": ["batch_ref", "date", "supplier_code", "supplier_dn", "supplier_invoice", "invoice_date", "default_warehouse_code", "default_project_code", "default_unit_code", "spk", "receiver", "notes", "item_code", "qty", "uom_code", "warehouse_code", "project_code", "unit_code", "condition", "line_notes", "source_po_no"],
        "required": ["batch_ref", "date", "item_code", "qty", "source_po_no"],
    },
    "mi": {
        "label": "MI — Pengeluaran Barang",
        "path": "/api/mi",
        "columns": ["batch_ref", "date", "division_code", "default_warehouse_code", "default_project_code", "default_unit_code", "spk", "receiver", "department", "source_type", "notes", "item_code", "qty", "uom_code", "warehouse_code", "project_code", "unit_code", "line_notes", "source_mro_no"],
        "required": ["batch_ref", "date", "item_code", "qty"],
    },
    "transfer": {
        "label": "Transfer Antar Gudang",
        "path": "/api/transfers",
        "columns": ["batch_ref", "date", "from_warehouse_code", "to_warehouse_code", "project_code", "notes", "item_code", "qty", "uom_code", "line_project_code", "unit_code", "line_notes"],
        "required": ["batch_ref", "date", "from_warehouse_code", "to_warehouse_code", "item_code", "qty"],
    },
    "loan": {
        "label": "Pinjam Antar Gudang",
        "path": "/api/loans",
        "columns": ["batch_ref", "date", "from_warehouse_code", "to_warehouse_code", "due_date", "project_code", "requester", "notes", "item_code", "qty", "uom_code", "line_project_code", "unit_code", "line_notes"],
        "required": ["batch_ref", "date", "from_warehouse_code", "to_warehouse_code", "item_code", "qty"],
    },
    "loan_return": {
        "label": "Pengembalian Pinjaman",
        "columns": ["batch_ref", "date", "source_loan_no", "notes", "item_code", "qty"],
        "required": ["batch_ref", "date", "source_loan_no", "item_code", "qty"],
    },
    "adjustment": {
        "label": "Penyesuaian Stok",
        "path": "/api/adjustments",
        "columns": ["batch_ref", "date", "warehouse_code", "division_code", "project_code", "adj_type", "reason", "notes", "item_code", "adjustment", "line_reason"],
        "required": ["batch_ref", "date", "warehouse_code", "item_code", "adjustment"],
    },
    "opname": {
        "label": "Stock Opname",
        "columns": ["batch_ref", "date", "warehouse_code", "division_code", "mode", "scope", "notes", "item_code", "counted"],
        "required": ["batch_ref", "date", "warehouse_code", "item_code", "counted"],
    },
}

ALL_DATASETS = {**MASTER_DATASETS, **TRANSACTION_DATASETS}


def _text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Ya" if v else "Tidak"
    return str(v).strip()


def _bool(v, default=False):
    if v is None or _text(v) == "":
        return default
    return _text(v).lower() in ("1", "true", "ya", "y", "yes", "aktif", "include", "included")


def _float(v, default=0.0):
    if v is None or _text(v) == "":
        return default
    try:
        return float(v)
    except Exception:
        return default


def _date(v):
    if v is None:
        return None
    if hasattr(v, "date"):
        try:
            return v.date().isoformat()
        except Exception:
            pass
    s = _text(v)
    return s[:10] if s else None


def _find_route(app, path, method="POST"):
    for route in list(app.router.routes):
        if getattr(route, "path", None) == path and method in (getattr(route, "methods", set()) or set()):
            return route
    return None


async def _call(app, path, body, user, method="POST"):
    route = _find_route(app, path, method)
    if not route:
        raise HTTPException(500, f"Endpoint transaksi tidak tersedia: {path}")
    return await route.endpoint(body, user)


def _template_bytes(key, spec):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    headers = spec["columns"]
    ws.append(headers)
    if spec.get("example"):
        ws.append(spec["example"])
    elif key in TRANSACTION_DATASETS:
        sample = []
        for h in headers:
            val = ""
            if h == "batch_ref": val = "BATCH-001"
            elif h == "date": val = "2026-09-17"
            elif h in ("qty", "price", "discount", "adjustment", "counted"): val = 1 if h != "price" else 100000
            elif h == "item_code": val = "BRG-00001"
            elif h.endswith("warehouse_code"): val = "GDG-00001"
            elif h == "from_warehouse_code": val = "GDG-00001"
            elif h == "to_warehouse_code": val = "GDG-00002"
            elif h == "supplier_code": val = "SUP-00001"
            elif h == "currency": val = "IDR"
            elif h == "mode": val = "live"
            elif h == "scope": val = "all"
            elif h == "tax_inclusive": val = "Tidak"
            sample.append(val)
        ws.append(sample)

    fill = PatternFill("solid", fgColor="1E3A5F")
    for c in ws[1]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for idx, h in enumerate(headers, 1):
        width = max(14, min(34, len(h) + 5))
        ws.column_dimensions[get_column_letter(idx)].width = width

    guide = wb.create_sheet("Petunjuk")
    guide.append(["TEMPLATE", spec["label"]])
    guide.append(["Dataset", key])
    guide.append(["Kolom wajib", ", ".join(spec.get("required", []))])
    guide.append(["Aturan", "Jangan mengubah nama header. Gunakan kode master (bukan ID internal)."])
    if key in TRANSACTION_DATASETS:
        guide.append(["batch_ref", "Baris dengan batch_ref yang sama digabung menjadi 1 transaksi dengan banyak item."])
        guide.append(["Nomor transaksi", "Nomor resmi tetap dibuat otomatis oleh sistem saat import."])
        guide.append(["Sumber", "Isi source_mro_no/source_ro_no/source_po_no bila transaksi harus menjaga traceability."])
    if key == "stock_limits":
        guide.append(["current_stock", "Opsional. Kosongkan bila hanya mengubah Min/Max. Isi hanya saat setup saldo/stock awal yang sudah disetujui."])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 100
    guide["A1"].font = Font(bold=True)
    guide["B1"].font = Font(bold=True)
    out = BytesIO(); wb.save(out); return out.getvalue()


def _read_rows(data: bytes):
    try:
        wb = load_workbook(BytesIO(data), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"File Excel tidak dapat dibaca: {e}")
    if "Data" not in wb.sheetnames:
        raise HTTPException(400, "Sheet 'Data' tidak ditemukan. Gunakan template dari sistem.")
    ws = wb["Data"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [_text(x) for x in rows[0]]
    out = []
    for excel_row, values in enumerate(rows[1:], 2):
        if not any(v is not None and _text(v) != "" for v in values):
            continue
        row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers)) if headers[i]}
        row["__row__"] = excel_row
        out.append(row)
    return headers, out


async def _maps(server):
    async def cmap(coll):
        docs = await coll.find({}, {"_id": 0}).to_list(10000)
        return {_text(d.get("code")).upper(): d for d in docs if d.get("code")}
    return {
        "items": await cmap(server.db.items),
        "warehouses": await cmap(server.db.warehouses),
        "projects": await cmap(server.db.projects),
        "units": await cmap(server.db.units),
        "divisions": await cmap(server.db.divisions),
        "suppliers": await cmap(server.db.suppliers),
        "uoms": await cmap(server.db.uoms),
        "item_categories": await cmap(server.db.item_categories),
        "supplier_categories": await cmap(server.db.supplier_categories),
        "taxes": await cmap(server.db.taxes),
        "contacts": await cmap(server.db.contacts),
    }


def _ref(maps, kind, code, row, field, errors, required=False):
    c = _text(code).upper()
    if not c:
        if required:
            errors.append(f"Baris {row}: {field} wajib diisi")
        return None
    rec = maps.get(kind, {}).get(c)
    if not rec:
        errors.append(f"Baris {row}: {field} '{code}' tidak ditemukan")
        return None
    return rec.get("id")


def _check_headers(headers, spec):
    missing = [x for x in spec.get("required", []) if x not in headers]
    if missing:
        raise HTTPException(400, "Header wajib tidak ada: " + ", ".join(missing))


async def _import_master(server, key, rows, user):
    maps = await _maps(server)
    errors = []
    prepared = []
    for r in rows:
        n = r["__row__"]
        for f in MASTER_DATASETS[key].get("required", []):
            if _text(r.get(f)) == "": errors.append(f"Baris {n}: {f} wajib diisi")
        if key == "items":
            prepared.append({
                "code": _text(r.get("code")), "name": _text(r.get("name")), "alias": _text(r.get("alias")) or None,
                "category_id": _ref(maps, "item_categories", r.get("category_code"), n, "category_code", errors) if _text(r.get("category_code")) else None,
                "division_id": _ref(maps, "divisions", r.get("division_code"), n, "division_code", errors) if _text(r.get("division_code")) else None,
                "base_uom_id": _ref(maps, "uoms", r.get("base_uom_code"), n, "base_uom_code", errors, True),
                "brand": _text(r.get("brand")) or None, "part_number": _text(r.get("part_number")) or None,
                "spec": _text(r.get("spec")) or None, "is_active": _bool(r.get("is_active"), True),
            })
        elif key == "units":
            prepared.append({
                "code": _text(r.get("code")), "name": _text(r.get("name")), "type": _text(r.get("type")) or None,
                "plate_no": _text(r.get("plate_no")) or None, "asset_no": _text(r.get("asset_no")) or None,
                "brand": _text(r.get("brand")) or None, "model": _text(r.get("model")) or None,
                "year": int(_float(r.get("year"), 0)) or None,
                "division_id": _ref(maps, "divisions", r.get("division_code"), n, "division_code", errors) if _text(r.get("division_code")) else None,
                "is_active": _bool(r.get("is_active"), True),
            })
        elif key == "suppliers":
            contact = None
            if _text(r.get("contact_name")):
                contact = {"id": server.gid(), "name": _text(r.get("contact_name")), "role": _text(r.get("contact_role")), "phone": _text(r.get("contact_phone")), "email": _text(r.get("contact_email")), "is_primary": True}
            bank = None
            if _text(r.get("bank_name")) or _text(r.get("bank_account_no")):
                bank = {"id": server.gid(), "bank_name": _text(r.get("bank_name")), "account_no": _text(r.get("bank_account_no")), "account_name": _text(r.get("bank_account_name")), "branch": _text(r.get("bank_branch")), "currency": _text(r.get("bank_currency")) or "IDR", "is_primary": True}
            prepared.append({
                "code": _text(r.get("code")), "name": _text(r.get("name")),
                "supplier_category_id": _ref(maps, "supplier_categories", r.get("supplier_category_code"), n, "supplier_category_code", errors) if _text(r.get("supplier_category_code")) else None,
                "phone": _text(r.get("phone")), "email": _text(r.get("email")), "legal_name": _text(r.get("legal_name")),
                "supplier_type": _text(r.get("supplier_type")) or "Lokal", "address": _text(r.get("address")), "city": _text(r.get("city")),
                "province": _text(r.get("province")), "country": _text(r.get("country")) or "Indonesia", "npwp": _text(r.get("npwp")), "nib": _text(r.get("nib")),
                "pkp": _bool(r.get("pkp")), "payment_term": _text(r.get("payment_term")), "currency": _text(r.get("currency")) or "IDR",
                "default_tax_id": _ref(maps, "taxes", r.get("default_tax_code"), n, "default_tax_code", errors) if _text(r.get("default_tax_code")) else None,
                "lead_time_days": int(_float(r.get("lead_time_days"), 0)), "min_order": _float(r.get("min_order"), 0),
                "delivery_days": _text(r.get("delivery_days")), "service_area": _text(r.get("service_area")),
                "contacts": [contact] if contact else [], "banks": [bank] if bank else [], "is_active": _bool(r.get("is_active"), True),
            })
        elif key == "stock_limits":
            prepared.append({
                "item_id": _ref(maps, "items", r.get("item_code"), n, "item_code", errors, True),
                "warehouse_id": _ref(maps, "warehouses", r.get("warehouse_code"), n, "warehouse_code", errors, True),
                "min_stock": _float(r.get("min_stock"), 0), "max_stock": _float(r.get("max_stock"), 0),
                "current_stock": None if _text(r.get("current_stock")) == "" else _float(r.get("current_stock"), 0),
            })
    if errors:
        return {"ok": False, "imported": 0, "errors": errors[:200]}

    imported = 0; created = 0; updated = 0
    if key == "stock_limits":
        for doc in prepared:
            filt = {"item_id": doc["item_id"], "warehouse_id": doc["warehouse_id"]}
            existing = await server.db.item_warehouse.find_one(filt)
            patch = {"item_id": doc["item_id"], "warehouse_id": doc["warehouse_id"], "min_stock": doc["min_stock"], "max_stock": doc["max_stock"]}
            if doc["current_stock"] is not None: patch["current_stock"] = doc["current_stock"]
            await server.db.item_warehouse.update_one(filt, {"$set": patch}, upsert=True)
            imported += 1; updated += 1 if existing else 0; created += 0 if existing else 1
    else:
        coll = getattr(server.db, key)
        for doc in prepared:
            existing = await coll.find_one({"code": doc["code"]}, {"_id": 0})
            if existing:
                await coll.update_one({"id": existing["id"]}, {"$set": doc})
                await server.audit(user, "excel_update", key, existing["id"], doc.get("code"))
                updated += 1
            else:
                doc.update({"id": server.gid(), "created_at": server.now_iso()})
                await coll.insert_one(doc)
                await server.audit(user, "excel_create", key, doc["id"], doc.get("code"))
                created += 1
            imported += 1
    return {"ok": True, "imported": imported, "created": created, "updated": updated, "errors": []}


async def _source_line(server, module, doc_no, item_id):
    cfg = {"mro": (server.db.mro, server.db.mro_lines, "mro_id"), "ro": (server.db.ro, server.db.ro_lines, "ro_id"), "po": (server.db.po, server.db.po_lines, "po_id")}[module]
    head = await cfg[0].find_one({"no": _text(doc_no)}, {"_id": 0})
    if not head: return None, None
    lines = await cfg[1].find({cfg[2]: head["id"], "item_id": item_id}, {"_id": 0}).to_list(10)
    if len(lines) != 1: return head, None
    return head, lines[0]


async def _transaction_payload(server, key, group, maps, errors):
    first = group[0]; n0 = first["__row__"]
    def rid(kind, field, required=False, row=None):
        rr = row or first
        return _ref(maps, kind, rr.get(field), rr["__row__"], field, errors, required) if _text(rr.get(field)) or required else None
    h = {"date": _date(first.get("date")), "notes": _text(first.get("notes")) or None}
    lines = []
    for r in group:
        n = r["__row__"]
        item_id = _ref(maps, "items", r.get("item_code"), n, "item_code", errors, True)
        uom_id = _ref(maps, "uoms", r.get("uom_code"), n, "uom_code", errors) if _text(r.get("uom_code")) else None
        qty = _float(r.get("qty"), 0)
        if key not in ("adjustment", "opname") and qty <= 0: errors.append(f"Baris {n}: qty harus lebih dari 0")
        line = {"item_id": item_id, "qty": qty, "uom_id": uom_id, "unit": _text(r.get("uom_code")) or None, "notes": _text(r.get("line_notes")) or None}
        if _text(r.get("warehouse_code")): line["warehouse_id"] = rid("warehouses", "warehouse_code", row=r)
        if _text(r.get("project_code")): line["project_id"] = rid("projects", "project_code", row=r)
        if _text(r.get("unit_code")): line["unit_id"] = rid("units", "unit_code", row=r)
        if key == "ro" and _text(r.get("source_mro_no")) and item_id:
            dh, sl = await _source_line(server, "mro", r.get("source_mro_no"), item_id)
            if not dh: errors.append(f"Baris {n}: source_mro_no '{r.get('source_mro_no')}' tidak ditemukan")
            elif not sl: errors.append(f"Baris {n}: item pada MRO sumber tidak unik/tidak ditemukan")
            else: line["sources"] = [{"mro_id": dh["id"], "line_id": sl["id"], "qty": qty}]
        if key == "po":
            line.update({"price": _float(r.get("price"), 0), "discount": _float(r.get("discount"), 0)})
            if _text(r.get("tax_code")):
                tax_id = rid("taxes", "tax_code", row=r); tax = maps.get("taxes", {}).get(_text(r.get("tax_code")).upper(), {})
                line.update({"tax_id": tax_id, "tax": _float(tax.get("rate"), 0), "tax_name": tax.get("name")})
            if _text(r.get("source_ro_no")) and item_id:
                dh, sl = await _source_line(server, "ro", r.get("source_ro_no"), item_id)
                if not dh: errors.append(f"Baris {n}: source_ro_no '{r.get('source_ro_no')}' tidak ditemukan")
                elif not sl: errors.append(f"Baris {n}: item pada RO sumber tidak unik/tidak ditemukan")
                else: line["sources"] = [{"ro_id": dh["id"], "line_id": sl["id"], "qty": qty}]
        if key == "do":
            line["condition"] = _text(r.get("condition")) or "Baik"
            if _text(r.get("source_po_no")) and item_id:
                dh, sl = await _source_line(server, "po", r.get("source_po_no"), item_id)
                if not dh: errors.append(f"Baris {n}: source_po_no '{r.get('source_po_no')}' tidak ditemukan")
                elif not sl: errors.append(f"Baris {n}: item pada PO sumber tidak unik/tidak ditemukan")
                else: line["sources"] = [{"po_id": dh["id"], "line_id": sl["id"], "qty": qty}]
        if key == "mi" and _text(r.get("source_mro_no")) and item_id:
            dh, sl = await _source_line(server, "mro", r.get("source_mro_no"), item_id)
            if not dh: errors.append(f"Baris {n}: source_mro_no '{r.get('source_mro_no')}' tidak ditemukan")
            elif not sl: errors.append(f"Baris {n}: item pada MRO sumber tidak unik/tidak ditemukan")
            else: line.update({"mro_line_id": sl["id"], "sources": [{"mro_id": dh["id"], "line_id": sl["id"], "qty": qty}]})
        if key in ("transfer", "loan"):
            if _text(r.get("line_project_code")): line["project_id"] = rid("projects", "line_project_code", row=r)
        if key == "adjustment": line = {"item_id": item_id, "adjustment": _float(r.get("adjustment"), 0), "reason": _text(r.get("line_reason")) or None}
        if key == "opname": line = {"item_id": item_id, "counted": _float(r.get("counted"), 0)}
        lines.append(line)

    if key in ("mro", "ro", "po", "do", "mi"):
        h.update({
            "division_id": rid("divisions", "division_code"), "default_warehouse_id": rid("warehouses", "default_warehouse_code"),
            "default_project_id": rid("projects", "default_project_code"), "default_unit_id": rid("units", "default_unit_code"),
            "spk": _text(first.get("spk")) or None,
        })
    if key in ("mro", "ro"):
        h.update({"requester": _text(first.get("requester")) or None, "department": _text(first.get("department")) or None})
    if key == "mro": h["need_date"] = _date(first.get("need_date"))
    if key == "po":
        h.update({"supplier_id": rid("suppliers", "supplier_code", True), "buyer_contact_id": rid("contacts", "buyer_contact_code"), "payment_term": _text(first.get("payment_term")), "delivery_term": _text(first.get("delivery_term")), "currency": _text(first.get("currency")) or "IDR", "eta": _date(first.get("eta")), "tax_inclusive": _bool(first.get("tax_inclusive")), "supplier_notes": _text(first.get("supplier_notes")), "internal_notes": _text(first.get("internal_notes"))})
    elif key == "do":
        h.update({"supplier_id": rid("suppliers", "supplier_code"), "supplier_dn": _text(first.get("supplier_dn")), "supplier_invoice": _text(first.get("supplier_invoice")), "invoice_date": _date(first.get("invoice_date")), "receiver": _text(first.get("receiver"))})
    elif key == "mi": h.update({"receiver": _text(first.get("receiver")), "department": _text(first.get("department")), "source_type": _text(first.get("source_type"))})
    elif key in ("transfer", "loan"):
        h.update({"from_warehouse_id": rid("warehouses", "from_warehouse_code", True), "to_warehouse_id": rid("warehouses", "to_warehouse_code", True), "project_id": rid("projects", "project_code")})
        if key == "loan": h.update({"due_date": _date(first.get("due_date")), "requester": _text(first.get("requester"))})
    elif key == "adjustment": h.update({"warehouse_id": rid("warehouses", "warehouse_code", True), "division_id": rid("divisions", "division_code"), "project_id": rid("projects", "project_code"), "adj_type": _text(first.get("adj_type")) or "Koreksi", "reason": _text(first.get("reason"))})
    elif key == "opname": h.update({"warehouse_id": rid("warehouses", "warehouse_code", True), "division_id": rid("divisions", "division_code"), "mode": _text(first.get("mode")) or "live", "scope": _text(first.get("scope")) or "all"})
    h["lines"] = lines
    return h


async def _import_transactions(server, key, rows, user):
    maps = await _maps(server); errors = []
    groups = OrderedDict()
    for r in rows:
        batch = _text(r.get("batch_ref"))
        if not batch: errors.append(f"Baris {r['__row__']}: batch_ref wajib diisi"); continue
        groups.setdefault(batch, []).append(r)
    prepared = []
    if key == "loan_return":
        for batch, group in groups.items(): prepared.append((batch, group))
    else:
        for batch, group in groups.items(): prepared.append((batch, await _transaction_payload(server, key, group, maps, errors)))
    if errors: return {"ok": False, "imported": 0, "errors": errors[:200]}

    imported = 0; created_docs = []
    for batch, payload in prepared:
        try:
            if key == "loan_return":
                first = payload[0]
                loan = await server.db.loans.find_one({"no": _text(first.get("source_loan_no"))}, {"_id": 0})
                if not loan: raise HTTPException(400, f"Loan sumber '{first.get('source_loan_no')}' tidak ditemukan")
                loan_lines = await server.db.loan_lines.find({"loan_id": loan["id"]}, {"_id": 0}).to_list(1000)
                by_item = defaultdict(list)
                for ll in loan_lines: by_item[ll.get("item_id")].append(ll)
                ret_lines = []
                for r in payload:
                    item_id = _ref(maps, "items", r.get("item_code"), r["__row__"], "item_code", [], True)
                    choices = by_item.get(item_id, [])
                    if len(choices) != 1: raise HTTPException(400, f"Baris {r['__row__']}: item pinjaman tidak unik/tidak ditemukan")
                    ret_lines.append({"loan_line_id": choices[0]["id"], "qty": _float(r.get("qty"), 0)})
                route = _find_route(server.app, "/api/loans/{did}/return", "POST")
                if not route: raise HTTPException(500, "Endpoint return loan tidak tersedia")
                result = await route.endpoint(loan["id"], {"date": _date(first.get("date")), "notes": _text(first.get("notes")), "lines": ret_lines}, user)
            elif key == "opname":
                counted = payload.pop("lines")
                result = await _call(server.app, "/api/opname", payload, user)
                did = result.get("id")
                existing = await server.db.opname_lines.find({"opname_id": did}, {"_id": 0}).to_list(10000)
                by_item = {x["item_id"]: x for x in existing}
                count_lines = [{"line_id": by_item[x["item_id"]]["id"], "counted": x["counted"]} for x in counted if x.get("item_id") in by_item]
                route = _find_route(server.app, "/api/opname/{did}/count", "PUT")
                if route and count_lines: result = await route.endpoint(did, {"lines": count_lines, "status": "Review"}, user)
            else:
                result = await _call(server.app, TRANSACTION_DATASETS[key]["path"], payload, user)
            imported += 1
            created_docs.append({"batch_ref": batch, "id": result.get("id") if isinstance(result, dict) else None, "no": result.get("no") if isinstance(result, dict) else None})
        except Exception as e:
            detail = getattr(e, "detail", None) or str(e)
            errors.append(f"Batch {batch}: {detail}")
    await server.audit(user, "excel_import", key, "batch", after={"imported": imported, "errors": len(errors)})
    return {"ok": len(errors) == 0, "imported": imported, "created_docs": created_docs, "errors": errors[:200]}


def install(server):
    @server.app.get("/api/excel/datasets", tags=["excel-import"])
    async def datasets(user=Depends(server.current_user)):
        server.require(user, "view")
        return {
            "masters": [{"key": k, "label": v["label"], "required": v.get("required", [])} for k, v in MASTER_DATASETS.items()],
            "transactions": [{"key": k, "label": v["label"], "required": v.get("required", [])} for k, v in TRANSACTION_DATASETS.items()],
        }

    @server.app.get("/api/excel/template/{dataset}", tags=["excel-import"])
    async def template(dataset: str, user=Depends(server.current_user)):
        server.require(user, "view")
        spec = ALL_DATASETS.get(dataset)
        if not spec: raise HTTPException(404, "Template tidak dikenal")
        data = _template_bytes(dataset, spec)
        return Response(content=data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="template_{dataset}.xlsx"'})

    @server.app.post("/api/excel/import/{dataset}", tags=["excel-import"])
    async def import_excel(dataset: str, file: UploadFile = File(...), user=Depends(server.current_user)):
        server.require(user, "create")
        spec = ALL_DATASETS.get(dataset)
        if not spec: raise HTTPException(404, "Dataset import tidak dikenal")
        if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
            raise HTTPException(400, "Gunakan file Excel .xlsx")
        data = await file.read()
        if len(data) > 20 * 1024 * 1024: raise HTTPException(400, "File terlalu besar (maksimal 20 MB)")
        headers, rows = _read_rows(data)
        _check_headers(headers, spec)
        if not rows: raise HTTPException(400, "Tidak ada data pada sheet Data")
        if dataset in MASTER_DATASETS:
            result = await _import_master(server, dataset, rows, user)
        else:
            result = await _import_transactions(server, dataset, rows, user)
        result.update({"dataset": dataset, "filename": file.filename, "rows": len(rows)})
        return result
