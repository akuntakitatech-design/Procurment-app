"""Stock Opname — workflow, valuasi, Freeze/Live, posting atomik, import/export & cetak (KHUSUS OPNAME).

Aturan (lihat PR Stock Opname Hardening):
- 1 dokumen = 1 gudang (header). Gudang dikunci setelah snapshot.
- 1 gudang = 1 Opname aktif (Counting/Review/Waiting Approval). Lock atomik via PK deterministik di
  koleksi strict-pk `opname_locks` (aman terhadap request bersamaan). Posted/Cancelled melepas lock.
- Workflow ditegakkan server: Counting -> Review -> Waiting Approval -> Posted; Return/Reject (alasan wajib)
  -> Counting; Cancel (alasan wajib) -> Cancelled. Status TIDAK pernah dibaca dari payload.
- Freeze: SEMUA tulis saldo/nilai `item_warehouse` pada gudang ber-lock freeze ditolak (titik tulis tunggal
  persediaan: post_movement + jalur saldo awal/valuasi/import), kecuali oleh posting Opname pemilik lock.
- Live: stok sistem dicatat per baris saat hitungan disimpan (system_at_count + counted_at); selisih =
  fisik - stok sistem saat hitung. Mutasi yang diinput setelah hitung dengan tanggal <= tanggal hitung
  menandai baris "perlu hitung ulang" (submit/posting diblokir sampai dihitung ulang).
- Posting atomik: klaim status (SELECT FOR UPDATE), pra-validasi seluruh baris, posting via engine existing
  (post_movement, MWA tidak diubah), gagal di tengah -> reversal kompensasi seluruh movement dokumen; retry
  memakai source_key ber-attempt sehingga tidak ada double posting maupun baris terlewat.
- Harga Moving Average & nilai selisih diredaksi server untuk user tanpa view_purchase_price.
- Dokumen legacy (tanpa workflow_version) tetap terbaca: selisih = fisik - snapshot (perilaku lama).
"""
import asyncio
import contextvars
import hashlib
import io
import math
import re

from fastapi import Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response

WF_VERSION = 2
ACTIVE = ("Counting", "Review", "Waiting Approval")
EDITABLE = ("Counting", "Review")
POSTING_CTX: contextvars.ContextVar = contextvars.ContextVar("opname_posting_doc", default=None)
COMPENSATING: contextvars.ContextVar = contextvars.ContextVar("opname_compensating", default=False)
STOCK_KEYS = {"current_stock", "avg_cost", "total_value", "_ver", "opening_qty", "opening_value", "opening_avg_cost"}
PRICE_KEYS = ("unit_price", "value", "approved_unit_cost", "cost_reason", "avg_cost", "posted_unit_cost",
              "posted_value", "cost_overridden")
SUMMARY_PRICE_KEYS = ("surplus_value", "shortage_value", "net_value")
EPS = 1e-9


class FrozenWarehouseError(HTTPException):
    pass


# Retry konflik lock InnoDB pada klaim CAS status dokumen Opname. HANYA 1213 (deadlock) & 1205 (lock wait timeout);
# error lain (termasuk 1062 / koneksi / validasi bisnis) TIDAK di-retry. Maks. 3 percobaan termasuk yang pertama.
LOCK_CONFLICT_ERRNOS = (1213, 1205)
CAS_MAX_ATTEMPTS = 3
CAS_RETRY_DELAY = 0.05


def lock_conflict_errno(exc):
    """Errno 1213/1205 bila exc adalah konflik lock MariaDB (pymysql OperationalError), selain itu None."""
    try:
        from pymysql.err import OperationalError
    except Exception:  # noqa: BLE001
        return None
    args = getattr(exc, "args", None) or ()
    code = args[0] if args else None
    return code if isinstance(exc, OperationalError) and code in LOCK_CONFLICT_ERRNOS else None


async def retry_lock_conflict(op, attempts=CAS_MAX_ATTEMPTS, sleep=asyncio.sleep, delay=CAS_RETRY_DELAY):
    """Jalankan `op()` (SATU unit transaksi utuh, mis. find_one_and_update = BEGIN; SELECT..FOR UPDATE; UPDATE; COMMIT).
    Bila InnoDB memilih transaksi ini sebagai korban 1213/1205, adapter (_Tx.__aexit__) sudah ROLLBACK seluruhnya
    sebelum exception sampai di sini, sehingga pengulangan aman & tidak pernah menerapkan perubahan dua kali.
    Percobaan habis -> HTTP 409 (bukan 500), tanpa perubahan data."""
    attempts = max(1, min(int(attempts), CAS_MAX_ATTEMPTS))
    for i in range(attempts):
        try:
            return await op()
        except Exception as exc:
            if lock_conflict_errno(exc) is None:
                raise
            if i == attempts - 1:
                raise HTTPException(409, "Dokumen sedang diproses bersamaan oleh pengguna lain. Muat ulang lalu coba lagi.") from exc
            await sleep(delay * (i + 1))


def _f(v):
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return x if math.isfinite(x) else None


def lock_id(warehouse_id: str) -> str:
    return "OPL" + hashlib.sha1(str(warehouse_id).encode()).hexdigest()[:16]


def frozen_message(no):
    return (f"Gudang sedang dalam proses Stock Opname {no}. "
            "Mutasi stok tidak dapat dilakukan hingga proses selesai.")


def _date10(v):
    return str(v or "")[:10]


def _minutes_ago(server, minutes):
    from datetime import datetime, timedelta, timezone
    ref = str(server.now_iso())
    dt = datetime.now(timezone.utc) - timedelta(minutes=minutes)
    return dt.isoformat() if "+" in ref or ref.endswith("Z") else dt.replace(tzinfo=None).isoformat()


def factor_of(item: dict, uom_id) -> float | None:
    """Faktor ke satuan dasar. None bila satuan bukan milik barang."""
    if not uom_id or uom_id == item.get("base_uom_id") or uom_id == f"legacy:{item.get('unit')}":
        return 1.0
    for r in item.get("uoms") or []:
        if r.get("uom_id") == uom_id:
            fx = _f(r.get("factor"))
            return fx if fx and fx > 0 else None
    return None


def compute_line(doc: dict, line: dict, iw: dict | None):
    """Angka per baris dari data dokumen (dipakai detail, ringkasan, cetak, pra-validasi posting)."""
    counted = _f(line.get("counted"))
    v2 = doc.get("workflow_version") == WF_VERSION
    posted = doc.get("status") == "Posted" and line.get("posted_delta") is not None
    current = float((iw or {}).get("current_stock") or 0)
    if posted:
        system = float(line.get("variance_base") if line.get("variance_base") is not None else line.get("snapshot") or 0)
    elif not v2 or doc.get("mode") == "freeze":
        system = float(line.get("snapshot") or 0)
    elif counted is not None and line.get("system_at_count") is not None:
        system = float(line.get("system_at_count"))
    else:
        system = current
    variance = None if counted is None else counted - system
    avg = _f((iw or {}).get("avg_cost")) or 0.0
    has_avg = current > EPS and avg > EPS
    approved = _f(line.get("approved_unit_cost"))
    if posted:
        price = _f(line.get("posted_unit_cost"))
    elif variance is not None and variance > EPS and approved is not None:
        price = approved
    else:
        price = avg if has_avg else (approved if approved is not None else None)
    value = None
    if posted and line.get("posted_value") is not None:
        value = float(line.get("posted_value"))
    elif variance is not None and price is not None:
        value = variance * price
    if counted is None:
        status = "uncounted"
    elif abs(variance) <= EPS:
        status = "match"
    else:
        status = "plus" if variance > 0 else "minus"
    needs_cost = (not posted and variance is not None and variance > EPS and not has_avg and (approved is None or approved <= 0))
    return {"counted": counted, "system_qty": system, "variance": variance, "status": status,
            "avg_cost": avg if has_avg else None, "unit_price": price, "value": value, "needs_cost": needs_cost,
            "reason_missing": variance is not None and abs(variance) > EPS and not str(line.get("reason") or "").strip(),
            "current_stock": current}


def summarize(rows):
    s = {"total": 0, "counted": 0, "uncounted": 0, "match": 0, "plus": 0, "minus": 0, "needs_recount": 0,
         "reason_missing": 0, "needs_cost": 0, "surplus_value": 0.0, "shortage_value": 0.0, "net_value": 0.0}
    for r in rows:
        s["total"] += 1
        s[r["status"]] = s.get(r["status"], 0) + 1
        if r["status"] != "uncounted":
            s["counted"] += 1
        s["needs_recount"] += 1 if r.get("needs_recount") else 0
        s["reason_missing"] += 1 if r.get("reason_missing") else 0
        s["needs_cost"] += 1 if r.get("needs_cost") else 0
        if r.get("value") is not None:
            if r["value"] > 0:
                s["surplus_value"] += r["value"]
            elif r["value"] < 0:
                s["shortage_value"] += r["value"]
    s["net_value"] = s["surplus_value"] + s["shortage_value"]
    for k in SUMMARY_PRICE_KEYS:
        s[k] = round(s[k], 4)
    return s


def readiness_errors(rows, posted_check=False):
    out = []
    n_unc = sum(1 for r in rows if r["status"] == "uncounted")
    if not rows:
        out.append("Stock Opname tidak memiliki baris barang")
    if n_unc:
        out.append(f"Masih ada {n_unc} barang yang belum dihitung")
    n_reason = sum(1 for r in rows if r.get("reason_missing"))
    if n_reason:
        out.append(f"{n_reason} barang selisih belum memiliki alasan")
    n_rc = sum(1 for r in rows if r.get("needs_recount"))
    if n_rc:
        out.append(f"{n_rc} barang perlu dihitung ulang karena ada mutasi setelah penghitungan")
    n_cost = sum(1 for r in rows if r.get("needs_cost"))
    if n_cost:
        out.append(f"{n_cost} barang surplus tanpa rata-rata persediaan belum memiliki Harga Satuan yang disetujui")
    return out


def install(server):
    app = server.app

    class _LiveDb:  # selalu resolve server.db saat dipanggil (tenant-isolation layer mengganti server.db belakangan)
        def __getattr__(self, name):
            return getattr(server.db, name)
    db = _LiveDb()
    import mariadb_motor as MM
    import transfer_lines as TL
    import attachment_integrity_guard_layer as AIG
    import doc_warehouse as DW

    # ---- registrasi draft lampiran modul opname (infrastruktur draft existing; modul lain tidak berubah)
    TL.DRAFT_ENTITY_OF.setdefault("opname", "opname_draft")
    TL.DRAFT_HEAD_COLL.setdefault("opname", "opname")
    TL.CLAIM_STAMP_MODULES.add("opname")
    AIG.DRAFT_ENTITIES.setdefault("opname_draft", "opname")
    for name in dir(AIG):
        obj = getattr(AIG, name)
        if isinstance(obj, dict) and obj.get("adjustment_draft") == "attachment_drafts":
            obj.setdefault("opname_draft", "attachment_drafts")
    MM.STRICT_PK_COLLECTIONS.add("opname_locks")

    # ------------------------------------------------------------------ helpers
    def is_admin(user):
        return (user or {}).get("role") == "admin"

    def can(user, act):
        if is_admin(user):
            return True
        if act == "post" and server.has_perm(user, "post_stock_opname"):
            return True
        return bool(server.has_perm(user, f"opname.{act}"))

    def need(user, act, msg):
        if not can(user, act):
            raise HTTPException(403, msg)

    def price_ok(user):
        return is_admin(user) or bool(server.has_perm(user, "view_purchase_price"))

    async def scope(did, user):
        fn = getattr(server, "require_doc_access", None)
        if fn:
            await fn("opname", did, user)

    async def assert_division(div_id):
        if div_id in (None, ""):
            return  # wajib-divisi ditegakkan layer validasi transaksi existing
        dv = await db.divisions.find_one({"id": div_id}, {"_id": 0, "id": 1, "is_active": 1})
        if not dv:
            raise HTTPException(400, "Divisi tidak ditemukan")
        if dv.get("is_active") is False:
            raise HTTPException(400, "Divisi tidak aktif")

    async def get_doc(did):
        d = await db.opname.find_one({"id": did}, {"_id": 0})
        if not d:
            raise HTTPException(404, "Stock Opname tidak ditemukan")
        return d

    async def all_lines(did):
        return await db.opname_lines.find({"opname_id": did}, {"_id": 0}).to_list(None)

    async def items_by_id(ids):
        ids = list({i for i in ids if i})
        if not ids:
            return {}
        rows = await db.items.find({"id": {"$in": ids}}, {"_id": 0}).to_list(None)
        return {r["id"]: r for r in rows}

    async def iw_map(wh):
        rows = await db.item_warehouse.find({"warehouse_id": wh}, {"_id": 0}).to_list(None)
        return {r.get("item_id"): r for r in rows}

    async def uom_labels():
        rows = await db.uoms.find({}, {"_id": 0}).to_list(None)
        return {r["id"]: (r.get("symbol") or r.get("code") or r.get("name")) for r in rows}

    def base_unit(it, um):
        return um.get(it.get("base_uom_id")) or it.get("unit") or ""

    async def history(did, doc, action, frm, to, user, reason=None, extra=None):
        ev = {"action": action, "from": frm, "to": to, "by": user.get("email"), "by_name": user.get("name"),
              "at": server.now_iso(), "reason": reason}
        hist = list(doc.get("history") or []) + [ev]
        patch = {"history": hist, **(extra or {})}
        await db.opname.update_one({"id": did}, {"$set": patch})
        return ev

    async def release_lock(did, wh=None):
        q = {"opname_id": did}
        await db.opname_locks.delete_many(q)
        await db.opname.update_one({"id": did}, {"$set": {"freeze_active": False}})

    async def enrich(doc, lines, user):
        items = await items_by_id([l.get("item_id") for l in lines])
        iwm = await iw_map(doc.get("warehouse_id"))
        um = await uom_labels()
        rows = []
        for l in lines:
            it = items.get(l.get("item_id"), {})
            c = compute_line(doc, l, iwm.get(l.get("item_id")))
            r = {**l, **c, "item_code": it.get("code"), "item_name": it.get("name"), "unit": base_unit(it, um),
                 "base_uom_id": it.get("base_uom_id"), "needs_recount": bool(l.get("needs_recount")),
                 "counted_uom_label": um.get(l.get("counted_uom_id")) if l.get("counted_uom_id") else None}
            r["variance_legacy"] = r["variance"] if r["variance"] is not None else 0  # kompatibel field lama
            rows.append(r)
        rows.sort(key=lambda r: (str(r.get("item_code") or ""), str(r.get("item_name") or "")))
        return rows

    def redact(rows, summary, user):
        if price_ok(user):
            return
        for r in rows:
            for k in PRICE_KEYS:
                r.pop(k, None)
        for k in SUMMARY_PRICE_KEYS:
            summary.pop(k, None)

    def permissions(doc, user):
        st = doc.get("status")
        return {"count": st in EDITABLE and can(user, "edit"), "add_item": st == "Counting" and can(user, "edit"),
                "review": st == "Counting" and can(user, "edit"), "submit": st in EDITABLE and can(user, "edit"),
                "return": (st == "Review" and can(user, "edit")) or (st == "Waiting Approval" and can(user, "post")),
                "reject": st == "Waiting Approval" and can(user, "post"),
                "approve": st in ("Waiting Approval", "Gagal Posting") and can(user, "post"),
                "cancel": st in ACTIVE and (can(user, "cancel") or can(user, "delete")),
                "import": st in EDITABLE and can(user, "edit"), "print": can(user, "print") or can(user, "view"),
                "price": price_ok(user), "override_cost": price_ok(user) and can(user, "edit") and st in EDITABLE}

    async def payload(doc, user, page=None, page_size=50, q=None, status=None):
        lines = await all_lines(doc["id"])
        rows = await enrich(doc, lines, user)
        summary = summarize(rows)
        wh = await db.warehouses.find_one({"id": doc.get("warehouse_id")}, {"_id": 0}) or {}
        dv = await db.divisions.find_one({"id": doc.get("division_id")}, {"_id": 0}) if doc.get("division_id") else None
        out = {**doc, "warehouse_name": wh.get("name"), "warehouse_code": wh.get("code"),
               "division_name": (dv or {}).get("name"), "legacy": doc.get("workflow_version") != WF_VERSION,
               "summary": summary, "permissions": permissions(doc, user),
               "readiness": readiness_errors(rows) if doc.get("status") in ACTIVE else []}
        sel = rows
        if q:
            ql = str(q).strip().lower()
            sel = [r for r in sel if ql in str(r.get("item_code") or "").lower() or ql in str(r.get("item_name") or "").lower()]
        if status and status != "all":
            if status == "variance":
                sel = [r for r in sel if r["status"] in ("plus", "minus")]
            elif status == "recount":
                sel = [r for r in sel if r.get("needs_recount")]
            elif status == "attention":
                sel = [r for r in sel if r.get("needs_recount") or r.get("reason_missing") or r.get("needs_cost")]
            else:
                sel = [r for r in sel if r["status"] == status]
        out["line_total"] = len(sel)
        if page is not None:
            ps = max(1, min(int(page_size or 50), 500))
            pg = max(1, int(page))
            out["page"], out["page_size"] = pg, ps
            out["page_count"] = max(1, math.ceil(len(sel) / ps))
            sel = sel[(pg - 1) * ps: pg * ps]
        out["lines"] = sel
        redact(sel, summary, user)
        return out

    # ------------------------------------------------------------------ freeze + live guard (item_warehouse)
    async def frozen_lock(wh):
        if not wh or not isinstance(wh, str) or COMPENSATING.get():
            return None
        lk = await db.opname_locks.find_one({"warehouse_id": wh, "mode": "freeze"}, {"_id": 0})
        if lk and POSTING_CTX.get() != lk.get("opname_id"):
            return lk
        return None

    async def any_freeze():
        if COMPENSATING.get():
            return False
        return bool(await db.opname_locks.find_one({"mode": "freeze"}, {"_id": 0, "id": 1}))

    async def resolve_whs(coll, flt):
        """Filter tanpa warehouse_id (mis. hapus per item_id): tentukan gudang dari baris yang cocok.
        Filter sudah ber-scope tenant (tenant layer membungkus di atas koleksi mentah)."""
        if not await any_freeze():
            return set()
        try:
            rows = await MM.MariaCollection.find(coll, flt or {}, {"_id": 0, "warehouse_id": 1}).to_list(None)
        except Exception:  # noqa: BLE001
            return set()
        return {r.get("warehouse_id") for r in rows if isinstance(r.get("warehouse_id"), str)}

    async def stock_unchanged(coll, flt, update):
        """Update yang hanya menyetel ulang nilai stok yang sama (mis. ubah min/max) bukan mutasi stok."""
        sets = (update or {}).get("$set") if isinstance(update, dict) else None
        if not isinstance(sets, dict) or set(update) - {"$set", "$setOnInsert"}:
            return False
        ff = flat(flt)
        if not isinstance(ff.get("item_id"), str) or not isinstance(ff.get("warehouse_id"), str):
            return False
        cur = await MM.MariaCollection.find_one(coll, flt, {"_id": 0})
        if not cur:
            return False
        for k in set(sets) & STOCK_KEYS:
            a, b = _f(sets.get(k)), _f(cur.get(k))
            if (a if a is not None else 0.0) != (b if b is not None else 0.0):
                return False
        return True

    async def guard_whs(coll, flt, update=None, docs=None):
        whs = wh_candidates(flt, update, docs)
        if not whs and docs is None:
            whs = await resolve_whs(coll, flt)
        await assert_not_frozen(whs)

    def flat(flt):
        """Gabungkan field level-atas + isi $and (tenant layer membungkus filter: {"$and": [asli, {tenant_id}]})."""
        out = {}
        if not isinstance(flt, dict):
            return out
        for k, v in flt.items():
            if k == "$and" and isinstance(v, (list, tuple)):
                for sub in v:
                    for kk, vv in flat(sub).items():
                        out.setdefault(kk, vv)
            elif not k.startswith("$"):
                out.setdefault(k, v)
        return out

    def wh_candidates(flt, update=None, docs=None):
        out = set()
        for src in [flat(flt)] + list(docs or []):
            v = (src or {}).get("warehouse_id")
            if isinstance(v, str):
                out.add(v)
            elif isinstance(v, dict) and isinstance(v.get("$in"), (list, tuple)):
                out |= {x for x in v["$in"] if isinstance(x, str)}
        if isinstance(update, dict):
            v = (update.get("$set") or {}).get("warehouse_id")
            if isinstance(v, str):
                out.add(v)
        return out

    def touches_stock(update):
        if not isinstance(update, dict):
            return True
        keys = set()
        for op, val in update.items():
            if op.startswith("$") and isinstance(val, dict):
                keys |= set(val)
            elif not op.startswith("$"):
                keys.add(op)
        return bool(keys & STOCK_KEYS)

    async def assert_not_frozen(whs):
        for wh in whs:
            lk = await frozen_lock(wh)
            if lk:
                raise FrozenWarehouseError(409, frozen_message(lk.get("opname_no")))

    async def cas_doc(flt, update):
        """Compare-and-set status dokumen Opname (klaim Approve & Post / transisi workflow). Klaim bersamaan dapat
        membuat InnoDB memilih salah satu transaksi sebagai korban 1213/1205 -> diulang via retry_lock_conflict.
        Return dokumen bila menang; None bila kalah CAS (pemanggil menjawab 409) — tidak pernah 500."""
        return await retry_lock_conflict(lambda: db.opname.find_one_and_update(flt, update))

    Coll = MM.MariaCollection
    if not getattr(Coll, "_opname_freeze_guard", False):
        o_ins1, o_insm, o_up1, o_upm = Coll.insert_one, Coll.insert_many, Coll.update_one, Coll.update_many
        o_rep, o_fau, o_del1, o_delm = Coll.replace_one, Coll.find_one_and_update, Coll.delete_one, Coll.delete_many
        o_fad, o_bulk = Coll.find_one_and_delete, Coll.bulk_write

        async def g_insert_one(self, document, *a, **k):
            if self.name == "item_warehouse":
                await guard_whs(self, None, docs=[document])
            elif self.name == "valuation_replays" and (document or {}).get("status") == "applying":
                # Revaluasi (mengubah nilai pool) dicek SELURUH pool sebelum satu pun ditulis -> tanpa partial.
                await assert_not_frozen({p.get("warehouse_id") for it in (document.get("items") or [])
                                         for p in (it.get("pools") or []) if isinstance(p.get("warehouse_id"), str)})
            res = await o_ins1(self, document, *a, **k)
            if self.name == "stock_ledger":
                await live_flag(document)
                await freeze_race_flag(document)
            return res

        async def g_insert_many(self, documents, *a, **k):
            documents = list(documents)
            if self.name == "item_warehouse":
                await guard_whs(self, None, docs=documents)
            return await o_insm(self, documents, *a, **k)

        def mk_upd(orig, single=False):
            async def g(self, filter, update, *a, **k):
                if self.name == "item_warehouse" and touches_stock(update):
                    if not (single and await any_freeze() and await stock_unchanged(self, filter, update)):
                        await guard_whs(self, filter, update)
                return await orig(self, filter, update, *a, **k)
            return g

        async def g_replace(self, filter, replacement, *a, **k):
            if self.name == "item_warehouse":
                await guard_whs(self, filter, docs=[replacement])
            return await o_rep(self, filter, replacement, *a, **k)

        def mk_del(orig):
            async def g(self, filter, *a, **k):
                if self.name == "item_warehouse":
                    await guard_whs(self, filter)
                return await orig(self, filter, *a, **k)
            return g

        async def g_bulk(self, requests, *a, **k):
            if self.name == "item_warehouse" and await any_freeze():
                raise FrozenWarehouseError(409, "Operasi massal saldo gudang ditolak selama ada gudang dalam Stock Opname Freeze.")
            return await o_bulk(self, requests, *a, **k)

        Coll.insert_one, Coll.insert_many = g_insert_one, g_insert_many
        Coll.update_one, Coll.update_many = mk_upd(o_up1, single=True), mk_upd(o_upm)
        Coll.find_one_and_update = mk_upd(o_fau)
        Coll.replace_one, Coll.delete_one, Coll.delete_many = g_replace, mk_del(o_del1), mk_del(o_delm)
        Coll.find_one_and_delete, Coll.bulk_write = mk_del(o_fad), g_bulk
        Coll._opname_freeze_guard = True

    async def freeze_race_flag(entry):
        """Mutasi yang lolos cek Freeze tepat sebelum lock dibuat (race) -> baris ditandai hitung ulang
        sehingga snapshot tidak pernah dianggap valid diam-diam."""
        try:
            wh, item = (entry or {}).get("warehouse_id"), (entry or {}).get("item_id")
            if not wh or not item:
                return
            lk = await db.opname_locks.find_one({"warehouse_id": wh, "mode": "freeze"}, {"_id": 0})
            if not lk or entry.get("doc_id") == lk.get("opname_id") or POSTING_CTX.get() == lk.get("opname_id"):
                return
            await db.opname_lines.update_many({"opname_id": lk["opname_id"], "item_id": item}, {"$set": {
                "needs_recount": True,
                "recount_reason": f"Mutasi {entry.get('doc_type') or ''} {entry.get('doc_no') or ''} terjadi bersamaan "
                                  "dengan snapshot Freeze"}})
        except Exception:  # noqa: BLE001
            return

    # ---- reversal/pembatalan dokumen lain: cek SEMUA gudang dokumen sebelum satu pun movement dibalik
    _orig_reverse = server.reverse_document_valuation
    if not getattr(_orig_reverse, "_opname_guard", False):
        async def reverse_guarded(doc_id, *a, **k):
            reason = k.get("reason") if "reason" in k else (a[1] if len(a) > 1 else "reversal")
            if reason in ("post_failed", "correction_failed"):
                tok = COMPENSATING.set(True)  # kompensasi kegagalan = memulihkan keadaan sebelum transaksi
                try:
                    return await _orig_reverse(doc_id, *a, **k)
                finally:
                    COMPENSATING.reset(tok)
            if await any_freeze():
                rows = await db.valuation_ledger.find({"doc_id": doc_id, "is_reversal": {"$ne": True}, "reversed": {"$ne": True}},
                                                      {"_id": 0, "warehouse_id": 1}).to_list(None)
                await assert_not_frozen({r.get("warehouse_id") for r in rows if isinstance(r.get("warehouse_id"), str)})
            return await _orig_reverse(doc_id, *a, **k)
        reverse_guarded._opname_guard = True
        server.reverse_document_valuation = reverse_guarded

    # ---- import Excel massal yang menulis saldo: pra-cek seluruh baris (tanpa partial import)
    async def precheck_rows(rows, only_with=None):
        if not await any_freeze():
            return
        codes = set()
        for r in rows or []:
            if only_with and str(r.get(only_with) if r.get(only_with) is not None else "").strip() == "":
                continue
            c = str(r.get("warehouse_code") or "").strip()
            if c:
                codes.add(c)
        if not codes:
            return
        whs = await db.warehouses.find({"code": {"$in": list(codes)}}, {"_id": 0, "id": 1}).to_list(None)
        await assert_not_frozen({w["id"] for w in whs})

    try:
        import excel_import_layer as XL
        _orig_im = XL._import_master
        if not getattr(_orig_im, "_opname_guard", False):
            async def import_master_guarded(server_arg, key, rows, user):
                if key == "stock_limits":
                    await precheck_rows(rows, only_with="current_stock")
                elif key == "opening_inventory":
                    await precheck_rows(rows)
                return await _orig_im(server_arg, key, rows, user)
            import_master_guarded._opname_guard = True
            XL._import_master = import_master_guarded
    except Exception:  # noqa: BLE001
        pass

    # ---- pra-cek tingkat DOKUMEN untuk endpoint mutasi stok: SEBELUM satu baris pun ditulis, seluruh gudang yang
    # disentuh (payload baru + movement aktif dokumen lama) dicek. Mencegah partial/ledger kompensasi pada dokumen
    # multi-gudang (mis. Loan Return keluar dari B lalu masuk ke A yang Freeze). Dipasang saat startup agar
    # membungkus route FINAL (setelah semua layer lain terpasang); hanya mengganti dependant.call.
    GUARDED = {("POST", "/api/do"), ("POST", "/api/mi"), ("POST", "/api/loans"), ("POST", "/api/transfers"),
               ("POST", "/api/adjustments"), ("POST", "/api/loans/{did}/return"), ("PUT", "/api/loan-returns/{rid}"),
               ("DELETE", "/api/loan-returns/{rid}"), ("PUT", "/api/transactions/{module}/{did}"),
               ("DELETE", "/api/transactions/{module}/{did}")}

    async def doc_level_whs(kw):
        found = set()

        def walk(o):
            if isinstance(o, dict):
                for kk, vv in o.items():
                    if isinstance(vv, str) and kk.endswith("warehouse_id"):
                        found.add(vv)
                    elif isinstance(vv, (dict, list)):
                        walk(vv)
            elif isinstance(o, list):
                for vv in o:
                    walk(vv)
        for k, v in kw.items():
            if isinstance(v, (dict, list)) and k != "user":
                walk(v)
            elif k in ("did", "rid") and isinstance(v, str):
                rows = await db.valuation_ledger.find({"doc_id": v, "is_reversal": {"$ne": True}, "reversed": {"$ne": True}},
                                                      {"_id": 0, "warehouse_id": 1}).to_list(None)
                found |= {r.get("warehouse_id") for r in rows if isinstance(r.get("warehouse_id"), str)}
        return found

    def patch_routes():
        for r in app.router.routes:
            p, ms = getattr(r, "path", None), getattr(r, "methods", set()) or set()
            dep = getattr(r, "dependant", None)
            if not dep or not any((m, p) in GUARDED for m in ms) or getattr(dep.call, "_opname_doc_guard", False):
                continue
            orig = dep.call

            async def guarded(*a, __orig=orig, **kw):
                if kw.get("module") != "opname" and await any_freeze():
                    await assert_not_frozen(await doc_level_whs(kw))
                return await __orig(*a, **kw)
            guarded._opname_doc_guard = True
            dep.call = guarded

    app.router.on_startup.append(patch_routes)
    server.opname_patch_mutation_routes = patch_routes

    async def live_flag(entry):
        """Mode Live: mutasi diinput SETELAH baris dihitung dengan tanggal transaksi <= tanggal hitung
        -> baris ditandai perlu hitung ulang (rekonsiliasi kronologis tidak bisa dipastikan)."""
        try:
            wh, item = (entry or {}).get("warehouse_id"), (entry or {}).get("item_id")
            if not wh or not item:
                return
            lk = await db.opname_locks.find_one({"warehouse_id": wh, "mode": "live"}, {"_id": 0})
            if not lk or entry.get("doc_id") == lk.get("opname_id"):
                return
            line = await db.opname_lines.find_one({"opname_id": lk["opname_id"], "item_id": item}, {"_id": 0})
            if not line or line.get("counted") is None or not line.get("counted_at"):
                return
            txn = str(entry.get("txn_at") or "")
            cat = str(line["counted_at"])
            earlier = (txn < cat) if len(txn) > 10 else (_date10(txn) <= _date10(cat))
            if earlier:
                await db.opname_lines.update_one({"id": line["id"]}, {"$set": {
                    "needs_recount": True,
                    "recount_reason": f"Mutasi {entry.get('doc_type') or ''} {entry.get('doc_no') or ''} "
                                      f"(tgl {_date10(txn)}) diinput setelah penghitungan"}})
        except Exception:  # noqa: BLE001 — penanda tidak boleh menggagalkan transaksi lain
            return

    # ------------------------------------------------------------------ routes
    def replace_route(path, method, endpoint, **kw):
        for r in list(app.router.routes):
            if getattr(r, "path", None) == path and method in (getattr(r, "methods", set()) or set()):
                app.router.routes.remove(r)
        app.add_api_route(path, endpoint, methods=[method], tags=["opname-workflow"], **kw)

    async def ep_get(did: str, user=Depends(server.current_user)):
        need(user, "view", "Anda tidak memiliki izin melihat Stock Opname")
        await scope(did, user)
        return await payload(await get_doc(did), user)

    async def ep_lines(did: str, page: int = Query(1), page_size: int = Query(50), q: str | None = Query(None),
                       status: str | None = Query(None), user=Depends(server.current_user)):
        need(user, "view", "Anda tidak memiliki izin melihat Stock Opname")
        await scope(did, user)
        return await payload(await get_doc(did), user, page, page_size, q, status)

    async def ep_create(body: dict, user=Depends(server.current_user)):
        need(user, "create", "Anda tidak memiliki izin membuat Stock Opname")
        body = body or {}
        wh_id = body.get("warehouse_id")
        wh = await db.warehouses.find_one({"id": wh_id}, {"_id": 0}) if wh_id else None
        if not wh:
            raise HTTPException(400, "Gudang wajib dipilih")
        if wh.get("is_active") is False:
            raise HTTPException(400, f"Gudang {wh.get('name')} tidak aktif")
        mode = body.get("mode") or "live"
        if mode not in ("live", "freeze"):
            raise HTTPException(400, "Mode penghitungan tidak valid")
        scope_v = body.get("scope") or "all"
        if scope_v not in ("all", "division"):
            raise HTTPException(400, "Scope barang tidak valid")
        if scope_v == "division" and not body.get("division_id"):
            raise HTTPException(400, "Scope Per Divisi memerlukan Divisi")
        await assert_division(body.get("division_id"))
        legacy_active = await db.opname.find_one({"warehouse_id": wh_id, "status": {"$in": list(ACTIVE)}}, {"_id": 0, "no": 1})
        if legacy_active:
            raise HTTPException(409, f"Gudang {wh.get('name')} masih memiliki Stock Opname aktif {legacy_active.get('no')}. "
                                     "Selesaikan atau batalkan dokumen tersebut terlebih dahulu.")
        did = server.gid()
        no = await server.next_number("OPN")
        lk = {"id": lock_id(wh_id), "warehouse_id": wh_id, "opname_id": did, "opname_no": no, "mode": mode,
              "created_at": server.now_iso(), "created_by": user.get("email")}
        for attempt in (1, 2):
            try:
                await db.opname_locks.insert_one(lk)
                break
            except Exception as exc:  # noqa: BLE001
                if not MM._is_duplicate_pk(exc):
                    raise
                cur = await db.opname_locks.find_one({"id": lk["id"]}, {"_id": 0}) or {}
                owner = await db.opname.find_one({"id": cur.get("opname_id")}, {"_id": 0, "status": 1, "no": 1})
                # Basi HANYA bila pemilik sudah final, atau pemilik tidak ada DAN lock sudah lama (bukan sedang dibuat).
                age_ok = str(cur.get("created_at") or "") < _minutes_ago(server, 10)
                if attempt == 1 and ((owner and owner.get("status") not in ACTIVE + ("Posting",)) or (not owner and age_ok)):
                    await db.opname_locks.delete_many({"id": lk["id"], "opname_id": cur.get("opname_id")})  # lock basi
                    continue
                raise HTTPException(409, f"Gudang {wh.get('name')} masih memiliki Stock Opname aktif "
                                         f"{(owner or {}).get('no') or cur.get('opname_no')}.")
        draft_id = body.get("attachment_draft_id")
        claimed = False
        try:
            if draft_id:
                claimed = await TL.claim_draft(server, draft_id, user, did, module="opname")
            now = server.now_iso()
            doc = {"id": did, "no": no, "date": body.get("date") or now[:10], "warehouse_id": wh_id,
                   "division_id": body.get("division_id"), "mode": mode, "scope": scope_v, "notes": body.get("notes"),
                   "counter_name": (body.get("counter_name") or "").strip() or user.get("name"),
                   "status": "Counting", "workflow_version": WF_VERSION, "freeze_active": mode == "freeze",
                   "snapshot_at": now, "post_attempt": 1, "created_by": user.get("email"),
                   "created_by_name": user.get("name"), "created_at": now,
                   "history": [{"action": "create", "from": None, "to": "Counting", "by": user.get("email"),
                                "by_name": user.get("name"), "at": now, "reason": None}]}
            await db.opname.insert_one(doc)
            iws = await db.item_warehouse.find({"warehouse_id": wh_id}, {"_id": 0}).to_list(None)
            items = await items_by_id([r.get("item_id") for r in iws])
            recs = []
            for iw in iws:
                it = items.get(iw.get("item_id"))
                if not it:
                    continue
                if scope_v == "division" and it.get("division_id") != body.get("division_id"):
                    continue
                recs.append({"id": server.gid(), "opname_id": did, "item_id": iw["item_id"],
                             "snapshot": float(iw.get("current_stock") or 0), "counted": None})
            for i in range(0, len(recs), 500):
                await db.opname_lines.insert_many(recs[i:i + 500])
        except Exception:
            await db.opname_lines.delete_many({"opname_id": did})
            await db.opname.delete_one({"id": did})
            await db.opname_locks.delete_many({"opname_id": did})
            if claimed:
                await TL.release_draft(server, draft_id, did)
            raise
        if claimed:
            await TL.bind_draft_attachments(server, draft_id, did, user, module="opname")
        await server.audit(user, "create", "opname", did, no, after={"warehouse_id": wh_id, "mode": mode, "lines": len(recs)})
        return await payload(await get_doc(did), user, page=1)

    async def apply_counts(doc, raw_lines, user, source="manual"):
        did = doc["id"]
        if doc.get("status") not in EDITABLE:
            raise HTTPException(409, f"Hasil hitung tidak dapat diubah pada status {doc.get('status')}")
        by_id = {l["id"]: l for l in await all_lines(did)}
        items = await items_by_id([l.get("item_id") for l in by_id.values()])
        seen, updates = set(), []
        for raw in raw_lines or []:
            lid = (raw or {}).get("line_id") or (raw or {}).get("id")
            if not lid or lid in seen:
                raise HTTPException(400, "Baris Stock Opname tidak valid atau duplikat")
            seen.add(lid)
            old = by_id.get(lid)
            if not old:
                raise HTTPException(400, "Baris tidak berasal dari Stock Opname ini")
            it = items.get(old.get("item_id"), {})
            label = it.get("code") or it.get("name") or old.get("item_id")
            patch = {}
            if raw.get("clear"):
                patch.update({"counted": None, "counted_qty": None, "counted_uom_id": None, "conversion_factor": None,
                              "counted_at": None, "counted_by": None, "system_at_count": None, "needs_recount": False})
            elif "qty" in raw or "counted" in raw:
                uom_id = raw.get("uom_id") or it.get("base_uom_id")
                fx = factor_of(it, uom_id)
                if fx is None:
                    raise HTTPException(400, f"{label}: satuan hitung tidak valid untuk barang ini")
                if "qty" in raw:
                    if raw.get("qty") in (None, ""):
                        raise HTTPException(400, f"{label}: qty fisik kosong — gunakan 'kosongkan' bila belum dihitung")
                    qty = _f(raw.get("qty"))
                    base = None if qty is None else qty * fx
                else:
                    if raw.get("counted") in (None, ""):
                        continue  # format lama: counted kosong = tidak diubah
                    base = _f(raw.get("counted"))
                    qty = None if base is None else base / fx
                if base is None or base < -EPS:
                    raise HTTPException(400, f"{label}: qty fisik harus angka >= 0")
                patch.update({"counted": round(base, 6), "counted_qty": qty, "counted_uom_id": uom_id,
                              "conversion_factor": fx, "counted_at": server.now_iso(), "counted_by": user.get("email"),
                              "counted_by_name": user.get("name"), "needs_recount": False, "recount_reason": None})
                if doc.get("workflow_version") == WF_VERSION and doc.get("mode") == "live":
                    patch["system_at_count"] = float(await server.stock_balance(old["item_id"], doc["warehouse_id"]) or 0)
            if "reason" in raw:
                patch["reason"] = (str(raw.get("reason") or "").strip() or None)
            if "approved_unit_cost" in raw:
                ac = raw.get("approved_unit_cost")
                if ac in (None, ""):
                    if old.get("approved_unit_cost") is not None and not price_ok(user):
                        raise HTTPException(403, "Anda tidak memiliki izin mengubah Harga Satuan")
                    patch.update({"approved_unit_cost": None, "cost_reason": None})
                else:
                    if not (price_ok(user) and can(user, "edit")):
                        raise HTTPException(403, "Anda tidak memiliki izin mengisi Harga Satuan (Moving Average)")
                    acv = _f(ac)
                    if acv is None or acv <= 0:
                        raise HTTPException(400, f"{label}: Harga Satuan harus lebih dari 0")
                    cr = str(raw.get("cost_reason") or old.get("cost_reason") or "").strip()
                    if not cr:
                        raise HTTPException(400, f"{label}: alasan Harga Satuan wajib diisi")
                    patch.update({"approved_unit_cost": acv, "cost_reason": cr, "cost_set_by": user.get("email"),
                                  "cost_set_at": server.now_iso()})
            if patch:
                updates.append((old, patch))
        for old, patch in updates:
            await db.opname_lines.update_one({"id": old["id"], "opname_id": did}, {"$set": patch})
            if "approved_unit_cost" in patch and patch.get("approved_unit_cost") != old.get("approved_unit_cost"):
                # nilai harga TIDAK ditulis ke audit log (audit dapat dibaca user tanpa izin harga)
                await server.audit(user, "override_cost", "opname", did, doc.get("no"),
                                   before={"item_id": old["item_id"], "had_cost": old.get("approved_unit_cost") is not None},
                                   after={"item_id": old["item_id"], "cost_set": patch.get("approved_unit_cost") is not None})
        if updates:
            await server.audit(user, "count", "opname", did, doc.get("no"), after={"lines": len(updates), "source": source})
        return len(updates)

    async def ep_count(did: str, body: dict, user=Depends(server.current_user)):
        need(user, "edit", "Anda tidak memiliki izin mengisi hasil hitung")
        await scope(did, user)
        doc = await get_doc(did)
        body = body or {}
        n = await apply_counts(doc, body.get("lines"), user)
        hp = {k: body[k] for k in ("notes", "counter_name") if k in body and doc.get("status") in EDITABLE}
        if hp:
            await db.opname.update_one({"id": did}, {"$set": hp})
        out = await payload(await get_doc(did), user, page=1, page_size=1)
        out["saved"] = n
        return out

    async def ep_add_item(did: str, body: dict, user=Depends(server.current_user)):
        need(user, "edit", "Anda tidak memiliki izin menambah barang")
        await scope(did, user)
        doc = await get_doc(did)
        if doc.get("status") != "Counting":
            raise HTTPException(409, "Barang hanya dapat ditambahkan pada status Counting")
        item_id = (body or {}).get("item_id")
        reason = str((body or {}).get("reason") or "").strip()
        if not reason:
            raise HTTPException(400, "Alasan penambahan barang wajib diisi")
        it = await db.items.find_one({"id": item_id}, {"_id": 0}) if item_id else None
        if not it:
            raise HTTPException(404, "Barang tidak ditemukan di Master Barang")
        if it.get("is_active") is False:
            raise HTTPException(400, f"Barang {it.get('code')} tidak aktif")
        if doc.get("scope") == "division" and it.get("division_id") != doc.get("division_id"):
            raise HTTPException(400, "Barang berada di luar scope divisi Stock Opname ini")
        if await db.opname_lines.find_one({"opname_id": did, "item_id": item_id}, {"_id": 0, "id": 1}):
            raise HTTPException(409, f"Barang {it.get('code')} sudah ada di Stock Opname ini")
        rec = {"id": server.gid(), "opname_id": did, "item_id": item_id,
               "snapshot": float(await server.stock_balance(item_id, doc["warehouse_id"]) or 0), "counted": None,
               "added_manually": True, "added_reason": reason, "added_by": user.get("email"), "added_at": server.now_iso()}
        await db.opname_lines.insert_one(rec)
        await server.audit(user, "add_item", "opname", did, doc.get("no"), reason=reason, after={"item_id": item_id})
        return {"ok": True, "line_id": rec["id"]}

    async def rows_for(doc, user):
        return await enrich(doc, await all_lines(doc["id"]), user)

    async def transition(did, user, action, reason=None):
        doc = await get_doc(did)
        st = doc.get("status")
        reason = str(reason or "").strip() or None
        if action in ("review", "submit"):
            need(user, "edit", "Anda tidak memiliki izin mengajukan Stock Opname")
            allowed_from = ("Counting",) if action == "review" else EDITABLE
            if st not in allowed_from:
                raise HTTPException(409, f"Aksi tidak sah dari status {st}")
            if not doc.get("division_id"):
                raise HTTPException(400, "Divisi wajib diisi sebelum diajukan")
            errs = readiness_errors(await rows_for(doc, user))
            if errs:
                raise HTTPException(400, "; ".join(errs))
            to = "Review" if action == "review" else "Waiting Approval"
            extra = ({"counted_completed_by": user.get("email"), "counted_completed_at": server.now_iso()} if action == "review"
                     else {"reviewed_by": user.get("email"), "reviewed_by_name": user.get("name"), "reviewed_at": server.now_iso()})
            if action == "submit" and st == "Counting":
                extra.update({"counted_completed_by": user.get("email"), "counted_completed_at": server.now_iso()})
        elif action in ("return", "reject"):
            if not reason:
                raise HTTPException(400, "Alasan wajib diisi")
            if action == "reject" or st == "Waiting Approval":
                need(user, "post", "Anda tidak memiliki izin approval Stock Opname")
            else:
                need(user, "edit", "Anda tidak memiliki izin mengembalikan Stock Opname")
            allowed_from = ("Waiting Approval",) if action == "reject" else ("Review", "Waiting Approval")
            if st not in allowed_from:
                raise HTTPException(409, f"Aksi tidak sah dari status {st}")
            to = "Counting"
            extra = {"last_return": {"action": action, "reason": reason, "by": user.get("email"), "at": server.now_iso()}}
        elif action == "cancel":
            if not (can(user, "cancel") or can(user, "delete")):
                raise HTTPException(403, "Anda tidak memiliki izin membatalkan Stock Opname")
            if not reason:
                raise HTTPException(400, "Alasan pembatalan wajib diisi")
            if st not in ACTIVE:
                raise HTTPException(409, f"Stock Opname berstatus {st} tidak dapat dibatalkan")
            to = "Cancelled"
            extra = {"cancelled_by": user.get("email"), "cancelled_at": server.now_iso(), "cancel_reason": reason}
        else:
            raise HTTPException(400, "Aksi workflow tidak dikenal")
        won = await cas_doc({"id": did, "status": st}, {"$set": {"status": to}})
        if not won:
            raise HTTPException(409, "Status dokumen berubah. Muat ulang lalu coba lagi.")
        await history(did, doc, action, st, to, user, reason, extra)
        if to == "Cancelled":
            await release_lock(did)
        await server.audit(user, {"review": "review", "submit": "submit", "return": "return", "reject": "reject",
                                  "cancel": "cancel"}[action], "opname", did, doc.get("no"), reason=reason,
                           before={"status": st}, after={"status": to})
        return await payload(await get_doc(did), user, page=1, page_size=1)

    async def ep_workflow(did: str, body: dict, user=Depends(server.current_user)):
        await scope(did, user)
        return await transition(did, user, (body or {}).get("action"), (body or {}).get("reason"))

    async def ep_submit(did: str, user=Depends(server.current_user)):
        await scope(did, user)
        return await transition(did, user, "submit")

    # ---- posting
    async def prepare_effects(doc, lines, user, override=None):
        """Pra-validasi SELURUH baris sebelum satu pun movement ditulis. override: {line_id: patch} (koreksi)."""
        wh = doc["warehouse_id"]
        items = await items_by_id([l.get("item_id") for l in lines])
        iwm = await iw_map(wh)
        effects, errs = [], []
        for l in lines:
            l = {**l, **((override or {}).get(l["id"]) or {})}
            label = (items.get(l["item_id"]) or {}).get("code") or l["item_id"]
            if l.get("counted") is None:
                errs.append(f"{label}: belum dihitung")
                continue
            if l.get("needs_recount"):
                errs.append(f"{label}: perlu dihitung ulang")
                continue
            base_doc = {**doc, "status": "Waiting Approval"}  # hitung dari data hitung, bukan nilai posted lama
            c = compute_line(base_doc, l, iwm.get(l["item_id"]))
            delta = c["variance"]
            if abs(delta) > EPS and not str(l.get("reason") or "").strip():
                errs.append(f"{label}: alasan selisih wajib diisi")
            effects.append((l, c["system_qty"], delta, c))
        return effects, errs, iwm, items

    async def post_effects(doc, effects, user, key_suffix):
        wh, did, no = doc["warehouse_id"], doc["id"], doc.get("no")
        results = []
        for l, base, delta, c in effects:
            if abs(delta) <= EPS:
                results.append((l, base, 0.0, None, 0.0))
                continue
            sk = f"OPN::{did}::{l['id']}{key_suffix}"
            if delta > 0:
                uc, ov = await DW._resolve_in_cost(l["item_id"], wh, l.get("approved_unit_cost"), l.get("cost_reason") or l.get("reason"))
                res = await server.post_movement("Stock Opname Adjustment", no, did, l["item_id"], wh, delta, 0, user=user,
                                                 line_id=l["id"], unit_cost_in=uc, require_cost=True, source_key=sk,
                                                 txn_at=doc.get("date"), division_id=doc.get("division_id"))
                results.append((l, base, delta, res.get("unit_cost") or uc, float(res.get("value_in") or delta * uc)))
            else:
                res = await server.post_movement("Stock Opname Adjustment", no, did, l["item_id"], wh, 0, -delta, user=user,
                                                 line_id=l["id"], source_key=sk, txn_at=doc.get("date"),
                                                 division_id=doc.get("division_id"))
                vo = float(res.get("value_out") or 0)
                results.append((l, base, delta, res.get("unit_cost"), -vo))
        return results

    async def save_results(results, override=None):
        for l, base, delta, uc, val in results:
            patch = {"variance_base": base, "posted_delta": delta, "posted_unit_cost": uc, "posted_value": val}
            patch.update((override or {}).get(l["id"]) or {})
            await db.opname_lines.update_one({"id": l["id"]}, {"$set": patch})

    async def ep_post(did: str, user=Depends(server.current_user)):
        need(user, "post", "Anda tidak memiliki izin Approve & Posting Stock Opname")
        await scope(did, user)
        doc = await get_doc(did)
        if doc.get("status") == "Posted":
            raise HTTPException(409, "Stock Opname sudah diposting")
        stale = doc.get("status") == "Posting" and str(doc.get("posting_started_at") or "") < _minutes_ago(server, 10)
        if doc.get("status") == "Gagal Posting" or stale:
            # Pemulihan: posting sebelumnya terhenti (crash/rollback gagal). Bersihkan sisa movement dokumen lalu
            # kembalikan ke Waiting Approval dengan attempt baru -> retry tidak pernah double/terlewat.
            await server.reverse_document_valuation(did, user=user, reason="post_failed", block_negative=False)
            won = await cas_doc({"id": did, "status": doc.get("status")}, {"$set": {
                "status": "Waiting Approval", "post_attempt": int(doc.get("post_attempt") or 1) + 1,
                "recovered_at": server.now_iso()}})
            if not won:
                raise HTTPException(409, "Status dokumen berubah. Muat ulang lalu coba lagi.")
            await server.audit(user, "post_recover", "opname", did, doc.get("no"), after={"from": doc.get("status")})
            doc = await get_doc(did)
        if doc.get("status") != "Waiting Approval":
            raise HTTPException(409, f"Posting hanya dari status Waiting Approval (status saat ini {doc.get('status')})")
        lines = await all_lines(did)
        effects, errs, iwm, items = await prepare_effects(doc, lines, user)
        if errs:
            raise HTTPException(400, "; ".join(errs[:20]))
        for l, base, delta, c in effects:  # stok tidak boleh negatif & biaya surplus harus valid SEBELUM menulis apa pun
            if delta < -EPS and c["current_stock"] + delta < -1e-6:
                raise HTTPException(409, f"Posting opname {(items.get(l['item_id']) or {}).get('code')} akan membuat stok negatif ({c['current_stock'] + delta:g})")
            if delta > EPS:
                await DW._resolve_in_cost(l["item_id"], doc["warehouse_id"], l.get("approved_unit_cost"), l.get("cost_reason") or l.get("reason"))
        claimed = await cas_doc({"id": did, "status": "Waiting Approval"},
                                {"$set": {"status": "Posting", "posting_started_at": server.now_iso(),
                                          "posting_by": user.get("email")}})
        if not claimed:
            raise HTTPException(409, "Stock Opname sedang/sudah diproses posting oleh pengguna lain")
        attempt = int(doc.get("post_attempt") or 1)
        suffix = "" if attempt <= 1 else f"::r{attempt}"
        token = POSTING_CTX.set(did)
        try:
            results = await post_effects(doc, effects, user, suffix)
        except Exception as exc:
            ok = True
            try:
                await server.reverse_document_valuation(did, user=user, reason="post_failed", block_negative=True)
            except Exception:  # noqa: BLE001
                ok = False
            msg = getattr(exc, "detail", None) or str(exc) or exc.__class__.__name__
            await db.opname.update_one({"id": did}, {"$set": {
                "status": "Waiting Approval" if ok else "Gagal Posting", "post_attempt": attempt + 1,
                "last_post_error": str(msg)[:500], "last_post_error_at": server.now_iso()}})
            await server.audit(user, "post_failed", "opname", did, doc.get("no"), reason=str(msg)[:300],
                               after={"rolled_back": ok, "attempt": attempt})
            if isinstance(exc, HTTPException):
                raise HTTPException(exc.status_code, f"Posting dibatalkan seluruhnya (tidak ada perubahan stok): {msg}")
            raise HTTPException(500, f"Posting dibatalkan seluruhnya (tidak ada perubahan stok): {msg}")
        finally:
            POSTING_CTX.reset(token)
        await save_results(results)
        now = server.now_iso()
        await db.opname.update_one({"id": did}, {"$set": {
            "status": "Posted", "approved_by": user.get("name"), "approved_by_email": user.get("email"),
            "approved_at": now, "posted_at": now, "last_post_error": None}})
        await history(did, {**doc, "history": (await get_doc(did)).get("history")}, "approve_post", "Waiting Approval", "Posted", user)
        await release_lock(did)
        await server.audit(user, "approve", "opname", did, doc.get("no"),
                           after={"lines": len(results), "moved": sum(1 for r in results if abs(r[2]) > EPS)})
        return await payload(await get_doc(did), user, page=1, page_size=1)

    # ---- edit (koreksi Posted via reversal) & delete
    import transaction_mutation_layer as TM
    tx_put = next((r for r in app.router.routes if getattr(r, "path", None) == "/api/transactions/{module}/{did}"
                   and "PUT" in (getattr(r, "methods", set()) or set())), None)
    tx_del = next((r for r in app.router.routes if getattr(r, "path", None) == "/api/transactions/{module}/{did}"
                   and "DELETE" in (getattr(r, "methods", set()) or set())), None)
    orig_put = tx_put.endpoint if tx_put else None
    orig_del = tx_del.endpoint if tx_del else None

    async def correct_posted(doc, body, user):
        did = doc["id"]
        need(user, "edit", "Anda tidak memiliki izin mengoreksi Stock Opname")
        if body.get("warehouse_id") not in (None, doc.get("warehouse_id")):
            raise HTTPException(400, "Gudang Stock Opname dikunci setelah snapshot")
        cap = await TM._capability(server, "opname", did, user)
        if cap.get("blockers"):
            raise HTTPException(409, cap.get("reason") or "Stock Opname tidak dapat dikoreksi")
        lines = await all_lines(did)
        by_id = {l["id"]: l for l in lines}
        items = await items_by_id([l.get("item_id") for l in lines])
        raw = body.get("lines") or []
        ids = [r.get("id") or r.get("line_id") for r in raw]
        if len(ids) != len(set(ids)) or set(ids) != set(by_id):
            raise HTTPException(400, "Koreksi Stock Opname Posted harus memuat seluruh baris yang sama (tidak boleh menambah/menghapus/mengganti barang)")
        override = {}
        for r in raw:
            old = by_id[r.get("id") or r.get("line_id")]
            if r.get("item_id") not in (None, old.get("item_id")):
                raise HTTPException(400, "Barang Stock Opname Posted tidak boleh diganti")
            if r.get("snapshot") is not None and abs(float(r["snapshot"]) - float(old.get("snapshot") or 0)) > EPS:
                raise HTTPException(400, "Snapshot Stock Opname Posted tidak boleh diubah")
            it = items.get(old["item_id"], {})
            if "qty" in r:
                fx = factor_of(it, r.get("uom_id") or it.get("base_uom_id"))
                q = _f(r.get("qty"))
                base = None if (q is None or fx is None) else q * fx
            else:
                base = _f(r.get("counted"))
            if base is None or base < -EPS:
                raise HTTPException(400, f"{it.get('code')}: qty hitung wajib diisi (>= 0) pada koreksi")
            p = {"counted": round(base, 6), "reason": (str(r.get("reason") if "reason" in r else old.get("reason") or "").strip() or None),
                 "needs_recount": False}
            if r.get("approved_unit_cost") not in (None, "") and _f(r.get("approved_unit_cost")) != _f(old.get("approved_unit_cost")):
                if not price_ok(user):
                    raise HTTPException(403, "Anda tidak memiliki izin mengisi Harga Satuan")
                if not str(r.get("cost_reason") or "").strip():
                    raise HTTPException(400, f"{it.get('code')}: alasan Harga Satuan wajib diisi")
                p.update({"approved_unit_cost": _f(r.get("approved_unit_cost")), "cost_reason": str(r.get("cost_reason")).strip()})
            override[old["id"]] = p
        # basis selisih tetap basis posting awal (snapshot/stok sistem saat hitung) — dibaca via variance_base
        fixed = [{**l, "snapshot": l.get("variance_base") if l.get("variance_base") is not None else l.get("snapshot"),
                  "system_at_count": None} for l in lines]
        legacy_doc = {**doc, "workflow_version": None}  # paksa basis = snapshot(=variance_base)
        new_eff, errs, iwm, _ = await prepare_effects(legacy_doc, fixed, user, override)
        if errs:
            raise HTTPException(400, "; ".join(errs[:20]))
        for l, base, delta, c in new_eff:
            old_delta = float(by_id[l["id"]].get("posted_delta") if by_id[l["id"]].get("posted_delta") is not None
                              else (float(by_id[l["id"]].get("counted") or 0) - float(by_id[l["id"]].get("snapshot") or 0)))
            after = c["current_stock"] - old_delta + delta
            if after < -1e-6:
                raise HTTPException(409, f"Koreksi {(items.get(l['item_id']) or {}).get('code')} akan membuat stok negatif ({after:g})")
        claimed = await cas_doc({"id": did, "status": "Posted"}, {"$set": {"status": "Posting"}})
        if not claimed:
            raise HTTPException(409, "Stock Opname sedang diproses pengguna lain")
        n = int(doc.get("correction_count") or 0) + 1
        old_eff = [(l, float(l.get("variance_base") if l.get("variance_base") is not None else l.get("snapshot") or 0),
                    float(l.get("posted_delta") if l.get("posted_delta") is not None else float(l.get("counted") or 0) - float(l.get("snapshot") or 0)),
                    {}) for l in lines]
        token = POSTING_CTX.set(did)
        try:
            await server.reverse_document_valuation(did, user=user, reason="opname_correction", block_negative=True)
            try:
                results = await post_effects(doc, new_eff, user, f"::c{n}")
            except Exception as exc:
                await server.reverse_document_valuation(did, user=user, reason="correction_failed", block_negative=False)
                restored = await post_effects(doc, old_eff, user, f"::c{n}x")
                await save_results(restored)
                await db.opname.update_one({"id": did}, {"$set": {"status": "Posted", "correction_count": n}})
                raise HTTPException(getattr(exc, "status_code", 500),
                                    f"Koreksi dibatalkan, efek posting sebelumnya dipulihkan: {getattr(exc, 'detail', exc)}")
        finally:
            POSTING_CTX.reset(token)
        await save_results(results, override)
        await db.opname.update_one({"id": did}, {"$set": {"status": "Posted", "correction_count": n,
                                                          "notes": body.get("notes", doc.get("notes")),
                                                          "updated_by": user.get("email"), "updated_at": server.now_iso()}})
        await history(did, await get_doc(did), "correction", "Posted", "Posted", user, body.get("reason"))
        safe = [{k: v for k, v in l.items() if k not in PRICE_KEYS} for l in lines]
        await server.audit(user, "edit", "opname", did, doc.get("no"), before={"lines": safe},
                           after={"correction": n, "lines": len(results)})
        return {"ok": True, "id": did, "no": doc.get("no")}

    async def ep_tx_put(module: str, did: str, body: dict, user=Depends(server.current_user)):
        if module != "opname":
            if await any_freeze():  # edit = reversal + posting ulang: cek gudang lama & baru sebelum menulis apa pun
                found = set()

                def walk(o):
                    if isinstance(o, dict):
                        for kk, vv in o.items():
                            if isinstance(vv, str) and kk.endswith("warehouse_id"):
                                found.add(vv)
                            else:
                                walk(vv)
                    elif isinstance(o, list):
                        for vv in o:
                            walk(vv)
                walk(body or {})
                rows = await db.valuation_ledger.find({"doc_id": did, "is_reversal": {"$ne": True}, "reversed": {"$ne": True}},
                                                      {"_id": 0, "warehouse_id": 1}).to_list(None)
                found |= {r.get("warehouse_id") for r in rows if isinstance(r.get("warehouse_id"), str)}
                await assert_not_frozen(found)
            return await orig_put(module, did, body, user)
        await scope(did, user)
        doc = await get_doc(did)
        body = body or {}
        if doc.get("status") == "Posted":
            return await correct_posted(doc, body, user)
        need(user, "edit", "Anda tidak memiliki izin mengubah Stock Opname")
        if doc.get("status") not in EDITABLE:
            raise HTTPException(409, f"Stock Opname berstatus {doc.get('status')} tidak dapat diubah")
        if body.get("warehouse_id") not in (None, "", doc.get("warehouse_id")):
            raise HTTPException(400, "Gudang Stock Opname dikunci setelah snapshot. Batalkan dokumen dan buat Stock Opname baru.")
        for k in ("mode", "scope"):
            if body.get(k) not in (None, "", doc.get(k)):
                raise HTTPException(400, "Mode/Scope tidak dapat diubah setelah snapshot")
        patch = {k: body[k] for k in ("date", "notes", "counter_name") if k in body}
        if "division_id" in body and body.get("division_id"):
            await assert_division(body["division_id"])
            patch["division_id"] = body["division_id"]
        if patch:
            await db.opname.update_one({"id": did}, {"$set": {**patch, "updated_by": user.get("email"), "updated_at": server.now_iso()}})
        if body.get("lines"):
            await apply_counts(await get_doc(did), body.get("lines"), user, source="edit")
        await server.audit(user, "edit", "opname", did, doc.get("no"), after=patch)
        return {"ok": True, "id": did, "no": doc.get("no")}

    async def ep_tx_del(module: str, did: str, user=Depends(server.current_user)):
        if module != "opname":
            return await orig_del(module, did, user)
        doc = await db.opname.find_one({"id": did}, {"_id": 0}) or {}
        if doc.get("status") == "Posting":
            raise HTTPException(409, "Stock Opname sedang diproses posting")
        out = await orig_del(module, did, user)
        await db.opname_locks.delete_many({"opname_id": did})
        return out

    # ---- Excel template/import (dokumen yang sama) & data cetak
    HEAD_ROW = 6
    COLS = ["No", "Kode Barang", "Nama Barang", "Satuan", "Stok Sistem", "Stok Fisik", "Alasan Selisih", "ID Baris"]

    async def ep_template(did: str, user=Depends(server.current_user)):
        need(user, "view", "Anda tidak memiliki izin melihat Stock Opname")
        await scope(did, user)
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        doc = await get_doc(did)
        p = await payload(doc, user)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Hitung Fisik"
        meta = [("No. Opname", doc.get("no")), ("Gudang", f"{p.get('warehouse_code') or ''} — {p.get('warehouse_name') or ''}"),
                ("Tanggal", doc.get("date")), ("Mode", doc.get("mode")),
                ("Petunjuk", "Isi kolom Stok Fisik (satuan sesuai kolom Satuan). Kosong = belum dihitung; 0 = fisik nol. "
                             "Alasan wajib bila ada selisih. Jangan ubah Kode Barang / ID Baris.")]
        for i, (k, v) in enumerate(meta, start=1):
            ws.cell(row=i, column=1, value=k).font = Font(bold=True)
            ws.cell(row=i, column=2, value=v)
        for j, h in enumerate(COLS, start=1):
            c = ws.cell(row=HEAD_ROW, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(horizontal="center")
        for i, r in enumerate(p["lines"], start=1):
            ws.append([i, r.get("item_code"), r.get("item_name"), r.get("unit"), r.get("system_qty"),
                       None if r.get("counted") is None else r.get("counted"), r.get("reason"), r.get("id")])
        for col, w in zip("ABCDEFGH", (6, 18, 40, 10, 12, 12, 36, 38)):
            ws.column_dimensions[col].width = w
        ws.column_dimensions["H"].hidden = True
        ws.freeze_panes = ws.cell(row=HEAD_ROW + 1, column=1)
        buf = io.BytesIO()
        wb.save(buf)
        fname = re.sub(r"[^A-Za-z0-9_-]+", "_", str(doc.get("no"))) + "_lembar_hitung.xlsx"
        return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{fname}"'})

    async def parse_import(doc, data, user):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception:  # noqa: BLE001
            raise HTTPException(400, "File bukan Excel (.xlsx) yang valid")
        ws = wb.worksheets[0]
        p = await payload(doc, user)
        errors, rows = [], []
        meta = {str(ws.cell(row=i, column=1).value or "").strip(): ws.cell(row=i, column=2).value for i in range(1, HEAD_ROW)}
        if str(meta.get("No. Opname") or "").strip() != str(doc.get("no")):
            errors.append(f"Nomor Opname di file ({meta.get('No. Opname') or '-'}) tidak sama dengan dokumen {doc.get('no')}")
        wh_txt = str(meta.get("Gudang") or "")
        if p.get("warehouse_code") and not wh_txt.startswith(str(p.get("warehouse_code"))):
            errors.append(f"Gudang di file ({wh_txt or '-'}) berbeda dengan gudang dokumen ({p.get('warehouse_code')})")
        header = [str(ws.cell(row=HEAD_ROW, column=j).value or "").strip() for j in range(1, len(COLS) + 1)]
        idx = {h: j for j, h in enumerate(header)}
        for need_col in ("Kode Barang", "Stok Fisik"):
            if need_col not in idx:
                errors.append(f"Kolom '{need_col}' tidak ditemukan pada baris {HEAD_ROW}")
        if errors and any("Kolom" in e for e in errors):
            return {"errors": errors, "rows": [], "applied": 0}
        by_code = {str(r.get("item_code") or "").strip().upper(): r for r in p["lines"]}
        by_line = {r["id"]: r for r in p["lines"]}
        items_all = {}
        seen = {}
        skipped = 0
        um = await uom_labels()
        for rn, vals in enumerate(ws.iter_rows(min_row=HEAD_ROW + 1, values_only=True), start=HEAD_ROW + 1):
            vals = list(vals) + [None] * len(COLS)
            get = lambda h: vals[idx[h]] if h in idx else None  # noqa: E731
            code = str(get("Kode Barang") or "").strip()
            qty_raw = get("Stok Fisik")
            if not code and qty_raw in (None, "") and not get("ID Baris"):
                continue
            if not code:
                errors.append(f"Baris {rn}: Kode Barang kosong")
                continue
            line = by_line.get(str(get("ID Baris") or "").strip()) or by_code.get(code.upper())
            if not line:
                if not items_all:
                    items_all = {str(i.get("code") or "").upper(): i for i in await db.items.find({}, {"_id": 0, "code": 1}).to_list(None)}
                errors.append(f"Baris {rn}: kode {code} " + ("tidak ada di snapshot Stock Opname ini (gunakan Tambah Barang)"
                                                             if code.upper() in items_all else "tidak ditemukan di Master Barang"))
                continue
            if str(line.get("item_code") or "").strip().upper() != code.upper():
                errors.append(f"Baris {rn}: kode {code} tidak cocok dengan ID baris dokumen")
                continue
            if line["id"] in seen:
                errors.append(f"Baris {rn}: barang {code} duplikat (sudah ada di baris {seen[line['id']]})")
                continue
            seen[line["id"]] = rn
            if qty_raw in (None, ""):
                skipped += 1
                continue
            qty = _f(qty_raw)
            if qty is None or qty < 0:
                errors.append(f"Baris {rn}: Stok Fisik '{qty_raw}' harus angka >= 0")
                continue
            sat = str(get("Satuan") or "").strip()
            uom_id = line.get("base_uom_id")
            if sat and sat.upper() != str(line.get("unit") or "").upper():
                it = (await items_by_id([line["item_id"]])).get(line["item_id"], {})
                match = [r.get("uom_id") for r in it.get("uoms") or [] if str(um.get(r.get("uom_id")) or "").upper() == sat.upper()]
                if not match:
                    errors.append(f"Baris {rn}: satuan '{sat}' tidak valid untuk {code}")
                    continue
                uom_id = match[0]
            reason = str(get("Alasan Selisih") or "").strip()
            rows.append({"row": rn, "line_id": line["id"], "item_code": line.get("item_code"), "item_name": line.get("item_name"),
                         "qty": qty, "uom_id": uom_id, "unit": sat or line.get("unit"), "reason": reason,
                         "previous": line.get("counted"), "system_qty": line.get("system_qty")})
        return {"errors": errors, "rows": rows, "skipped_blank": skipped}

    async def ep_import(did: str, request: Request, file: UploadFile = File(...), mode: str = Query("preview"),
                        user=Depends(server.current_user)):
        need(user, "edit", "Anda tidak memiliki izin import hasil hitung")
        await scope(did, user)
        doc = await get_doc(did)
        if doc.get("status") not in EDITABLE:
            raise HTTPException(409, f"Stock Opname berstatus {doc.get('status')} tidak dapat diubah lewat import")
        data = await file.read()
        if len(data) > 10 * 1024 * 1024:
            raise HTTPException(400, "Ukuran file import maksimal 10 MB")
        res = await parse_import(doc, data, user)
        res["mode"] = mode
        res["valid_rows"] = len(res["rows"])
        if mode != "commit":
            return res
        if res["errors"]:
            raise HTTPException(400, {"msg": "Import dibatalkan: perbaiki error terlebih dahulu", "errors": res["errors"][:200]})
        raw = [{"line_id": r["line_id"], "qty": r["qty"], "uom_id": r["uom_id"], **({"reason": r["reason"]} if r["reason"] else {})}
               for r in res["rows"]]
        res["applied"] = await apply_counts(doc, raw, user, source="excel_import") if raw else 0
        res["summary"] = (await payload(await get_doc(did), user, page=1, page_size=1))["summary"]
        return res

    async def ep_print(did: str, user=Depends(server.current_user)):
        need(user, "view", "Anda tidak memiliki izin melihat Stock Opname")
        await scope(did, user)
        out = await payload(await get_doc(did), user)
        out["printed_by"] = user.get("name")
        out["printed_at"] = server.now_iso()
        return out

    replace_route("/api/opname/{did}", "GET", ep_get)
    replace_route("/api/opname/{did}/lines", "GET", ep_lines)
    replace_route("/api/opname", "POST", ep_create)
    replace_route("/api/opname/{did}/count", "PUT", ep_count)
    replace_route("/api/opname/{did}/submit", "POST", ep_submit)
    replace_route("/api/opname/{did}/post", "POST", ep_post)
    replace_route("/api/opname/{did}/workflow", "POST", ep_workflow)
    replace_route("/api/opname/{did}/add-item", "POST", ep_add_item)
    replace_route("/api/opname/{did}/template.xlsx", "GET", ep_template)
    replace_route("/api/opname/{did}/import", "POST", ep_import)
    replace_route("/api/opname/{did}/print-data", "GET", ep_print)
    if orig_put:
        replace_route("/api/transactions/{module}/{did}", "PUT", ep_tx_put)
    if orig_del:
        replace_route("/api/transactions/{module}/{did}", "DELETE", ep_tx_del)
