"""Detailed item-level MRO traceability reporting.

Keeps one row per MRO + item while combining multiple downstream RO/PO/DO/MI
references into arrays. Quantities stay in the canonical/base inventory UOM.
PO monetary values are aggregated per row. DO value is calculated from PO DPP per
base unit multiplied by the quantity actually received.
"""
from io import BytesIO

from fastapi import Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def _find_route(app, path, method):
    for route in list(app.router.routes):
        if getattr(route, "path", None) == path and method in (getattr(route, "methods", set()) or set()):
            return route
    return None


def _add_ref(bucket: dict, doc: dict | None):
    if not doc or not doc.get("id"):
        return
    bucket[doc["id"]] = {
        "id": doc["id"],
        "no": doc.get("no"),
        "date": doc.get("date"),
    }


def _refs(bucket: dict):
    return sorted(
        bucket.values(),
        key=lambda x: ((x.get("date") or ""), (x.get("no") or "")),
    )


def _f(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _join_refs(refs, key):
    values = []
    for ref in refs or []:
        value = ref.get(key)
        if value:
            values.append(str(value)[:10] if key == "date" else str(value))
    return "\n".join(values)


import contextvars

_CHUNK = 5000
_LAST_BATCH: contextvars.ContextVar = contextvars.ContextVar("trace_last_batch", default=None)


class TraceBatch:
    """P7 — pengganti N+1 query traceability: data dimuat batch (`$in`) lalu dibaca per kunci dari memori.

    `find(coll, field, key, extra)` == `db[coll].find({field: key, **extra}).to_list(None)` (urutan dasar created_at,pk per
    kunci sama karena satu kunci selalu berada dalam satu chunk); `one(coll, id)` == `find_one({"id": id})` (dokumen
    pertama urutan dasar). Kunci yang belum dimuat / bukan string -> query asli (hasil identik, hanya lebih lambat)."""

    def __init__(self, server):
        self.db = server.db
        self.groups = {}

    @staticmethod
    def _tag(coll, field, extra):
        return coll, field, tuple(sorted((extra or {}).items()))

    async def load(self, coll, field, keys, extra=None):
        tag = self._tag(coll, field, extra)
        loaded, data = self.groups.setdefault(tag, (set(), {}))
        todo = [k for k in dict.fromkeys(keys) if isinstance(k, str) and k not in loaded]
        for i in range(0, len(todo), _CHUNK):
            part = todo[i:i + _CHUNK]
            pset = set(part)
            for d in await getattr(self.db, coll).find({field: {"$in": part}, **(extra or {})}, {"_id": 0}).to_list(None):
                v = d.get(field)
                for k in ([v] if isinstance(v, str) else [x for x in v if x in pset] if isinstance(v, list) else []):
                    if k in pset:
                        data.setdefault(k, []).append(d)
            loaded.update(part)

    async def find(self, coll, field, key, extra=None, cap=None):
        loaded, data = self.groups.get(self._tag(coll, field, extra), (set(), {}))
        if isinstance(key, str) and key in loaded:
            rows = data.get(key, [])
            return list(rows[:cap] if cap else rows)
        return await getattr(self.db, coll).find({field: key, **(extra or {})}, {"_id": 0}).to_list(cap)

    async def one(self, coll, key):
        loaded, data = self.groups.get(self._tag(coll, "id", None), (set(), {}))
        if isinstance(key, str) and key in loaded:
            rows = data.get(key)
            return rows[0] if rows else None
        return await getattr(self.db, coll).find_one({"id": key}, {"_id": 0})


async def trace_batch(server, mro_ids, line_filter=None):
    """Muat seluruh rantai MRO line -> RO -> PO -> DO / MI untuk `mro_ids` dalam ±10 query batch (bukan per baris)."""
    b = TraceBatch(server)
    await b.load("mro_lines", "mro_id", mro_ids)
    lines = [ln for k in mro_ids for ln in (b.groups[b._tag("mro_lines", "mro_id", None)][1].get(k) or [])]
    lids = [ln.get("id") for ln in lines]
    await b.load("allocations", "source_line_id", lids, {"target_type": "ro"})
    await b.load("allocations", "source_line_id", lids, {"target_type": "mi"})
    ro_a = [a for k in lids for a in await b.find("allocations", "source_line_id", k, {"target_type": "ro"})]
    mi_a = [a for k in lids for a in await b.find("allocations", "source_line_id", k, {"target_type": "mi"})]
    await b.load("allocations", "source_line_id", [a.get("target_line_id") for a in ro_a], {"target_type": "po"})
    po_a = [a for r in ro_a for a in await b.find("allocations", "source_line_id", r.get("target_line_id"), {"target_type": "po"})]
    await b.load("allocations", "source_line_id", [a.get("target_line_id") for a in po_a], {"target_type": "do"})
    do_a = [a for p in po_a for a in await b.find("allocations", "source_line_id", p.get("target_line_id"), {"target_type": "do"})]
    await b.load("ro", "id", [a.get("target_doc_id") for a in ro_a])
    await b.load("po", "id", [a.get("target_doc_id") for a in po_a])
    await b.load("po_lines", "id", [a.get("target_line_id") for a in po_a])
    await b.load("do", "id", [a.get("target_doc_id") for a in do_a])
    await b.load("mi", "id", [a.get("target_doc_id") for a in mi_a])
    return b


async def _build_rows(server, date_from=None, date_to=None, user=None):
    """Baris MRO Traceability (dipakai endpoint lama JSON/Excel DAN Pusat Laporan -> satu sumber; scope divisi &
    nilai finansial dipasang oleh report_control_scope_layer/report_financial_fix yang membungkus fungsi ini).
    Periode = tanggal MRO bisnis WIB inklusif. Tanpa batas jumlah MRO/baris (tidak ada pemotongan diam-diam)."""
    from reporting.scope import local_day

    show_financial = bool(user and server.has_perm(user, "view_purchase_price"))
    mros = await server.db.mro.find({}, {"_id": 0}).sort("date", -1).to_list(None)
    if date_from or date_to:
        mros = [m for m in mros if (not date_from or local_day(m.get("date")) >= date_from)
                and (not date_to or local_day(m.get("date")) <= date_to)]
    items = {i["id"]: i for i in await server.db.items.find({}, {"_id": 0}).to_list(None)}
    categories = {
        c["id"]: c for c in await server.db.item_categories.find({}, {"_id": 0}).to_list(None)
    }

    out = []
    B = await trace_batch(server, [mh["id"] for mh in mros])  # P7: rantai dimuat batch (dulu ±6 query per baris)
    for mh in mros:
        lines = await B.find("mro_lines", "mro_id", mh["id"])
        grouped = {}

        for line in lines:
            item_id = line.get("item_id")
            it = items.get(item_id, {})
            category_name = (
                categories.get(it.get("category_id"), {}).get("name")
                or it.get("category")
                or "-"
            )
            row = grouped.setdefault(item_id, {
                "mro_id": mh["id"],
                "mro_no": mh.get("no"),
                "mro_date": mh.get("date"),
                "item_id": item_id,
                "item_code": it.get("code"),
                "item_name": it.get("name"),
                "category": category_name,
                "unit": it.get("unit") or line.get("unit"),
                "request": 0.0,
                "qty_ro": 0.0,
                "qty_po": 0.0,
                "qty_received": 0.0,
                "qty_mi": 0.0,
                "po_gross": 0.0 if show_financial else None,
                "po_discount": 0.0 if show_financial else None,
                "po_dpp": 0.0 if show_financial else None,
                "po_tax": 0.0 if show_financial else None,
                "po_total": 0.0 if show_financial else None,
                "do_total": 0.0 if show_financial else None,
                "mro_refs": {},
                "ro_refs": {},
                "po_refs": {},
                "do_refs": {},
                "mi_refs": {},
            })

            row["request"] += _f(line.get("qty"))
            _add_ref(row["mro_refs"], mh)

            ro_allocs = await B.find("allocations", "source_line_id", line["id"], {"target_type": "ro"})
            row["qty_ro"] += sum(_f(a.get("qty")) for a in ro_allocs)

            for ra in ro_allocs:
                rh = await B.one("ro", ra.get("target_doc_id"))
                _add_ref(row["ro_refs"], rh)

                po_allocs = await B.find("allocations", "source_line_id", ra.get("target_line_id"), {"target_type": "po"})
                for pa in po_allocs:
                    allocated_po_qty = _f(pa.get("qty"))
                    row["qty_po"] += allocated_po_qty
                    ph = await B.one("po", pa.get("target_doc_id"))
                    _add_ref(row["po_refs"], ph)

                    po_line = await B.one("po_lines", pa.get("target_line_id")) or {}
                    po_line_qty = _f(po_line.get("qty"))
                    gross_line = po_line_qty * _f(po_line.get("price"))
                    discount_line = _f(po_line.get("discount"))
                    dpp_line = gross_line - discount_line
                    tax_line = dpp_line * _f(po_line.get("tax")) / 100.0
                    total_line = _f(po_line.get("total")) if po_line.get("total") is not None else dpp_line + tax_line
                    share = (allocated_po_qty / po_line_qty) if po_line_qty > 0 else 0.0

                    if show_financial:
                        row["po_gross"] += gross_line * share
                        row["po_discount"] += discount_line * share
                        row["po_dpp"] += dpp_line * share
                        row["po_tax"] += tax_line * share
                        row["po_total"] += total_line * share

                    do_allocs = await B.find("allocations", "source_line_id", pa.get("target_line_id"), {"target_type": "do"})
                    unit_dpp = (dpp_line / po_line_qty) if po_line_qty > 0 else 0.0
                    for da in do_allocs:
                        received_qty = _f(da.get("qty"))
                        row["qty_received"] += received_qty
                        if show_financial:
                            row["do_total"] += unit_dpp * received_qty
                        dh = await B.one("do", da.get("target_doc_id"))
                        _add_ref(row["do_refs"], dh)

            mi_allocs = await B.find("allocations", "source_line_id", line["id"], {"target_type": "mi"})
            row["qty_mi"] += sum(_f(a.get("qty")) for a in mi_allocs)
            for ma in mi_allocs:
                mih = await B.one("mi", ma.get("target_doc_id"))
                _add_ref(row["mi_refs"], mih)

        for row in grouped.values():
            row["outstanding"] = max(0.0, row["request"] - row["qty_mi"])
            row["status"] = (
                "Completed" if row["request"] > 0 and row["qty_mi"] >= row["request"]
                else "Partial" if row["qty_mi"] > 0
                else "Open"
            )
            for key in ("mro_refs", "ro_refs", "po_refs", "do_refs", "mi_refs"):
                row[key] = _refs(row[key])
            out.append(row)

    out_rows = out
    _LAST_BATCH.set(B)  # dipakai ulang wrapper report_financial_fix (request sama) -> tanpa query ulang
    return out_rows


def _excel(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "MRO Traceability"
    headers = [
        "Kode Barang", "Nama Barang", "Kategori", "Satuan",
        "Tanggal MRO", "No. MRO", "Qty MRO",
        "Tanggal RO", "No. RO", "Qty RO",
        "Tanggal PO", "No. PO", "Qty PO",
        "Nilai PO", "Diskon PO", "DPP PO", "Pajak PO", "Total PO",
        "Tanggal DO", "No. DO", "Qty DO", "Total DO (DPP x Qty Diterima)",
        "Tanggal MI", "No. MI", "Qty MI", "Outstanding", "Status",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        ws.append([
            row.get("item_code"), row.get("item_name"), row.get("category"), row.get("unit"),
            _join_refs(row.get("mro_refs"), "date"), _join_refs(row.get("mro_refs"), "no"), row.get("request"),
            _join_refs(row.get("ro_refs"), "date"), _join_refs(row.get("ro_refs"), "no"), row.get("qty_ro"),
            _join_refs(row.get("po_refs"), "date"), _join_refs(row.get("po_refs"), "no"), row.get("qty_po"),
            row.get("po_gross"), row.get("po_discount"), row.get("po_dpp"), row.get("po_tax"), row.get("po_total"),
            _join_refs(row.get("do_refs"), "date"), _join_refs(row.get("do_refs"), "no"), row.get("qty_received"), row.get("do_total"),
            _join_refs(row.get("mi_refs"), "date"), _join_refs(row.get("mi_refs"), "no"), row.get("qty_mi"), row.get("outstanding"), row.get("status"),
        ])

    money_cols = {14, 15, 16, 17, 18, 22}
    qty_cols = {7, 10, 13, 21, 25, 26}
    for row in ws.iter_rows(min_row=2):
        for idx, cell in enumerate(row, start=1):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if idx in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'
            elif idx in qty_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.####'

    widths = [15, 28, 18, 12, 14, 24, 12, 14, 24, 12, 14, 24, 12, 16, 16, 16, 16, 16, 14, 24, 12, 22, 14, 24, 12, 14, 14]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def install(server):
    app = server.app
    path = "/api/reports/mro-traceability"
    old = _find_route(app, path, "GET")
    if old and old in app.router.routes:
        app.router.routes.remove(old)

    async def mro_traceability_detail(
        date_from: str | None = Query(None),
        date_to: str | None = Query(None),
        user=Depends(server.current_user),
    ):
        server.require(user, "view")
        return await _build_rows(server, date_from, date_to, user)

    app.add_api_route(path, mro_traceability_detail, methods=["GET"], tags=["reports"])

    @app.get("/api/reports/mro-traceability/export.xlsx", tags=["reports"])
    async def export_mro_traceability(
        date_from: str | None = Query(None),
        date_to: str | None = Query(None),
        user=Depends(server.current_user),
    ):
        server.require(user, "export")
        rows = await _build_rows(server, date_from, date_to, user)
        stream = _excel(rows)
        suffix = f"_{date_from or 'awal'}_{date_to or 'akhir'}"
        filename = f"mro-traceability{suffix}.xlsx"
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
