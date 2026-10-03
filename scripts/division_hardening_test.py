"""Tahap 2 final hardening: lookup separation, division scope on lookup/print/export/report/dashboard/search (one tenant)."""
import io
import sys
import uuid

import requests

sys.path.insert(0, "/app/backend/tests")
import receipt_control_test as T  # noqa: E402

call, check, API = T.call, T.check, T.API


class U:
    def __init__(self, email):
        self.S = requests.Session()
        tok = self.S.post(f"{API}/auth/login", json={"email": email, "password": "TestPass123!"}).json().get("token")
        self.S.headers.update({"Authorization": f"Bearer {tok}"})

    def __call__(self, method, path, body=None, raw=False):
        r = self.S.request(method, f"{API}/{path}", json=body)
        if raw:
            return r.status_code, r.content
        try:
            return r.status_code, r.json()
        except Exception:
            return r.status_code, {}


def chain(div, item, wh, proj, qty, price, sup):
    sc, mro = call("POST", "mro", {"no": f"MRO-{uuid.uuid4().hex[:6]}", "division_id": div, "submitted": True, "lines": [
        {"item_id": item, "qty": qty, "warehouse_id": wh, "project_id": proj}]}, 200)
    call("POST", f"mro/{mro['id']}/submit", {})
    sc, ro = call("POST", "ro", {"division_id": div, "submitted": True, "lines": [{"item_id": item, "qty": qty, "warehouse_id": wh,
                  "project_id": proj, "sources": [{"mro_id": mro["id"], "line_id": mro["lines"][0]["id"], "qty": qty}]}]}, 200)
    call("POST", f"ro/{ro['id']}/submit", {})
    return mro, ro


def mk_po(div, lines, sup):
    sc, po = call("POST", "po", {"supplier_id": sup, "division_id": div, "lines": lines}, 200)
    call("POST", f"po/{po['id']}/submit", {})
    if call("GET", f"po/{po['id']}")[1].get("status") != "Approved":
        call("POST", f"po/{po['id']}/approve", {})
    return po


def po_line(ro, item, wh, proj, qty, price):
    return {"item_id": item, "qty": qty, "price": price, "warehouse_id": wh, "project_id": proj,
            "sources": [{"ro_id": ro["id"], "line_id": ro["lines"][0]["id"], "qty": qty}]}


def ids(u, path):
    sc, r = u("GET", path)
    rows = r if isinstance(r, list) else (r or {}).get("items") or []
    return {x.get("id") for x in rows}


def main():
    M = T.setup()
    u = uuid.uuid4().hex[:5]
    dA = M["div"]
    sc, dB = call("POST", "master/divisions", {"code": f"DB{u}", "name": "Divisi B"}, 200)
    R = {}
    for tag, d in (("A", dA), ("B", dB)):
        for key, coll, body in (("it", "items", {"code": f"I{tag}{u}", "name": f"Barang {tag}", "unit": "PCS"}),
                                ("wh", "warehouses", {"code": f"W{tag}{u}", "name": f"Gudang Div {tag}"}),
                                ("pr", "projects", {"code": f"P{tag}{u}", "name": f"Proyek Div {tag}"}),
                                ("un", "units", {"code": f"U{tag}{u}", "name": f"Unit Div {tag}"})):
            R[key + tag] = call("POST", f"master/{coll}", {**body, "division_id": d["id"], "is_active": True}, 200)[1]
        R["spk" + tag] = call("POST", "spk", {"spk_number": f"SPK-{tag}{u}", "project_name": f"Proyek {tag}", "spk_value": 10,
                                              "procurement_budget": 5, "status": "active", "division_id": d["id"]}, 200)[1]
    sc, uom = call("POST", "master/uoms", {"code": f"G{u}", "name": f"Global{u}"}, 200)
    sc, tax = call("POST", "master/taxes", {"code": f"TX{u}", "name": f"PPN{u}", "rate": 11}, 200)
    sup = M["supX"]["id"]
    mA, rA = chain(dA["id"], R["itA"]["id"], R["whA"]["id"], R["prA"]["id"], 2, 1000, sup)
    mB, rB = chain(dB["id"], R["itB"]["id"], R["whB"]["id"], R["prB"]["id"], 3, 1000, sup)
    mA2, rA2 = chain(dA["id"], R["itA"]["id"], R["whA"]["id"], R["prA"]["id"], 1, 1000, sup)
    mB2, rB2 = chain(dB["id"], R["itB"]["id"], R["whB"]["id"], R["prB"]["id"], 1, 1000, sup)
    X, Y, Z = 2 * 1000, 3 * 1000, 1000 + 1000
    poA = mk_po(dA["id"], [po_line(rA, R["itA"]["id"], R["whA"]["id"], R["prA"]["id"], 2, 1000)], sup)
    poB = mk_po(dB["id"], [po_line(rB, R["itB"]["id"], R["whB"]["id"], R["prB"]["id"], 3, 1000)], sup)
    poAB = mk_po(dA["id"], [po_line(rA2, R["itA"]["id"], R["whA"]["id"], R["prA"]["id"], 1, 1000),
                            po_line(rB2, R["itB"]["id"], R["whB"]["id"], R["prB"]["id"], 1, 1000)], sup)
    sc, adjB = call("POST", "adjustments", {"warehouse_id": R["whB"]["id"], "reason": "Saldo awal", "lines": [
        {"item_id": R["itB"]["id"], "adjustment": 5, "approved_unit_cost": 100, "reason": "Saldo awal"}]})
    check("setup Penyesuaian Div B created", sc == 200, f"{sc} {adjB}")
    view_all = {f"{m}.view" for m in ("mro", "ro", "po", "do", "mi", "transfer", "loan", "adjustment", "opname")}
    call("PUT", "access/roles/manager", {"permissions": sorted(view_all | {"po.print", "mro.print", "ro.print", "do.print", "transfer.print", "adjustment.print", "export",
                                         "items.view", "warehouses.view", "projects.view", "units.view", "spk.view", "suppliers.view", "po.create", "view_purchase_price"}),
                                         "division_scope": {"mode": "all"}}, 200)
    call("PUT", "access/roles/warehouse", {"permissions": ["mro.view", "mro.create"], "division_scope": {"mode": "all"}}, 200)
    us = {}
    for tag, role, divs in (("A", "manager", [dA["id"]]), ("B", "manager", [dB["id"]]), ("AB", "manager", [dA["id"], dB["id"]]),
                            ("LK", "warehouse", [dA["id"]])):
        email = f"hd_{tag}_{uuid.uuid4().hex[:5]}@example.com"
        sc, nu = call("POST", "users", {"email": email, "password": "TestPass123!", "name": tag, "role": role}, 200)
        call("PUT", f"access/users/{nu['id']}", {"overrides": {}, "division_override": {"mode": "selected", "divisions": divs}}, 200)
        us[tag] = U(email)
    A, B, AB, LK = us["A"], us["B"], us["AB"], us["LK"]

    # A. lookup separation
    sc, r = LK("GET", "master/items")
    check("T1 Barang.Lihat=Tidak -> full master list 403", sc == 403 and "tidak memiliki izin" in str(r), f"{sc}")
    sc, lk = LK("GET", "lookup/items")
    got = {x["id"] for x in lk} if sc == 200 else set()
    check("T1 MRO.Tambah -> Barang lookup works, only Divisi A (+ tanpa divisi)", sc == 200 and R["itA"]["id"] in got and R["itB"]["id"] not in got, f"{sc}")
    allowed = {"id", "code", "name", "is_active", "division_id", "label", "unit", "base_uom_id", "uom_id", "uom_name",
               "uom_conversions", "conversions", "category_id", "specification", "brand", "item_type"}
    check("T1 lookup minimal fields (no tenant/audit metadata)", sc == 200 and all(set(x) <= allowed for x in lk), [set(x) - allowed for x in lk][:1])
    sc, r = LK("GET", "lookup/taxes")
    check("lookup without functional reason -> 403", sc == 403 and "memilih data Pajak" in str(r), f"{sc} {r}")
    sc, r = LK("GET", "master/warehouses")
    check("T1 Gudang full master 403 for lookup-only user", sc == 403)
    # B. division on lookup
    for name, key in (("items", "it"), ("warehouses", "wh"), ("projects", "pr"), ("units", "un"), ("spk", "spk")):
        a, b = ids(A, f"lookup/{name}"), ids(B, f"lookup/{name}")
        check(f"T2/T3 lookup {name}: A sees A not B, B sees B not A", R[key + "A"]["id"] in a and R[key + "B"]["id"] not in a
              and R[key + "B"]["id"] in b and R[key + "A"]["id"] not in b)
        q = R[key + "B"].get("code") or R[key + "B"].get("spk_number")
        check(f"T2 lookup {name} search/direct query for B code -> empty for A", not ids(A, f"lookup/{name}?q={q}&active_only=false&limit=5"))
    check("T4 global Satuan visible to A and B", uom["id"] in ids(A, "lookup/uoms") and uom["id"] in ids(B, "lookup/uoms"))
    check("T4 global Pajak visible to A and B (not division-restricted)", tax["id"] in ids(A, "lookup/taxes") and tax["id"] in ids(B, "lookup/taxes"))
    # C. print
    pr = lambda us_, dt, did: us_("POST", "verify/generate", {"doc_type": dt, "doc_id": did, "doc_no": "x"})[0]  # noqa: E731
    check("T5 PO A print: A allowed", pr(A, "PO", poA["id"]) == 200 and A("GET", f"po/{poA['id']}")[0] == 200)
    check("T5 PO A print: B blocked (detail + print)", pr(B, "PO", poA["id"]) == 403 and B("GET", f"po/{poA['id']}")[0] == 403)
    check("T6 PO A+B: A blocked", pr(A, "PO", poAB["id"]) == 403 and A("GET", f"po/{poAB['id']}")[0] == 403)
    check("T6 PO A+B: B blocked", pr(B, "PO", poAB["id"]) == 403 and B("GET", f"po/{poAB['id']}")[0] == 403)
    check("T6 PO A+B: AB allowed", pr(AB, "PO", poAB["id"]) == 200 and AB("GET", f"po/{poAB['id']}")[0] == 200)
    check("print MRO B by A blocked", pr(A, "MRO", mB["id"]) == 403 and A("GET", f"mro/{mB['id']}")[0] == 403)
    check("print RO B by A blocked", pr(A, "RO", rB["id"]) == 403)
    if adjB.get("id"):
        check("print Penyesuaian (Gudang Div B) by A blocked", pr(A, "ADJUSTMENT", adjB["id"]) == 403 and A("GET", f"adjustments/{adjB['id']}")[0] in (403, 404))
        check("Penyesuaian Div B visible for B", adjB["id"] in ids(B, "adjustments"))
    poBd = call("GET", f"po/{poB['id']}")[1]
    sc, doB = call("POST", "do", T.do_body(M, poBd, 1, "supX"))
    check("setup DO Div B created", sc == 200, f"{sc} {doB}")
    if sc == 200:
        check("print DO B: A blocked, B allowed", pr(A, "DO", doB["id"]) == 403 and A("GET", f"do/{doB['id']}")[0] == 403
              and pr(B, "DO", doB["id"]) == 200 and doB["id"] not in ids(A, "do"))
    sc, trB = call("POST", "transfers", {"from_warehouse_id": R["whB"]["id"], "to_warehouse_id": M["wh"]["id"],
                                         "lines": [{"item_id": R["itB"]["id"], "qty": 1, "unit": "PCS"}], "notes": "uji"})
    check("setup Transfer from Gudang Div B created", sc == 200, f"{sc} {trB}")
    if sc == 200:
        check("print Transfer B: A blocked, B allowed", pr(A, "TRANSFER", trB["id"]) == 403 and A("GET", f"transfers/{trB['id']}")[0] in (403, 404)
              and pr(B, "TRANSFER", trB["id"]) == 200 and trB["id"] not in ids(A, "transfers"))
    check("SPK B detail by A blocked", A("GET", f"spk/{R['spkB']['id']}")[0] == 403 and pr(A, "SPK", R["spkB"]["id"]) == 403)
    # D. export
    from openpyxl import load_workbook
    sc, blob = A("GET", "reports/mro-traceability/export.xlsx", raw=True)
    txt = ""
    if sc == 200:
        wb = load_workbook(io.BytesIO(blob), read_only=True)
        txt = " ".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None)
    check("T7 export xlsx as A: zero Divisi B rows", sc == 200 and mA["no"] in txt and mB["no"] not in txt and mB2["no"] not in txt and "Divisi B" not in txt, f"{sc}")
    # E. reports / dashboard totals
    def prem(us_):
        sc, d = us_("GET", "dashboard-premium")
        return (d.get("po") or {}).get("po_value"), (d.get("po") or {}).get("total_docs"), d
    va, na, _ = prem(A)
    vb, nb, _ = prem(B)
    vab, nab, _ = prem(AB)
    check(f"T8 premium PO total A == X ({X}) not X+Y", va == X and na == 1, (va, na))
    check(f"T8 premium PO total B == Y ({Y})", vb == Y and nb == 1, (vb, nb))
    check(f"T8 premium PO total AB == X+Y+Z ({X + Y + Z})", vab == X + Y + Z and nab == 3, (vab, nab))
    sc, rows = A("GET", "reports/mro-traceability")
    nos = {r.get("mro_no") for r in rows} if sc == 200 else set()
    check("report rows A: only Divisi A MRO", mA["no"] in nos and mB["no"] not in nos and mB2["no"] not in nos, nos)
    for p in ("reports/lead-time", "reports/unit-usage", "inventory/position"):
        sc, rows = A("GET", p)
        s = str(rows)
        check(f"report {p} A has no Divisi B data", sc == 200 and R["itB"]["code"] not in s and mB["no"] not in s and R["unB"]["code"] not in s, f"{sc}")
    da, db_, dab = A("GET", "dashboard")[1], B("GET", "dashboard")[1], AB("GET", "dashboard")[1]
    check("T9 dashboard po_outstanding A=1, B=1, AB=3", (da.get("po_outstanding"), db_.get("po_outstanding"), dab.get("po_outstanding")) == (1, 1, 3),
          (da.get("po_outstanding"), db_.get("po_outstanding"), dab.get("po_outstanding")))
    check("T9 dashboard mro_open A=2, B=2", da.get("mro_open") == db_.get("mro_open") == 2, (da.get("mro_open"), db_.get("mro_open")))
    # search
    for q, rid in ((mB["no"], mB["id"]), (poB["no"], poB["id"]), (R["itB"]["code"], R["itB"]["id"]), (R["prB"]["code"], R["prB"]["id"]),
                   (R["unB"]["code"], R["unB"]["id"])):
        sc, rows = A("GET", f"search?q={q}")
        check(f"search A for B '{q}' hidden", sc == 200 and rid not in {x.get("id") for x in rows}, rows)
    sc, rows = A("GET", "pull/ro-for-po")
    check("source picker A hides RO Divisi B", sc == 200 and rB["id"] not in {x.get("ro_id") for x in rows})
    # G. direct API A -> B
    for path in (f"po/{poB['id']}", f"mro/{mB['id']}", f"ro/{rB['id']}", f"traceability/{mB['id']}",
                 f"transactions/po/{poB['id']}/capability", f"po/{poB['id']}/email-context"):
        sc, r = A("GET", path)
        check(f"T10 direct A -> B {path.split('/')[0]} blocked", sc in (403, 404) and mB["no"] not in str(r) and poB["no"] not in str(r), f"{sc}")
    # zero-division user
    lk_id = next(x["id"] for x in call("GET", "users")[1] if x["email"].lower().startswith("hd_lk_"))
    call("DELETE", f"users/{lk_id}", None, 200)
    sc, nz = call("POST", "users", {"email": f"hd_z_{u}@example.com", "password": "TestPass123!", "name": "Z", "role": "manager"}, 200)
    call("PUT", f"access/users/{nz['id']}", {"overrides": {}, "division_override": {"mode": "selected", "divisions": []}}, 200)
    Zu = U(f"hd_z_{u}@example.com")
    vz, nzc, _ = prem(Zu)
    dz = Zu("GET", "dashboard")[1]
    check("no effective divisions -> zero dashboard/PO/list", (vz or 0) == 0 and not nzc and not dz.get("po_outstanding") and not ids(Zu, "po"), (vz, nzc))
    check("admin (Semua Divisi) sees all 3 POs", {poA["id"], poB["id"], poAB["id"]} <= ids(lambda m, p, b=None: call(m, p, b), "po"))

    passed = sum(1 for _, ok in T.RESULTS if ok)
    print(f"\n{passed}/{len(T.RESULTS)} passed")
    sys.exit(0 if passed == len(T.RESULTS) else 1)


if __name__ == "__main__":
    main()
